"""Structural checks for the concise requirements specification and user manual."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "software_delivery"
WORK = ROOT / ".codex_work" / "compact_docs"


def effective_font(run, paragraph):
    rpr = run._element.rPr
    fonts = None if rpr is None else rpr.rFonts
    east = None if fonts is None else fonts.get(qn("w:eastAsia"))
    ascii_name = None if fonts is None else fonts.get(qn("w:ascii"))
    size = run.font.size
    bold = run.bold
    style = paragraph.style
    while style is not None:
        srpr = style._element.rPr
        sfonts = None if srpr is None else srpr.rFonts
        if east is None and sfonts is not None:
            east = sfonts.get(qn("w:eastAsia"))
        if ascii_name is None and sfonts is not None:
            ascii_name = sfonts.get(qn("w:ascii"))
        if size is None:
            size = style.font.size
        if bold is None:
            bold = style.font.bold
        style = style.base_style
    return east, ascii_name, None if size is None else size.pt, bold


def check_document(path: Path, pdf: Path, expected_pages: int, body_heading: str,
                   required: tuple[str, ...], min_tables: int, min_headings: int) -> None:
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    pages = len(PdfReader(str(pdf)).pages)
    assert pages == expected_pages, (path.name, pages, expected_pages)
    assert f"共{pages}页" in text
    assert all(term in text for term in required)
    assert not any(term in text for term in ("共XX页", "版本记录", "修订记录", "V1.0"))
    assert len(doc.tables) >= min_tables

    body_started = False
    headings = 0
    font_errors = []
    heading_errors = []
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == body_heading:
            body_started = True
        if not body_started:
            continue
        is_heading = paragraph.style is not None and paragraph.style.name.startswith("Heading")
        if is_heading:
            headings += 1
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            east, ascii_name, size, bold = effective_font(run, paragraph)
            if (east, ascii_name, size) != ("楷体", "Times New Roman", 14.0):
                font_errors.append((index, run.text[:20], east, ascii_name, size))
            if is_heading and bold is not True:
                heading_errors.append((index, paragraph.text))
    assert headings >= min_headings, headings
    assert not font_errors, font_errors[:10]
    assert not heading_errors, heading_errors[:10]

    fill_errors = []
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
                fill = None if shd is None else shd.get(qn("w:fill"))
                if row_index == 0 and fill != "D9D9D9":
                    fill_errors.append((table_index, row_index, cell_index, fill))
                if row_index > 0 and fill not in (None, "auto", "FFFFFF"):
                    fill_errors.append((table_index, row_index, cell_index, fill))
    assert not fill_errors, fill_errors[:10]
    print(f"PASS {path.name} pages={pages} tables={len(doc.tables)} headings={headings}")


check_document(
    DOCS / "端到端协同跟踪原型软件_需求规格说明书.docx",
    WORK / "requirements" / "requirements_word.pdf",
    23,
    "1 范围",
    ("V8兼容视图", "40 Hz", "10 Hz", "target_valid_mask", "IMM-JPDA", "Learned-K"),
    18,
    24,
)
check_document(
    DOCS / "端到端协同跟踪原型软件_使用说明书.docx",
    WORK / "manual" / "manual_word.pdf",
    18,
    "1 软件概述与使用准备",
    ("V8四十赫兹数据使用", "40 Hz", "10 Hz", "IMM-JPDA", "INTERFACE.md"),
    11,
    23,
)
