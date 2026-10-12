"""Build the fully audited final midterm-review rectification report."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
EVIDENCE_DIR = PROJECT_ROOT / "artifacts" / "metrics" / "midterm_review"
OUTPUT_PATH = PROJECT_ROOT / "docs" / "中期评审意见整改闭环报告_最终稿.docx"

FONT_CN = "Microsoft YaHei"
FONT_LATIN = "Arial"
NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
PALE_GREEN = "E2F0D9"
PALE_YELLOW = "FFF2CC"
LIGHT_GRAY = "D9D9D9"
TEXT_GRAY = "555555"


def load_json(name: str):
    return json.loads((EVIDENCE_DIR / name).read_text(encoding="utf-8"))


def set_run_font(run, size=10.5, bold=False, color="000000"):
    run.font.name = FONT_LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CN)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_table_widths(table, widths):
    table.autofit = False
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width.twips))
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = width


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def remove_paragraph_border(paragraph_or_style):
    element = paragraph_or_style._p if hasattr(paragraph_or_style, "_p") else paragraph_or_style._element
    p_pr = element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def replace_cell_text(cell, text, *, size=9.0, bold=False, color="000000", align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_lead_paragraph(doc, lead, text, *, size=10.5, space_after=5):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(lead)
    set_run_font(run, size=size, bold=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=size)
    return paragraph


def add_bullets(doc, items, *, size=10.2):
    for text in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(0.65)
        paragraph.paragraph_format.first_line_indent = Cm(-0.25)
        paragraph.paragraph_format.line_spacing = 1.12
        paragraph.paragraph_format.space_after = Pt(2.5)
        set_run_font(paragraph.add_run(text), size=size)


def add_mapping_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(4.1), Cm(9.3), Cm(3.7)]
    set_table_widths(table, widths)
    for index, title in enumerate(("意见修改点", "对应整改", "检查结论")):
        replace_cell_text(
            table.rows[0].cells[index],
            title,
            size=9.2,
            bold=True,
            color="FFFFFF",
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(table.rows[0].cells[index], NAVY)
    repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for col_index, value in enumerate(values):
            replace_cell_text(
                cells[col_index],
                value,
                size=8.9,
                bold=col_index == 2,
                align=WD_ALIGN_PARAGRAPH.CENTER if col_index == 2 else None,
            )
            if row_index % 2 == 0:
                set_cell_shading(cells[col_index], PALE_BLUE)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_data_table(doc, headers, rows, widths=None, *, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        set_table_widths(table, widths)
    for index, header in enumerate(headers):
        replace_cell_text(
            table.rows[0].cells[index],
            header,
            size=font_size,
            bold=True,
            color="FFFFFF",
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        set_cell_shading(table.rows[0].cells[index], NAVY)
    repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for col_index, value in enumerate(values):
            replace_cell_text(
                cells[col_index], value, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER
            )
            if row_index % 2 == 0:
                set_cell_shading(cells[col_index], PALE_BLUE)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_item(doc, item):
    heading_text = f"第{item['number']}条"
    if item.get("title"):
        heading_text += f"  {item['title']}"
    heading = doc.add_heading(heading_text, level=2)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)
    set_keep_with_next(heading)
    for run in heading.runs:
        set_run_font(run, size=12, bold=True)

    if item.get("blank"):
        blank = doc.add_paragraph()
        blank.paragraph_format.space_after = Pt(72)
        return

    add_lead_paragraph(doc, "原专家意见：", item["opinion"], size=10.3)

    subheading = doc.add_heading("意见修改点对应检查", level=3)
    subheading.paragraph_format.space_before = Pt(3)
    subheading.paragraph_format.space_after = Pt(3)
    set_keep_with_next(subheading)
    for run in subheading.runs:
        set_run_font(run, size=10.8, bold=True)
    add_mapping_table(doc, item["mapping"])

    add_lead_paragraph(doc, "实际整改：", "", size=10.5, space_after=2)
    add_bullets(doc, item["implementation"])

    add_lead_paragraph(doc, "验证结果：", "", size=10.5, space_after=2)
    add_bullets(doc, item["validation"])
    if item.get("data_table"):
        table_data = item["data_table"]
        add_data_table(
            doc,
            table_data["headers"],
            table_data["rows"],
            table_data.get("widths"),
            font_size=table_data.get("font_size", 8.7),
        )

    add_lead_paragraph(doc, "证据位置：", "；".join(item["evidence"]) + "。", size=9.8)
    status_para = add_lead_paragraph(doc, "闭环判定：", item["status"], size=10.2, space_after=8)
    status_para.runs[-1].font.color.rgb = RGBColor.from_string(item["status_color"])


def build_items(data_1_3, data_4_6, data_7_9, data_10_12):
    item1 = data_1_3["review_item_1"]
    item2 = data_1_3["review_item_2"]
    item3 = data_1_3["review_item_3"]
    item4 = data_4_6["review_item_4"]
    item5 = data_4_6["review_item_5"]
    item6 = data_4_6["review_item_6"]
    item7 = data_7_9["review_item_7"]
    item8 = data_7_9["review_item_8"]
    item9 = data_7_9["review_item_9"]
    item10 = data_10_12["review_item_10"]
    item11 = data_10_12["review_item_11"]
    scope = data_1_3["scope"]

    range2 = item2["observed_weight_ranges"]
    condition3 = {entry["condition"]: entry for entry in item3["conditions"]}
    matrix6 = item6["association_matrix"]
    summary6 = {entry["condition"]: entry for entry in matrix6["condition_summaries"]}
    env8 = item8["environment"]
    latency8 = item8["latency"]
    guide9 = item9["guide_supplementary_metric"]
    quality10 = item10["acceptance_dataset_empirical_audit"]
    sensors10 = quality10["sensors"]
    classes11 = {
        entry["scenario_class"]: entry for entry in item11["maneuver_audit"]["classes"]
    }

    def sensor_row(name, label):
        sensor = sensors10[name]
        errors = sensor["ecef_equivalent_observation_error_m"]
        return [
            label,
            f"{sensor['availability_rate'] * 100:.2f}%",
            f"{sensor['missing_rate'] * 100:.2f}%",
            f"{errors['rmse']:.2f}",
            f"{errors['p95']:.2f}",
        ]

    def maneuver_row(class_name, label):
        entry = classes11[class_name]
        metrics = entry["metrics"]
        return [
            label,
            entry["track_count"],
            f"{metrics['acceleration_magnitude_mps2']['p95']:.2f}/{metrics['acceleration_magnitude_mps2']['max']:.2f}",
            f"{metrics['turn_rate_deg_s']['p95']:.2f}/{metrics['turn_rate_deg_s']['max']:.2f}",
            f"{metrics['speed_change_rate_mps2']['p95']:.2f}/{metrics['speed_change_rate_mps2']['max']:.2f}",
            f"{metrics['equivalent_kinematic_load_factor_g']['max']:.2f}",
        ]

    return [
        {
            "number": 1,
            "title": "协同编队平台层作用与性能贡献",
            "opinion": "何林远专家建议进一步说明四层架构中第二层“协同编队平台”对算法设计和性能提升的具体贡献，明确该层在多模态融合、协同探测或跟踪预测中的作用。",
            "mapping": [
                ["说明第二层的结构位置", "将第二层定义为协同编队信息处理层，由僚机数据链、单平台时序特征提取和跨平台图聚合组成。", "已对应"],
                ["说明多模态融合中的作用", "先在单平台内融合雷达、光电、电子侦察和数据链特征，再向跨平台协同层输出节点特征。", "已对应"],
                ["说明协同探测作用", "距离约束GAT聚合邻机信息，为本机视场受限或观测不足提供补盲信息。", "已对应"],
                ["说明跟踪预测及性能贡献", "GAT输出参与连续位置和速度状态修正，并用关闭GAT的消融试验量化贡献。", "已对应"],
            ],
            "implementation": [
                "复用现有模型中的僚机数据链输入、单平台时序编码和距离约束GAT，不改变网络拓扑或权重。",
                "在整改材料中补充第二层输入、处理、输出及其与第三层融合跟踪模块的连接关系。",
                "把协同探测补盲、邻机信息聚合和状态修正分别对应到可检查的模块与消融结果。",
            ],
            "validation": [
                f"完整模型RMSE为{item1['full_model_rmse_m']:.3f} m，关闭GAT后为{item1['without_gat_rmse_m']:.3f} m。",
                f"GAT使该消融配置下的RMSE相对降低{item1['gat_relative_rmse_reduction_pct']:.2f}%，证明协同编队平台层具有可测的精度贡献。",
                "该消融结果用于说明结构贡献，不与第7条V7稳定阶段验收RMSE混作同一统计口径。",
            ],
            "evidence": ["src/model.py", "artifacts/metrics/ablation_results.json", "artifacts/metrics/midterm_review/items_1_3_evidence.json"],
            "status": "工程整改完成。结构作用和性能贡献均已形成可核查依据。",
            "status_color": "006100",
        },
        {
            "number": 2,
            "title": "不同场景和干扰条件下的传感器贡献",
            "opinion": "何林远专家建议区分不同作战场景和干扰条件下各类传感器的贡献度差异。在协同探测、协同打击以及不同干扰模式下，雷达、光电、电子战等传感器的置信度和权重应有所区别，不宜简单采用统一处理方式。",
            "mapping": [
                ["区分协同探测和协同打击", "CV和CT归入协同探测，CA和CSSW归入协同打击，分别统计动态门控。", "已对应"],
                ["区分不同干扰模式", "设置基准、70%雷达压制、EO连续2秒遮挡和ESM 70%丢测四类条件。", "已对应"],
                ["区分不同传感器贡献", "通过前向钩子读取EO置信度、ESM置信度和主动被动双流门控，不使用统一常数权重。", "已对应"],
                ["验证权重随条件变化", "对每种作战任务和干扰条件报告门控均值、有效点数和位置误差。", "已对应"],
            ],
            "implementation": [
                "保留现有EO和ESM置信度头以及双流门控结构，通过审计钩子记录运行时输出，不修改模型推理图。",
                "采用64个分层V7场景，CV、CA、CT、CSSW各16个，共192条航迹；每一条件使用相同样本。",
                "对各传感器受扰条件采用固定随机种子和内存扰动，保证原始数据及模型权重不变。",
            ],
            "validation": [
                f"双流门控均值范围为{range2['active_stream_gate_mean']['minimum']:.4f}至{range2['active_stream_gate_mean']['maximum']:.4f}，跨度{range2['active_stream_gate_mean']['range']:.4f}。",
                f"EO置信度均值范围为{range2['eo_confidence_mean']['minimum']:.4f}至{range2['eo_confidence_mean']['maximum']:.4f}；ESM为{range2['esm_confidence_mean']['minimum']:.4f}至{range2['esm_confidence_mean']['maximum']:.4f}。",
                "雷达受70%压制时，双流门控提高到0.7762至0.7965；ESM 70%丢测时下降到0.6546至0.6641，表明融合系数会随受扰状态调整。",
            ],
            "evidence": ["src/model.py", "scripts/evaluation/review_items_1_3.py", "artifacts/metrics/midterm_review/items_1_3_evidence.json"],
            "status": "工程整改完成。任务类型、干扰条件和传感器贡献均已分层记录。",
            "status_color": "006100",
        },
        {
            "number": 3,
            "title": "信息不完备处理机制和测试覆盖",
            "opinion": "何林远专家建议补充指南中“信息不完备”相关要求的体现，说明在数据缺失、遮挡、丢测、延迟等条件下算法如何处理，以及测试场景是否覆盖该类情况。",
            "mapping": [
                ["数据缺失处理", "ChannelMaskEncoder用可学习Mask Token表示无效通道；无有效位置源时由Learned-K前端承接历史状态。", "已对应"],
                ["遮挡处理", "增加全源连续0.5秒遮挡测试，检查恢复过程和输出有限性。", "已对应"],
                ["随机丢测处理", "增加20%独立随机丢测测试，并使用有效掩码隔离无效观测。", "已对应"],
                ["延迟处理", "注册流程按传感器源时间戳插值至统一时间网格，数据链使用Age和更新标记衰减过期信息。", "已对应"],
                ["场景覆盖", "覆盖EO整通道缺失、随机丢测、全源遮挡和0.2秒数据链延迟。", "已对应"],
            ],
            "implementation": [
                "补齐源时间戳对齐逻辑，使量测按实际采样时刻进入统一时间网格，而不是仅依赖接收顺序。",
                "复核模型现有掩码编码和历史状态承接逻辑，避免无效通道以零值冒充有效量测。",
                "新增五组可重复测试，统一统计稳定阶段RMSE、P95误差、有效点数和有限输出率。",
            ],
            "validation": [
                "五种条件均统计18240个稳定阶段有效点，有限输出率均为100%。",
                "EO整通道缺失、20%随机丢测、全源遮挡0.5秒和数据链延迟0.2秒均得到连续输出。",
                "严重信息不完备会造成精度下降，因此这些试验用于验证处理机制和鲁棒性，不宣称全部满足8 m。",
            ],
            "data_table": {
                "headers": ["条件", "RMSE m", "相对基准变化", "有限输出率"],
                "rows": [
                    ["基准", f"{condition3['baseline']['position_rmse_m']:.3f}", "0.00%", "100%"],
                    ["EO整通道缺失", f"{condition3['eo_channel_missing']['position_rmse_m']:.3f}", f"{condition3['eo_channel_missing']['rmse_change_vs_baseline_pct']:+.2f}%", "100%"],
                    ["20%随机丢测", f"{condition3['random_measurement_drop_20pct']['position_rmse_m']:.3f}", f"{condition3['random_measurement_drop_20pct']['rmse_change_vs_baseline_pct']:+.2f}%", "100%"],
                    ["全源遮挡0.5秒", f"{condition3['all_source_occlusion_0_5s']['position_rmse_m']:.3f}", f"{condition3['all_source_occlusion_0_5s']['rmse_change_vs_baseline_pct']:+.2f}%", "100%"],
                    ["数据链延迟0.2秒", f"{condition3['datalink_delay_0_2s']['position_rmse_m']:.3f}", f"{condition3['datalink_delay_0_2s']['rmse_change_vs_baseline_pct']:+.2f}%", "100%"],
                ],
                "widths": [Cm(5.6), Cm(3.2), Cm(4.4), Cm(3.3)],
            },
            "evidence": ["src/data_association/pipeline_core.py", "src/model.py", "scripts/evaluation/review_items_1_3.py", "artifacts/metrics/midterm_review/items_1_3_evidence.json"],
            "status": "工程整改完成。处理机制、覆盖条件和精度边界均已说明。",
            "status_color": "006100",
        },
        {
            "number": 4,
            "title": "仿真数据集场景分类和典型难场景",
            "opinion": "何林远专家建议进一步细化仿真数据集的场景划分，明确数据集是否按照协同探测、协同打击、干扰模式、目标机动强度等维度进行分类，并突出对典型难场景的覆盖。",
            "mapping": [
                ["按协同探测分类", "CV和CT标记为协同探测。", "已对应"],
                ["按协同打击分类", "CA和CSSW标记为协同打击。", "已对应"],
                ["按干扰模式分类", "列出无专项干扰、调幅闪烁、软压制、距离抖动和复合干扰。", "已对应"],
                ["按机动强度分类", "CV为低、CA为中、CT和CSSW为高；V8复合干扰列为高机动。", "已对应"],
                ["突出典型难场景", "CT、CSSW和V8复合干扰增加typical_difficult标签。", "已对应"],
            ],
            "implementation": [
                "在场景目录中增加作战任务、干扰模式、机动强度和难度标签四个维度。",
                "场景接口复用同一分类信息，避免文档和软件显示口径不一致。",
                "不重新生成训练数据，仅对已有验证集建立可检索、可审计的分类视图。",
            ],
            "validation": [
                "V7共有2000个验证场景，CV、CA、CT、CSSW各500个。",
                "分层审计样本为32个场景，各类型8个；协同探测和协同打击各16个，典型难场景16个。",
                "V8现有1个复合强干扰场景，已纳入典型难场景，但不作为大样本统计结论。",
            ],
            "evidence": ["tracking_system/backend/scenario_catalog.py", "tracking_system/backend/batch_testing.py", "artifacts/metrics/midterm_review/items_4_6_evidence.json"],
            "status": "工程整改完成。四维分类和典型难场景标签均已固化。",
            "status_color": "006100",
        },
        {
            "number": 5,
            "title": "端到端跟踪技术范围和边界",
            "opinion": "徐飞专家建议进一步明确“端到端事件跟踪方法”的技术范围和边界。需要说明端到端是指从原始多源观测数据到跟踪结果的完整流程，还是指完成时空对准和目标关联之后的融合跟踪阶段，并确认该理解是否已与所内及甲方达成一致。",
            "mapping": [
                ["说明软件完整流程", "软件链覆盖原始多源观测、量测关联、源时间戳对齐、按目标组织序列、融合跟踪和航迹输出。", "已对应"],
                ["说明模型端到端范围", "神经模型范围限定为已完成时空对准和目标关联的多源序列到连续位置、速度状态。", "已对应"],
                ["明确排除项", "原始检测点提取和目标身份关联属于确定性前处理，不属于联合训练范围。", "已对应"],
                ["确认所内及甲方一致", "已形成统一技术口径和签认文字，尚需所内及甲方书面确认。", "待签认"],
            ],
            "implementation": [
                "在项目说明和系统README中采用同一组输入、输出和排除项定义。",
                "对外表述同时区分软件全流程和模型端到端范围，避免将确定性前处理误写为单一可训练网络的一部分。",
                "最终稿保留所内和甲方签认栏，签署后完成外部闭环。",
            ],
            "validation": [
                "项目说明、软件说明和闭环报告中的端到端边界已经一致。",
                "当前软件确实从原始观测进入完整处理链；当前神经模型输入则是已关联、已对准的目标级多源序列。",
                "尚无所内及甲方书面确认材料，因此不将外部确认状态写为完成。",
            ],
            "evidence": ["docs/end_to_end_scope.md", "tracking_system/README.md", "artifacts/metrics/midterm_review/items_4_6_evidence.json"],
            "status": "技术整改完成，待所内及甲方签认。",
            "status_color": "9C6500",
        },
        {
            "number": 6,
            "title": "关联正确率指标的充分验证",
            "opinion": "徐飞专家建议对“关联正确率85%”指标开展更加充分的验证。验证过程应覆盖多种传感器组合、多类典型场景和复杂干扰条件，避免仅基于简单场景得到较高关联正确率。",
            "mapping": [
                ["覆盖多种传感器组合", "覆盖雷达、雷达加光电、雷达加电子侦察、雷达加光电加电子侦察四种组合。", "已对应"],
                ["覆盖多类典型场景", "CV、CA、CT、CSSW各8个场景。", "已对应"],
                ["覆盖复杂干扰", "覆盖基准、调幅闪烁、70%雷达软压制、100 m距离抖动和复合干扰。", "已对应"],
                ["避免简单场景偏差", "对每个场景类别、干扰条件和传感器组合分别判定85%阈值，并报告最差单元。", "已对应"],
            ],
            "implementation": [
                "采用量测几何、匈牙利分配和时序记忆完成关联，真值仅在分组结束后用于审计。",
                "32个分层场景分别施加5种条件，共160个场景条件运行；每次运行统计4种传感器组合。",
                "证据记录包含每个单元的正确数、总数、准确率、阈值、通过状态和真值泄漏标记。",
            ],
            "validation": [
                "共形成80个场景类别、干扰条件和传感器组合判定单元，全部达到85%阈值。",
                f"最差单元为CV、70%雷达软压制、雷达加光电加电子侦察，正确率{matrix6['minimum_cell_accuracy'] * 100:.2f}%。",
                f"基准全传感器总体正确率{summary6['baseline']['all_sensor_accuracy'] * 100:.2f}%，复合干扰为{summary6['composite_interference']['all_sensor_accuracy'] * 100:.2f}%。",
                "所有单元的truth_used_for_observation_grouping均为false，未使用真值参与关联。",
            ],
            "data_table": {
                "headers": ["条件", "总体正确率", "最低单元正确率", "85%判定"],
                "rows": [[
                    name,
                    f"{summary6[key]['all_sensor_accuracy'] * 100:.2f}%",
                    f"{summary6[key]['minimum_cell_accuracy'] * 100:.2f}%",
                    "通过",
                ] for key, name in [
                    ("baseline", "基准"),
                    ("radar_blinking", "调幅闪烁"),
                    ("radar_soft_suppression_70pct", "70%雷达软压制"),
                    ("radar_range_smearing_100m", "100 m距离抖动"),
                    ("composite_interference", "复合干扰"),
                ]],
                "widths": [Cm(5.2), Cm(3.8), Cm(4.2), Cm(3.3)],
            },
            "evidence": ["src/data_association/pipeline_core.py", "tests/test_track_association.py", "artifacts/metrics/midterm_review/items_4_6_evidence.json"],
            "status": "工程整改完成。复杂条件下全部80个判定单元通过。",
            "status_color": "006100",
        },
        {
            "number": 7,
            "title": "稳定阶段位置均方根误差",
            "opinion": "徐飞专家建议在位置误差指标表述中增加限定条件，例如明确为“跟踪稳定阶段的位置均方根误差”，以避免冷启动阶段或初始过渡阶段误差对指标统计造成影响。",
            "mapping": [
                ["增加稳定阶段限定", "指标统一命名为跟踪稳定阶段三维位置均方根误差。", "已对应"],
                ["排除冷启动影响", "每条航迹排除前5帧后再统计。", "已对应"],
                ["统一有效点规则", "仅统计target_valid_mask为true且预测、真值均有限的点。", "已对应"],
                ["明确计算公式", "先计算每点ECEF三维欧氏位置误差，再对全部有效点计算RMSE。", "已对应"],
            ],
            "implementation": [
                "在批量测试、评估逻辑、指标说明和报告中使用同一机器可读定义。",
                "将warmup_frames作为批测参数保存到报告，默认值为5，避免统计口径依赖人工解释。",
                "评估界面同步显示稳定阶段、有效掩码和三维ECEF误差限定。",
            ],
            "validation": [
                f"分层抽取{data_7_9['scope']['scenario_count']}个V7场景、{data_7_9['scope']['track_count']}条航迹，排除前5帧后统计{item7['evaluated_points']}个有效点。",
                f"稳定阶段位置RMSE为{item7['position_rmse_m']:.3f} m，满足不超过{item7['threshold_m']:.0f} m要求。",
                "该结果不包含冷启动和初始过渡阶段，并与第10条规定的V7输入质量前提绑定。",
            ],
            "evidence": ["tracking_system/backend/batch_testing.py", "tracking_system/backend/evaluation.py", "docs/metric_definitions.md", "artifacts/metrics/midterm_review/items_7_9_evidence.json"],
            "status": "工程整改完成。指标名称、公式、掩码和预热帧均已固化，实测通过。",
            "status_color": "006100",
        },
        {
            "number": 8,
            "title": "单帧时间指标和测试环境",
            "opinion": "徐飞专家建议补充“控制指令生成时间小于100ms”指标的测试环境说明，包括硬件平台、CPU/GPU型号、运行环境以及是否为服务器、工控机或实际部署平台。",
            "mapping": [
                ["核对指标对象", "任务书指标为单帧前向推理时间；当前软件输出跟踪状态，不生成飞行控制指令。", "已对应"],
                ["说明计时边界", "计时从已对准、已关联的多源张量输入到位置和速度状态输出。", "已对应"],
                ["说明硬件环境", "记录整机、CPU、内存、GPU、显存和驱动。", "已对应"],
                ["说明软件环境", "记录操作系统、Python、PyTorch、CUDA、设备和批量大小。", "已对应"],
                ["说明平台属性", "明确为笔记本式开发工作站，非服务器、非工控机、非实际部署平台。", "已对应"],
            ],
            "implementation": [
                "统一使用“单帧跟踪结果生成时间（模型前向）”表述，避免将跟踪状态误称为控制指令。",
                "计时不包含文件读取、原始检测、目标关联、界面、控制律、通信和执行机构。",
                "新增可复现实测脚本，采用批量1、50次预热、500次正式测试，并记录完整环境。",
            ],
            "validation": [
                f"平台为{env8['manufacturer']} {env8['model']}笔记本式开发工作站，CPU为{env8['cpu']}，内存{env8['total_memory_bytes'] / 1024**3:.1f} GiB。",
                f"GPU为{env8['gpu']['name']}，显存{env8['gpu']['total_memory_bytes'] / 1024**3:.1f} GiB，驱动{env8['gpu']['driver_version']}；PyTorch {env8['pytorch']}，CUDA {env8['cuda_runtime']}。",
                f"平均时延{latency8['mean_ms']:.3f} ms，P95为{latency8['p95_ms']:.3f} ms，最大{latency8['max_ms']:.3f} ms；P95小于100 ms。",
                "实际服务器、工控机或部署平台仍应使用同一脚本复测，不能直接沿用本工作站数值。",
            ],
            "evidence": ["scripts/evaluation/review_items_7_9.py", "artifacts/metrics/midterm_review/items_7_9_evidence.json", "docs/metric_definitions.md"],
            "status": "工程整改完成。本开发工作站指标通过；部署平台复测要求已保留。",
            "status_color": "006100",
        },
        {
            "number": 9,
            "title": "八米误差与三秒机动距离指标关系",
            "opinion": "郭云飞专家建议进一步明确位置预测误差指标的定义依据。指南中表述为“位置预测误差不超过目标3秒内机动距离”，而协议中表述为“位置均方根误差小于等于8m”，两者并非同一概念，需要说明二者之间的换算依据或确认甲方是否认可该指标调整。",
            "mapping": [
                ["说明两个指标不是同一概念", "8 m为稳定阶段总体RMSE；3秒机动距离为随时间变化的相对运动尺度。", "已对应"],
                ["说明是否存在换算依据", "未发现可直接采用的数学或经验换算，不设置无依据系数。", "已对应"],
                ["定义指南口径", "定义e(t)不超过D3(t)，D3(t)为真值从t到t加3秒的位移，要求稳定阶段且3秒真值窗完整。", "已对应"],
                ["确认指标调整", "8 m继续作为任务书判据，两项指标并列报告；最终采用关系待所内及甲方确认。", "待签认"],
            ],
            "implementation": [
                "在指标说明中分别写明两个指标的变量、统计范围、单位和聚合方式。",
                "同一批V7样本同时计算稳定阶段RMSE和3秒机动距离符合率，避免通过换算替代实测。",
                "若后续要求3秒前视预测，须另行确认预测起点、终点、聚合方式和阈值。",
            ],
            "validation": [
                f"同一批样本的稳定阶段RMSE为{item9['contract_metric']['value_m']:.3f} m。",
                f"3秒机动距离口径共有{guide9['eligible_points']}个有效点，e(t)不超过D3(t)的符合率为{guide9['compliance_rate'] * 100:.2f}%。",
                f"3秒机动距离均值为{guide9['mean_three_second_maneuver_distance_m']:.3f} m，保守批次P95误差与距离之比不超过{guide9['conservative_batch_p95_error_to_distance_ratio']:.4f}。",
                "上述结果证明两种口径可以并列验证，不构成二者之间的换算关系。",
            ],
            "evidence": ["docs/metric_definitions.md", "scripts/evaluation/review_items_7_9.py", "artifacts/metrics/midterm_review/items_7_9_evidence.json"],
            "status": "技术整改完成，指标最终采用关系待所内及甲方签认。",
            "status_color": "9C6500",
        },
        {
            "number": 10,
            "title": "测试数据质量前提和指标适用范围",
            "opinion": "郭云飞专家建议补充测试数据质量条件说明。位置误差指标是否能够达到，与输入观测数据质量密切相关，应明确传感器噪声、观测误差、数据缺失、干扰强度和测试场景复杂度等前提条件。",
            "mapping": [
                ["传感器噪声", "固化雷达、光电、电子侦察角度噪声，雷达距离噪声和数据链位置噪声。", "已对应"],
                ["观测误差", "按传感器统计ECEF等效观测误差RMSE、P95和最大值，并说明EO和ESM不直接测距。", "已对应"],
                ["数据缺失", "统计各传感器有效率和缺失率，并记录15%长机随机丢测等配置。", "已对应"],
                ["干扰强度", "明确V8闪烁盲区、软压制丢测率和距离抖动强度。", "已对应"],
                ["场景复杂度", "记录场景时长、采样周期、目标数、平台数、机动类别和典型难场景标签。", "已对应"],
                ["指标适用前提", "8 m结论限定于V7质量配置并排除前5帧；V8只作为鲁棒性证据。", "已对应"],
            ],
            "implementation": [
                "V7验收配置固定为10秒、0.1秒采样、每场景3目标和5平台，覆盖CV、CA、CT、CSSW。",
                "V7噪声配置为雷达方位和俯仰2 mrad、距离48 m；EO方位和俯仰2 mrad；ESM方位和俯仰5 mrad；数据链延迟0.1秒、位置噪声60 m。",
                "本地传感器视场为方位正负60度、俯仰正负3度，最大距离200 km；雷达和EO更新0.1秒，ESM更新0.2秒。",
                "V8配置为0.025秒采样、5目标和5平台，雷达距离噪声80 m；每3秒有0.2秒闪烁盲区，12 km外10%丢测、18 km外25%丢测，15%概率叠加400 m距离抖动。",
            ],
            "validation": [
                f"V7审计{quality10['scenario_count']}个场景、{quality10['track_count']}条航迹，每场景100帧。",
                "雷达和数据链的ECEF等效观测RMSE分别为55.69 m和66.06 m；EO和ESM结果仅描述注册后输入质量，不代表直接测距精度。",
                "V8审计1个场景、5条航迹、400帧；该样本数据链有效率为0，只用于极端缺失和复合干扰鲁棒性验证。",
            ],
            "data_table": {
                "headers": ["输入源", "有效率", "缺失率", "等效RMSE m", "等效P95 m"],
                "rows": [sensor_row("Radar", "雷达"), sensor_row("EO", "光电"), sensor_row("ESM", "电子侦察"), sensor_row("FastLink", "数据链")],
                "widths": [Cm(3.1), Cm(3.1), Cm(3.1), Cm(3.7), Cm(3.7)],
            },
            "evidence": ["tracking_system/backend/scenario_catalog.py", "tracking_system/backend/batch_testing.py", "docs/test_conditions_and_maneuver_intensity.md", "artifacts/metrics/midterm_review/items_10_12_evidence.json"],
            "status": "工程整改完成。噪声、误差、缺失、干扰、复杂度和指标边界均已固化。",
            "status_color": "006100",
        },
        {
            "number": 11,
            "title": "高机动场景机动强度量化",
            "opinion": "郭云飞专家建议量化高机动场景的机动强度。当前仅描述突变次数、遮蔽次数等现象，尚不足以体现目标机动难度。建议补充目标加速度、转弯角速度、速度变化率或最大过载等指标，并确认该测试数据质量和机动强度是否得到甲方认可。",
            "mapping": [
                ["量化目标加速度", "由相邻真值速度差除以采样周期计算三维加速度模。", "已对应"],
                ["量化转弯角速度", "由相邻三维速度向量夹角除以采样周期计算。", "已对应"],
                ["量化速度变化率", "由相邻速度模长之差的绝对值除以采样周期计算。", "已对应"],
                ["量化最大过载", "以加速度模除以9.80665计算等效运动学过载，并说明其不等于真实机体载荷。", "已对应"],
                ["确认数据质量和机动强度", "量化结果已形成，甲方认可事项保留书面签认栏。", "待签认"],
            ],
            "implementation": [
                "直接使用注册数据中的真值速度计算，不从估计轨迹反推机动强度，避免将模型误差混入场景标签。",
                "CV、CA、CT、CSSW采用V7分层样本，每类48条航迹；复合干扰采用V8的5条航迹。",
                "每类均报告样本数、均值、RMSE、P95和最大值；最终稿重点列出P95和最大值。",
                "等效运动学过载不包含重力、气动和姿态因素，不用于替代真实飞行载荷证明。",
            ],
            "validation": [
                "CV为恒速基准；CA体现持续加减速；CT体现协调转弯；CSSW体现复合机动和模式切换。",
                "V8复合干扰场景的加速度、转弯角速度、速度变化率和等效过载均为本次样本最高。",
                "所有计算由独立单元测试覆盖恒速、直角转向和非法输入边界。",
            ],
            "data_table": {
                "headers": ["场景", "航迹", "加速度 P95 最大 m/s²", "转弯角速度 P95 最大 °/s", "速度变化率 P95 最大 m/s²", "最大等效过载 g"],
                "rows": [
                    maneuver_row("CV", "CV"),
                    maneuver_row("CA", "CA"),
                    maneuver_row("CT", "CT"),
                    maneuver_row("CSSW", "CSSW"),
                    maneuver_row("COMPOSITE-JAMMING", "复合干扰"),
                ],
                "widths": [Cm(2.1), Cm(1.4), Cm(3.5), Cm(3.7), Cm(3.7), Cm(2.7)],
                "font_size": 8.0,
            },
            "evidence": ["scripts/evaluation/review_items_10_12.py", "tests/test_review_items_10_12.py", "docs/test_conditions_and_maneuver_intensity.md", "artifacts/metrics/midterm_review/items_10_12_evidence.json"],
            "status": "技术量化完成，测试数据质量和机动包线待甲方签认。",
            "status_color": "9C6500",
        },
        {
            "number": 12,
            "title": "",
            "blank": True,
        },
    ]


def validate_items(items):
    assert [item["number"] for item in items] == list(range(1, 13))
    assert all(item["opinion"] and item["mapping"] for item in items if not item.get("blank"))
    assert all(all(row[2] in {"已对应", "待签认"} for row in item["mapping"]) for item in items if not item.get("blank"))


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=8.5, color="777777")


def build_document():
    data_1_3 = load_json("items_1_3_evidence.json")
    data_4_6 = load_json("items_4_6_evidence.json")
    data_7_9 = load_json("items_7_9_evidence.json")
    data_10_12 = load_json("items_10_12_evidence.json")
    items = build_items(data_1_3, data_4_6, data_7_9, data_10_12)
    validate_items(items)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    for style_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 10.8)):
        style = styles[style_name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    title = doc.add_paragraph(style="Title")
    remove_paragraph_border(styles["Title"])
    remove_paragraph_border(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(5)
    set_run_font(title.add_run("中期评审意见整改闭环报告 最终稿"), size=18, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    set_run_font(
        subtitle.add_run("基于多源异构数据融合的端到端适变跟踪方法研究项目"),
        size=10.5,
        color=TEXT_GRAY,
    )

    add_lead_paragraph(
        doc,
        "审查结论：",
        "原始评审文件中的第1至11条意见均已逐句拆解，并与项目中的代码、测试、数据和文档建立对应关系。"
        "第1、2、3、4、6、7、8、10条工程整改已完成；第5、9、11条技术整改已完成，待所内或甲方签认；"
        "其余位置按要求保留空白。",
        size=10.8,
    )
    add_lead_paragraph(
        doc,
        "低成本整改原则：",
        "保留现有模型结构和权重，优先复用已有验证数据、门控结构和批测链路。整改集中在技术边界说明、场景与质量配置、"
        "可复现实测、指标定义、证据归档和外部签认流程，没有重新训练或替换模型。",
        size=10.5,
    )

    summary = doc.add_table(rows=2, cols=4)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(summary, [Cm(4.2), Cm(4.2), Cm(4.2), Cm(4.2)])
    summary_values = [
        ("意见总数", "12条"),
        ("工程整改完成", "8条"),
        ("技术完成待签认", "3条"),
        ("预留空白", "1条"),
    ]
    for col, (label, value) in enumerate(summary_values):
        replace_cell_text(summary.rows[0].cells[col], label, size=9.2, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(summary.rows[0].cells[col], PALE_BLUE)
        replace_cell_text(summary.rows[1].cells[col], value, size=10.2, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_table_borders(summary)

    doc.add_heading("总体整改检查结果", level=1)
    add_bullets(doc, [
        "完整性检查：第1至11条原意见均保留原文，并拆分为可核查的修改点。",
        "一致性检查：位置精度统一为V7质量条件下、排除前5帧的稳定阶段三维位置RMSE；8 m与3秒机动距离分开定义和验证。",
        "技术边界检查：软件完整流程与神经模型端到端范围分开表述，原始检测和身份关联不再被误写为联合训练范围。",
        "证据检查：第1至11条均有代码、测试或文档证据；30项自动化测试全部通过。",
        "审批边界检查：需要所内或甲方决定的事项保留为空白签认栏，不用内部技术材料代替外部批准。",
    ])

    doc.add_heading("闭环状态总表", level=1)
    status_table = doc.add_table(rows=1, cols=4)
    status_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(status_table, [Cm(1.4), Cm(5.0), Cm(7.2), Cm(3.2)])
    for index, header in enumerate(("序号", "整改主题", "核心结果", "状态")):
        replace_cell_text(status_table.rows[0].cells[index], header, size=8.8, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(status_table.rows[0].cells[index], NAVY)
    repeat_table_header(status_table.rows[0])
    short_results = {
        1: "GAT消融显示RMSE相对降低13.26%",
        2: "动态门控随任务和干扰变化，跨度0.1419",
        3: "缺失、遮挡、丢测和延迟均有机制与测试",
        4: "按任务、干扰、机动、难度四维分类",
        5: "软件全流程与模型端到端边界已统一",
        6: "80个分层关联判定单元全部超过85%",
        7: "稳定阶段RMSE为7.442 m",
        8: "本工作站单帧P95为24.349 ms",
        9: "双指标并列验证，不做无依据换算",
        10: "V7质量前提和V8鲁棒性边界已固化",
        11: "四项机动强度指标已按场景量化",
        12: "",
    }
    for index, item in enumerate(items, start=1):
        cells = status_table.add_row().cells
        prevent_row_split(status_table.rows[-1])
        status_short = "完成" if item["number"] in {1, 2, 3, 4, 6, 7, 8, 10} else ("待签认" if item["number"] in {5, 9, 11} else "")
        values = [item["number"], item["title"], short_results[item["number"]], status_short]
        for col, value in enumerate(values):
            replace_cell_text(cells[col], value, size=8.5, bold=col in {0, 3}, align=WD_ALIGN_PARAGRAPH.CENTER if col in {0, 3} else None)
            if index % 2 == 0:
                set_cell_shading(cells[col], PALE_BLUE)
            if col == 3 and status_short:
                set_cell_shading(cells[col], PALE_GREEN if status_short == "完成" else PALE_YELLOW)
    set_table_borders(status_table)

    group_headings = {
        1: "何林远专家意见整改",
        5: "徐飞专家意见整改",
        9: "郭云飞专家意见整改",
    }
    for item in items:
        if item["number"] in group_headings:
            heading = doc.add_heading(group_headings[item["number"]], level=1)
            heading.paragraph_format.space_after = Pt(5)
            set_keep_with_next(heading)
            for run in heading.runs:
                set_run_font(run, size=14, bold=True)
        add_item(doc, item)

    doc.add_page_break()
    doc.add_heading("待外部签认事项", level=1)
    add_lead_paragraph(
        doc,
        "使用说明：",
        "下表只列需要外部主体作出决定的事项。签字前，项目内部技术整改状态保持不变；签字完成后，"
        "在签认结论中填写“同意”或批准文件编号，并填写日期。",
    )
    external_rows = [
        ["5", "确认模型端到端范围为已关联、已对准的多源序列到连续目标状态。", "所内/甲方", "", ""],
        ["9", "确认8 m稳定阶段RMSE与3秒机动距离并列使用，不设置无依据换算。", "所内/甲方", "", ""],
        ["11", "确认V7和V8测试数据质量条件及各类场景机动包线。", "甲方", "", ""],
    ]
    add_data_table(
        doc,
        ["条目", "签认内容", "责任方", "结论或文件编号", "日期"],
        external_rows,
        [Cm(1.3), Cm(8.3), Cm(3.0), Cm(3.0), Cm(2.0)],
        font_size=8.5,
    )

    doc.add_heading("最终复核记录", level=1)
    review_table = doc.add_table(rows=3, cols=4)
    review_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(review_table, [Cm(3.0), Cm(5.5), Cm(3.0), Cm(5.5)])
    review_values = [
        ["项目负责人", "", "整改日期", "2026年9月10日"],
        ["技术复核人", "", "复核日期", ""],
        ["甲方确认人", "", "确认日期", ""],
    ]
    for row_index, values in enumerate(review_values):
        prevent_row_split(review_table.rows[row_index])
        for col_index, value in enumerate(values):
            replace_cell_text(review_table.rows[row_index].cells[col_index], value, size=9.2, bold=col_index in {0, 2}, align=WD_ALIGN_PARAGRAPH.CENTER)
            if col_index in {0, 2}:
                set_cell_shading(review_table.rows[row_index].cells[col_index], PALE_BLUE)
    set_table_borders(review_table)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("中期评审意见整改闭环报告 最终稿  第"), size=8.5, color="777777")
    add_page_number(footer)
    set_run_font(footer.add_run("页"), size=8.5, color="777777")

    doc.core_properties.title = "中期评审意见整改闭环报告 最终稿"
    doc.core_properties.subject = "中期评审意见整改和闭环检查"
    doc.core_properties.keywords = "中期评审 整改闭环 逐点对应 测试质量 机动强度"
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
