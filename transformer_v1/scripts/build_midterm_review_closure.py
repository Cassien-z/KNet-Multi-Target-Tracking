"""Build the cumulative midterm-review closure table for items 1-12."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
EVIDENCE_1_3_PATH = (
    PROJECT_ROOT
    / "artifacts"
    / "metrics"
    / "midterm_review"
    / "items_1_3_evidence.json"
)
EVIDENCE_4_6_PATH = (
    PROJECT_ROOT
    / "artifacts"
    / "metrics"
    / "midterm_review"
    / "items_4_6_evidence.json"
)
EVIDENCE_7_9_PATH = (
    PROJECT_ROOT
    / "artifacts"
    / "metrics"
    / "midterm_review"
    / "items_7_9_evidence.json"
)
EVIDENCE_10_12_PATH = (
    PROJECT_ROOT
    / "artifacts"
    / "metrics"
    / "midterm_review"
    / "items_10_12_evidence.json"
)
OUTPUT_PATH = PROJECT_ROOT / "docs" / "中期评审意见整改闭环表_第1至12条.docx"

FONT_CN = "Microsoft YaHei"
FONT_LATIN = "Arial"
NAVY = "1F4E78"
PALE_BLUE = "DDEBF7"
PALE_GREEN = "E2F0D9"
LIGHT_GRAY = "D9D9D9"
TEXT_GRAY = "4F4F4F"


def set_run_font(run, size=10, bold=False, color="000000"):
    run.font.name = FONT_LATIN
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CN)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=65, start=90, bottom=65, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_table_widths(table, widths):
    """Apply fixed column widths to both the table grid and every cell."""
    table.autofit = False
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width.twips))
    tbl_pr = table._tbl.tblPr
    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.append(tbl_width)
    tbl_width.set(qn("w:type"), "dxa")
    tbl_width.set(qn("w:w"), str(sum(width.twips for width in widths)))
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = width


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def remove_paragraph_border(paragraph_or_style):
    element = paragraph_or_style._p if hasattr(paragraph_or_style, "_p") else paragraph_or_style._element
    p_pr = element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_cell_paragraph(cell, text, *, bold_lead=None, size=9.0, color="000000", space_after=1.5):
    paragraph = cell.add_paragraph() if cell.paragraphs[0].text else cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.05
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, size=size, bold=True, color=color)
        tail = paragraph.add_run(text[len(bold_lead):])
        set_run_font(tail, size=size, color=color)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, color=color)
    return paragraph


def replace_cell_text(cell, text, *, size=9.5, bold=False, color="000000", align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def format_body_cell(cell):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def build_document():
    data_1_3 = json.loads(EVIDENCE_1_3_PATH.read_text(encoding="utf-8"))
    data_4_6 = json.loads(EVIDENCE_4_6_PATH.read_text(encoding="utf-8"))
    data_7_9 = json.loads(EVIDENCE_7_9_PATH.read_text(encoding="utf-8"))
    data_10_12 = json.loads(EVIDENCE_10_12_PATH.read_text(encoding="utf-8"))
    item1 = data_1_3["review_item_1"]
    item2 = data_1_3["review_item_2"]
    item3 = data_1_3["review_item_3"]
    item4 = data_4_6["review_item_4"]
    item5 = data_4_6["review_item_5"]
    item6 = data_4_6["review_item_6"]
    item7 = data_7_9["review_item_7"]
    item8 = data_7_9["review_item_8"]
    item9 = data_7_9["review_item_9"]
    item10 = data_10_12["review_item_10"]
    item11 = data_10_12["review_item_11"]
    item12 = data_10_12["review_item_12"]
    scope = data_1_3["scope"]

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.05)
    section.bottom_margin = Cm(1.0)
    section.left_margin = Cm(1.15)
    section.right_margin = Cm(1.15)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    title = doc.add_paragraph(style="Title")
    remove_paragraph_border(styles["Title"])
    remove_paragraph_border(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("中期评审意见整改闭环表")
    set_run_font(title_run, size=18, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(5)
    subtitle_run = subtitle.add_run("基于多源异构数据融合的端对端适变跟踪方法研究项目  第1至12条")
    set_run_font(subtitle_run, size=10.5, color=TEXT_GRAY)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(5)
    intro.paragraph_format.line_spacing = 1.1
    intro_run = intro.add_run(
        "本表累计记录中期评审意见第1至12条的实际整改、验证范围和可复核证据。"
        "第10至12条完成测试质量前提固化、机动强度量化和协议经费科目核对；"
        "复用现有模型与数据，不修改模型权重。需甲方或财务签认的事项单列，不以内部整改代替外部批准。"
    )
    set_run_font(intro_run, size=10)

    summary = doc.add_table(rows=2, cols=5)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    summary_widths = [Cm(3.0), Cm(4.0), Cm(4.0), Cm(4.0), Cm(4.0)]
    set_table_widths(summary, summary_widths)
    summary_values = [
        ("整改范围", "专家意见第1至12条"),
        ("质量审计", "V7 64场景 / V8 1场景"),
        ("稳定阶段RMSE", f"{item7['position_rmse_m']:.3f} m"),
        ("最大机动强度", "84.30 m/s² / 8.60 g"),
        ("整改边界", "内部完成 / 外部待签认"),
    ]
    for col, (label, value) in enumerate(summary_values):
        for row in range(2):
            format_body_cell(summary.rows[row].cells[col])
        replace_cell_text(summary.rows[0].cells[col], label, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(summary.rows[0].cells[col], PALE_BLUE)
        replace_cell_text(summary.rows[1].cells[col], value, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_table_borders(summary)

    headers = ["序号", "专家意见", "整改措施", "验证结果与证据", "状态"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(1.2), Cm(5.6), Cm(7.3), Cm(10.0), Cm(1.9)]
    set_table_widths(table, widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=120, bottom=120)
        set_cell_shading(cell, NAVY)
        replace_cell_text(cell, header, size=9.5, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_table_header(table.rows[0])

    opinions = [
        "何林远专家建议进一步说明四层架构中第二层“协同编队平台”对算法设计和性能提升的具体贡献，明确该层在多模态融合、协同探测或跟踪预测中的作用。",
        "何林远专家建议区分不同作战场景和干扰条件下各类传感器的贡献度差异。在协同探测、协同打击以及不同干扰模式下，雷达、光电、电子战等传感器的置信度和权重应有所区别，不宜简单采用统一处理方式。",
        "何林远专家建议补充指南中“信息不完备”相关要求的体现，说明在数据缺失、遮挡、丢测、延迟等条件下算法如何处理，以及测试场景是否覆盖该类情况。",
    ]

    opinions.extend(
        [
            "何林远专家建议进一步细化仿真数据集的场景划分，明确数据集是否按照协同探测、协同打击、干扰模式、目标机动强度等维度进行分类，并突出对典型难场景的覆盖。",
            "徐飞专家建议进一步明确“端到端事件跟踪方法”的技术范围和边界。需要说明端到端是指从原始多源观测数据到跟踪结果的完整流程，还是指完成时空对准和目标关联之后的融合跟踪阶段，并确认该理解是否已与所内及甲方达成一致。",
            "徐飞专家建议对“关联正确率85%”指标开展更加充分的验证。验证过程应覆盖多种传感器组合、多类典型场景和复杂干扰条件，避免仅基于简单场景得到较高关联正确率。",
        ]
    )

    opinions.extend(
        [
            "郭云飞专家建议补充测试数据质量条件说明。位置误差指标是否能够达到，与输入观测数据质量密切相关，应明确传感器噪声、观测误差、数据缺失、干扰强度和测试场景复杂度等前提条件。",
            "郭云飞专家建议量化高机动场景的机动强度。当前仅描述突变次数、遮蔽次数等现象，尚不足以体现目标机动难度。建议补充目标加速度、转弯角速度、速度变化率或最大过载等指标，并确认该测试数据质量和机动强度是否得到甲方认可。",
            "胡静秋专家建议核对协议经费明细与相关财务文件要求的一致性。根据18号文，经费科目中似乎未设置“设计费”这一项，当前协议中出现该科目可能存在不规范之处。",
        ]
    )

    opinions.extend(
        [
            "徐飞专家建议在位置误差指标表述中增加限定条件，例如明确为“跟踪稳定阶段的位置均方根误差”，以避免冷启动阶段或初始过渡阶段误差对指标统计造成影响。",
            "徐飞专家建议补充“控制指令生成时间小于100ms”指标的测试环境说明，包括硬件平台、CPU/GPU型号、运行环境以及是否为服务器、工控机或实际部署平台。",
            "郭云飞专家建议进一步明确位置预测误差指标的定义依据。指南中表述为“位置预测误差不超过目标3秒内机动距离”，而协议中表述为“位置均方根误差小于等于8m”，两者并非同一概念，需要说明二者之间的换算依据或确认甲方是否认可该指标调整。",
        ]
    )

    actions = [
        [
            "架构映射：将第二层明确为协同编队信息处理层，由僚机数据链输入、单平台时序特征提取和距离约束GAT跨平台聚合组成。",
            "算法作用：多模态特征先在各平台内融合，再由GAT根据节点特征相关性聚合邻机信息，实现协同探测补盲和跟踪状态修正。",
            "修改方式：复用现有GAT结构和消融脚本，不改变网络参数及模型权重。",
        ],
        [
            "新增动态权重审计：通过前向钩子观测现有EO置信度头、ESM置信度头及双流门控α，不改变推理图和权重。",
            "任务划分：CV/CT归入协同探测，CA/CSSW归入协同打击。",
            "受扰条件：基准、雷达70%压制、EO连续2秒遮挡、ESM 70%丢测。报告分别统计各条件和任务场景下的门控均值及有效样本数。",
        ],
        [
            "缺失处理：现有ChannelMaskEncoder使用可学习Mask Token表示无效通道；无有效位置源时，Learned-K前端承接上一时刻观测状态。",
            "延迟处理：注册流程现已读取传感器源时间戳，并按源采样时刻插值到统一时间网格；数据链继续使用Age和更新标记进行时效衰减。",
            "新增测试：EO整通道缺失、20%随机丢测、全源连续0.5秒遮挡、数据链0.2秒延迟，全部采用固定随机种子并仅在内存中扰动测试输入。",
        ],
    ]

    actions.extend(
        [
            [
                "新增四维场景分类：作战任务、干扰模式、机动强度和典型难度标签；CV/CT归入协同探测，CA/CSSW归入协同打击。",
                "机动强度按生成模型分为低、中、高三级；CT和CSSW标记为典型难场景，V8复合强干扰场景单独列示。",
                "场景目录和审计脚本统一复用现有验证数据，不重新生成训练数据。",
            ],
            [
                "统一口径：软件全流程覆盖原始多源观测、量测关联、源时间戳对齐、融合跟踪和航迹输出。",
                "模型端到端范围限定为：已完成时空对准和目标关联的多源序列，直接映射为连续位置与速度状态。",
                "明确排除：原始检测点提取和目标身份关联不属于神经网络联合训练范围；上述口径已写入项目说明，所内及甲方仍需在复核栏签署确认。",
            ],
            [
                "重做分层关联审计：CV、CA、CT、CSSW各8个场景，覆盖基准、调幅闪烁、70%雷达软压制、100 m距离抖动和复合干扰。",
                "覆盖雷达、雷达+光电、雷达+电子侦察、雷达+光电+电子侦察四种组合，共80个场景类型、干扰和传感器组合判定单元。",
                "关联阶段仅使用量测几何、匈牙利分配和时序记忆；真值只在分组完成后用于审计，不参与目标关联。",
            ],
        ]
    )

    actions.extend(
        [
            [
                "固化V7验收前提：传感器噪声、视场、量测周期、15%长机随机丢测、0.1秒数据链延迟、目标/平台数量及场景类别均写入数据集目录和接口。",
                "分离指标边界：8 m仅用于V7并排除前5帧；V8复合干扰及严重信息不完备条件仅作为鲁棒性证据。",
                "新增实证审计：按传感器统计有效率、缺失率和ECEF等效观测误差，结果保存为机器可读JSON。",
            ],
            [
                "从注册真值速度按有限差分计算加速度模、三维转弯角速度、速度变化率和等效运动学过载。",
                "对CV、CA、CT、CSSW和V8复合干扰场景分别报告样本数、P95和最大值，不再只用突变或遮蔽次数描述难度。",
                "等效过载只表示轨迹运动学强度；数据质量与机动包线另设甲方签认栏。",
            ],
            [
                "核对现行技术协议原表及文件哈希，确认存在“设计费11万元”，支出合计30万元。",
                "因当前材料未包含18号文及财务解释，不擅自改挂科目或拆分金额。",
                "建立受控变更步骤：取得依据、财务确认替代科目、复算总额、甲乙双方书面审批并归档。",
            ],
        ]
    )

    actions.extend(
        [
            [
                "统一名称为“跟踪稳定阶段三维位置均方根误差”，明确ECEF三维欧氏位置误差的RMSE公式。",
                "统一统计掩码：仅使用target_valid_mask有效点，并排除每条航迹前5帧冷启动及初始过渡阶段。",
                "批测JSON新增机器可读指标定义，评估界面和项目说明同步显示稳定阶段限定条件。",
            ],
            [
                "核对任务书原文：项目指标为单帧前向推理时间；软件输出跟踪状态，不生成飞行控制指令，故统一称为“单帧跟踪结果生成时间（模型前向）”。",
                "固化计时边界：从已对准和关联的多源张量输入到位置、速度状态输出；不含文件读取、原始检测、关联、界面、控制律和通信执行链路。",
                "新增可复现实测脚本，记录整机、CPU、内存、GPU、驱动、操作系统、Python、PyTorch、CUDA、预热和迭代次数。",
            ],
            [
                "不设置换算系数：8 m RMSE与目标3秒机动距离分别作为绝对统计精度和相对机动尺度，二者并列报告。",
                "指南补充口径定义为e(t)不超过D3(t)，其中e(t)为当前输出位置误差，D3(t)为真值从t到t+3 s的位移；仅统计稳定阶段且3秒真值窗完整的点。",
                "8 m继续作为任务书判据；若甲方要求3秒前视预测口径，须书面确认预测起点、终点、聚合方式和阈值。",
            ],
        ]
    )

    # Keep the cumulative table in the expert opinion's original 1-12 order.
    opinions = opinions[:6] + opinions[9:] + opinions[6:9]
    actions = actions[:6] + actions[9:] + actions[6:9]

    item1_results = [
        f"完整模型RMSE：{item1['full_model_rmse_m']:.3f} m。",
        f"关闭GAT后RMSE：{item1['without_gat_rmse_m']:.3f} m。",
        f"GAT带来相对RMSE降低：{item1['gat_relative_rmse_reduction_pct']:.2f}%。",
        "证据：artifacts/metrics/ablation_results.json；src/model.py。",
    ]

    ranges = item2["observed_weight_ranges"]
    item2_results = [
        "64个分层场景、192条航迹完成动态审计。",
        f"双流门控α范围：{ranges['active_stream_gate_mean']['minimum']:.4f}至{ranges['active_stream_gate_mean']['maximum']:.4f}，跨度{ranges['active_stream_gate_mean']['range']:.4f}。",
        f"EO置信度均值范围：{ranges['eo_confidence_mean']['minimum']:.4f}至{ranges['eo_confidence_mean']['maximum']:.4f}。",
        f"ESM置信度均值范围：{ranges['esm_confidence_mean']['minimum']:.4f}至{ranges['esm_confidence_mean']['maximum']:.4f}。",
        "结论：融合系数随任务类别和传感器受扰状态变化，并非统一固定权重。",
    ]

    condition_names = {
        "baseline": "基准",
        "eo_channel_missing": "EO整通道缺失",
        "random_measurement_drop_20pct": "20%随机丢测",
        "all_source_occlusion_0_5s": "全源遮挡0.5秒",
        "datalink_delay_0_2s": "数据链延迟0.2秒",
    }
    item3_results = ["各条件均评估18240个稳定阶段位置点，有限输出率均为100%。"]
    for result in item3["conditions"]:
        item3_results.append(
            f"{condition_names[result['condition']]}：RMSE {result['position_rmse_m']:.3f} m，"
            f"相对基准变化{result['rmse_change_vs_baseline_pct']:+.2f}%。"
        )
    item3_results.extend(
        [
            "结论：缺失、遮挡、丢测和延迟均已有明确处理路径及可重复测试覆盖；严重信息不完备会造成精度退化，相关测试前提将在后续第10条整改中统一限定。",
            "证据文件：items_1_3_evidence.json（artifacts/metrics/midterm_review）。",
        ]
    )

    inventory = item4["full_inventory"]
    sample4 = item4["audit_sample"]
    item4_results = [
        f"V7验证集共{inventory['v7']['scenario_count']}个场景：CV、CA、CT、CSSW各500个。",
        f"关联审计样本32个，各类8个；协同探测与协同打击各16个，典型难场景{sample4['typical_difficult_scenario_count']}个。",
        "干扰维度覆盖无专项干扰、调幅闪烁、软压制、距离抖动和复合干扰。",
        "证据：scenario_catalog.py；items_4_6_evidence.json。",
    ]

    item5_results = [
        "已在项目README和技术范围说明中统一模型输入、输出、联合学习范围及确定性前处理边界。",
        "结论：模型端到端是关联与对准后的多源序列到目标状态；软件层面仍保留从原始观测到航迹输出的完整流程。",
        f"仓库口径状态：{item5['repository_baseline_status']}。外部确认仍需所内及甲方签署，不虚构已确认结论。",
        "证据：docs/end_to_end_scope.md；tracking_system/README.md。",
    ]

    item5_results[2] = "仓库技术口径已固化并记录。外部确认仍需所内及甲方签署，不虚构已确认结论。"

    matrix6 = item6["association_matrix"]
    summaries6 = {item["condition"]: item for item in matrix6["condition_summaries"]}
    worst6 = matrix6["worst_cell"]
    item6_results = [
        f"完成{item6['scenario_condition_runs']}个场景条件运行和80个分层判定单元，全部达到85%阈值。",
        f"最差单元：{worst6['scenario_class']}、70%雷达压制、三传感器组合，正确率{worst6['accuracy'] * 100:.2f}%。",
        f"基准全传感器总体正确率{summaries6['baseline']['all_sensor_accuracy'] * 100:.2f}%；复合干扰为{summaries6['composite_interference']['all_sensor_accuracy'] * 100:.2f}%。",
        "全部审计文件的真值泄漏标记为否；原始验证集和模型权重未修改。",
        "证据文件：items_4_6_evidence.json（artifacts/metrics/midterm_review）。",
    ]

    env8 = item8["environment"]
    gpu8 = env8["gpu"]
    latency8 = item8["latency"]
    guide9 = item9["guide_supplementary_metric"]
    item7_results = [
        f"完成{data_7_9['scope']['scenario_count']}个分层场景、{data_7_9['scope']['track_count']}条航迹测试；排除前5帧后统计{item7['evaluated_points']}个有效点。",
        f"稳定阶段位置RMSE为{item7['position_rmse_m']:.3f} m，满足不超过{item7['threshold_m']:.0f} m要求。",
        "证据：batch_testing.py；evaluation.py；metric_definitions.md；items_7_9_evidence.json。",
    ]
    item8_results = [
        f"平台：{env8['platform_category']}，{env8['manufacturer']} {env8['model']}；非服务器、非工控机、非实际部署平台。",
        f"CPU：{env8['cpu']}，{env8['logical_cpu_count']}逻辑处理器；内存{env8['total_memory_bytes'] / 1024**3:.1f} GiB。",
        f"GPU：{gpu8['name']}，显存{gpu8['total_memory_bytes'] / 1024**3:.1f} GiB，驱动{gpu8['driver_version']}；PyTorch {env8['pytorch']}，CUDA {env8['cuda_runtime']}。",
        f"单帧批量1、预热{latency8['warmup_iterations']}次、正式{latency8['test_iterations']}次：平均{latency8['mean_ms']:.3f} ms，P95 {latency8['p95_ms']:.3f} ms，最大{latency8['max_ms']:.3f} ms；P95小于100 ms。",
        "适用边界：本结果只代表该开发工作站，实际部署平台须用同一脚本复测。",
    ]
    item9_results = [
        "结论：不存在可直接采用的数学或经验换算关系，未将8 m RMSE改写为3秒机动距离。",
        f"同一批样本稳定阶段RMSE为{item9['contract_metric']['value_m']:.3f} m；指南补充口径共{guide9['eligible_points']}个有效点，符合率{guide9['compliance_rate'] * 100:.2f}%。",
        f"目标3秒机动距离均值为{guide9['mean_three_second_maneuver_distance_m']:.3f} m；保守批次P95误差/距离比不超过{guide9['conservative_batch_p95_error_to_distance_ratio']:.4f}。",
        "仓库内定义和双指标测试已闭环；协议与指南最终采用关系仍待所内及甲方书面确认。",
        "证据：metric_definitions.md；items_7_9_evidence.json。",
    ]

    v7_quality = item10["acceptance_dataset_empirical_audit"]
    v8_quality = item10["robustness_dataset_empirical_audit"]
    v7_sensors = v7_quality["sensors"]
    item10_results = [
        f"V7审计{v7_quality['scenario_count']}场景、{v7_quality['track_count']}航迹；每场景100帧、3目标、5平台，采样周期0.1秒。",
        "V7有效率/缺失率："
        f"雷达{v7_sensors['Radar']['availability_rate'] * 100:.2f}%/{v7_sensors['Radar']['missing_rate'] * 100:.2f}%，"
        f"光电{v7_sensors['EO']['availability_rate'] * 100:.2f}%/{v7_sensors['EO']['missing_rate'] * 100:.2f}%，"
        f"电子侦察{v7_sensors['ESM']['availability_rate'] * 100:.2f}%/{v7_sensors['ESM']['missing_rate'] * 100:.2f}%，"
        f"数据链{v7_sensors['FastLink']['availability_rate'] * 100:.2f}%/{v7_sensors['FastLink']['missing_rate'] * 100:.2f}%。",
        f"雷达/数据链ECEF等效观测RMSE为{v7_sensors['Radar']['ecef_equivalent_observation_error_m']['rmse']:.2f}/{v7_sensors['FastLink']['ecef_equivalent_observation_error_m']['rmse']:.2f} m；EO/ESM不直接测距。",
        f"V8复合干扰审计{v8_quality['scenario_count']}场景、{v8_quality['track_count']}航迹，仅作为鲁棒性证据。",
        "证据：scenario_catalog.py；test_conditions_and_maneuver_intensity.md；items_10_12_evidence.json。",
    ]

    maneuver_rows = {
        row["scenario_class"]: row for row in item11["maneuver_audit"]["classes"]
    }

    def maneuver_summary(class_name):
        metrics = maneuver_rows[class_name]["metrics"]
        return (
            f"{class_name}：加速度P95/最大{metrics['acceleration_magnitude_mps2']['p95']:.2f}/"
            f"{metrics['acceleration_magnitude_mps2']['max']:.2f} m/s²，转弯角速度P95/最大"
            f"{metrics['turn_rate_deg_s']['p95']:.2f}/{metrics['turn_rate_deg_s']['max']:.2f} °/s，"
            f"速度变化率P95/最大{metrics['speed_change_rate_mps2']['p95']:.2f}/"
            f"{metrics['speed_change_rate_mps2']['max']:.2f} m/s²，"
            f"最大等效过载{metrics['equivalent_kinematic_load_factor_g']['max']:.2f} g。"
        )

    item11_results = [
        maneuver_summary("CV") + " " + maneuver_summary("CA"),
        maneuver_summary("CT") + " " + maneuver_summary("CSSW"),
        maneuver_summary("COMPOSITE-JAMMING"),
        "量化计算和证据已完成；测试数据质量及机动包线仍需甲方书面认可。",
        "证据：test_conditions_and_maneuver_intensity.md；items_10_12_evidence.json。",
    ]

    agreement = item12["agreement_audit"]
    design_fee = agreement["design_fee_entries"][0]["amount_10k_cny"]
    item12_results = [
        f"核对确认：现行技术协议含设计费{design_fee:.0f}万元，支出合计{agreement['reported_total_10k_cny']:.0f}万元。",
        f"已记录协议SHA-256：{agreement['agreement_sha256']}，确保后续变更对应同一版本。",
        "当前提供材料未包含18号文，无法据此确定合规替代科目；未擅自改挂或拆分11万元。",
        "已形成5步受控变更清单；取得财务解释并完成甲乙双方书面变更后，方可标记外部闭环。",
        "证据：test_conditions_and_maneuver_intensity.md；items_10_12_evidence.json。",
    ]

    result_blocks = [
        item1_results,
        item2_results,
        item3_results,
        item4_results,
        item5_results,
        item6_results,
        item7_results,
        item8_results,
        item9_results,
        item10_results,
        item11_results,
        item12_results,
    ]
    statuses = ["已完成", "已完成", "已完成\n覆盖与机制验证"]

    statuses.extend(
        ["已完成", "边界已固化\n待外部签署", "已完成\n全部单元通过"]
    )
    statuses.extend(
        ["已完成\n稳定阶段通过", "已完成\n时延通过", "内部闭环\n待甲方确认"]
    )
    statuses.extend(
        ["已完成\n条件已固化", "量化已完成\n待甲方确认", "核对已完成\n待财务变更"]
    )

    for row_index in range(len(opinions)):
        cells = table.add_row().cells
        for index, width in enumerate(widths):
            format_body_cell(cells[index])
        set_table_widths(table, widths)
        prevent_row_split(table.rows[-1])
        replace_cell_text(cells[0], str(row_index + 1), size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        replace_cell_text(cells[1], opinions[row_index], size=8.9)
        cells[2].text = ""
        for paragraph_text in actions[row_index]:
            add_cell_paragraph(cells[2], paragraph_text, size=8.8, space_after=2)
        cells[3].text = ""
        for paragraph_text in result_blocks[row_index]:
            add_cell_paragraph(cells[3], paragraph_text, size=8.8, space_after=1.5)
        replace_cell_text(cells[4], statuses[row_index], size=8.8, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cells[4], PALE_GREEN)

    set_table_borders(table)

    doc.add_heading("整改结论", level=1)
    heading = doc.paragraphs[-1]
    heading.paragraph_format.space_before = Pt(5)
    heading.paragraph_format.space_after = Pt(2)
    for run in heading.runs:
        set_run_font(run, size=11, bold=True)

    conclusion = doc.add_paragraph()
    conclusion.paragraph_format.line_spacing = 1.05
    conclusion.paragraph_format.space_after = Pt(3)
    run = conclusion.add_run(
        "第1至12条均已形成仓库内整改措施、测试或核对记录和可追溯证据。第10条把8 m结果与V7输入质量条件绑定，"
        "并将V8复合干扰单列为鲁棒性证据；第11条已按场景量化加速度、转弯角速度、速度变化率和等效运动学过载；"
        "第12条确认协议存在设计费11万元、总额30万元，并建立受控变更流程。所有技术整改均复用现有权重。"
    )
    set_run_font(run, size=9.2)

    boundary = doc.add_paragraph()
    boundary.paragraph_format.line_spacing = 1.05
    boundary.paragraph_format.space_after = Pt(4)
    lead = boundary.add_run("结果使用边界：")
    set_run_font(lead, size=9.2, bold=True)
    tail = boundary.add_run(
        "第8条时延仅代表本次开发工作站；第9条指标关系、第11条数据质量与机动包线仍需甲方书面认可；"
        "第12条在取得适用18号文及财务解释前不得擅自改挂科目，须经甲乙双方书面变更。"
        "因此上述外部确认项不以仓库内技术整改代替批准。"
    )
    set_run_font(tail, size=9.2)

    doc.add_heading("待外部签认事项", level=1)
    external_heading = doc.paragraphs[-1]
    external_heading.paragraph_format.space_before = Pt(4)
    external_heading.paragraph_format.space_after = Pt(2)
    for heading_run in external_heading.runs:
        set_run_font(heading_run, size=11, bold=True)

    external_rows = [
        ["条目", "签认内容", "责任方", "签认结论", "日期"],
        ["5", "确认模型端到端范围为关联、对准后的多源序列至连续目标状态。", "所内/甲方", "", ""],
        ["9", "确认8 m稳定阶段RMSE与3秒机动距离采用并列口径，不设置无依据换算。", "所内/甲方", "", ""],
        ["11", "确认V7/V8测试数据质量条件及CV、CA、CT、CSSW、复合干扰机动包线。", "甲方", "", ""],
        ["12", "依据18号文确认11万元合规科目及是否拆分，并完成双方书面变更。", "财务/甲乙双方", "", ""],
    ]
    external = doc.add_table(rows=len(external_rows), cols=5)
    external.alignment = WD_TABLE_ALIGNMENT.CENTER
    external_widths = [Cm(1.4), Cm(12.0), Cm(4.0), Cm(4.2), Cm(3.0)]
    set_table_widths(external, external_widths)
    for row_index, values in enumerate(external_rows):
        prevent_row_split(external.rows[row_index])
        for col_index, value in enumerate(values):
            cell = external.rows[row_index].cells[col_index]
            format_body_cell(cell)
            set_cell_margins(cell, top=45, bottom=45)
            replace_cell_text(
                cell,
                value,
                size=8.7,
                bold=row_index == 0,
                color="FFFFFF" if row_index == 0 else "000000",
                align=WD_ALIGN_PARAGRAPH.CENTER if col_index != 1 else None,
            )
            if row_index == 0:
                set_cell_shading(cell, NAVY)
    repeat_table_header(external.rows[0])
    set_table_borders(external)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)

    signoff = doc.add_table(rows=2, cols=4)
    signoff.alignment = WD_TABLE_ALIGNMENT.CENTER
    signoff_data = [
        ["项目负责人", "", "整改日期", "2026年9月10日"],
        ["复核确认人", "", "确认日期", ""],
    ]
    signoff_widths = [Cm(3.2), Cm(7.0), Cm(3.2), Cm(7.0)]
    set_table_widths(signoff, signoff_widths)
    for row_index, values in enumerate(signoff_data):
        for col_index, value in enumerate(values):
            cell = signoff.rows[row_index].cells[col_index]
            format_body_cell(cell)
            set_cell_margins(cell, top=45, bottom=45)
            is_label = col_index in (0, 2)
            replace_cell_text(
                cell,
                value,
                size=9.0,
                bold=is_label,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
            if is_label:
                set_cell_shading(cell, PALE_BLUE)
    set_table_borders(signoff)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("中期评审意见整改闭环表  第1至12条")
    set_run_font(footer_run, size=8.5, color="777777")

    doc.core_properties.title = "中期评审意见整改闭环表"
    doc.core_properties.subject = "中期评审意见第1至12条整改记录"
    doc.core_properties.keywords = "中期评审 整改闭环 测试质量 机动强度 经费核对"
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
