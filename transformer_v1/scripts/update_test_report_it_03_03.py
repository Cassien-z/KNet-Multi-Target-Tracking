from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "software_delivery" / "端到端协同跟踪算法软件_测试报告.docx"
RUN_DIR = ROOT / "artifacts" / "reports" / "software_test" / "IT-03-03" / "20260925_run01"
FIG_DIR = RUN_DIR / "figures"
RESULT = json.loads((RUN_DIR / "result.json").read_text(encoding="utf-8"))
BACKUP = RUN_DIR / "端到端协同跟踪算法软件_测试报告_写入IT-03-03前备份.docx"


def set_run_font(run, bold: bool | None = None) -> None:
    run.font.name = "楷体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "楷体")
    run.font.size = Pt(14)
    if bold is not None:
        run.bold = bold


def add_paragraph(cell, text: str = "", *, bold: bool = False, center: bool = False, after: float = 5.0):
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    if text:
        set_run_font(paragraph.add_run(text), bold=bold)
    return paragraph


def add_figure(cell, filename: str, caption: str) -> None:
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(FIG_DIR / filename), width=Cm(10.2))
    add_paragraph(cell, caption, center=True, after=7)


assert RESULT["overall_passed"] is True
assert RESULT["passed_count"] == RESULT["check_count"]
assert RESULT["formal_metrics"]["improvement_pct"] >= 15.0
for figure in RESULT["figures"]:
    assert (FIG_DIR / figure).is_file(), figure

if not BACKUP.exists():
    shutil.copy2(REPORT, BACKUP)

doc = Document(REPORT)
table = doc.tables[14]
table.autofit = False

label_cell = table.cell(5, 0)
label_cell._element.clear_content()
label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
add_paragraph(label_cell, "测试结果", bold=True, center=True, after=0)

cell = table.cell(5, 1)
cell._element.clear_content()
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

add_paragraph(
    cell,
    "2026年9月25日在Windows环境使用Python 3.11.5执行IT-03-03。正式证据绑定V7完整批次"
    "46fada4af5ef4c70a03fa56ef521f8c8；Transformer结果读取同一批测报告，IMM-JPDA使用"
    "scripts/evaluation/test_rmse_improvement.py对同一场景清单重新计算。全部验收检查通过。",
)

add_paragraph(cell, "1. 比较范围和基线配置", bold=True)
add_paragraph(
    cell,
    "两种方法均覆盖V7完整2000个文件、6000个目标和570000个目标帧点。每个文件包含3个目标，序列长度为"
    "100帧，均排除前5帧并保留3×95=285个点；文件顺序、目标数量、时间窗和有效点数逐项一致。两种总体RMSE"
    "均按逐文件平方误差和有效点数加权汇总，未将逐文件百分比或逐文件RMSE直接平均。",
)
add_paragraph(
    cell,
    "IMM-JPDA使用NumPy实现的CV、CA、左转CT和右转CT四模型交互多模型滤波，时间步长dt=0.1 s；"
    "观测源为长机Radar和僚机FastLink，平方Mahalanobis门限为25.0，漏检权重为0.15。每个目标使用首帧"
    "真值初始化，该设置对基线有利并已在证据方法名中显式登记。V8原始40 Hz场景未混入本次V7比较。",
)

add_paragraph(cell, "2. 完整V7正式结果", bold=True)
add_paragraph(
    cell,
    "Transformer稳定阶段点加权三维位置RMSE为7.488174721 m，IMM-JPDA在完全相同的570000个点上的RMSE为"
    "189.127525101 m。按(189.127525101-7.488174721)/189.127525101×100%计算，RMSE降低率为"
    "96.040674293%，高于15%阈值81.040674293个百分点；两项RMSE均为有限正数。",
)
add_figure(cell, "figure_1_overall_rmse_improvement.png", "图4-13-1  V7完整集Transformer与IMM-JPDA总体RMSE及降低率")

add_paragraph(cell, "3. 运动类型与逐文件结果", bold=True)
add_paragraph(
    cell,
    "CA、CSSW、CT和CV每类均包含500个文件和142500个点。四类Transformer/IMM-JPDA RMSE及降低率依次为："
    "CA 7.040874/184.257961 m、96.178795%；CSSW 8.515877/197.417479 m、95.686361%；"
    "CT 7.281349/198.873837 m、96.338710%；CV 7.012769/174.932623 m、95.991160%。四类均显著高于15%要求。",
)
add_figure(cell, "figure_2_motion_class_comparison.png", "图4-13-2  不同运动类型的Transformer与IMM-JPDA RMSE")
add_paragraph(
    cell,
    "进一步在每个文件内部按相同3个目标和285个点计算降低率，2000/2000个文件均不低于15%。逐文件最低"
    "降低率为72.091354%（track_cssw_0322.mat），中位数为88.975018%，5百分位为82.320920%。"
    "因此总体结论并非仅由少数IMM-JPDA长尾失效场景支配。",
)
add_paragraph(
    cell,
    "IMM-JPDA逐文件RMSE中位数为62.051781 m、95百分位为469.675896 m、最大值为1103.425134 m；"
    "每文件300个可能目标更新帧中实际更新287～300帧。无论某帧是否获得新观测，基线仍按预测状态在同一"
    "285个稳态目标帧点上计误差，不删除失配或漏检后的预测帧。",
)
add_figure(cell, "figure_3_per_file_improvement.png", "图4-13-3  2000个文件的RMSE降低率分布与累计比例")

add_paragraph(cell, "4. 公式、边界、异常与复现检查", bold=True)
add_paragraph(
    cell,
    "独立复算IMM-JPDA平方误差和、总体RMSE及降低率，与证据JSON一致。构造边界中14.99%判为不通过，"
    "15.00%和15.01%判为通过；基线RMSE为0或负数、任一RMSE为NaN以及Transformer RMSE为负数时均返回"
    "无效结论，不允许形式通过。判定使用未四舍五入值，指标含义为相对基线的误差降低百分比。",
)
add_paragraph(
    cell,
    "本轮完整证据与2026年9月17日使用相同脚本、相同V7清单和相同参数生成的历史证据逐项复核："
    "IMM-JPDA总体RMSE、降低率、570000个点以及2000条逐文件结果完全一致。两份证据因所属批次不同而"
    "job_signature不同，本轮只使用与当前批次签名一致的新证据进行正式判定。",
)
add_figure(cell, "figure_4_reproducibility_and_guards.png", "图4-13-4  完整基线结果复现及阈值与异常输入防护")

add_paragraph(cell, "5. 判定", bold=True)
add_paragraph(
    cell,
    "本用例通过。V7完整范围内Transformer位置RMSE为7.488174721 m，IMM-JPDA基线为189.127525101 m，"
    "相对降低96.040674293%，满足不低于15%的技术指标。场景、目标、时间窗、预热帧、有效点和加权方式"
    "一致，基线设置、真值初始化、异常防护和证据签名均已记录；2000个逐文件结果亦全部达到15%。",
)

for paragraph in cell.paragraphs:
    for run in paragraph.runs:
        if run._element.xpath(".//a:blip"):
            continue
        set_run_font(run)

doc.save(REPORT)
print(REPORT)

