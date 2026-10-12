from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "software_delivery" / "端到端协同跟踪算法软件_测试报告.docx"
RUN_DIR = ROOT / "artifacts" / "reports" / "software_test" / "IT-03-01" / "20260924_run01"
FIG_DIR = RUN_DIR / "figures"
RESULT = json.loads((RUN_DIR / "result.json").read_text(encoding="utf-8"))
BACKUP = RUN_DIR / "端到端协同跟踪算法软件_测试报告_写入IT-03-01前备份.docx"


def set_run_font(run, bold: bool | None = None) -> None:
    run.font.name = "楷体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "楷体")
    run.font.size = Pt(14)
    if bold is not None:
        run.bold = bold


def add_paragraph(cell, text: str = "", *, bold: bool = False, center: bool = False, after: float = 5.0):
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    if text:
        set_run_font(paragraph.add_run(text), bold=bold)
    return paragraph


def add_figure(cell, filename: str, caption: str) -> None:
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(FIG_DIR / filename), width=Cm(10.2))
    add_paragraph(cell, caption, center=True, after=7)


assert RESULT["overall_passed"] is True
assert RESULT["check_count"] == RESULT["passed_count"] == 43
for figure in RESULT["figures"]:
    assert (FIG_DIR / figure).is_file(), figure

if not BACKUP.exists():
    shutil.copy2(REPORT, BACKUP)

doc = Document(REPORT)
table = doc.tables[12]
table.autofit = False

label_cell = table.cell(5, 0)
label_cell._element.clear_content()
label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
add_paragraph(label_cell, "测试结果", bold=True, center=True, after=0)

cell = table.cell(5, 1)
cell._element.clear_content()
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

add_paragraph(
    cell,
    "2026年9月24日在Windows环境执行IT-03-01。正式验证使用Python 3.11.5、PyTorch 2.1.0和"
    "NVIDIA GeForce RTX 3060 Laptop GPU，批次编号为46fada4af5ef4c70a03fa56ef521f8c8。"
    "加载冻结V7权重，SHA-256为d22427274a8b680c477f2e7074d91efd03a89a5347f33862a6c6de7ce35dc51a；"
    "V7验证数据SHA-256为df64557ec6f34018bd05c8025d4d4dfa75aa7e215fa32741056f96c6a59fb78a。"
    "共执行43项范围、统计、阈值、异常和复现检查，43项通过。",
)

add_paragraph(cell, "1. 全量范围与正式验收结果", bold=True)
add_paragraph(
    cell,
    "正式批测完整覆盖V7验证集2000个场景、6000条航迹和570000个稳态有效点；每条航迹序列长度为100帧，"
    "剔除前5帧后保留95帧，max_scenarios为空，未作有利抽样。CA、CSSW、CT和CV各包含500个场景、"
    "1500条航迹及142500个有效点。正式指标按逐航迹位置RMSE等权平均计算，结果为7.063701790 m，"
    "小于等于8 m，绝对裕量0.936298210 m，约占阈值的11.70%。",
)
add_paragraph(
    cell,
    "作为补充口径，将全部稳态点的平方误差汇总后开方，点加权位置RMSE为7.488174721 m，也小于等于8 m，"
    "绝对裕量0.511825279 m，约占阈值的6.40%。独立脚本复算的两项结果分别为7.063701790 m和"
    "7.488174721 m，与批测报告的差异不超过1×10⁻¹² m。",
)
add_figure(cell, "figure_1_overall_rmse.png", "图4-11-1  V7全量航迹等权与点加权位置RMSE")

add_paragraph(cell, "2. 运动类型分层结果", bold=True)
add_paragraph(
    cell,
    "按运动类型分层后，CA、CSSW、CT和CV的航迹等权平均RMSE依次为6.647521 m、8.125481 m、"
    "6.875572 m和6.606234 m；点加权RMSE依次为7.040874 m、8.515877 m、7.281349 m和7.012769 m。"
    "四类样本数量完全相同，因此V7总体结果没有因类别数量不均而被偏置。",
)
add_paragraph(
    cell,
    "CSSW分层的航迹等权平均RMSE为8.125481 m，单独高于8 m，说明强转弯/强切换类场景仍是当前模型的"
    "薄弱项。合同验收判据事先规定为V7完整验证集的航迹等权总体RMSE不大于8 m，未规定每一运动类型均需"
    "单独不大于8 m，因此该分层结果作为风险项披露，不改变总体判定。V8强干扰附加集未并入本用例V7总体统计。",
)
add_figure(cell, "figure_2_motion_class_rmse.png", "图4-11-2  不同运动类型的位置RMSE分层结果")

add_paragraph(cell, "3. 航迹分布与尾部误差", bold=True)
add_paragraph(
    cell,
    "6000条航迹中，航迹RMSE不大于8 m的比例为71.80%。全部稳态点的位置误差中位数为5.467232 m，"
    "95百分位为13.686215 m，最大值为61.465439 m。总体均值达到指标，但仍存在少量长尾误差，后续优化应"
    "优先关注CSSW场景和高分位样本。",
)
add_figure(cell, "figure_3_track_rmse_distribution.png", "图4-11-3  V7逐航迹RMSE分布与累计比例")

add_paragraph(cell, "4. 统计公式、边界与复现性检查", bold=True)
add_paragraph(
    cell,
    "使用两条人工航迹验证统计公式：逐航迹RMSE分别为3.535534 m和6.928203 m，航迹等权平均为"
    "5.231869 m，点加权结果为5.813777 m，证明两种口径计算链路相互独立且均可由原始平方误差复算。"
    "阈值边界7.999 m和8.000 m判为通过，8.001 m判为不通过；含NaN的合成输入会触发失败，不会被静默删除。",
)
add_paragraph(
    cell,
    "新批次与历史完整批次2ea4f9eb0f2845b599fcf4837d71a018使用相同模型和数据摘要，两项RMSE均在"
    "1×10⁻¹² m内一致。代表样例track_ca_0000.pkl_target_0连续运行3次，RMSE均为5.622602 m，"
    "输出最大差和RMSE极差均为0。",
)
add_figure(cell, "figure_4_reproducibility_and_threshold.png", "图4-11-4  全量结果复现与8 m阈值边界检查")

add_paragraph(cell, "5. 判定", bold=True)
add_paragraph(
    cell,
    "本用例通过。冻结V7模型在完整V7验证集上的航迹等权平均位置RMSE为7.063701790 m，满足不大于8 m的"
    "技术指标；补充点加权RMSE为7.488174721 m，同样满足8 m阈值。正式范围、统计公式、阈值边界、"
    "非有限值处理和复现性检查均符合测试大纲要求。CSSW分层结果高于8 m以及尾部误差已作为限制项单列记录。",
)

for paragraph in cell.paragraphs:
    for run in paragraph.runs:
        if run._element.xpath(".//a:blip"):
            continue
        set_run_font(run)

doc.save(REPORT)
print(REPORT)

