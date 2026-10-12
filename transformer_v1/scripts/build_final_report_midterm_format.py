"""Build a completion-report draft by filling the original midterm-report DOCX."""

from __future__ import annotations

import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from build_v7_final_report_docx import REPORT_FIGURES
from build_v7_technical_summary_docx import set_east_asia_font


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "基于多源异构数据融合的端对端适变跟踪方法研究_中期报告20260525.docx"
SOURCE_MD = ROOT / "docs" / "基于多源异构数据融合的端对端适变跟踪方法研究_结题报告初稿.md"
SOURCE_DOCX = ROOT / "docs" / "基于多源异构数据融合的端对端适变跟踪方法研究_结题报告初稿.docx"
OUTPUT = ROOT / "docs" / "基于多源异构数据融合的端对端适变跟踪方法研究_结题报告_中期报告格式初稿.docx"


def find_paragraph(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"未找到段落：{text}")


def clear_runs(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph_text(paragraph, text: str, *, bold: bool | None = None) -> None:
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    if bold is not None:
        run.bold = bold


def remove_between(start_paragraph, end_paragraph) -> None:
    body = start_paragraph._p.getparent()
    children = list(body)
    start_index = children.index(start_paragraph._p)
    end_index = children.index(end_paragraph._p)
    for child in children[start_index + 1 : end_index]:
        body.remove(child)


def copy_ppr(target, template) -> None:
    ppr = target._p.get_or_add_pPr()
    for child in list(ppr):
        ppr.remove(child)
    source = template._p.pPr
    if source is not None:
        for child in source:
            if child.tag != qn("w:sectPr"):
                ppr.append(deepcopy(child))


def add_inline_text(paragraph, text: str, *, default_bold: bool = False) -> None:
    pattern = re.compile(r"(\*\*.*?\*\*|`.*?`)")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            run.bold = default_bold
        token = match.group(0)
        run = paragraph.add_run(token[2:-2] if token.startswith("**") else token[1:-1])
        run.bold = token.startswith("**") or default_bold
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        run.bold = default_bold


def insert_paragraph(anchor, text: str, template, *, bold: bool = False):
    paragraph = anchor.insert_paragraph_before()
    copy_ppr(paragraph, template)
    add_inline_text(paragraph, text, default_bold=bold)
    return paragraph


def insert_heading(anchor, text: str, template):
    paragraph = anchor.insert_paragraph_before()
    copy_ppr(paragraph, template)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    set_east_asia_font(run, "仿宋_GB2312")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def _set_cell_text(cell, text: str, *, size: float = 10.5, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    # The midterm template numbers some first-column paragraphs automatically.
    # Remove that inherited list property before writing an explicit serial number.
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is not None:
        ppr.remove(num_pr)
    clear_runs(paragraph)
    lines = text.split("\n")
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.size = Pt(size)
        set_east_asia_font(run, "宋体")
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def fill_table(table: Table, rows: list[list[str]], *, header: bool = True) -> None:
    if len(rows) != len(table.rows) or any(len(row) != len(table.columns) for row in rows):
        raise ValueError("表格维度与母版不一致")
    seen = set()
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            key = cell._tc
            if key in seen:
                continue
            seen.add(key)
            _set_cell_text(cell, value, bold=header and row_index == 0)


def keep_table_rows_together(table: Table) -> None:
    """Keep a short table on one page when it fits there."""
    for row_index, row in enumerate(table.rows):
        keep = row_index < len(table.rows) - 1
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = tr_pr.find(qn("w:cantSplit"))
        if cant_split is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = keep


def insert_cloned_table(anchor, table_xml) -> Table:
    cloned = deepcopy(table_xml)
    anchor._p.addprevious(cloned)
    return Table(cloned, anchor._parent)


def insert_markdown_table(doc: Document, anchor, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            _set_cell_text(table.cell(row_index, col_index), value, size=9.5, bold=row_index == 0)
    anchor._p.addprevious(table._tbl)


def insert_picture(doc: Document, anchor, number: int, caption: str, caption_template) -> None:
    image_paragraph = anchor.insert_paragraph_before()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.first_line_indent = Cm(0)
    image_paragraph.paragraph_format.space_before = Pt(5)
    image_paragraph.paragraph_format.space_after = Pt(2)
    image_paragraph.paragraph_format.keep_with_next = True
    image_paragraph.add_run().add_picture(str(REPORT_FIGURES[number]), width=Cm(14.4))
    caption_paragraph = anchor.insert_paragraph_before()
    copy_ppr(caption_paragraph, caption_template)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.first_line_indent = Cm(0)
    caption_paragraph.paragraph_format.space_after = Pt(6)
    run = caption_paragraph.add_run(f"图{number}  {caption}")
    run.font.size = Pt(10.5)
    set_east_asia_font(run, "宋体")


def insert_reserved_figure(doc: Document, anchor, number: int, caption: str, caption_template) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.width = Cm(14.4)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_text(cell, "模型总体架构图预留位\n待插入项目组自绘的 V7 最终架构图", size=11)
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    row = table.rows[0]
    row.height = Cm(6.8)
    anchor._p.addprevious(table._tbl)
    caption_paragraph = anchor.insert_paragraph_before()
    copy_ppr(caption_paragraph, caption_template)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.first_line_indent = Cm(0)
    run = caption_paragraph.add_run(f"图{number}  {caption}")
    run.font.size = Pt(10.5)
    set_east_asia_font(run, "宋体")


def iter_docx_blocks(parent):
    """Yield paragraphs and tables in their document order."""
    parent_element = parent.element.body if isinstance(parent, Document.__class__) else None
    if hasattr(parent, "_element") and hasattr(parent._element, "body"):
        parent_element = parent._element.body
    elif isinstance(parent, _Cell):
        parent_element = parent._tc
    if parent_element is None:
        raise TypeError("不支持的DOCX容器")
    for child in parent_element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def parse_docx_sections() -> dict[str, list[tuple]]:
    """Recover the technical-summary blocks directly from the retained Word draft."""
    source = Document(SOURCE_DOCX)
    sections: dict[str, list[tuple]] = {}
    current: str | None = None
    for block in iter_docx_blocks(source):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if re.match(r"^[34]\.\d+\s+", text):
                current = text
                sections[current] = []
                continue
            if re.match(r"^[1-6](?:\s|\.)", text) and not re.match(r"^[34]\.\d+\s+", text):
                current = None
                continue
            if current is None or not text:
                continue
            if match := re.match(r"图\s*(\d+)\s{2,}(.+)", text):
                number = int(match.group(1))
                kind = "reserved" if number == 7 else "figure"
                sections[current].append((kind, number, match.group(2)))
            else:
                sections[current].append(("paragraph", text))
        elif current is not None:
            joined = "\n".join(cell.text for row in block.rows for cell in row.cells)
            if "模型总体架构图预留位" in joined:
                continue
            rows = [[cell.text.strip() for cell in row.cells] for row in block.rows]
            sections[current].append(("table", rows))
    return sections


def parse_markdown_sections() -> dict[str, list[tuple]]:
    if not SOURCE_MD.is_file():
        return parse_docx_sections()
    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    sections: dict[str, list[tuple]] = {}
    current: str | None = None
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1
        if line.startswith("### "):
            current = line[4:]
            sections[current] = []
            continue
        if line.startswith("## "):
            current = None
            continue
        if current is None or not line:
            continue
        if line.startswith("|"):
            table_rows = []
            while True:
                cells = [item.strip() for item in line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", item) for item in cells):
                    table_rows.append(cells)
                if index >= len(lines) or not lines[index].strip().startswith("|"):
                    break
                line = lines[index].strip()
                index += 1
            sections[current].append(("table", table_rows))
            continue
        if match := re.match(r"\*\*图\s*(\d+)\s*预留\*\*\s*(.+)", line):
            sections[current].append(("reserved", int(match.group(1)), match.group(2)))
            continue
        if match := re.match(r"\*\*图\s*(\d+)\*\*\s*(.+)", line):
            sections[current].append(("figure", int(match.group(1)), match.group(2)))
            continue
        sections[current].append(("paragraph", line))
    return sections


def normalize_imported_text(text: str) -> str:
    replacements = {
        "第 4 章": "后续验证结果部分",
        "第4章": "后续验证结果部分",
        "第 3 章": "本节前述技术内容",
        "第3章": "本节前述技术内容",
        "本报告只": "本节只",
        "本稿": "本报告",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def replace_static_toc(doc: Document) -> None:
    """Replace the copied midterm TOC fields after body bookmarks are rebuilt."""
    toc_lines = [
        "目录",
        "1.项目研究目标\t1",
        "（一）研究目标\t1",
        "（二）研究内容\t1",
        "（三）研究内容完成情况\t5",
        "（四）研究成效\t6",
        "2.研究工作进程\t6",
        "2.1 多模态异构数据的统一表征与融合机理\t6",
        "2.2 端到端多模态信息融合与适变跟踪框架设计\t10",
        "2.3 标准数据集生成与算法验证\t14",
        "3.取得的主要成果及应用前景分析\t22",
        "3.1取得的主要成果\t22",
        "3.2应用前景分析\t71",
        "4.文档资料情况\t76",
        "5.申报专利情况\t76",
        "6.国防科技报告清单（如有请填写）\t78",
    ]
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    sdts = doc._element.body.findall(qn("w:sdt"))
    if len(sdts) != 1:
        raise ValueError(f"目录内容控件数量异常：{len(sdts)}")
    paragraphs = sdts[0].findall(".//w:sdtContent/w:p", namespace)
    if len(paragraphs) != len(toc_lines):
        raise ValueError(f"目录行数异常：{len(paragraphs)}")
    for paragraph, value in zip(paragraphs, toc_lines):
        for child in list(paragraph):
            if child.tag != qn("w:pPr"):
                paragraph.remove(child)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = value
        run.append(text)
        paragraph.append(run)


def insert_source_sections(doc: Document, anchor, section_names: list[str], sections, body_template, minor_heading_template, caption_template) -> None:
    for section_name in section_names:
        title = re.sub(r"^\d+\.\d+\s*", "", section_name)
        insert_heading(anchor, title, minor_heading_template)
        for block in sections[section_name]:
            kind = block[0]
            if kind == "paragraph":
                insert_paragraph(anchor, normalize_imported_text(block[1]), body_template)
            elif kind == "table":
                insert_markdown_table(doc, anchor, block[1])
            elif kind == "figure":
                insert_picture(doc, anchor, block[1], block[2], caption_template)
            elif kind == "reserved":
                insert_reserved_figure(doc, anchor, block[1], block[2], caption_template)


def build() -> Path:
    shutil.copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)

    # Capture reusable format and table templates before removing body ranges.
    body_template = doc.paragraphs[168]
    minor_heading_template = doc.paragraphs[83]
    caption_template = doc.paragraphs[34]
    table_templates = {index: deepcopy(doc.tables[index]._tbl) for index in (1, 2)}
    persistent_tables = {index: doc.tables[index] for index in range(3, 11)}

    section1_completion = find_paragraph(doc, "（三）研究内容完成情况")
    section2_heading = find_paragraph(doc, "2.研究工作进程")
    target_intro = find_paragraph(doc, "按照课题研究目标要求，研究目标实现情况如下表所示：")
    target_caption = find_paragraph(doc, "表3  课题目标完成情况表")
    section3_heading = find_paragraph(doc, "3.取得的主要成果及应用前景分析")
    section4_heading = find_paragraph(doc, "4.文档资料情况")

    remove_between(section3_heading, section4_heading)
    remove_between(section2_heading, target_intro)
    remove_between(section1_completion, section2_heading)

    # Section 1 completion and final metric table.
    section1_paragraphs = [
        "项目围绕多源异构观测统一表征、端到端多模态融合跟踪、多平台协同状态修正、标准数据集生成和原型软件验证等任务完成了研究与实现。结题阶段锁定 V7 模型和完整 V7 注册验证集作为最终性能主口径，覆盖 CV、CA、CT、CSSW 四类机动，各类500个场景，共2000个场景、6000条航迹和570000个稳定阶段有效目标点。",
        "在最终主批次上，总体三维位置RMSE为7.488米；按合同关联对口径独立复算的正确率为99.90%，召回率为96.60%，F1为98.22%；与现有IMM-JPDA参考实现进行同场景比较得到96.04%的RMSE降幅；开发工作站T=1模型前向P95为20.532毫秒。相关指标均保存任务号、模型摘要、数据摘要和独立证据文件，能够由原始结果重新计算。",
        "项目同时形成了多源异构仿真数据生成与注册流程、端到端融合跟踪模型、自动化原型验证与效能评估系统，以及需求、设计、测试和使用等配套文档。现有结论适用于已定义的V7仿真分布和开发工作站测试条件，现场部署、真实传感器数据和完整软件链路实时性仍需在后续工程阶段验证。",
        "表1  项目技术指标完成情况",
    ]
    for text in section1_paragraphs:
        paragraph = insert_paragraph(section2_heading, text, body_template)
        if text.startswith("表1"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
    metric_table = insert_cloned_table(section2_heading, table_templates[1])
    fill_table(
        metric_table,
        [
            ["指标", "预期指标", "结题指标"],
            ["航迹平均位置均方根误差", "≤8米", "7.488米"],
            ["目标正确关联率", "≥85%", "99.90%（召回率96.60%）"],
            ["传感器模态", "≥3种", "3种（雷达、光电、ESM）"],
            ["典型机载作战想定", "≥2种", "2种（协同探测、协同打击）"],
            ["干扰模式", "≥3种", "3种（调幅闪烁、随距压制、距离欺骗/抖动）"],
            ["平均位置RMSE降低", "≥15%", "96.04%"],
            ["单帧前向推理时间", "≤100毫秒", "P95为20.532毫秒"],
        ],
    )
    insert_heading(section2_heading, "（四）研究成效", minor_heading_template)
    insert_paragraph(
        section2_heading,
        "项目形成了从场景建模、数据生成、时空注册、目标关联、端到端融合跟踪到批量测试和证据归档的闭环技术链。相较中期阶段，结题阶段完成了V7完整集验证、指标口径审计、参考实现复算、模型导出一致性检查和缺测条件压力测试，使成果由阶段性演示转化为可追溯、可重算的结题证据。",
        body_template,
    )

    # Section 2 maps the work summary to the three research tasks in the midterm outline.
    progress_intro = [
        "项目实施期为2025年9月至2026年9月。研究工作按照技术方案设计、算法与数据研发、系统集成和结题验证三个阶段推进，并围绕中期评审提出的指标口径、数据覆盖、对比条件和材料一致性问题进行了集中整改。最终形成的V7模型、完整V7注册验证集和配套证据文件构成结题验收的统一对象。",
        "研究进程以三项合同研究内容为主线组织：第一项完成多模态异构数据的统一表征、目标关联、野值处理和缺测补全；第二项完成可学习状态更新、主动与被动双流融合、时序Transformer和多平台图注意力协同；第三项完成典型场景数据生成、标准化验证集、原型软件和完整指标验证。各阶段形成的文档、模型、工具、论文和专利材料与后续清单保持一致。",
    ]
    for text in progress_intro:
        insert_paragraph(target_intro, text, body_template)

    progress_sections = [
        (
            "2.1 多模态异构数据的统一表征与融合机理",
            [
                "项目完成了雷达、光电、电子侦察和协同链路信息的时间整理、空间注册、有效性标注和目标级序列组织。针对不同来源采样周期和时间戳不一致问题，注册程序按统一时间网格选择PCHIP、线性插值、单点保持或无效处理，并以最大插值间隔限制不可靠外推；针对多平台局部坐标差异，程序依据平台位置姿态将量测转换到统一ECEF坐标。",
                "目标关联采用量测驱动的门控、匈牙利分配和时间记忆机制，目标真值只在关联完成后用于评分。结题阶段从原始V7场景重新组织候选量测并开展独立审计，避免直接使用注册数据中的目标身份作为分组答案。短时缺测通过有效性掩码和历史状态承接表示，观测恢复后再由模型进行状态修正。",
                "相关模块通过IT-01-02至IT-01-04受控用例复核：关联门限边界、交叉航迹历史连续性、插值分支、最大插值间隔、多平台空间注册和异常值门限均形成可重复记录。多平台注册用例中，同一目标在两个局部坐标中的数值距离约7023.71米，转换到统一ECEF坐标后的结果间距约为1.14×10^-12米，说明坐标变换实现与预设几何关系一致。该结果只验证转换程序，不代表模型跟踪精度。",
                "通过上述工作，项目形成多源异构观测时空统一表征模型、航迹与目标关联模型、野值剔除与缺失数据补全模型，为端到端融合跟踪提供来源、时间、坐标和有效性含义明确的目标级输入。",
                "**研究任务组织与阶段推进。** 项目将统一表征工作分为场景与传感器条件整理、量测注册、目标关联、异常处理和序列封装五个相互衔接的环节。前一环节的输出同时作为后一环节的输入和检查对象，使平台状态、目标真值、原始量测、关联结果及模型输入之间能够建立对应关系。该组织方式避免将数据生成、注册和关联分别作为孤立模块描述，也为后续检查误差究竟来自输入条件、关联错误还是模型估计提供了基础。",
                "项目早期首先完成雷达、光电、电子侦察和协同链路数据的字段梳理，统一场景编号、平台编号、目标编号、时间戳、坐标系和有效性标记。中期以后，工作重点由“数据能够被模型读取”转向“每一个输入值能够说明来源和含义”，补充了更新周期、视场范围、测量噪声、缺失状态和链路时效等质量条件。结题阶段进一步将这些条件与V7场景清单、注册结果和逐航迹评价记录对应，形成从原始场景到评价样本的可追溯链路。",
                "**中期评审意见的落实。** 针对评审提出的端到端范围、输入质量、协同平台作用和关联充分性问题，项目明确软件全流程从结构化观测接入到航迹输出，神经网络模型从已经完成时空注册和目标关联的目标级序列开始计算。量测门控、匈牙利分配及历史记忆属于软件前处理，不作为网络联合学习的组成部分。该边界已经同步到技术说明、软件设计和测试口径中，避免把位置精度结果反向解释为关联算法已经得到验证。",
                "关联模块的整改重点是防止目标真值参与分组过程。项目重新从原始V7场景组织候选量测，只利用量测几何关系、门控条件和历史连续性完成关联，在关联结束后再使用真值标签统计正确数、输出关联对数和真值可关联对数。由此分别得到正确率、召回率和F1，既反映错误关联，也反映漏关联；旧的观测事件纯度仅保留为过程性审计指标，不再替代合同关联口径。",
                "**异常与缺测处理。** 对超出物理范围、违反时间连续性或不满足有效标记的输入，程序先执行质量检查和野值处理，再将通道有效性掩码与数值共同传入模型。短时缺测由历史状态承接，观测恢复后通过新的有效量测修正；长时间全源不可用时，系统只能依赖状态传播，误差会随遮挡时间增长。结题材料因此同时保留标准条件结果和信息不完备压力结果，不把程序能够连续输出等同于所有缺测条件下均满足精度要求。",
                "**质量控制与证据追溯。** 统一表征相关检查由固定受控用例和真实V7场景审计两部分构成。受控用例用于确认插值分支、最大时间间隔、坐标变换和门限边界是否按设计执行；场景审计用于检查在批量数据中是否存在标记缺失、时间顺序异常、目标数量变化或坐标范围异常。两类检查解决的问题不同，前者证明实现逻辑，后者说明最终数据对象的实际状态，均不能由单一示意图替代。",
                "项目对过程图和验证图采用不同管理方式。数据流程、关联机制和坐标注册图用于说明处理顺序及接口关系，图注标明其示意性质；带有性能数字的图必须能够追溯到V7任务记录或独立统计文件。结题阶段清理了来源不明确、样本范围不清或手工填写数字的旧图，保留能够解释技术路线且不承担验收判定的过程图，并统一采用中文标题、图例和坐标标注。",
                "**阶段成果衔接。** 统一表征研究的直接成果不仅是若干预处理程序，还包括数据字段约定、质量条件、注册后的目标级序列以及对应的测试记录。这些成果分别支撑技术方案中的处理流程、设计说明书中的接口定义、测试大纲中的检查项目和测试报告中的执行结果。正式交付时如调整传感器字段、坐标基准或时间同步规则，需要同步更新注册结果和后续指标，不能只修改说明文字。",
                "**时间统一处理。** 对更新频率不同的传感器，注册程序依据源时间戳与目标时间网格的关系选择插值、单点保持或无效处理。插值方法本身不能弥补过长时间缺测，因此同时设置最大允许间隔，并把超出间隔的通道标记为不可用。该规则使模型能够区分“经过合理时间对齐的观测”和“因缺测而不应继续外推的数值”，也为压力试验中的随机丢测和连续遮挡建立了一致输入含义。",
                "**空间统一处理。** 多平台观测首先保留各自平台位置、姿态和局部测量含义，再转换到统一地固坐标系。转换过程分别检查平台基准、角度单位、轴向定义和输出范围，避免数值能够计算但坐标含义发生偏移。受控用例验证数学变换，V7场景检查批量数据范围，两者共同支撑多平台信息进入图协同模块；坐标注册误差如未被发现，会同时影响关联距离和最终位置评价。",
                "**关联连续性管理。** 单帧几何距离只能说明当前时刻候选量测的接近程度，目标交叉或量测稀疏时容易发生身份交换。项目在门控和匈牙利分配之外保留历史关联记忆，并通过交叉航迹用例检查目标接近、分离过程中的连续性。对无法可靠匹配的量测，程序允许暂不输出关联而不是强制配对，因此最终评价同时报告正确率和召回率。",
                "统一表征工作的完成度以接口可用、逻辑受控、批量数据可审计和后续指标可重算四个层面判断。仅能生成注册文件不足以说明任务完成，还需要确认字段定义已进入文档、受控用例能够复现、完整场景没有明显结构异常，并能由原始数据重新得到关联统计。结题阶段的补充工作主要围绕后三个层面展开，使阶段性程序转化为可供验收复核的研究成果。",
            ],
        ),
        (
            "2.2 端到端多模态信息融合与适变跟踪框架设计",
            [
                "项目完成了以V7权重为最终版本的端到端融合跟踪模型。模型以物理状态传播和可学习增益为前端，将先验状态、观测残差和质量信息结合；随后分别编码主动雷达全维观测和光电、电子侦察测角观测，通过双向交叉注意力和动态门控形成单机多模态特征。",
                "融合特征进入时序Transformer提取运动趋势，再由图注意力网络聚合多平台信息，残差回归头输出位置和速度修正。结题测试对八种模态有效性组合、融合门控、图邻接边界、注意力归一化、无效节点隔离以及PyTorch与ONNX Runtime六维输出一致性进行了检查，验证模型结构和软件实现按设计运行。",
                "版本演进过程中，项目先后处理了模态缺失表示、关联真值泄漏风险、批次混用、时延统计起止点和参考基线输入条件等问题。最终模型只以V7完整验证集结果作为结题主结论，V5、V6、V8和抽样测试仅保留为过程或边界证据。",
                "项目同步完成模型部署、批量任务、逐航迹结果保存、多维评估和报告导出的原型软件闭环。模型导出一致性测试中，PyTorch与ONNX Runtime六维状态输出均未超过位置1米、速度0.2米每秒的测试阈值，为后续软件调用提供了数值一致性依据。",
                "**框架形成过程。** 模型研发并非一次完成，而是经历输入语义确认、单平台多模态融合、长时序建模、多平台协同和部署验证等阶段。早期版本主要用于打通训练与推理链路，中期版本用于检查融合结构和界面可用性，结题阶段则围绕V7权重固定评价对象，补齐完整集批测、模块消融、参考实现对照和部署一致性证据。阶段版本用于说明研发过程，最终指标只引用V7，避免把不同模型和不同数据条件的结果合并。",
                "V5、V6相关记录反映了网络规模、输入组织和训练配置的探索过程；V8复合强干扰样本用于观察超出标准质量条件后的退化现象。选择V7作为结题版本的依据是其已经形成权重、注册验证集、任务配置和独立评价文件相互绑定的完整证据链，并在统一口径下达到项目总体数值要求。现有资料没有对所有历史版本开展完全相同协议的全量复测，因此报告不作历史版本性能的简单排名。",
                "**物理先验与学习模块协同。** 框架首先依据运动状态和时间间隔形成先验预测，再由网络结合观测残差和质量信息估计修正量。这样既保留状态传播的物理含义，又允许模型学习不同模态、不同运动阶段下的修正强度。最终输出为位置和速度状态，不直接产生控制指令；控制系统如何使用航迹结果属于后续系统集成范围。该边界同时用于解释模型前向时延的起止位置。",
                "主动雷达提供距离、方位等较完整的观测信息，光电和电子侦察主要提供角度或辐射特征。模型没有将三类通道视为数值意义完全相同的输入，而是分别编码后再通过交叉注意力交换信息，并结合通道有效性控制融合权重。被动通道失效时，相应特征受到掩码约束；雷达持续可用时，模型仍可依靠主动观测和历史状态运行，但其性能条件与全模态标准配置不同。",
                "**时序与协同建模。** 时序Transformer用于在连续序列中提取速度变化、转向和机动持续性等特征，图注意力网络则在平台节点之间聚合对同一目标的观测和状态信息。平台之间的连接关系、无效节点隔离和注意力归一化均通过受控测试检查。推理时关闭时序模块或图协同模块会改变既有V7权重的工作条件，因此消融结果用于说明当前模型对这些模块的依赖，不等同于各变体重新训练后的严格结构比较。",
                "项目将协同平台贡献从架构示意转化为可执行检查：一方面确认多平台输入能够经过统一坐标和时间处理进入图结构，另一方面在完整验证集上关闭GAT观察误差变化。该检查说明已训练模型利用了跨平台特征，但不能单独量化通信链路、平台数量和编队构型各自的净贡献。若后续需要确定最小平台配置，应在固定训练和测试条件下开展专门的节点数量与链路拓扑试验。",
                "**软件集成与任务管理。** 原型软件围绕数据集选择、模型注册、参数配置、推理执行、结果回放和报告导出组织功能。批量任务保存任务号、模型摘要、数据摘要、场景清单和逐航迹结果，使同一批次的总体指标、分类指标和图件能够由统一对象生成。该机制减少了人工复制数字造成的版本混用，并允许在发现统计口径问题时定位到具体任务重新计算。",
                "模型部署验证重点检查研究环境与软件运行环境之间的数值一致性。项目将V7模型导出为ONNX格式，使用固定输入同时执行PyTorch和ONNX Runtime推理，并比较六维状态输出。测试阈值用于判断导出过程是否引入不可接受的数值差异，不代表实际目标平台已经完成性能验收；目标硬件上的算力、并发、内存和通信条件仍需在工程阶段重新测量。",
                "**问题收敛与版本冻结。** 结题阶段发现的主要问题包括模态缺失标记含义不统一、部分关联统计可能受到真值字段影响、不同批次图片混用、时延边界表述不清以及参考实现输入条件不完全一致。项目分别通过接口统一、原始场景重算、任务摘要核对、独立计时和对比条件说明进行处理。问题处理结果进入代码、证据文件和文档三处，防止只修改报告措辞而实现与结果仍保持旧状态。",
                "V7冻结后，项目原则上不再以零散试验覆盖主批次结论。新增分析应引用相同权重和数据摘要，并说明是完整集、分层样本还是受控用例；如重新训练权重、更换验证集或改变前5帧排除规则，应建立新的任务身份并重新生成位置精度、关联、基线和时延证据。该规则是保证最终报告、测试报告和软件记录保持一致的基础。",
                "**模块接口与状态流转。** 模型输入由目标维、平台维、时间维和通道维共同组成，并配套有效性及质量信息。前端状态更新产生单平台的初步状态特征，双流模块融合主动与被动观测，时序模块汇总历史变化，图协同模块交换平台信息，最后由残差回归头输出状态修正。按这一顺序组织接口，便于在测试中分别构造通道屏蔽、时间长度变化和节点隔离条件，并将异常定位到相应模块。",
                "**训练结果与软件配置的对应。** 原型软件加载模型时同时读取输入维度、平台数量、时间长度和输出定义等配置信息，避免仅凭权重文件名推断网络结构。批测任务记录实际加载的模型摘要和数据摘要，报告图件再引用任务身份。该对应关系对V7结题版本尤其重要：同名模型如重新导出或重新训练，其摘要发生变化后应视为新的被测对象，不能继续沿用旧任务结果。",
                "**功能验证与性能验证分工。** 软件功能测试主要确认模型注册、数据查询、任务创建、进度管理、结果保存、回放和报告导出能够按流程运行；算法性能验证则在固定输入与统计规则下判断精度、关联、基线和时延。功能可用不自动代表性能达标，性能数字正确也不能替代软件操作链检查。结题报告将两类结果汇总说明，详细步骤仍由测试大纲和测试报告承载。",
                "模型研发阶段对异常情况采用显式记录方式。输入维度不一致、模型配置缺失、场景文件异常或推理结果非有限值时，任务应给出错误信息并停止相应批次，而不是继续生成不完整统计。对允许存在的通道缺失，则通过有效性掩码进入正常推理和压力测试。将“非法输入”和“任务允许的缺测”区分处理，有助于提高批量验证的可解释性。",
            ],
        ),
        (
            "2.3 标准数据集生成与算法验证",
            [
                "项目形成面向协同探测和协同打击任务的多源异构仿真数据生成流程，覆盖CV、CA、CT、CSSW四类机动、雷达/光电/电子侦察三类传感器模态，以及调幅闪烁、随距压制和距离欺骗/抖动三类干扰生成机制。V7原始验证集和注册验证集各包含2000个场景，各类500个场景，每场景3个目标、100帧、5个平台。",
                "结题阶段冻结V7权重、V7注册验证集、前5帧排除规则和有效点掩码，完成2000场景、6000航迹、570000有效点的完整推理。总体位置RMSE为7.488米，四类机动结果和逐航迹分布均保存；独立关联对审计得到99.90%的正确率、96.60%的召回率和98.22%的F1；同场景IMM-JPDA参考实现对比得到96.04%的RMSE降幅。",
                "开发工作站T=1模型前向独立计时采用50次预热和300次同步正式测量，P95为20.532毫秒。该指标只覆盖已准备张量进入模型至状态输出，不含检测、关联、通信、文件读写、界面和控制；非因果完整序列模型的在线逐帧精度仍需后续验证。",
                "项目还完成64场景信息不完备压力测试、完整集模块消融、场景级误差分布、高误差尾部贡献、原始计时分布和输入可用性分析。测试结果既说明V7标准条件下总体指标达到项目内部要求，也显示全源遮挡、随机丢测和额外链路延迟会造成明显退化。",
                "阶段成果已落实为技术方案、需求规格说明书、设计说明书、测试大纲、测试报告、使用说明书、结题报告初稿和中期评审整改闭环表；同时形成SCI论文初稿1篇和发明专利申请材料1套。论文尚未发表，专利尚未取得受理号，相关状态在研究进程、成果说明和申报表格中采用同一口径。",
                "**场景与数据集建设过程。** 数据生成工作从任务想定、平台编队、目标运动、传感器配置和干扰条件五个层面展开。协同探测与协同打击用于组织任务背景，CV、CA、CT和CSSW用于描述目标运动类别，三类传感器和三类干扰机制用于形成观测条件。不同层级分别承担检索、生成和评价作用，报告不把任务想定名称直接等同于某一种固定机动或某一种干扰。",
                "每个场景均保存平台和目标的时间序列、传感器量测及必要的有效性信息。注册流程再按照统一时间轴和ECEF坐标组织目标级序列，为模型推理提供固定的数据结构。项目保留原始V7场景和注册V7场景两个层次：原始场景用于关联审计和数据条件复核，注册场景用于模型批量推理。二者场景数量相同，但用途不同，不能用注册文件中的目标身份代替原始关联过程。",
                "**从阶段演示到完整验证。** 中期阶段主要通过样例和抽样结果说明技术路线已经能够运行，结题阶段则要求固定被测模型、数据范围、评价规则和运行环境。为此，项目建立2000场景完整批次，保存6000条逐航迹记录，并在统计时统一排除前5帧、仅使用有效目标点。此前来源、版本或样本范围不清的结果不再进入最终指标表，过程图只承担方法说明作用。",
                "完整验证不仅计算一个总体RMSE，还保留四类机动结果、逐航迹分布、场景级误差和高误差尾部贡献。总体指标用于合同阈值判断，分类与分布结果用于定位性能差异。例如CSSW类别误差高于其他类别，说明总体达标并不意味着所有机动条件具有相同裕量。报告同时呈现总体结论和类别差异，避免只选择有利统计量。",
                "**验证矩阵与证据层级。** 项目把验证证据分为完整集主批次、分层补充试验和受控功能用例。完整集主批次承担最终位置精度和参考实现对比；32场景关联矩阵、64场景信息不完备试验承担条件变化分析；受控用例检查坐标注册、门限、图邻接和模型导出等实现逻辑。三类证据在样本规模和结论范围上不同，正文按其实际作用引用，不将小样本压力试验扩展为完整集结论。",
                "关联评价采用正确输出关联对、全部输出关联对和全部真值可关联对三个计数，分别形成正确率、召回率和F1。位置评价采用稳定阶段有效点的三维欧氏误差，并以有效点数加权计算总体RMSE。参考实现与V7在相同场景和有效点规则下比较，但输入模态与初始化条件仍存在差异，因此96.04%的降幅作为当前配置的可复核结果，不进一步表述为严格同输入算法净增益。",
                "**时延与运行条件。** 时延测试将模型输入张量准备完成作为起点，以输出位置和速度状态作为终点，采用50次预热和300次正式同步计时。T=1结果用于对应单帧模型前向指标，T=100完整序列结果和高负载结果用于说明不同输入与负载下的计算变化。测试未包含原始检测、关联注册、文件读取、界面刷新、通信和控制环节，因此不能直接代表完整软件链或机载平台端到端时延。",
                "信息不完备试验按照光电缺失、随机丢测、全源短时遮挡和额外链路延迟等条件组织。其目的不是重新定义V7标准验证集，而是观察输入质量下降时误差变化和软件能否保持连续运行。结果表明不同缺失方式的影响并不相同，尤其全源遮挡会使模型更多依赖状态传播。后续工程应用需要根据任务允许的中断时间设置告警、降级和重新捕获策略。",
                "**结果复核与图件生成。** 结题阶段将验证图统一改为中文标题、图例和坐标轴，并由保存的JSON、逐航迹记录或批测报告重新绘制。绘图程序在生成前检查场景数量、机动类别、样本分母和任务摘要，图注说明数据来源及适用范围。过程图、场景真值图和性能验证图分别标识，避免把场景轨迹或模型结构示意误认为算法输出证据。",
                "**文档与交付准备。** 技术方案说明总体路线，需求规格说明书定义软件功能，设计说明书记录模块和接口，测试大纲规定测试内容，测试报告保存执行结果，使用说明书说明操作流程，结题报告汇总工作和技术结论。中期报告与整改闭环表作为过程资料保留。上述文档与模型、数据、工具清单采用相同成果名称和数量，论文及专利只按当前已有凭证填写，不以草稿文件替代正式发表或受理状态。",
                "项目对计算性成果提出最小复核集合要求：位置精度需保留权重、注册验证集、场景清单、预热与掩码规则、批测配置和逐航迹结果；关联指标需保留原始场景、关联规则和计数文件；时延需保留输入形态、硬件条件和原始计时样本；图件需保留数据摘要和生成程序。正式移交时还应由项目组核定介质、版本、日期、责任人和签收状态。",
                "**完成情况自评。** 从现有项目内部证据看，三项研究内容均已形成对应的模型、数据、软件或文档，并完成V7完整验证和补充边界测试。能够由当前文件直接重算的数值作为明确结果；由消融或分层样本支持的内容保留条件说明；论文发表、专利受理、现场软件验收和外部签认等事项仅按现有正式凭证填写。该分级原则贯穿研究进程、主要成果和后续表格。",
                "**批次执行流程。** 每次正式验证先确认权重、数据集、场景范围、预热帧数和有效点规则，再创建批量任务；任务完成后检查场景数、航迹数和异常记录，随后计算总体与分类指标。关联、参考实现和前向时延由独立程序生成证据，再依据模型摘要和数据条件与主批次对应。最后才生成图件和结题表格，防止先画图后补统计口径。",
                "**指标复核流程。** 位置RMSE由逐航迹有效点数和误差平方重新汇总，不对四类RMSE直接作算术平均；关联正确率由正确输出关联对除以全部输出关联对，同时以召回率检查漏关联；时延由原始同步样本计算P95，不以完整序列耗时简单除以帧数替代T=1测量。上述规则分别写入测试材料和总结报告，确保不同文档引用同一计算含义。",
                "**高误差样本处理。** 项目没有从总体统计中删除CSSW或其他高误差航迹，而是保留原始结果并通过场景级分布和尾部贡献图分析其影响。对异常值的处理只针对输入或计算错误，真实存在的困难样本仍进入总体评价。后续如对高误差场景进行数据修订或模型优化，应建立新版本重新评价，并保留与当前V7结论的清晰区分。",
                "**资料审查与一致性检查。** 结题材料形成后，项目按研究任务、成果名称、指标数字和状态表述四个方面交叉核对。研究进程中的完成事项必须能够在文档资料或模型工具清单中找到对应项；论文和专利的叙述必须与发表、受理和授权表格一致；测试报告中的V7数字必须与结题报告的表格和图件一致。当前版本按照这一原则统一填写。",
                "**测试结果的分级使用。** 对能够由完整批次直接重算的结果，报告给出明确数值和样本范围；对模块消融、输入质量和分层压力试验，只说明在对应样本与配置下观察到的变化；对真实部署、现场功能和极端干扰条件，则作为后续验证事项。采用分级表述能够避免将受控用例、局部样本和完整集结果混为同一证据强度。",
                "**评审问题闭环。** 中期评审意见并非只在文字层面答复，而是分别转化为数据检查、受控功能用例、补充压力试验或边界说明。能够通过当前代码和数据验证的问题形成运行记录；需要外部确认的指标解释和机动包线事项保留签认条件。研究进程、技术总结和整改材料因此采用相同结论，不把内部完成状态写成外部已经批准。",
                "**成果归档层次。** 项目资料可按任务依据、技术方案、实现设计、测试记录和总结结论五个层次整理。结题报告位于证据链末端，用于汇总工作和解释结果，不替代设计说明与测试报告中的细节。归档时应同时保存可阅读文档和必要的机器可读结果，使审查人员既能了解技术路线，也能在需要时复核关键统计。",
                "**验收准备工作的衔接。** 在正式提交前，还需由项目组依据甲方最终要求核对封面信息、签字盖章、交付介质和成果凭证。该类管理信息与技术计算分开维护：补充项目编号、责任人或签收状态不改变V7主批次数字；如甲方改变评价场景、指标解释或被测平台，则需要重新确定验证任务并更新相关图表。通过区分版式完善、状态更新和技术口径变化，可以避免在成稿过程中无意改变已经固定的统计结论。",
            ],
        ),
    ]
    for heading, paragraphs in progress_sections:
        insert_heading(target_intro, heading, minor_heading_template)
        for text in paragraphs:
            insert_paragraph(target_intro, text, body_template)

    progress_caption = insert_paragraph(target_intro, "表2  研究进度安排及完成情况", body_template)
    progress_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    progress_caption.paragraph_format.page_break_before = True
    progress_table = insert_cloned_table(target_intro, table_templates[2])
    fill_table(
        progress_table,
        [
            ["节点", "进度", "进度要求", "完成情况"],
            ["1", "2025.9—2025.12", "完成多模态异构数据、端到端融合跟踪和标准数据集生成技术调研，形成总体框架和技术方案。", "完成需求分析、总体技术路线和模块划分；完成《基于多源异构数据融合的端到端适变跟踪方法研究技术方案》。"],
            ["2", "2026.1—2026.5", "设计端到端智能信息融合算法，研究典型任务场景下的数据生成与验证方法，完成软件测试大纲。", "完成统一表征、目标关联、双流融合、时序Transformer、图注意力协同和数据生成模块；完成阶段模型、验证流程及软件测试大纲。"],
            ["3", "2026.6—2026.9", "完成算法软件功能与性能测试、验收文档、SCI论文1篇和专利申请书1篇，完成软件配套文档、测试报告和验收技术报告。", "完成V7模型与完整验证集测试；完成需求、设计、测试大纲、测试报告、使用说明书和结题报告初稿；形成SCI论文初稿1篇和发明专利申请材料1套，论文尚未发表、专利尚未取得受理号；结题材料正在汇总完善。"],
        ],
    )
    keep_table_rows_together(progress_table)
    set_paragraph_text(target_intro, "合同研究工作进程已按计划完成，结题验收材料正在统一整理。")
    set_paragraph_text(target_caption, "表3  课题目标完成情况表")

    target_table = persistent_tables[3]
    fill_table(
        target_table,
        [
            ["课题名称", "研究内容", "目标实现情况", "成果说明"],
            ["基于多源异构数据融合的端到端适变跟踪方法研究", "研究内容1", "完成雷达、光电、电子侦察和协同链路观测的时间整理、空间注册、目标关联、异常处理与缺测补全；通过受控用例和原始V7场景关联审计验证实现边界。", "模型、数据库：\n1套多源异构观测时空统一表征模型；1套航迹与目标关联模型；1套野值剔除与缺失数据补全模型。"],
            ["基于多源异构数据融合的端到端适变跟踪方法研究", "研究内容2", "完成可学习状态更新、主动/被动双流融合、时序Transformer和多平台图注意力协同构成的V7模型；完整集RMSE为7.488米，关联对正确率99.90%，参考实现对比降幅96.04%，T=1前向P95为20.532毫秒。", "模型、工具：\n1套端到端多模态智能融合与适变跟踪核心模型；1套基于时空神经网络的深度学习模型；1套复杂机动与强干扰下的高精度状态估计模型；1套融合跟踪自动化原型验证与效能评估系统。"],
            ["基于多源异构数据融合的端到端适变跟踪方法研究", "研究内容3", "形成协同探测和协同打击两类想定的数据生成流程；V7原始与注册验证集各2000场景，覆盖四类机动、三类传感器模态和三类干扰机制；完成完整集、压力条件、消融、基线和时延验证。", "模型、数据库：\n1套机载协同场景下的多源异构仿真数据库及标准化生成流程；形成V7完整验证集、逐航迹结果和独立评价证据。"],
        ],
    )

    # Section 3 maps the former technical summary and V7 validation into 3.1.
    sections = parse_markdown_sections()
    section1_tech = [name for name in sections if re.match(r"3\.(?:[1-4])\s", name)]
    section2_tech = [name for name in sections if re.match(r"3\.(?:[5-9]|1[0-3])\s", name)]
    validation = [name for name in sections if re.match(r"4\.(?:[1-9]|1[0-9])\s", name)]

    insert_heading(section4_heading, "3.1取得的主要成果", minor_heading_template)
    insert_heading(section4_heading, "3.1.1 技术研发", minor_heading_template)
    insert_paragraph(
        section4_heading,
        "围绕多源异构观测、复杂机动目标和多平台协同跟踪问题，项目形成了从场景与数据条件、观测注册和目标关联，到双流融合、时序建模、图协同和状态输出的完整技术方案。以下内容按最终V7实现和结题证据重新归纳。",
        body_template,
    )
    insert_source_sections(doc, section4_heading, section1_tech, sections, body_template, minor_heading_template, caption_template)

    insert_heading(section4_heading, "3.1.2 算法与软件开发", minor_heading_template)
    insert_source_sections(doc, section4_heading, section2_tech + validation, sections, body_template, minor_heading_template, caption_template)

    insert_heading(section4_heading, "3.1.3 报告与论文专利", minor_heading_template)
    reports_and_ip = [
        "围绕项目研究、软件交付和验收复核，已形成《基于多源异构数据融合的端到端适变跟踪方法研究技术方案》《端到端协同跟踪算法软件需求规格说明书》《端到端协同跟踪算法软件设计说明书》《端到端协同跟踪算法软件测试大纲》《端到端协同跟踪算法软件测试报告》《端到端协同跟踪算法软件使用说明书》以及本结题报告初稿。中期报告和中期评审意见整改闭环表继续作为过程资料归档。",
        "论文方面，已形成《多无人机异构数据端到端适变协同跟踪网络》SCI论文初稿1篇，内容围绕物理先验状态更新、主动与被动双流交叉注意力、长时序Transformer和协同图注意力网络展开。截至本稿统计日期，该论文尚未正式发表，因此不计入论文发表数量。",
        "专利方面，已形成“一种基于物理先验与时空图注意力的多无人机异构数据端到端协同跟踪方法”发明专利申请材料1套，包括权利要求书、说明书、说明书摘要和附图等文件。截至本稿统计日期尚未取得受理号和授权号，因此在专利受理表和授权表中不列为已受理或已授权成果。",
        "上述状态与表2研究进度、表4研究报告清单、专利申报情况和论文发表情况保持一致：已完成的材料按“已形成”表述，只有取得正式编号的成果才计入受理、授权、登记或发表数量。",
    ]
    for text in reports_and_ip:
        insert_paragraph(section4_heading, text, body_template)

    insert_heading(section4_heading, "3.2应用前景分析", minor_heading_template)
    application_paragraphs = [
        "本项目形成的多源异构观测统一表征、端到端融合跟踪和多平台协同状态修正方法，可作为机载协同感知链路中的目标级信息处理模块。其直接应用方向包括预警机与无人机编队的多源航迹融合、非共视条件下的目标状态估计，以及局部传感器受限时的跨平台信息补充。",
        "在算法验证和装备论证阶段，项目形成的数据生成、批量测试和多维评估工具可用于传感器配置比较、任务场景推演、干扰条件构造和候选算法筛选。通过固定模型、数据、统计口径和证据摘要，可降低不同方案因测试范围不一致而产生的比较偏差，为后续半实物仿真和目标设备移植提供基准。",
        "**面向协同探测任务的应用。** 在多平台分布式探测条件下，不同载机的视场、测量精度和通信时效可能存在差异。项目形成的统一表征与图协同方法可用于把各平台的目标级观测组织到共同时间和坐标基准下，再利用跨平台信息修正单个平台的局部估计。该能力适合在仿真和半实物环境中研究平台编队、传感器组合与信息共享策略对航迹质量的影响。",
        "**面向协同打击任务的应用。** 对持续机动目标，航迹结果需要同时反映位置、速度和状态变化趋势。项目模型通过物理先验、时序特征和多平台协同输出连续状态，可为后续任务规划或火控系统提供目标级状态输入。当前软件输出并非控制指令，尚未与具体武器、制导或控制系统完成接口和闭环验证，因此应用前景应理解为可供集成的跟踪算法基础。",
        "**装备论证与方案比较。** 仿真数据生成和自动批测能力可以服务于方案早期论证。在保持场景、样本范围和评价规则一致的条件下，可比较不同传感器组合、平台数量、链路延迟和候选算法的性能变化。逐航迹记录、分类统计和高误差样本定位可帮助论证人员识别总体平均值掩盖的局部风险，并为传感器配置调整和测试资源分配提供依据。",
        "原型系统还可作为算法迭代的回归验证工具。模型或数据处理模块更新后，通过固定V7基准任务重新运行，可以检查总体精度、关联、时延和关键功能是否发生非预期变化；对新增场景则建立独立任务和摘要，避免覆盖既有结论。这种版本化评估方式有助于把研究代码逐步转化为可维护的软件部件。",
        "在工程化应用前，仍需把当前研究原型扩展为因果或滑动窗口推理形式，并在目标处理平台上测量从原始量测接入、关联注册、模型前向到航迹发布的全链路时延。还需引入真实传感器误差、通信中断、虚警和欺骗目标等条件，验证模型在仿真分布之外的稳定性和失效提示机制。",
        "**工程转化的第一阶段是接口适配。** 需要依据目标平台的传感器消息、坐标基准、时钟同步、目标编号和链路协议，重新确认数据字段与更新策略。仿真环境中已经明确的有效性掩码、插值上限和异常门限，在真实系统中必须结合设备特性校准。若接口语义发生变化，应重新生成注册数据并评估模型，不宜仅通过格式转换直接沿用既有精度结论。",
        "**第二阶段是在线化与资源适配。** 当前完整序列模型主要用于离线验证，工程实现需要设计因果递推或滑动窗口推理，明确缓存长度、状态初始化、目标新生与消亡以及通信中断后的恢复方式。同时应在目标CPU、GPU或专用计算平台上测量模型前向、数据搬移、关联注册和完整任务调度的时间与资源占用，并在多目标并发条件下检查最坏时延。",
        "**第三阶段是半实物和实测验证。** 应逐步引入设备噪声、系统偏差、时间同步误差、虚警、漏检、目标交叉和复杂电磁干扰，检查模型在训练分布之外的退化方式。测试不仅需要记录平均精度，还应关注高误差尾部、航迹中断、错误关联持续时间和恢复能力。对失效条件应提供可观测的质量标记，便于上层系统采取降级或人工复核措施。",
        "现有结果显示CSSW类别的误差裕量小于其他机动类别，全源遮挡、随机丢测和额外链路延迟也会造成不同程度退化。后续优化宜先定位高误差片段与机动突变、通道可用性和关联置信度的关系，再决定是否采用数据增强、动态窗口、置信度建模或协同策略调整。该路径比只追求总体平均值下降更有利于提高任务可靠性。",
        "参考实现对比可为后续算法选择提供基准，但当前IMM-JPDA与V7的输入模态和初始化条件不完全相同。工程论证阶段应在相同传感器输入、初始化、预处理和干扰条件下复测，并可增加其他成熟模块化算法作为对照。这样才能区分多源信息增加带来的收益与模型结构本身带来的收益，形成更适合型号决策的证据。",
        "相关技术也可扩展到低空协同监管、无人机集群巡检、海上目标监测和重点区域安防等多源感知任务。扩展时需要依据新领域的传感器类型、运动规律、坐标基准和评价要求重新训练与验证，不能直接沿用V7仿真集上的米级精度结论。",
        "在低空监管场景中，可研究雷达、光电、无线电侦测和多站共享信息的联合航迹维护；在无人机集群巡检中，可研究不同视角对同一目标或设施的持续观测；在海上目标监测中，可结合岸基、舰载和无人平台信息改善遮挡及远距离条件下的状态连续性。这些方向与本项目的数据组织和协同融合方法具有共通性，但任务目标、误差尺度和合规要求均需重新定义。",
        "项目形成的数据、模型、软件和证据链还具有科研复用价值。标准化场景生成流程可支持新算法的对照试验，逐航迹与分层评价方法可支持误差诊断，中文图件和报告生成流程可用于技术评审。复用时应固定版本并保留来源说明，避免把为本项目设置的V7条件误认为通用行业基准。",
        "**成果转化建议。** 可按照“研究原型完善—半实物验证—目标平台移植—现场任务验证”分阶段推进。每个阶段分别定义输入条件、功能范围、性能指标和退出准则，并保存与模型、数据和平台绑定的测试记录。论文和专利材料可围绕已形成的物理先验、双流融合、时序建模和图协同方法继续凝练，但成果状态仍应以正式投稿、受理或授权凭证为准。",
        "在研究原型完善阶段，重点应放在代码结构、配置管理、异常处理和自动回归测试上，使模型、数据和评价程序能够在受控环境中重复运行。半实物阶段再接入更接近实际设备的数据格式与时间节拍，观察同步误差、通信波动和计算负载对结果的共同影响。两个阶段的评价对象不同，应分别建立基线和记录，不用离线仿真数字直接代替半实物结论。",
        "目标平台移植阶段需要在算力、内存、功耗和软件生态约束下选择模型精度、时间窗口与并发能力的平衡点。可根据任务优先级采用模型压缩、算子优化、分批调度或异构计算，但任何优化都应重新检查状态输出一致性和完整链路时延。若压缩或量化改变精度，应将其作为新的工程版本评价，而不是继续使用研究环境下的V7性能数字。",
        "现场任务验证应覆盖正常运行、部分传感器不可用、通信延迟、目标交叉和高机动等代表性条件，并预先规定数据记录、故障处置和结果复盘方式。除精度与时延外，还应关注连续运行稳定性、任务恢复、状态质量标记和人工可解释信息。通过这些工作，才能判断研究成果在具体装备体系中的适用范围和保障需求。",
        "从长期建设角度看，项目形成的价值还包括一套面向多源协同跟踪的验证方法：先固定输入与被测对象，再区分完整集、压力试验和受控用例，最后以可追溯图表表达结论。该方法可随新传感器、新平台和新模型扩展，帮助后续项目减少评价口径反复变化造成的重复工作，并为跨阶段结果比较保留必要依据。",
        "工程应用还需要建立模型版本和数据版本的长期维护机制。真实任务中传感器配置、平台软件和作战条件可能持续变化，模型应在发现输入分布漂移或性能下降时触发复核，而不能长期沿用一次离线验证结论。通过保留任务摘要、典型样本和回归指标，可以为后续升级提供可比较的历史基线。",
        "对于需要较高可信度的任务，可在学习模型之外保留物理约束、输入质量检查和传统估计方法作为监控或降级手段。当输入超出训练范围、模型输出出现异常或计算资源不足时，系统应能够报告状态并切换到预先定义的处理策略。该方向有助于提高工程系统的可解释性和运行韧性，也是后续从研究原型走向任务系统时需要重点设计的内容。",
        "应用推广过程中应继续坚持结论与证据范围一致。V7结果证明当前模型在既定仿真分布和开发工作站条件下达到项目内部验收口径，并提供了进一步集成的技术基础；新的平台、地区、传感器或任务模式均需要重新建立代表性数据和测试协议。只有在相应环境完成验证后，才能形成面向具体应用的性能承诺。",
        "从组织实施角度，可由算法、软件、数据和总体单位共同维护应用验证清单：算法侧负责模型版本和性能边界，软件侧负责接口、部署与运行日志，数据侧负责场景代表性和质量条件，总体单位负责确定任务指标及验收方法。各方以同一任务身份交换结果，可以减少模型参数、数据版本和报告口径之间的信息偏差，也便于在出现性能变化时快速定位责任环节。",
        "总体而言，现有成果已经具备继续开展软件集成、半实物验证和场景扩展的技术基础；其近期价值主要体现在算法原型、数据与测试工具链及可复核的性能基准，形成工程应用结论仍以目标平台和真实数据验证为前提。",
    ]
    for text in application_paragraphs:
        insert_paragraph(section4_heading, text, body_template)

    # Section 4 document lists, using actual current files.
    report_table = persistent_tables[4]
    report_rows = [["序号", "研究报告清单"]] + [
        ["1", "《基于多源异构数据融合的端到端适变跟踪方法研究技术方案》"],
        ["2", "《端到端协同跟踪算法软件需求规格说明书》"],
        ["3", "《端到端协同跟踪算法软件设计说明书》"],
        ["4", "《端到端协同跟踪算法软件测试大纲》"],
        ["5", "《端到端协同跟踪算法软件测试报告》"],
        ["6", "《端到端协同跟踪算法软件使用说明书》"],
        ["7", "《基于多源异构数据融合的端到端适变跟踪方法研究结题报告》"],
        ["8", "《基于多源异构数据融合的端到端适变跟踪方法研究中期报告》"],
        ["9", "《中期评审意见整改闭环表》"],
        ["", ""],
        ["", ""],
    ]
    fill_table(report_table, report_rows)
    fill_table(
        persistent_tables[5],
        [["序号", "模型清单"]] + [[str(i), value] for i, value in enumerate([
            "多源异构观测时空统一表征模型",
            "航迹与目标关联模型",
            "野值剔除与缺失数据补全模型",
            "端到端多模态智能融合与适变跟踪核心模型",
            "基于时空神经网络的深度学习模型",
            "复杂机动与强干扰下的高精度状态估计模型",
        ], 1)],
    )
    fill_table(
        persistent_tables[6],
        [["序号", "工具清单"], ["1", "融合跟踪自动化原型验证与效能评估系统"], ["", ""], ["", ""], ["", ""]],
    )

    # Official registration tables remain empty because no formal number exists.
    for table_index in (7, 8, 9, 10):
        table = persistent_tables[table_index]
        for row_index in range(1, len(table.rows)):
            for col_index in range(len(table.columns)):
                value = str(row_index) if col_index == 0 else ""
                _set_cell_text(table.cell(row_index, col_index), value)

    patent_text = find_paragraph(doc, "专利受理0个（正在整理专利1项），其中国防专利0个，发明专利0个。")
    set_paragraph_text(patent_text, "专利受理0个（已形成发明专利申请材料1套，尚未取得受理号），其中国防专利0个，发明专利受理0个。")
    paper_text = find_paragraph(doc, "论文发表0篇（已完成1篇论文撰写工作），其中0篇EI论文，1篇SCI论文。")
    set_paragraph_text(paper_text, "论文发表0篇（已形成SCI论文初稿1篇，尚未正式发表），其中EI论文0篇、SCI论文发表0篇。")
    paper_heading = find_paragraph(doc, "（4）论文发表：")
    paper_heading.paragraph_format.keep_with_next = True
    paper_text.paragraph_format.keep_with_next = True
    paper_text.paragraph_format.keep_together = True
    keep_table_rows_together(persistent_tables[10])

    replace_static_toc(doc)

    doc.core_properties.title = "基于多源异构数据融合的端到端适变跟踪方法研究结题报告"
    doc.core_properties.subject = "按中期报告格式编制的工作总结、技术总结与应用前景初稿"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
