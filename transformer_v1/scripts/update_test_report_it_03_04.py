from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "software_delivery" / "端到端协同跟踪算法软件_测试报告.docx"
RUN_DIR = ROOT / "artifacts" / "reports" / "software_test" / "IT-03-04" / "20260925_run01"
FIG_DIR = RUN_DIR / "figures"
RESULT = json.loads((RUN_DIR / "result.json").read_text(encoding="utf-8"))
BACKUP = RUN_DIR / "端到端协同跟踪算法软件_测试报告_写入IT-03-04前备份.docx"


def set_run_font(run, bold: bool | None = None) -> None:
    run.font.name = "楷体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "楷体")
    run.font.size = Pt(14)
    if bold is not None:
        run.bold = bold


def add_paragraph(cell, text: str = "", *, bold: bool = False, center: bool = False, after: float = 5.0):
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    if text:
        set_run_font(paragraph.add_run(text), bold=bold)
    return paragraph


def add_figure(cell, filename: str, caption: str) -> None:
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(FIG_DIR / filename), width=Cm(10.2))
    add_paragraph(cell, caption, center=True, after=7)


assert RESULT["overall_passed"] is True
assert RESULT["passed_count"] == RESULT["check_count"] == 46
assert RESULT["formal_metric"]["passed"] is True
assert RESULT["assessment"]["overall_status"] == "passed"
assert RESULT["backend_unit_tests"]["result"] == "14/14 passed"
for figure in RESULT["figures"]:
    assert (FIG_DIR / figure).is_file(), figure

if not BACKUP.exists():
    shutil.copy2(REPORT, BACKUP)

doc = Document(REPORT)
table = doc.tables[15]
table.autofit = False

label_cell = table.cell(5, 0)
label_cell._element.clear_content()
label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
add_paragraph(label_cell, "测试结果", bold=True, center=True, after=0)

cell = table.cell(5, 1)
cell._element.clear_content()
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

add_paragraph(
    cell,
    "2026年9月25日在Windows 10开发工作站执行IT-03-04。硬件为Dell Inspiron 16 7610、"
    "Intel Core i7-11800H、32 GB内存和NVIDIA GeForce RTX 3060 Laptop GPU（6 GB）；"
    "软件为Python 3.11.5、PyTorch 2.1.0、CUDA 12.1。测试期间无并行评估任务，桌面会话空闲。"
    "本轮46项验收检查全部通过，相关评估与批测后端单元测试14/14通过。",
)

add_paragraph(cell, "1. 正式时延口径与执行方法", bold=True)
add_paragraph(
    cell,
    "正式指标使用研究Python模型的B=1、T=1、N=5输入，从已注册张量输入到模型状态输出测量纯模型前向墙钟时间；"
    "不包含文件读写、原始检测提取、目标关联、界面绘制和飞行控制指令生成。模型置于eval和inference_mode，"
    "每轮CUDA前向后执行同步，先预热50次，再使用time.perf_counter连续测量500次并保留全部原始样本。"
    "验收值取未四舍五入的P95，要求不高于100 ms。",
)
add_paragraph(
    cell,
    "时延证据绑定V7完整批次46fada4af5ef4c70a03fa56ef521f8c8、任务签名、模型SHA-256和cuda:0设备；"
    "样例为track_ca_0000.pkl。另对同一模型的B=1、T=100、N=5完整序列独立执行同样的50次预热和500次测量。"
    "T=100结果只作完整序列性能说明，不除以100，不用于替代T=1正式判定；T=1测试也不证明该非因果时序模型可在线逐帧部署。",
)

add_paragraph(cell, "2. 前向时延结果", bold=True)
add_paragraph(
    cell,
    "T=1的500个样本均为有限正数，最小值9.402800 ms、均值13.238647 ms、P50为10.992300 ms、"
    "P95为23.136980 ms、P99为24.343569 ms、最大值24.705500 ms、标准差4.533429 ms，"
    "没有样本超过100 ms。P95低于指标76.863020 ms，正式时延指标通过。",
)
add_paragraph(
    cell,
    "T=100的最小值54.068000 ms、均值70.029367 ms、P50为63.093650 ms、P95为94.094790 ms、"
    "P99为101.929610 ms、最大值117.394100 ms、标准差13.834096 ms，其中6/500个样本超过100 ms。"
    "该组数据是完整100帧输入的一次前向时延分布，最大值超过100 ms不影响以T=1 P95定义的正式单帧指标判定。",
)
add_figure(cell, "figure_1_latency_distribution.png", "图4-14-1  T=1与T=100模型前向时延分布")
add_figure(cell, "figure_2_latency_statistics.png", "图4-14-2  T=1与T=100前向时延统计量")
add_figure(cell, "figure_3_cdf_and_margin.png", "图4-14-3  时延累计分布及T=1 P95指标裕量")

add_paragraph(cell, "3. 评估、状态与归档能力", bold=True)
add_paragraph(
    cell,
    "将当前批次的关联正确率、相对IMM-JPDA的RMSE降低率和本轮时延证据放入独立证据目录后，"
    "评估服务对位置RMSE、关联正确率、RMSE降低率和T=1前向P95四项均判为passed，总评为passed。"
    "时延项读取23.136980 ms，并在备注中单列T=100 P95=94.094790 ms。报告目录能够识别已完成任务，"
    "并成功导出包含原始任务、四项判据、评估结论和来源说明的assessment_bundle.json及assessment_criteria.csv。",
)
add_paragraph(
    cell,
    "异常分支检查结果符合预期：缺少关联和基线证据时总评为incomplete且两项不可用；关联或基线证据签名过期时拒绝证据并保持incomplete；"
    "时延证据签名过期时将时延项标为unavailable，不回退到完整序列耗时除法；失败任务被拒绝评估。"
    "阈值边界中99.999 ms和100.000 ms通过，100.001 ms不通过；零、负数、NaN和无穷大计时样本均拒绝。",
)
add_figure(cell, "figure_4_assessment_and_interfaces.png", "图4-14-4  评估状态防护及研究与交付入口输入范围")

add_paragraph(cell, "4. 交付接口与适用范围", bold=True)
add_paragraph(
    cell,
    "静态检查确认portable/artifacts/coop_tracker_v7.onnx含18个输入，其中16个时序输入固定为B=1、T=100，"
    "参考位置输入为[1,3]、时间步长输入为[1]，输出固定为[1,100,6]，并包含N=5传感器维。"
    "因此交付ONNX按T=100固定形状验证，不宣称支持T=1。MATLAB端在前序环境检查中因File system inconsistency未进入模型代码，"
    "C端黄金样例程序完成一致性验证但没有复用会话的纯前向计时输出；两端均无可用于本用例正式时延判定的证据，未将其写成通过。",
)
add_paragraph(
    cell,
    "本轮结果来自笔记本式开发工作站，并非实际部署平台。Python研究入口的正式判定不受MATLAB/C环境限制影响；"
    "实际部署平台仍应使用同一计时脚本和相同50次预热、500次重复、CUDA同步及P95口径复测。",
)

add_paragraph(cell, "5. 判定", bold=True)
add_paragraph(
    cell,
    "本用例通过。Python研究模型T=1纯前向时延P95为23.136980 ms，满足不高于100 ms的技术指标，"
    "全部原始样本、环境、负载、任务签名和模型摘要可追溯；T=100完整序列数据已单列且未折算。"
    "证据齐全时四项综合评估全部通过，缺失、过期、失败任务和异常数值均能被正确拦截，报告及CSV/JSON归档功能正常。",
)

for paragraph in cell.paragraphs:
    for run in paragraph.runs:
        if run._element.xpath(".//a:blip"):
            continue
        set_run_font(run)

doc.save(REPORT)
print(REPORT)
