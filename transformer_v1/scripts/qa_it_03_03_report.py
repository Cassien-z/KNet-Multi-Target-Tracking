from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


root = Path(__file__).resolve().parents[1]
report = root / "docs" / "software_delivery" / "端到端协同跟踪算法软件_测试报告.docx"
doc = Document(report)

assert len(doc.tables) == 16
assert len(doc.inline_shapes) == 50

table = doc.tables[14]
physical_cells = table.rows[5]._tr.tc_lst
grid_spans = [int(tc.tcPr.gridSpan.val) if tc.tcPr.gridSpan is not None else 1 for tc in physical_cells]
assert len(physical_cells) == 2
assert grid_spans == [1, 3]
assert table.cell(5, 0).text.strip() == "测试结果"

result_cell = table.cell(5, 1)
text = result_cell.text
required = [
    "7.488174721 m",
    "189.127525101 m",
    "96.040674293%",
    "2000/2000个文件均不低于15%",
    "最低降低率为72.091354%",
    "V8原始40 Hz场景未混入",
    "本用例通过",
]
for phrase in required:
    assert phrase in text, phrase

blips = result_cell._tc.xpath(".//a:blip")
assert len(blips) == 4

for paragraph in result_cell.paragraphs:
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        east_asia = run._element.get_or_add_rPr().get_or_add_rFonts().get(qn("w:eastAsia"))
        assert east_asia == "楷体", (run.text, east_asia)
        assert run.font.size is not None and round(run.font.size.pt, 4) == 14.0, (run.text, run.font.size)

print({
    "tables": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "row5_physical_cells": len(physical_cells),
    "row5_grid_spans": grid_spans,
    "result_figures": len(blips),
    "required_phrases": len(required),
    "font": "楷体 14pt",
})

