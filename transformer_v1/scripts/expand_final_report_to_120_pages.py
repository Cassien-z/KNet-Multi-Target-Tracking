"""Expand the completion report with verified process evidence and result explanations."""

from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "docs" / "基于多源异构数据融合的端对端适变跟踪方法研究_结题报告_中期报告格式初稿.docx"
OUTPUT = ROOT / "docs" / "基于多源异构数据融合的端对端适变跟踪方法研究_结题报告_扩充稿.docx"
NEW_FIG = ROOT / "artifacts" / "figures" / "final_report_expansion_120p"
TEST = ROOT / "artifacts" / "reports" / "software_test"


def set_east_asia_font(run, name: str) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def find_paragraph(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"未找到段落：{text}")


def copy_ppr(target, template) -> None:
    ppr = target._p.get_or_add_pPr()
    for child in list(ppr):
        ppr.remove(child)
    source = template._p.pPr
    if source is not None:
        for child in source:
            if child.tag != qn("w:sectPr"):
                ppr.append(deepcopy(child))


def clear_runs(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def replace_text(paragraph, text: str) -> None:
    clear_runs(paragraph)
    paragraph.add_run(text)


def add_inline_text(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*.*?\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        run = paragraph.add_run(match.group(0)[2:-2])
        run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_paragraph(anchor, text: str, body_template, *, keep_with_next: bool = False):
    paragraph = anchor.insert_paragraph_before()
    copy_ppr(paragraph, body_template)
    add_inline_text(paragraph, text)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    return paragraph


def add_heading(anchor, text: str, heading_template, *, page_break: bool = False):
    paragraph = anchor.insert_paragraph_before()
    copy_ppr(paragraph, heading_template)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    set_east_asia_font(run, "仿宋_GB2312")
    return paragraph


def add_figure(anchor, path: Path, label: str, caption: str, caption_template, *, width_cm: float = 14.4):
    if not path.is_file():
        raise FileNotFoundError(path)
    image_p = anchor.insert_paragraph_before()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.first_line_indent = Cm(0)
    image_p.paragraph_format.space_before = Pt(5)
    image_p.paragraph_format.space_after = Pt(2)
    image_p.paragraph_format.keep_with_next = True
    image_p.add_run().add_picture(str(path), width=Cm(width_cm))
    caption_p = anchor.insert_paragraph_before()
    copy_ppr(caption_p, caption_template)
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.first_line_indent = Cm(0)
    caption_p.paragraph_format.space_after = Pt(6)
    caption_p.paragraph_format.keep_with_next = True
    run = caption_p.add_run(f"图{label}  {caption}")
    run.font.size = Pt(10.5)
    set_east_asia_font(run, "宋体")
    return caption_p


def set_cell(cell, text: str, *, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    clear_runs(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) < 18 else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    set_east_asia_font(run, "宋体")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for side in ("top", "left", "bottom", "right"):
        node = mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), "90")
        node.set(qn("w:type"), "dxa")


def add_table(doc: Document, anchor, caption: str, rows: list[list[str]], body_template):
    cap = add_paragraph(anchor, caption, body_template, keep_with_next=True)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for i, values in enumerate(rows):
        row = table.rows[i]
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for j, value in enumerate(values):
            set_cell(table.cell(i, j), value, bold=i == 0)
    anchor._p.addprevious(table._tbl)
    spacer = anchor.insert_paragraph_before()
    spacer.paragraph_format.space_after = Pt(3)
    return table


def test_figure(code: str, filename: str) -> Path:
    runs = sorted((TEST / code).glob("*/figures"))
    if not runs:
        raise FileNotFoundError(code)
    return runs[0] / filename


def chapter_number_existing_figures(doc: Document) -> None:
    """Remove the inherited duplicate figure numbering without changing figure order."""
    section3_seen = False
    for paragraph in doc.paragraphs:
        text = paragraph.text
        stripped = text.strip()
        if stripped == "3.取得的主要成果及应用前景分析":
            section3_seen = True
            continue
        if not section3_seen:
            if stripped == "图1 项目总体研究内容":
                replace_text(paragraph, "图1-1  项目总体研究内容")
            elif "研究内容框架如图1所示" in text:
                replace_text(paragraph, text.replace("研究内容框架如图1所示", "研究内容框架如图1-1所示"))
            continue
        if re.search(r"图\s*\d+", text):
            updated = re.sub(r"图\s*(\d+)", lambda m: f"图3-{m.group(1)}", text)
            updated = updated.replace("中期报告图3-19", "中期报告图19")
            replace_text(paragraph, updated)


def update_static_toc(doc: Document) -> None:
    """Update the retained template TOC after the final Word pagination review."""
    toc_lines = [
        "目录",
        "1.项目研究目标\t1",
        "（一）研究目标\t1",
        "（二）研究内容\t1",
        "（三）研究内容完成情况\t5",
        "（四）研究成效\t6",
        "2.研究工作进程\t6",
        "2.1 多模态异构数据的统一表征与融合机理\t6",
        "2.2 端到端多模态信息融合与适变跟踪框架设计\t22",
        "2.3 标准数据集生成与算法验证\t38",
        "3.取得的主要成果及应用前景分析\t61",
        "3.1取得的主要成果\t61",
        "3.2应用前景分析\t110",
        "4.文档资料情况\t118",
        "5.申报专利情况\t119",
        "6.国防科技报告清单（如有请填写）\t120",
    ]
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    sdts = doc._element.body.findall(qn("w:sdt"))
    if len(sdts) != 1:
        raise ValueError(f"目录内容控件数量异常：{len(sdts)}")
    paragraphs = sdts[0].findall(".//w:sdtContent/w:p", namespace)
    if len(paragraphs) != len(toc_lines):
        raise ValueError(f"目录行数异常：{len(paragraphs)}")
    for paragraph, value in zip(paragraphs, toc_lines):
        for child in list(paragraph):
            if child.tag != qn("w:pPr"):
                paragraph.remove(child)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = value
        run.append(text)
        paragraph.append(run)


def add_section_21(doc: Document, anchor, body, heading, caption) -> None:
    n = 1
    def fig(path: Path, text: str, width=14.4):
        nonlocal n
        add_figure(anchor, path, f"2-{n}", text, caption, width_cm=width)
        n += 1

    add_heading(anchor, "2.1.1 输入接口与有效性语义", heading, page_break=True)
    for text in [
        "统一表征的第一项工作是把不同来源的字段含义固定下来。V7模型入口包含雷达、光电、电子侦察、平台姿态、参考位置、快慢链路、更新时间、缓存年龄和采样间隔等18类输入。检查内容不仅包括数组能否读入，还核对数据类型、张量形状、元素数量、ONNX别名和有限值要求，使训练、导出和软件推理使用同一套接口定义。",
        "接口核对把传感器数值和有效性掩码作为成对对象。数值为零不能自动解释为缺测，通道失效也不能仅靠填零表达；软件必须同时给出有效性标记、更新时间或缓存年龄。这样做可以区分真实零值、暂未更新、超出视场和链路失效，避免模型把占位值误认为实际观测。",
    ]:
        add_paragraph(anchor, text, body)
    fig(test_figure("IT-01-01", "图1_输入接口逐字段核对.png"), "模型输入接口逐字段核对结果")
    add_paragraph(anchor, "图2-1记录了18类输入在类型、形状、元素数、运行时别名和有限值方面的核对状态。该图属于接口一致性证据，证明软件调用条件与模型定义能够对应；它不承担位置精度或实时性判定。接口发生任何增删、排序变化或单位变化时，都应重新执行同类核对并重新生成模型摘要。", body)
    fig(test_figure("IT-01-01", "图2_模态缺失组合结果.png"), "典型模态缺失组合与有限输出检查")
    add_paragraph(anchor, "模态组合检查覆盖完整三模态、仅雷达、雷达与光电、雷达与电子侦察、全部慢链路缺失以及全部观测源缺失等条件。图2-2上半部分给出各来源的有效输入点数，下半部分核对有限输出比例。有限输出只能说明软件没有产生NaN或无穷值，不能据此认定严重缺测时仍满足8米精度要求；精度退化由后续分层压力试验单独说明。", body)
    fig(test_figure("IT-01-01", "图3_链路年龄边界结果.png"), "快慢链路缓存年龄边界检查")
    add_paragraph(anchor, "协同链路不是持续同步的新鲜量测。当前受控用例分别检查快链路2秒和慢链路5秒的缓存年龄边界：低于或等于门限时保持有效，超过门限后标记为无效。采用边界点前后各0.001秒的样本，可以确认比较符号、时间单位和缓存更新逻辑一致，防止临界值在不同程序中产生相反判断。", body)
    add_table(doc, anchor, "表2-1  模型输入接口及质量信息分类", [
        ["类别", "主要字段", "质量或有效性信息", "在模型中的作用"],
        ["主动观测", "雷达距离、方位、俯仰及派生量", "radar_mask、update、age", "提供较完整的空间约束"],
        ["被动观测", "光电和电子侦察测角及编码特征", "eo_mask、esm_mask、更新时间", "在视场允许时补充方向与辐射信息"],
        ["平台状态", "参考位置、姿态、速度和惯导信息", "有限值、坐标与角度单位", "完成局部观测注册和状态传播"],
        ["协同链路", "快链路与慢链路目标级状态", "缓存年龄、链路有效性", "提供跨平台补充状态"],
        ["运行参数", "采样间隔、序列长度、平台与目标维度", "配置摘要、模型摘要", "保证权重、数据和软件调用相互匹配"],
    ], body)

    add_heading(anchor, "2.1.2 异步采样与时间尺度对齐", heading, page_break=True)
    add_paragraph(anchor, "雷达、光电、电子侦察和协同链路可能在不同时间到达，且同一来源也可能因视场或通信条件产生间断。项目没有把所有输入简单复制到统一帧率，而是先保留源时间戳，再根据目标时间网格、邻近有效点数量和最大允许间隔选择插值、保持或无效处理。", body)
    fig(NEW_FIG / "01_多源异步采样与统一时间网格.png", "多源异步采样与统一时间网格示意")
    add_paragraph(anchor, "图2-4展示的是处理逻辑示意，而不是某一具体场景的实测时间戳。雷达和统一网格可以具有较密更新，光电、电子侦察及协同链路更新较稀；注册程序应以源时间为依据判断当前网格点可使用何种信息。图中的差异说明时间对齐会改变数据组织方式，但不能生成原始观测中不存在的新信息。", body)
    fig(test_figure("IT-01-03", "figure_1_interpolation_branches.png"), "不同有效点数量下的插值分支")
    add_paragraph(anchor, "当有效点不少于4个时采用PCHIP保持局部形状；只有2至3个点时采用线性插值；只有单点或没有可靠邻点时不进行跨区间外推。分支设计兼顾曲线平滑和小样本稳定性。PCHIP只用于时间轴上的数值整理，不改变传感器测量误差模型，也不应被解释为模型预测能力。", body)
    fig(test_figure("IT-01-03", "figure_2_gap_boundary.png"), "最大插值间隔0.45秒边界验证")
    fig(test_figure("IT-01-03", "figure_3_missing_masks.png"), "缺失保持与禁止外推时的输出掩码")
    add_paragraph(anchor, "最大插值间隔设置为0.45秒，距离最近真实观测小于或等于该值时允许使用相应分支，超过后将输出标记为无效。图2-6验证临界点的包含关系，图2-7进一步检查区间前、端点、区间内、末点和区间后的掩码变化。数值计算与掩码生成必须同时正确，否则即使插值曲线看起来连续，模型仍可能接收到不应使用的外推量。", body)
    fig(test_figure("IT-01-03", "figure_4_multisensor_validity.png"), "多传感器独立时间对齐后的有效点变化")
    add_paragraph(anchor, "不同传感器分别依据自身时间戳和有效点进行对齐，因此对齐后的有效点数量不必相同。受控样例中雷达、光电和电子侦察由原始9、6、4个有效点分别扩展为40、35、27个对齐有效点。该结果用于证明分来源处理和掩码传播能够工作，不能把对齐后的点数增加理解为传感器实际观测次数增加。", body)
    add_table(doc, anchor, "表2-2  时间尺度对齐分支与使用条件", [
        ["输入条件", "处理方式", "有效性处理", "主要风险控制"],
        ["有效点不少于4个且邻近间隔合规", "PCHIP插值", "按邻近真实观测与区间共同判定", "避免高阶插值跨越长缺测区间"],
        ["有效点2至3个且邻近间隔合规", "线性插值", "仅区间内或允许端点有效", "避免少量点构造不稳定曲线"],
        ["仅有单个有效点", "单点保持或仅端点有效", "禁止跨区间外推", "防止把一个历史值长期当作新观测"],
        ["无有效点或超过0.45秒", "不生成可用观测", "mask置为无效", "由其他来源和状态传播承担连续性"],
    ], body)

    add_heading(anchor, "2.1.3 目标关联与历史连续性", heading, page_break=True)
    add_paragraph(anchor, "目标关联位于原始量测与目标级序列之间。软件先依据几何距离和传感器条件形成候选，再通过门控剔除不合理配对，使用匈牙利算法得到一对一分配，并结合历史记忆维持交叉航迹的身份连续性。真值标签只在分组完成后用于离线评分，不进入关联决策。", body)
    fig(NEW_FIG / "02_关联代价矩阵与一对一分配.png", "关联代价矩阵与匈牙利一对一分配示例", 12.2)
    add_paragraph(anchor, "图2-9使用3条历史航迹和3个候选量测构造代价矩阵。匈牙利算法从全局组合中选择总代价较小的一对一分配，避免多个历史航迹同时占用同一候选量测。示例中的数字只用于解释算法关系，最终V7关联结果来自原始场景重算，不由该示例数值推导。", body)
    fig(NEW_FIG / "03_跨平台关联门限边界.png", "跨平台关联门限边界检查")
    fig(NEW_FIG / "04_交叉航迹历史连续性.png", "交叉航迹条件下的历史连续性检查")
    add_paragraph(anchor, "门限边界用1499米、1500米和1501米三个距离确认当前实现采用小于等于1500米的包含关系。交叉航迹用例则比较保留与清空历史记忆时的编号变化：保留记忆时两个目标身份保持连续，清空记忆后仅依据当前几何关系容易交换编号。这两项检查分别约束单帧候选范围和跨帧身份连续性。", body)
    add_table(doc, anchor, "表2-3  关联评价统计量及解释", [
        ["统计量", "计算分母", "反映的问题", "结题报告使用方式"],
        ["正确率", "全部输出关联对", "输出配对中有多少正确", "合同正确关联率主口径，V7为99.90%"],
        ["召回率", "全部真值可关联对", "应关联对象中有多少被找到", "检查漏关联，V7为96.60%"],
        ["F1", "正确率与召回率调和平均", "兼顾错关联与漏关联", "V7为98.22%"],
        ["观测事件纯度", "历史观测事件", "旧流程的事件级纯净程度", "仅作为过程审计，不替代合同口径"],
    ], body)

    add_heading(anchor, "2.1.4 空间注册与异常观测控制", heading, page_break=True)
    add_paragraph(anchor, "多平台观测首先处于各自局部坐标系。项目根据平台参考位置和姿态，把局部量测转换到统一ECEF坐标，再执行跨平台关联和模型输入组织。受控验证分别检查基础几何关系、旋转组合、矩阵行列式和同一目标在不同平台下的汇聚结果，用于排除轴向、角度单位和旋转顺序错误。", body)
    fig(test_figure("IT-01-04", "figure_2_geometry_validation.png"), "空间注册基础几何关系与旋转矩阵检查")
    add_paragraph(anchor, "六组人工几何关系的数值误差均低于受控用例容差，旋转矩阵行列式保持在1附近。该结果证明测试输入下的数学变换实现一致，但不代表真实平台姿态、时间同步或设备系统误差已经消除。批量场景仍需通过范围、连续性和最终航迹误差进行补充检查。", body)
    fig(test_figure("IT-01-04", "figure_3_threshold_boundary.png"), "角度与距离异常门限的严格超限规则")
    fig(test_figure("IT-01-04", "figure_4_outlier_suppression.png"), "多维异常点的中值基线与剔除结果")
    add_paragraph(anchor, "异常值判定采用严格超过门限的规则：等于0.05弧度或100米的样本保留，超过门限的样本剔除。图2-13验证门限比较符号，图2-14在方位、俯仰和距离三个维度注入同一时刻异常点，检查基于中值基线的剔除标记。剔除只针对明显不满足输入规则的量测，不能把真实存在的高机动或困难样本从最终评价中删除。", body)
    add_table(doc, anchor, "表2-4  统一表征关键边界及验证目的", [
        ["边界项目", "当前受控值", "验证目的", "结论限制"],
        ["快链路缓存年龄", "2秒", "确认临界值有效、超限失效", "不代表真实链路最大可用时延"],
        ["慢链路缓存年龄", "5秒", "确认缓存与更新时间逻辑", "工程门限需按设备重新标定"],
        ["最大插值间隔", "0.45秒", "限制跨长缺测区间外推", "只适用于当前注册配置"],
        ["角度异常门限", "0.05弧度", "确认严格超限剔除", "不是传感器标称精度"],
        ["距离异常门限", "100米", "确认多维异常处理", "与模型8米RMSE含义不同"],
        ["跨平台关联距离", "1500米", "确认候选门控包含关系", "不单独决定完整关联性能"],
    ], body)


def add_section_22(doc: Document, anchor, body, heading, caption) -> None:
    n = 15
    def fig(path: Path, text: str, width=14.4):
        nonlocal n
        add_figure(anchor, path, f"2-{n}", text, caption, width_cm=width)
        n += 1

    add_heading(anchor, "2.2.1 物理状态更新与可学习增益", heading, page_break=True)
    add_paragraph(anchor, "V7模型前端保留状态传播的物理含义，并利用网络学习观测修正强度。先验状态由上一时刻位置、速度和采样间隔形成，观测残差反映当前量测与先验之间的偏差，质量信息控制不同来源在当前时刻的可信程度。前端输出再进入双流、时序和图协同模块。", body)
    fig(test_figure("IT-02-01", "figure_1_information_matrix_fusion.png"), "信息矩阵融合的独立复算结果")
    add_paragraph(anchor, "图2-15分别构造仅雷达、雷达加快链路、雷达加慢链路以及雷达加快慢链路四种组合，将模块输出与独立的信息矩阵计算结果比较。各组合结果在受控数值精度内一致，说明来源增加时的信息叠加和协方差处理按预定公式执行。此处验证的是融合算子，不是完整神经网络在实际场景中的误差。", body)
    fig(test_figure("IT-02-01", "figure_2_cache_freshness_and_covariance.png"), "缓存新鲜度权重与雷达距离噪声关系")
    add_paragraph(anchor, "缓存权重随年龄衰减，且快慢链路采用不同时间尺度；雷达距离噪声随目标距离增长，体现远距离观测的不确定性增加。将两类关系并列检查，可以避免融合模块只比较数值大小而忽略信息来源的时效和测量条件。图中曲线来自受控公式检查，真实设备参数仍应在接口适配阶段标定。", body)
    fig(test_figure("IT-02-01", "figure_3_missing_recovery_and_velocity_clip.png"), "首帧缺失、连续缺失恢复与速度限幅")
    add_paragraph(anchor, "首帧缺失时模型不能依赖尚未形成的历史缓存，连续缺失期间则主要依靠状态传播。受控序列显示学习状态与量测状态在观测恢复后重新靠拢，并对位置跳变导致的瞬时速度设置限幅。限幅属于输入与状态稳定措施，只防止明显不合理的数值传播，不应替代对真实高机动状态的建模。", body)
    fig(test_figure("IT-02-01", "figure_4_feature_order_and_gain_range.png"), "特征顺序和六维学习增益范围核对")
    add_paragraph(anchor, "输入特征按固定顺序进入模型，六维学习增益在代表样本中保持有限并随维度、时间变化。固定顺序保证训练权重与软件推理读取的是同一物理量，增益范围检查用于发现饱和、非有限值或维度错位。观察到的增益范围只描述当前样本，不作为所有场景的理论上下界。", body)
    add_table(doc, anchor, "表2-5  状态更新模块的输入输出关系", [
        ["环节", "主要输入", "主要输出", "检查重点"],
        ["物理传播", "上一时刻位置、速度、采样间隔", "先验位置和速度", "时间单位、状态顺序与首帧初始化"],
        ["观测残差", "注册观测、先验状态、有效性", "各来源残差信息", "缺测通道不得形成伪残差"],
        ["信息融合", "协方差、缓存年龄、量测质量", "融合观测与不确定性", "独立复算与来源组合变化"],
        ["学习增益", "状态、残差和质量特征", "六维修正权重", "有限性、范围及维度顺序"],
        ["状态限制", "更新后位置和速度", "稳定状态表示", "非法跳变、限幅及恢复过程"],
    ], body)

    add_heading(anchor, "2.2.2 模态缺失感知与双流融合", heading, page_break=True)
    add_paragraph(anchor, "主动雷达与被动光电、电子侦察的测量含义和有效性模式不同。V7先对主动与被动来源分别编码，再通过双向交叉注意力交换信息，并利用有效性掩码、缺失标记和门控系数抑制无效通道。该设计允许模型在输入组合变化时继续运行，同时保留不同模态的独立语义。", body)
    fig(test_figure("IT-02-02", "figure_1_eight_mask_combinations.png"), "八种雷达、光电和电子侦察有效性组合响应")
    add_paragraph(anchor, "三类传感器有效或无效可以形成八种组合。图2-19上半部分给出代表样本的稳定阶段位置RMSE，下半部分确认八种组合均产生有限输出。只有被动通道可用而雷达失效时误差明显高于完整输入，说明当前模型对主动距离约束更敏感；该现象来自代表样本，不能替代完整集的通道消融。", body)
    fig(test_figure("IT-02-02", "figure_2_passive_weights_and_missing_token.png"), "无效被动通道权重归零与缺失标记作用")
    add_paragraph(anchor, "当光电或电子侦察通道无效时，相应注意力权重应被掩码归零；两类被动通道同时无效时，模型使用缺失标记保持输入结构而不是把占位数值当作观测。缺失标记响应随组合变化，说明模型能够区分不同不可用状态。权重大小只反映内部计算，不等同于对最终误差的因果贡献。", body)
    fig(test_figure("IT-02-02", "figure_4_noise_and_pose_response.png"), "被动测角噪声和平台姿态扰动响应")
    add_paragraph(anchor, "图2-21左侧只改变被动测角噪声，观察位置RMSE和融合门控变化；右侧采用零姿态、常量姿态和动态姿态三种输入，核对姿态信息是否实际进入计算。测试结果表明输出随受控扰动变化，没有被完全忽略。该检查用于发现接口失联或恒定分支，不能直接形成抗噪性能等级。", body)
    fig(NEW_FIG / "08_传感器信息受限条件误差变化.png", "传感器信息受限条件下的分层样本误差变化")
    add_paragraph(anchor, "64场景分层样本进一步比较标准输入、雷达量测抑制70%、光电遮挡2秒和电子侦察丢测70%四种条件。标准输入RMSE约7.45米，雷达量测大幅减少时上升到约21.48米，光电和电子侦察受限时分别约7.62米和7.85米。结果说明该样本中雷达距离约束影响更明显，同时也再次证明补充压力试验不能与2000场景主批次的7.488米混用。", body)
    add_table(doc, anchor, "表2-6  模态组合测试的结论边界", [
        ["证据类型", "样本范围", "可支持结论", "不可扩展结论"],
        ["八种掩码受控样例", "固定代表输入", "有效性掩码生效且输出有限", "不能代表各组合完整集精度"],
        ["噪声与姿态扰动", "固定输入的参数扫描", "相应输入确实参与模型计算", "不能给出实际设备抗扰等级"],
        ["64场景信息受限试验", "四类机动各16场景", "比较指定缺失条件下的退化趋势", "不能替代V7完整集验收值"],
        ["2000场景模块消融", "V7完整注册验证集", "说明既有权重对模块和通道的依赖", "不能代表重新训练后的最优变体"],
    ], body)

    add_heading(anchor, "2.2.3 长时序建模与多平台图协同", heading, page_break=True)
    add_paragraph(anchor, "单帧融合解决当前时刻的多模态组织问题，连续机动趋势仍需通过时序模块提取。V7使用时序Transformer在目标序列上形成上下文，再把各平台作为图节点，以距离门限和节点有效性构建邻接关系，通过图注意力聚合同一目标的跨平台信息。", body)
    fig(test_figure("IT-02-03", "figure_4_temporal_scope_and_repeatability.png"), "时序结构范围与固定输入重复运行检查")
    add_paragraph(anchor, "时序结构检查确认当前实现包含2层编码器，每层注意力头数为4；固定窗口在重复前向和批次顺序变化时保持数值稳定。结构数量来自V7模型配置，重复运行误差用于发现随机状态、批处理依赖或非确定性放大。该测试不评价更深网络或更长窗口是否更优。", body)
    fig(test_figure("IT-02-04", "figure_1_adjacency_boundary.png"), "20千米图邻接距离与节点有效性边界")
    add_paragraph(anchor, "图2-24分别展示所有节点有效和节点2无效两种情况。有效节点间距离不超过20千米时建立边，超过门限或任一节点无效时不建立跨节点边。自环用于保留节点自身信息。邻接矩阵只决定哪些节点可以交换信息，具体贡献还由注意力权重和特征内容决定。", body)
    fig(test_figure("IT-02-04", "figure_3_node_count_degradation.png"), "有效节点数量变化下的跨节点边与计算退化")
    add_paragraph(anchor, "随着有效节点从5个减少到1个，跨节点有效边数量相应下降，而节点自身计算仍可保留。测试以固定5槽位模拟节点逐一失效，说明软件能够在部分平台不可用时退化为较小协同图；其记录的池化往返次数只反映当前实现的执行行为，不是不同平台规模下的精度曲线。", body)
    add_table(doc, anchor, "表2-7  图协同模块的约束和检查项", [
        ["检查项", "输入变化", "预期行为", "工程含义"],
        ["距离边界", "19.9、20.0、20.1千米", "前两者保留边，后一项取消边", "拓扑与门限规则一致"],
        ["节点有效性", "单个或多个平台失效", "无效节点不与其他节点交换信息", "避免把占位节点传播到全图"],
        ["注意力归一化", "邻接位置和非邻接位置", "有效邻接按行归一化，非邻接为零", "确保权重只在允许边上分配"],
        ["节点数量退化", "有效节点5个降至1个", "协同边减少但保留自身状态", "支持部分平台退出时继续运行"],
        ["完整集关闭GAT", "统一V7权重与2000场景", "RMSE由7.488米增至8.853米", "说明既有模型使用跨平台特征"],
    ], body)

    add_heading(anchor, "2.2.4 状态输出与软件运行边界", heading, page_break=True)
    add_paragraph(anchor, "模型最终输出位置和速度六维状态。软件部署不仅需要核对PyTorch与ONNX的逐维差异，还要确认输出由学习状态、深度残差和参考位置按设计组合，参考位置变化只作用于位置分量，速度分量不因坐标平移而改变。", body)
    fig(test_figure("IT-02-05", "figure_2_state_composition.png"), "六维最终状态的组成量级")
    add_paragraph(anchor, "图2-26采用对数纵轴展示六维状态的三个组成部分。参考位置承担大尺度空间基准，学习状态和深度残差在其上形成目标状态修正。位置与速度量级不同，不能直接按柱高比较重要性；该图的作用是核对组合公式和量级关系，防止位置参考量被重复加到速度分量。", body)
    fig(test_figure("IT-02-05", "figure_3_reference_position_response.png"), "参考位置平移对位置与速度分量的响应")
    add_paragraph(anchor, "当参考位置只改变px、py、pz时，最终输出的三个位置分量按相同量变化，vx、vy、vz保持不变。图2-27用预期变化和实际变化并列核对这一性质，说明输出坐标平移逻辑符合状态定义。若目标平台更换坐标原点或参考基准，仍应重新检查坐标轴、单位和时间同步。", body)
    fig(test_figure("IT-02-05", "figure_4_runtime_status_and_margin.png"), "运行入口状态与黄金样例阈值余量")
    add_paragraph(anchor, "Python和C接口运行路径在当前环境中完成检查，MATLAB本机启动受环境条件限制，因此图2-28将其标记为环境受限，而不是写成已通过。右侧显示ONNX位置、速度和C接口位置、速度四类误差占相应阈值的比例，均保留余量。该结果证明当前黄金样例和调用路径满足测试阈值，不等同于目标装备平台已经完成部署验收。", body)
    add_table(doc, anchor, "表2-8  软件运行与性能结论的范围", [
        ["项目", "当前完成状态", "现有证据", "后续工作"],
        ["PyTorch模型前向", "已完成", "V7权重、完整集结果和原始计时", "目标硬件复测与在线化"],
        ["ONNX Runtime", "已完成黄金样例核对", "六维逐维差异低于测试阈值", "扩大样本并测量目标平台性能"],
        ["C接口调用", "已完成受控调用", "三次调用与状态差异记录", "纳入完整软件链和异常恢复"],
        ["MATLAB接口", "本机环境受限", "保留启动状态记录", "在具备完整运行环境后复测"],
        ["全链路实时性", "尚未形成验收结论", "当前仅有模型前向计时", "纳入读取、注册、关联、通信与显示"],
    ], body)


def add_section_23(doc: Document, anchor, body, heading, caption) -> None:
    n = 29
    def fig(path: Path, text: str, width=14.4):
        nonlocal n
        add_figure(anchor, path, f"2-{n}", text, caption, width_cm=width)
        n += 1

    add_heading(anchor, "2.3.1 数据集组织与场景覆盖", heading, page_break=True)
    add_paragraph(anchor, "标准数据集工作不仅包括生成场景文件，还包括固定任务想定、平台编队、目标运动、传感器配置、干扰条件、原始与注册数据层次以及评价清单。V7主批次以2000个场景作为统一对象，每类机动500个场景，每场景3个目标和5个平台，形成6000条目标航迹。", body)
    fig(ROOT / "artifacts" / "figures" / "数据集整体结构.png", "多源异构仿真数据集的层级组织")
    add_paragraph(anchor, "数据集按任务想定、场景、平台、目标、传感器和时间序列组织。原始层保存平台与目标真值、传感器量测和有效性信息；注册层保存统一时间与坐标下的目标级序列；评价层保存场景清单、逐航迹结果和统计证据。保持三层分离，既可以复核关联与注册，也能避免模型推理直接读取真值身份。", body)
    fig(test_figure("IT-01-05", "figure_1_dataset_inventory.png"), "V7四类机动文件清单与V8补充场景盘点")
    add_paragraph(anchor, "清单盘点显示V7原始MAT和注册PKL在CV、CA、CT、CSSW四类中各有500个场景文件，数量相互对应。另有1个V8复合干扰场景用于补充观察强干扰边界。V8不是V7验收数据集的一部分，其数量也不足以形成大样本统计，因此只在研究过程和边界分析中说明。", body)
    fig(test_figure("IT-01-05", "figure_2_scenario_coverage.png"), "两类任务想定的传感器、机动与干扰覆盖矩阵")
    add_paragraph(anchor, "覆盖矩阵把协同探测和高机动强干扰两类想定与传感器、协同链路、四类运动及典型干扰条件对应。矩阵用于说明场景构成，不表示每个V7场景同时具有所有干扰。V7主批次主要承担标准质量条件下的精度验收，强干扰条件通过补充样本和分层测试考察。", body)
    fig(NEW_FIG / "06_V7四类机动强度量化.png", "V7四类机动的加速度、转弯率与等效运动过载")
    add_paragraph(anchor, "机动强度由64场景分层样本的注册真值速度通过有限差分计算。CV接近恒速直线，CA体现持续加速，CT的速度大小基本不变而转弯角速度较高，CSSW同时包含速度变化和转向。等效运动过载按加速度模除以重力加速度计算，只表示运动学量级，不包含飞行器气动和结构载荷模型。", body)
    fig(NEW_FIG / "07_V7验收数据集输入质量条件.png", "V7验收数据集的通道可用性与更新周期")
    add_paragraph(anchor, "验收配置中雷达和光电更新周期为0.1秒，电子侦察为0.2秒；64场景审计中雷达、光电、电子侦察和协同链路有效输入比例约为78.8%、48.1%、47.2%和63.5%。这些比例综合反映视场、更新周期和配置缺测，说明模型输入本身具有明显不完整性。图中比例来自分层审计，不替代完整2000场景的逐槽统计。", body)
    add_table(doc, anchor, "表2-9  V7主批次数据规模与质量条件", [
        ["项目", "配置或规模", "用途", "说明"],
        ["场景数量", "2000个，四类各500个", "完整集精度和基线比较", "V7最终验收批次"],
        ["目标与平台", "每场景3目标、5平台", "多目标多平台协同", "形成6000条目标航迹"],
        ["序列长度", "10秒、0.1秒帧周期、100帧", "时序建模与评价", "评价统一排除前5帧"],
        ["传感器", "雷达、光电、电子侦察", "三类模态融合", "光电与电子侦察不直接提供距离"],
        ["链路", "0.1秒延迟、60米位置噪声标准差", "跨平台状态补充", "属于仿真配置"],
        ["评价有效点", "570000个", "总体三维位置RMSE", "6000条航迹×95帧"],
    ], body)

    add_heading(anchor, "2.3.2 干扰条件与兼容性边界", heading, page_break=True)
    add_paragraph(anchor, "为了检查数据生成与读取流程能否表达更严苛的观测条件，项目保留1个V8复合干扰场景。该场景包含雷达周期盲区、距离相关软压制和距离拖尾等机制，帧周期、目标数量和噪声配置也与V7主批次不同。它的作用是暴露边界和验证兼容性，不用于证明V7模型在强干扰下满足8米指标。", body)
    fig(test_figure("IT-01-05", "figure_3_interference_manifestation.png"), "V8复合干扰样本中的盲区和软压制表现")
    add_paragraph(anchor, "图2-34上半部分记录雷达有效观测数量随时间的变化，周期性低谷对应设定的盲区；下半部分按距离区间统计软压制后的新增缺失比例，远距离区间影响更明显。图中的百分比只描述这一个补充场景，不可推广为V7数据集或实际电磁环境的发生概率。", body)
    fig(test_figure("IT-01-05", "figure_4_smearing_and_compatibility.png"), "距离拖尾分布与400帧数据兼容性")
    add_paragraph(anchor, "距离拖尾检查比较雷达距离残差分布与设定范围，并逐帧核对400帧场景通过读取、注册和输出链路。右图的线性索引关系用于确认没有因较长序列发生截断、乱序或重复。兼容性通过只说明软件能处理该数据结构，不表示模型在该场景上的输出精度满足主批次阈值。", body)
    fig(NEW_FIG / "05_研发版本跨数据集RMSE对照.png", "研发模型与不同版本数据集的交叉验证RMSE")
    add_paragraph(anchor, "交叉验证显示V5、V6、V7模型在V7验收集上的RMSE分别约为10.380米、7.917米和7.488米，而三种模型在V8复合干扰样本上均出现明显退化。由于历史版本的形成过程、训练配置和评测协议并未全部重新统一，图2-36只用于说明版本演进和分布变化风险；最终验收结论仍限定为V7模型在V7完整验证集上的7.488米。", body)
    add_table(doc, anchor, "表2-10  数据集版本及在结题报告中的使用规则", [
        ["数据对象", "规模", "主要用途", "可形成的结论"],
        ["V7原始场景", "2000场景", "关联重算、场景与输入审计", "原始量测条件和关联统计"],
        ["V7注册验证集", "2000场景、6000航迹", "模型完整集推理与验收", "7.488米及配套完整集结果"],
        ["V7分层样本", "32或64场景", "关联矩阵、缺测和输入质量分析", "对应条件下的补充趋势"],
        ["V8复合干扰场景", "1场景、5航迹", "兼容性与极端边界观察", "仅能描述该代表样本"],
        ["V5、V6历史数据", "历次研发批次", "说明版本形成过程", "不得替代V7最终验收口径"],
    ], body)

    add_heading(anchor, "2.3.3 验证矩阵与证据归档", heading, page_break=True)
    add_paragraph(anchor, "结题阶段将验证证据分为完整集主批次、分层补充试验和受控功能用例。完整集主批次承担合同精度和同场景参考实现比较；分层试验解释关联条件、输入质量和缺测压力；受控用例检查插值、门限、坐标变换、图邻接和模型导出等实现逻辑。三类证据的样本规模和结论范围不同。", body)
    fig(NEW_FIG / "09_验证证据分层及结论使用范围.png", "项目验证证据分层及结论使用范围")
    add_paragraph(anchor, "图2-37把不同证据放在同一结构中：主批次给出最终数值，补充试验解释条件变化，功能用例证明程序按照设计执行。报告中的任何结论都应能够回到相应层级，不能用一个人工门限样例代替完整场景性能，也不能用总体RMSE反向证明每个接口分支均正确。", body)
    add_table(doc, anchor, "表2-11  研究任务、验证证据与成果文件对应关系", [
        ["研究任务", "主要验证对象", "代表证据", "成果文件"],
        ["统一表征", "字段、时间、坐标、关联、异常与缺测", "IT-01-01至IT-01-04、原始V7关联审计", "技术方案、设计说明书、测试报告"],
        ["融合跟踪模型", "状态更新、双流、时序、图协同和状态输出", "IT-02-01至IT-02-05、完整集消融", "设计说明书、模型权重、测试报告"],
        ["数据集生成", "场景数量、机动、传感器、干扰与兼容性", "IT-01-05、V7场景清单、分层质量审计", "数据集、数据说明和测试材料"],
        ["性能验收", "精度、关联、基线和模型前向时延", "V7完整集和独立指标JSON", "测试报告、结题报告"],
        ["工程原型", "模型注册、批量任务、结果保存和报告输出", "软件功能测试与运行记录", "需求、设计、使用说明书"],
    ], body)
    add_paragraph(anchor, "证据归档以任务号、模型摘要、数据摘要、场景清单、统计规则和输出文件为最小组合。位置精度保留逐航迹有效点和误差平方，关联保留正确对、输出对和真值可关联对计数，时延保留原始同步样本，图件保留数据摘要与生成程序。这样的组织方式使结题报告中的图表能够重新生成，也便于后续版本与V7基准进行回归比较。", body)

    add_heading(anchor, "2.3.4 正式批次执行与指标复核", heading, page_break=True)
    add_paragraph(anchor, "正式验证从固定被测对象开始，而不是从结果图开始。项目在任务创建前确认V7权重路径与摘要、V7注册验证集版本、2000个场景清单、前5帧排除规则、有效点定义和运行环境；任务完成后再检查场景数、航迹数、异常记录和逐航迹文件是否齐全。只有完整性检查通过，才进入总体与分类指标汇总。", body)
    fig(NEW_FIG / "10_V7正式验证批次执行流程.png", "V7正式验证批次的执行与证据归档流程")
    add_paragraph(anchor, "图2-38概括正式批次从对象冻结到结果归档的五个环节。各环节均保留可核对记录，使后续人员能够判断图表究竟来自哪一个模型、哪一份数据和哪一次运行。若模型权重、数据集、场景范围、预热规则或统计程序任一项变化，应建立新的任务身份，不能直接覆盖V7主批次。", body)
    add_paragraph(anchor, "完整性检查把“任务运行完成”和“结果可用于统计”分开。任务状态显示成功仍需确认输出文件数量、场景与航迹对应关系、有限值比例、每条航迹有效帧数以及异常日志；少量文件缺失会改变统计分母，即使总体数值看起来接近也不应与完整批次合并。对允许存在的传感器缺测，则通过掩码进入正常统计，不把缺测样本误删为执行异常。", body)
    fig(NEW_FIG / "11_结题主要指标计算边界.png", "结题主要指标的统计对象、分母与结果边界")
    add_paragraph(anchor, "位置、关联和时延三个指标使用不同统计对象。位置RMSE按570000个稳定阶段有效目标点的三维误差平方汇总；关联指标从原始场景重算关联对，分别使用输出对和真值可关联对作为分母；时延从准备好的输入张量进入模型开始，到六维状态输出结束。图2-39把三类指标的对象、分母和边界并列，防止在文档中只写一个数值而省略其含义。", body)
    add_table(doc, anchor, "表2-12  正式批次执行的质量门检查清单", [
        ["检查阶段", "检查项目", "通过条件", "异常处理"],
        ["任务创建前", "权重路径与SHA256", "与V7冻结摘要一致", "停止任务并重新选择被测模型"],
        ["任务创建前", "数据集摘要和场景清单", "V7完整注册集、2000场景", "建立新数据版本，禁止混入主批次"],
        ["任务创建前", "评价配置", "排除前5帧、有效点和维度定义一致", "修正配置后创建新任务"],
        ["任务执行后", "场景和航迹数量", "2000场景、6000航迹", "定位缺失文件并重新执行"],
        ["任务执行后", "有限输出与异常日志", "无未解释的非有限值和中断", "按场景记录并排除执行故障"],
        ["指标汇总前", "逐航迹有效点", "每条95帧，总计570000点", "核对掩码、序列长度和预热规则"],
        ["图表生成前", "指标JSON与任务身份", "摘要和分母与正文一致", "禁止手工复制其他批次数值"],
        ["材料归档前", "报告、图件和表格数字", "主指标逐项一致", "回到源记录重算并更新全部引用"],
    ], body)
    add_paragraph(anchor, "指标复核采用独立路径。位置精度由逐航迹有效点数和平方误差重新汇总，不对四类RMSE作简单算术平均；关联同时报告正确率、召回率和F1，避免高纯度掩盖漏关联；时延直接保存同步计时样本并计算P95，不以完整序列总时间除以帧数替代T=1测量。参考实现比较则按场景文件名和有效点范围配对，同时保留双方输入模态不同的说明。", body)

    add_heading(anchor, "2.3.5 版本绑定与高误差样本诊断", heading, page_break=True)
    add_paragraph(anchor, "可复核结果需要同时绑定模型、数据、运行和结果四类身份。权重文件名不能单独证明被测模型没有变化，数据目录名称也不能替代场景清单和摘要；同一份模型在不同评价配置、运行环境或场景范围下产生的结果应分别归档。项目使用摘要、任务号和结构化指标文件把这些对象联系起来。", body)
    fig(NEW_FIG / "12_模型数据运行结果身份绑定.png", "模型、数据、运行与结果的身份绑定关系")
    add_paragraph(anchor, "模型身份包含权重路径、SHA256和网络配置；数据身份包含版本、摘要、场景清单和预处理规则；运行身份包含任务号、硬件软件环境、随机种子与运行时间；结果身份包含原始输出、逐航迹记录、指标JSON和图件数据摘要。四类信息共同构成可追溯任务对象，任一部分无法对应时，结果只能作为待核记录，不能进入最终验收表。", body)
    add_paragraph(anchor, "版本绑定也用于解释历史结果。V5和V6运行记录保留其原始模型、数据与配置，不改写成V7结论；V8复合干扰样本单独归档，不与V7完整集混合计算。对V7主批次增加新的分布分析时，可以复用已经保存的逐航迹和场景结果，但必须注明是同一任务的二次统计，而不是重新运行后产生的新验收批次。", body)
    fig(NEW_FIG / "13_高误差场景诊断流程.png", "高误差场景的定位、复现与处置流程")
    add_paragraph(anchor, "高误差样本诊断从场景、航迹和时间片定位开始，再核对对应时刻的通道有效性、缓存年龄、关联状态和异常标记。随后区分误差是由高机动场景本身、时间或空间注册、目标关联、模型估计还是运行故障引起，并通过原任务复算、受控扰动或模块对照复现。只有确认输入或计算错误后才修订数据或结果，真实存在的困难样本继续保留在总体评价中。", body)
    add_paragraph(anchor, "诊断结论可以形成四类处置：数据定义或生成错误进入数据修订并建立新版本；模型在合规输入下出现系统性退化进入训练或结构优化；超出V7定义的极端条件进入适用边界和后续试验；软件执行异常进入日志、异常处理和回归测试。采用分类处置可以避免通过删除困难场景改善总体数字，也避免把接口错误归结为模型能力不足。", body)
    add_table(doc, anchor, "表2-13  高误差样本诊断信息与判定依据", [
        ["诊断层次", "需要检查的信息", "可能发现", "后续处理"],
        ["场景与运动", "机动类别、加速度、转弯率和阶段", "CSSW切换或转向突变", "增加同类样本或优化时序建模"],
        ["输入质量", "传感器有效性、缺测长度和缓存年龄", "长缺测、远距离或通道同时失效", "建立告警、降级与恢复策略"],
        ["注册与关联", "时间差、坐标范围、门控和身份连续性", "坐标偏差、漏关联或身份交换", "修订前处理并重新生成注册集"],
        ["模型估计", "残差、门控、注意力和状态修正", "权重饱和或历史偏差持续", "开展受控消融和针对性训练"],
        ["软件执行", "日志、非有限值、文件完整性和运行状态", "任务中断、结果缺失或环境差异", "修复软件后建立新任务复测"],
        ["统计与报告", "样本分母、任务号、摘要和图件来源", "批次混用或统计口径不一致", "从源记录复算并统一更新材料"],
    ], body)
    add_paragraph(anchor, "正式移交时，应将主批次任务摘要、关键指标JSON、逐航迹结果、原始计时样本、关联计数、图件生成程序和必要运行说明作为结题报告的支撑材料归档。报告正文负责解释结论和适用范围，机器可读记录负责支持复算，两者结合才能在后续模型升级、甲方复核或平台移植时保持结果连续性。", body)


def add_application_expansion(doc: Document, anchor, body, heading) -> None:
    add_heading(anchor, "3.2.1 应用对象与能力对应关系", heading, page_break=True)
    add_paragraph(anchor, "项目成果的可应用部分包括多源目标级数据组织、时空注册和关联程序、端到端融合跟踪模型、场景生成工具以及批量评价流程。不同应用方向可以复用其中一部分，但均需要重新确定传感器字段、坐标和时间基准、目标运动范围以及评价阈值，不能直接承诺V7仿真集上的米级精度。", body)
    add_table(doc, anchor, "表3-1  潜在应用方向与所需补充验证", [
        ["应用方向", "可复用能力", "新增输入与条件", "必须补充的验证"],
        ["机载协同探测", "多平台注册、图协同和连续航迹输出", "实装雷达、光电、电子侦察和链路协议", "设备误差、时钟同步、非共视与链路中断"],
        ["协同任务规划", "位置速度状态与机动趋势", "规划系统接口和状态质量标记", "闭环接口、更新频率和异常输出处置"],
        ["低空协同监管", "多站目标级融合和批量评价", "低空雷达、光电和无线电侦测数据", "密集目标、低慢小目标和合规要求"],
        ["无人机集群巡检", "跨视角信息互补和轨迹维护", "任务区域、目标或设施先验", "遮挡、通信带宽和平台数量变化"],
        ["海上目标监测", "异构观测统一和长时序状态估计", "岸基、舰载及无人平台数据", "海杂波、远距离误差和长时间缺测"],
        ["装备方案论证", "场景生成、固定批测和分层统计", "候选传感器与算法配置", "同输入、同初始化和同评价规则对比"],
    ], body)
    add_paragraph(anchor, "上述应用方向中的“可复用”是指方法和工具链具有迁移基础，不表示已经形成面向相应行业或装备的成品。实际转化应从接口适配和数据代表性检查开始，先建立新的基准任务，再决定模型是否可以直接迁移、微调或需要重新训练。", body)

    add_heading(anchor, "3.2.2 工程转化阶段与退出准则", heading, page_break=True)
    add_paragraph(anchor, "从研究原型到任务系统可划分为原型完善、半实物验证、目标平台移植和现场验证四个阶段。每一阶段都应固定输入条件、被测版本、性能指标和退出准则；上一阶段通过只说明具备进入下一阶段的条件，不自动形成最终装备应用结论。", body)
    add_table(doc, anchor, "表3-2  工程转化阶段、主要任务与退出准则", [
        ["阶段", "主要任务", "建议验证内容", "退出准则"],
        ["研究原型完善", "代码结构、配置、异常处理和自动回归", "固定V7基准、错误输入和重复运行", "模型数据任务可追溯，主要功能可重复"],
        ["半实物验证", "接入近真实格式、时间节拍和设备模型", "同步误差、通信波动、虚警漏检与负载", "形成半实物条件下的独立基线和失效记录"],
        ["目标平台移植", "算子、内存、功耗和并发适配", "精度一致性、全链路时延和资源占用", "满足目标平台功能、性能和稳定性要求"],
        ["现场任务验证", "真实场景、连续运行和故障恢复", "正常、缺测、交叉、高机动和干扰条件", "由任务方依据现场协议形成应用结论"],
    ], body)
    add_paragraph(anchor, "工程阶段的指标体系应从单一平均精度扩展到高误差尾部、错误关联持续时间、航迹中断与恢复、输入质量标记、资源占用和连续运行稳定性。特别是在线系统还需要区分模型前向时延、算法链路时延和从传感器到航迹发布的系统时延。", body)

    add_heading(anchor, "3.2.3 版本维护与风险控制", heading, page_break=True)
    add_paragraph(anchor, "模型版本、数据版本、软件版本和报告版本应保持绑定。传感器配置、平台软件、场景分布或统计规则发生变化时，应建立新的任务身份，并至少重新执行接口核对、固定基准回归和受影响指标测试。若改变训练权重或模型结构，应将其作为V7之后的新工程版本管理。", body)
    add_table(doc, anchor, "表3-3  工程应用风险与控制措施", [
        ["风险来源", "可能影响", "监测信息", "控制措施"],
        ["输入分布漂移", "标准精度下降或高误差尾部增加", "通道可用率、残差和状态质量", "建立代表性新数据并重新验证或训练"],
        ["时间与坐标配置错误", "系统性位置偏差或错误关联", "时间差、坐标范围和注册残差", "接口版本控制与受控几何回归"],
        ["长时间缺测或通信中断", "状态外推误差累积", "缓存年龄、连续缺测长度", "告警、降级、重新捕获和传统方法备份"],
        ["目标数量与平台规模变化", "图拓扑和计算负载变化", "有效节点数、边数和最坏时延", "规模压力测试与资源调度"],
        ["模型压缩或算子替换", "数值差异和精度损失", "逐维输出差异和完整集回归", "作为独立版本重新形成证据"],
        ["报告或图件批次混用", "验收数字与被测对象不一致", "模型摘要、数据摘要和任务号", "由任务对象自动生成图表并执行一致性审查"],
    ], body)
    add_paragraph(anchor, "对于高可信任务，可以在学习模型之外保留输入质量检查、物理约束和传统估计方法作为监控或降级手段。当输入超出训练范围、模型输出非有限、计算资源不足或状态质量持续下降时，系统应给出可观察状态并执行预先定义的处置流程。该能力需要在后续软件设计中明确，不宜仅依靠离线平均指标推断。", body)


def build() -> Path:
    if not BASE.is_file():
        raise FileNotFoundError(BASE)
    doc = Document(BASE)
    chapter_number_existing_figures(doc)

    body_template = find_paragraph(doc, "项目实施期为2025年9月至2026年9月。研究工作按照技术方案设计、算法与数据研发、系统集成和结题验证三个阶段推进，并围绕中期评审提出的指标口径、数据覆盖、对比条件和材料一致性问题进行了集中整改。最终形成的V7模型、完整V7注册验证集和配套证据文件构成结题验收的统一对象。")
    heading_template = find_paragraph(doc, "2.1 多模态异构数据的统一表征与融合机理")
    caption_template = find_paragraph(doc, "图3-1  项目总体技术路线与端到端范围。")

    anchor22 = find_paragraph(doc, "2.2 端到端多模态信息融合与适变跟踪框架设计")
    add_section_21(doc, anchor22, body_template, heading_template, caption_template)

    anchor23 = find_paragraph(doc, "2.3 标准数据集生成与算法验证")
    add_section_22(doc, anchor23, body_template, heading_template, caption_template)

    progress_caption = find_paragraph(doc, "表2  研究进度安排及完成情况")
    add_section_23(doc, progress_caption, body_template, heading_template, caption_template)

    section4 = find_paragraph(doc, "4.文档资料情况")
    add_application_expansion(doc, section4, body_template, heading_template)

    update_static_toc(doc)

    doc.core_properties.title = "基于多源异构数据融合的端到端适变跟踪方法研究结题报告"
    doc.core_properties.subject = "按中期报告格式扩充的结题报告稿"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
