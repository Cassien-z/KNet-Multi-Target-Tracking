from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "software_delivery" / "端到端协同跟踪算法软件_测试报告.docx"
RUN_DIR = ROOT / "artifacts" / "reports" / "software_test" / "IT-02-04" / "20260924_run01"
FIG_DIR = RUN_DIR / "figures"
RESULT = json.loads((RUN_DIR / "result.json").read_text(encoding="utf-8"))
BACKUP = RUN_DIR / "端到端协同跟踪算法软件_测试报告_写入IT-02-04前备份.docx"


def set_run_font(run, bold: bool | None = None) -> None:
    run.font.name = "楷体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    run.font.size = Pt(14)
    if bold is not None:
        run.bold = bold


def add_paragraph(cell, text: str = "", *, bold: bool = False, center: bool = False, after: float = 5.0):
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    if text:
        set_run_font(paragraph.add_run(text), bold=bold)
    return paragraph


def add_figure(cell, filename: str, caption: str) -> None:
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(FIG_DIR / filename), width=Cm(10.2))
    add_paragraph(cell, caption, center=True, after=7)


doc = Document(REPORT)
cell = doc.tables[10].cell(5, 0)
cell._element.clear_content()
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

add_paragraph(cell, "测试结果", bold=True, center=True, after=8)
add_paragraph(
    cell,
    "2026年9月24日使用Python 3.11.5、PyTorch 2.1.1在CPU环境执行IT-02-04。加载冻结V7模型best_model.pt，"
    "SHA-256为d22427274a8b680c477f2e7074d91efd03a89a5347f33862a6c6de7ce35dc51a。"
    "共执行48项构图、隔离、退化输入、置换、接口和证据一致性检查，48项通过。",
)

add_paragraph(cell, "1. 20 km邻接门限和节点有效性", bold=True)
add_paragraph(
    cell,
    "构图按“边两端节点均有有效传感器信息，且平台间欧氏距离不超过20000 m”建立跨节点边。双方有效时，19.9 km和20.0 km均建立双向边，"
    "20.1 km不建立跨节点边；在19.9 km下任一端无效或双方无效时，跨节点边均为0。所有5个固定槽位始终保留自环。边界快照中，节点0对距离19.9 km、20.0 km、20.1 km节点的邻接分别为1、1、0。",
)
add_paragraph(
    cell,
    "将节点2置为无效后，该节点对应行列除自环外均为0。图注意力各查询节点的权重和误差最大值为0，所有非邻接边注意力最大值为0，表明掩码后的归一化与邻接矩阵一致。",
)
add_figure(cell, "figure_1_adjacency_boundary.png", "图4-9-1  20 km门限、节点有效性与自环构图")
add_figure(cell, "figure_2_attention_and_isolation.png", "图4-9-2  图注意力权重与无效节点隔离")

add_paragraph(cell, "2. 无效节点隔离和节点置换", bold=True)
add_paragraph(
    cell,
    "仅在图模块输入处把无效节点2的有限特征替换为大幅随机值。该无效节点自身输出最大变化为7.126677，但所有有效节点输出最大变化为0，"
    "有效节点归一化池化结果最大变化也为0；无效节点池化权重为0。因此无效填充值不会通过图边传播到有效节点，也不进入最终池化。",
)
add_paragraph(
    cell,
    "在图模块边界同步置换5个节点的特征、邻接和池化权重，再按逆置换恢复顺序。节点级输出最大绝对差为2.384×10⁻⁷，"
    "全局池化特征最大绝对差为2.384×10⁻⁷，满足10⁻⁵容差。该结论只适用于图模块边界，不扩展为参考原点、目标关联及上游预处理无条件置换不变。",
)

add_paragraph(cell, "3. 节点数量变化和退化输入", bold=True)
add_paragraph(
    cell,
    "固定5个槽位，用mask依次模拟0至5个有效节点。有效节点数为0、1、2、3、4、5时，跨节点有向边数分别为0、0、2、6、12、20，"
    "自环数始终为5，池化归一化分母分别为0、1、2、3、4、5。各组图输出均为有限值；全无效图通过10⁻⁶稳定项得到零池化特征，不出现除零或NaN。"
    "单节点自环输出形状为[1,1,96]且有限；两个相距25 km的有效孤立节点仅保留各自自环。",
)
add_paragraph(
    cell,
    "节点数增加时池化特征L2范数依次为0、17.042、14.343、12.849、10.169、9.782。该数值只用于记录归一化响应，不把节点数增加或特征范数变化作为精度必然改善的判据。",
)
add_figure(cell, "figure_3_node_count_degradation.png", "图4-9-3  固定槽位下的节点失效、邻接和池化响应")

add_paragraph(cell, "4. GAT推理开关和完整V7消融证据", bold=True)
add_paragraph(
    cell,
    "在V7 CV代表样例上，开启和关闭GAT的输出形状均为[1,100,6]且全部有限，位置RMSE分别为5.795 m和5.945 m，"
    "两种输出最大绝对差为78.213516。关闭GAT只跳过CollaborativeGAT学习型节点交互，仍保留按有效节点进行的平台维平均。",
)
add_paragraph(
    cell,
    "复核同一冻结V7权重和完整2000场景的既有推理时消融证据：完整模型位置RMSE为7.488 m，关闭GAT后为8.853 m，"
    "两者均统计570000个有效点。该结果说明当前已训练权重使用了跨平台图聚合，但消融模型没有重新训练，不能解释为重新训练架构之间的净性能差。",
)
add_figure(cell, "figure_4_gat_switch_and_v7_ablation.png", "图4-9-4  GAT开关与V7完整集推理时消融")

add_paragraph(cell, "5. 判定", bold=True)
add_paragraph(
    cell,
    "本用例通过。跨节点边满足有效性与不超过20 km的距离条件，各槽位保留自环，无效节点池化权重为0；无效图特征不会影响有效节点和最终池化；"
    "图模块同步置换结果在容差内一致；单节点、孤立节点、节点失效和全无效图均无除零或NaN。适用边界及推理时消融口径已明确记录。",
)

for paragraph in cell.paragraphs:
    for run in paragraph.runs:
        if run._element.xpath(".//a:blip"):
            continue
        set_run_font(run)

if not BACKUP.exists():
    shutil.copy2(REPORT, BACKUP)
doc.save(REPORT)
print(REPORT)
