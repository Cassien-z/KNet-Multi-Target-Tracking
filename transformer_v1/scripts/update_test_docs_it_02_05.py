from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "docs" / "software_delivery"
OUTLINE = DOC_DIR / "端到端协同跟踪算法软件_测试大纲.docx"
REPORT = DOC_DIR / "端到端协同跟踪算法软件_测试报告.docx"
RUN_DIR = ROOT / "artifacts" / "reports" / "software_test" / "IT-02-05" / "20260924_run01"
FIG_DIR = RUN_DIR / "figures"
RESULT = json.loads((RUN_DIR / "result.json").read_text(encoding="utf-8"))


PURPOSE = (
    "验证软件从统一输入到六维状态输出的端到端链路、残差构成和坐标基准正确，并验证冻结模型在"
    "Python/PyTorch与Python ONNX Runtime下可部署、可调用且黄金样例输出一致；C接口作为补充入口验证。"
)

CONDITIONS = (
    "共用环境、冻结版本和数据隔离要求按第3章执行；Python/PyTorch与Python ONNX Runtime为本用例正式验收入口，"
    "C接口作为补充实测。MATLAB仅记录当前环境可用性，安装环境在启动阶段失败时不影响正式判定。"
    "参考ONNX导出形状固定为批量1、序列长度100、节点数5；形状改变时必须重新导出并重新开展一致性验证。"
    "使用portable/artifacts/golden_case中的随包黄金样例及其PyTorch参考输出，最大位置向量差≤1 m、"
    "最大速度分量差≤0.2 m/s。其他模态与边界组合由4.1及4.6至4.9核验。"
)

PROCEDURE = (
    "1）加载冻结V7权重和随包黄金样例，在Python/PyTorch下执行端到端推理，检查输出形状、有限性和状态顺序"
    "[px,vx,py,vy,pz,vz]。\n"
    "2）运行时捕获Learned-K六维状态、位置残差和速度残差，验证最终输出等于粗状态与交错残差之和；"
    "构造非零ref_pos变化，确认其仅叠加至位置分量。\n"
    "3）核对冻结权重、ONNX模型、黄金样例、清单和导出元数据的来源及SHA-256，验收时不覆盖原模型。\n"
    "4）使用Python ONNX Runtime加载冻结ONNX，对同一黄金输入连续推理3次，记录最大位置向量差、"
    "逐速度分量差、逐状态维最大绝对误差和超差样本数。\n"
    "5）使用ONNX模型检查器核对文件完整性、算子集、18项输入节点名、float32类型、固定轴和[1,100,6]输出。\n"
    "6）逐字段比较manifest、NPZ和小端float32二进制文件，核对元素数、字节数及行优先展开结果。\n"
    "7）补充执行C端verify_golden_case程序3次，记录返回码和相对同一PyTorch黄金输出的误差。\n"
    "8）尝试启动MATLAB入口并保存日志；若MATLAB安装环境在进入项目代码前失败，只记录环境限制，不纳入正式判定。\n"
    "9）本用例只验证固定黄金样例的一致性，不以该结果替代完整验证集位置RMSE指标。"
)

EXPECTED = (
    "六维状态顺序、残差合成、参考位置和输入内存布局正确；Python/PyTorch与Python ONNX Runtime均成功运行，"
    "相对同一PyTorch黄金输出的最大位置向量差≤1 m、最大速度分量差≤0.2 m/s，连续推理结果一致。"
    "C接口按相同阈值给出补充结果。模型、权重来源、样例和误差可追溯；其他形状不在当前接口范围内。"
    "仅证明所测固定样例的一致性，不替代完整验证集上的8 m指标。"
)


def set_run_font(run, bold: bool | None = None) -> None:
    run.font.name = "楷体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "楷体")
    run.font.size = Pt(14)
    if bold is not None:
        run.bold = bold


def replace_cell_text(cell, text: str) -> None:
    cell._element.clear_content()
    for index, line in enumerate(text.splitlines()):
        paragraph = cell.paragraphs[0] if index == 0 and cell.paragraphs else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(3)
        set_run_font(paragraph.add_run(line))


def add_paragraph(cell, text: str = "", *, bold: bool = False, center: bool = False, after: float = 5.0):
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    if text:
        set_run_font(paragraph.add_run(text), bold=bold)
    return paragraph


def add_figure(cell, filename: str, caption: str) -> None:
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(FIG_DIR / filename), width=Cm(10.2))
    add_paragraph(cell, caption, center=True, after=7)


def sync_case_definition(path: Path) -> None:
    doc = Document(path)
    table = doc.tables[11]
    table.autofit = False
    replace_cell_text(table.cell(1, 1), PURPOSE)
    replace_cell_text(table.cell(2, 1), CONDITIONS)
    replace_cell_text(table.cell(3, 1), PROCEDURE)
    replace_cell_text(table.cell(4, 1), EXPECTED)
    doc.save(path)


outline_backup = RUN_DIR / "端到端协同跟踪算法软件_测试大纲_写入IT-02-05前备份.docx"
report_backup = RUN_DIR / "端到端协同跟踪算法软件_测试报告_写入IT-02-05前备份.docx"
if not outline_backup.exists():
    shutil.copy2(OUTLINE, outline_backup)
if not report_backup.exists():
    shutil.copy2(REPORT, report_backup)

sync_case_definition(OUTLINE)
sync_case_definition(REPORT)

doc = Document(REPORT)
doc.tables[11].autofit = False
label_cell = doc.tables[11].cell(5, 0)
label_cell._element.clear_content()
label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
add_paragraph(label_cell, "测试结果", bold=True, center=True, after=0)

cell = doc.tables[11].cell(5, 1)
cell._element.clear_content()
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

add_paragraph(
    cell,
    "2026年9月24日在Windows CPU环境执行IT-02-05。正式验证环境为Python 3.11.5、PyTorch 2.1.0、"
    "ONNX Runtime 1.18.1和ONNX 1.16.2。冻结V7权重SHA-256为"
    "d22427274a8b680c 477f2e7074d91efd 03a89a5347f33862 a6c6de7ce35dc51a，"
    "交付ONNX模型SHA-256为42eb6db019592ac9 ac1a4020c6b472b0 3762950080a9aff4 a8ffee084005e0c3。"
    "正式范围共执行46项检查，46项通过。",
)

add_paragraph(cell, "1. 六维状态构成和参考坐标", bold=True)
add_paragraph(
    cell,
    "运行时捕获Learned-K六维状态、三维位置残差和三维速度残差，按[px,vx,py,vy,pz,vz]交错重构。"
    "重构输出与模型最终输出的最大绝对差为0，证明最终输出由Learned-K状态与深度残差逐元素相加构成。"
    "在ref_pos上施加[1024,-2048,4096] m的非零变化后，px、py、pz的输出变化与输入完全一致，"
    "位置最大误差为0，vx、vy、vz最大变化为0，证明参考位置只叠加至位置分量。",
)
add_figure(cell, "figure_2_state_composition.png", "图4-10-1  最终六维状态的粗状态、深度残差和参考位置构成")
add_figure(cell, "figure_3_reference_position_response.png", "图4-10-2  ref_pos对位置与速度分量的影响")

add_paragraph(cell, "2. Python/PyTorch与ONNX Runtime黄金样例一致性", bold=True)
add_paragraph(
    cell,
    "PyTorch使用冻结权重重新执行同一黄金输入，输出形状为[1,100,6]且全部有限。相对随包PyTorch黄金输出，"
    "最大位置向量差为0 m，最大速度分量差为1.526×10⁻⁵ m/s。Python ONNX Runtime关闭内存模式优化后连续推理3次，"
    "重复结果最大绝对差为0；相对黄金输出的最大位置向量差为0.5 m，最大速度分量差为0.096771 m/s，"
    "均满足1 m和0.2 m/s阈值。",
)
add_paragraph(
    cell,
    "ONNX逐维最大绝对误差依次为px 0.125 m、vx 0.030472 m/s、py 0 m、vy 0.096771 m/s、"
    "pz 0.5 m、vz 0.080139 m/s，各维超差样本数均为0。绝对ECEF位置处于百万米量级，0.125～0.5 m差异"
    "对应float32量化范围，不影响本用例判定。",
)
add_figure(cell, "figure_1_per_dimension_parity.png", "图4-10-3  ONNX Runtime相对PyTorch黄金输出的逐维误差")

add_paragraph(cell, "3. 模型签名和二进制布局", bold=True)
add_paragraph(
    cell,
    "交付ONNX通过模型完整性检查，算子集为17；18项输入名称与接口清单一致，数据类型均为tensor(float)，"
    "输出节点名为prediction、类型为tensor(float)、固定形状为[1,100,6]。逐字段核对18项输入和1项输出，"
    "共19个数组的manifest形状、元素数、4字节float32小端文件大小以及二进制与NPZ内容均一致。"
    "该验证覆盖固定B=1、T=100、N=5的交付布局，不宣称调用接口能够自动识别相同元素数但错误的形状、单位或字节序。",
)

add_paragraph(cell, "4. C接口补充验证和MATLAB环境记录", bold=True)
add_paragraph(
    cell,
    "C端verify_golden_case使用同一ONNX和同一组二进制输入连续完成3次推理，返回码为0；"
    "相对PyTorch黄金输出的最大位置向量差为0.5 m，最大速度分量差为0.047150 m/s，满足相同阈值。"
    "本机MATLAB可执行文件存在，但启动时即报“System Error: File system inconsistency”并返回1，"
    "未进入verify_golden_case或项目模型代码。该问题归类为MATLAB安装环境启动限制，按本用例修订口径保存日志但不参与正式判定。",
)
add_figure(cell, "figure_4_runtime_status_and_margin.png", "图4-10-4  各执行入口状态及黄金样例误差裕量")

add_paragraph(cell, "5. 判定", bold=True)
add_paragraph(
    cell,
    "本用例通过。Python/PyTorch和Python ONNX Runtime正式验证的46项检查全部通过；六维状态顺序、"
    "粗状态与残差合成、ref_pos叠加范围、固定输入输出签名及float32小端行优先布局均符合要求。"
    "ONNX和C补充实测的黄金样例误差均低于位置1 m、速度0.2 m/s阈值。该结论只证明所测固定黄金样例的"
    "部署一致性，不替代完整V7验证集上的8 m位置RMSE指标，也不覆盖B、T或N变化后的模型。",
)

for paragraph in cell.paragraphs:
    for run in paragraph.runs:
        if run._element.xpath(".//a:blip"):
            continue
        set_run_font(run)

doc.save(REPORT)
print(OUTLINE)
print(REPORT)
