"""Structural quality checks for the generated user manual."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DOCUMENT = ROOT / "docs" / "software_delivery" / "端到端协同跟踪原型软件_使用说明书.docx"


def main() -> None:
    doc = Document(DOCUMENT)
    all_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    all_text += "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )

    forbidden = ["版本记录", "文档版本", "编制日期", "V1.0", "共XX页", "目录将在打开文档时更新"]
    found = {term: all_text.count(term) for term in forbidden if term in all_text}
    assert not found, f"Forbidden document-control text remains: {found}"
    assert "共45页" in all_text, "Cover page count was not updated"

    body_started = False
    font_errors: list[tuple] = []
    heading_errors: list[tuple] = []
    heading_count = 0
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == "1 文档适用范围":
            body_started = True
        if not body_started:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        is_heading = style_name.startswith("Heading")
        if is_heading:
            heading_count += 1
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            r_pr = run._element.rPr
            r_fonts = r_pr.rFonts if r_pr is not None else None
            style = paragraph.style
            style_east_asia = None
            style_ascii = None
            style_size = None
            style_bold = None
            while style is not None:
                style_r_pr = style._element.get_or_add_rPr()
                style_fonts = style_r_pr.rFonts
                if style_fonts is not None:
                    style_east_asia = style_east_asia or style_fonts.get(qn("w:eastAsia"))
                    style_ascii = style_ascii or style_fonts.get(qn("w:ascii"))
                style_size = style_size or style.font.size
                if style_bold is None and style.font.bold is not None:
                    style_bold = style.font.bold
                style = style.base_style
            east_asia = r_fonts.get(qn("w:eastAsia")) if r_fonts is not None else None
            ascii_font = r_fonts.get(qn("w:ascii")) if r_fonts is not None else None
            east_asia = east_asia or style_east_asia
            ascii_font = ascii_font or style_ascii
            effective_size = run.font.size or style_size
            size = effective_size.pt if effective_size is not None else None
            if east_asia != "楷体" or ascii_font != "Times New Roman" or size != 14:
                font_errors.append((index, style_name, run.text[:30], east_asia, ascii_font, size))
            effective_bold = run.bold if run.bold is not None else style_bold
            if is_heading and effective_bold is not True:
                heading_errors.append((index, paragraph.text, run.text, run.bold))

    assert heading_count == 70, f"Unexpected heading count: {heading_count}"
    assert not font_errors, f"Body font errors: {font_errors[:10]}"
    assert not heading_errors, f"Heading emphasis errors: {heading_errors[:10]}"

    fill_errors: list[tuple] = []
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
                fill = shading.get(qn("w:fill")) if shading is not None else None
                if row_index == 0 and fill != "D9D9D9":
                    fill_errors.append((table_index, row_index, cell_index, fill, "header"))
                if row_index > 0 and fill not in (None, "auto", "FFFFFF"):
                    fill_errors.append((table_index, row_index, cell_index, fill, "body"))
    assert not fill_errors, f"Table fill errors: {fill_errors[:10]}"

    print(
        f"PASS paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} "
        f"headings={heading_count} pages=45"
    )


if __name__ == "__main__":
    main()
