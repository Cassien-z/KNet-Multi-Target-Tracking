"""Build the reviewable Word draft from the V7 technical-summary Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "结题报告_技术总结初稿_V7.md"
OUTPUT = ROOT / "docs" / "结题报告_技术总结初稿_V7.docx"
FIGURES = ROOT / "artifacts" / "reports" / "v7_final_validation" / "2ea4f9eb0f2845b599fcf4837d71a018"
FIGURE_FILES = {
    1: "01_V7分机动与总体位置精度.png",
    2: "02_V7逐航迹误差分布.png",
    3: "04_V7目标关联正确率.png",
    4: "07_V7干扰与传感器组合关联矩阵.png",
    5: "03_V7同场景基线对比.png",
    6: "05_V7模型前向时间.png",
    7: "08_V7模型模块消融.png",
    8: "06_V7信息不完备扰动实验.png",
}


def set_east_asia_font(style, name: str) -> None:
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def add_formatted_text(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1].replace("/", "/\u200b"))
            run.font.name = "Consolas"
            run.font.size = Pt(8.6)
        else:
            paragraph.add_run(part)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, code, separate, text, end):
        run._r.append(node)


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    set_east_asia_font(normal, "宋体")
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor(31, 39, 48)
    normal.paragraph_format.line_spacing = 1.38
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.first_line_indent = Cm(0.72)

    title = doc.styles["Title"]
    set_east_asia_font(title, "黑体")
    title.font.size = Pt(19)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(15)
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.keep_with_next = True
    title_ppr = title.element.get_or_add_pPr()
    for border in list(title_ppr.findall(qn("w:pBdr"))):
        title_ppr.remove(border)

    for name, size, before, after in (("Heading 1", 15, 18, 9), ("Heading 2", 12.5, 13, 6)):
        style = doc.styles[name]
        set_east_asia_font(style, "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("结题报告技术总结初稿").font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    add_page_number(footer)
    footer.add_run(" 页")
    for run in footer.runs:
        run.font.size = Pt(8)

    in_sources = False
    for raw in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        line = line.rstrip(" ")
        if line.endswith("  "):
            line = line[:-2]
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_formatted_text(paragraph, line[2:])
        elif line.startswith("## "):
            heading = line[3:]
            in_sources = heading == "证据与插图来源"
            doc.add_paragraph(heading, style="Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
        elif match := re.match(r"\*\*图\s*(\d+)\*\*\s*(.+)", line):
            number = int(match.group(1))
            figure_path = FIGURES / FIGURE_FILES[number]
            if not figure_path.is_file():
                raise FileNotFoundError(figure_path)
            img = doc.add_paragraph()
            img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img.paragraph_format.first_line_indent = Cm(0)
            img.paragraph_format.space_before = Pt(6)
            img.paragraph_format.space_after = Pt(3)
            img.paragraph_format.keep_with_next = True
            img.add_run().add_picture(str(figure_path), width=Cm(15.2))
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.first_line_indent = Cm(0)
            caption.paragraph_format.space_after = Pt(9)
            run = caption.add_run(f"图 {number}  {match.group(2)}")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(68, 75, 82)
        else:
            paragraph = doc.add_paragraph()
            if in_sources or line.startswith("**编制状态") or line.startswith("**统计截止") or line.startswith("**最终验证口径"):
                paragraph.paragraph_format.first_line_indent = Cm(0)
            if in_sources:
                paragraph.paragraph_format.space_after = Pt(4)
            add_formatted_text(paragraph, line)

    doc.core_properties.title = "基于多源异构数据融合的端对端适变跟踪方法研究技术总结初稿"
    doc.core_properties.subject = "V7模型及完整V7数据集结题技术总结"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
