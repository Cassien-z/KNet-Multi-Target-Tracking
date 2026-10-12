"""Build the combined V7 completion-report draft from its Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_v7_technical_summary_docx import (
    FIGURE_FILES,
    FIGURES,
    add_formatted_text,
    add_page_number,
    set_east_asia_font,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "基于多源异构数据融合的端对端适变跟踪方法研究_结题报告初稿.md"
OUTPUT = ROOT / "docs" / "基于多源异构数据融合的端对端适变跟踪方法研究_结题报告初稿.docx"
PROCESS_FIGURES = ROOT / "artifacts" / "figures" / "v7_process"
EXPANSION_FIGURES = ROOT / "artifacts" / "figures" / "v7_report_expansion"
SOFTWARE_TEST_FIGURES = ROOT / "artifacts" / "reports" / "software_test"
REPORT_FIGURES = {
    1: PROCESS_FIGURES / "01_V7项目总体技术路线.png",
    2: PROCESS_FIGURES / "02_V7多平台协同观测场景.png",
    3: PROCESS_FIGURES / "09_V7多平台多目标三维场景.png",
    4: PROCESS_FIGURES / "06_V7四类目标机动样本.png",
    5: PROCESS_FIGURES / "03_V7多源观测关联注册流程.png",
    6: SOFTWARE_TEST_FIGURES / "IT-01-04" / "20260923_run01" / "figures" / "figure_1_multiplatform_registration.png",
    8: PROCESS_FIGURES / "07_V7主动被动双流交互.png",
    9: SOFTWARE_TEST_FIGURES / "IT-02-02" / "20260923_run01" / "figures" / "figure_3_module_snapshot_and_fusion_gate.png",
    10: PROCESS_FIGURES / "08_V7时序建模与图协同.png",
    11: SOFTWARE_TEST_FIGURES / "IT-02-04" / "20260924_run01" / "figures" / "figure_2_attention_and_isolation.png",
    12: SOFTWARE_TEST_FIGURES / "IT-02-05" / "20260924_run01" / "figures" / "figure_1_per_dimension_parity.png",
    **{number + 12: FIGURES / filename for number, filename in FIGURE_FILES.items()},
    21: SOFTWARE_TEST_FIGURES / "IT-02-03" / "20260923_run01" / "figures" / "figure_3_missing_and_recovery.png",
    22: EXPANSION_FIGURES / "17_V7场景级误差累计分布.png",
    23: EXPANSION_FIGURES / "18_V7高误差航迹尾部贡献.png",
    24: EXPANSION_FIGURES / "19_V7前向原始计时分布.png",
    25: EXPANSION_FIGURES / "20_V7同场景配对基线分布.png",
    26: EXPANSION_FIGURES / "21_V7多源输入可用性.png",
}


def _shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def _borders(cell) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")


def _cell_margin(cell, top=95, left=110, bottom=95, right=110) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows or any(len(row) != len(rows[0]) for row in rows):
        raise ValueError("指标表列数不一致")
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    profiles = {
        "阶段": ((2.25, 3.15, 5.20, 5.20), ()),
        "条目": ((1.35, 3.15, 8.55, 2.75), (0,)),
        "成果类别": ((2.25, 4.45, 5.30, 3.80), ()),
        "输入来源": ((2.35, 2.20, 2.20, 9.05), (1, 2)),
        "场景类别": ((3.20, 2.00, 5.25, 5.35), (1, 2, 3)),
        "处理环节": ((3.20, 3.85, 4.85, 3.90), ()),
        "指标": ((3.25, 2.25, 2.65, 7.65), (1, 2)),
        "机动类别": ((3.20, 3.00, 4.10, 5.50), (1, 2, 3)),
        "主批次审计来源": ((3.10, 2.20, 2.65, 7.85), (1, 2)),
        "输入形态": ((3.20, 3.90, 5.20, 3.50), (3,)),
        "64 场景测试条件": ((3.80, 2.20, 2.65, 7.15), (1, 2)),
        "内容": ((2.25, 4.05, 4.80, 4.70), ()),
        "复核环节": ((2.25, 4.15, 4.70, 4.70), ()),
        "V7 机动类别": ((2.65, 1.65, 3.85, 3.85, 3.80), (1, 2, 3, 4)),
    }
    width_values, center_cols = profiles.get(rows[0][0], ((3.25, 2.25, 2.65, 7.65), (1, 2)))
    widths = [Cm(value) for value in width_values]
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _borders(cell)
            _cell_margin(cell)
            if i == 0:
                _shade(cell, "324B64")
            elif i % 2 == 0:
                _shade(cell, "F3F6F8")
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.16
            if i == 0:
                p.paragraph_format.keep_with_next = True
            if j in center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_formatted_text(p, value.strip())
            for run in p.runs:
                run.font.size = Pt(10.5)
                if i == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
        trpr = table.rows[i]._tr.get_or_add_trPr()
        split = OxmlElement("w:cantSplit")
        trpr.append(split)
        if i == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_reserved_figure(doc: Document, number: int, caption_text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(15.2)
    row = table.rows[0]
    row.height = Cm(7.7)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = row.cells[0]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _shade(cell, "F7F9FA")
    _borders(cell)
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    for side in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{side}"))
        edge.set(qn("w:val"), "dashed")
        edge.set(qn("w:color"), "AAB9C2")
        edge.set(qn("w:sz"), "7")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run("模型总体架构图预留位\n待插入项目组自绘的 V7 最终架构图")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(111, 127, 137)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.space_after = Pt(9)
    run = caption.add_run(f"图 {number}  {caption_text}")
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(68, 75, 82)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.1)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    set_east_asia_font(normal, "仿宋")
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(27, 35, 43)
    normal.paragraph_format.line_spacing = 1.42
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.first_line_indent = Cm(0.98)

    title = doc.styles["Title"]
    set_east_asia_font(title, "黑体")
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.space_after = Pt(0)
    title.paragraph_format.keep_with_next = True
    title_ppr = title.element.get_or_add_pPr()
    for border in list(title_ppr.findall(qn("w:pBdr"))):
        title_ppr.remove(border)

    for name, size, before, after in (("Heading 1", 16, 20, 11), ("Heading 2", 14.5, 14, 7)):
        style = doc.styles[name]
        set_east_asia_font(style, "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("基于多源异构数据融合的端对端适变跟踪方法研究  结题报告初稿")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(85, 93, 100)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    add_page_number(footer)
    footer.add_run(" 页")
    for item in footer.runs:
        item.font.size = Pt(8)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("基于多源异构数据融合的\n端对端适变跟踪方法研究")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Cm(0)
    subtitle.paragraph_format.space_before = Pt(25)
    run = subtitle.add_run("结题报告初稿")
    run.font.name = "黑体"
    run.font.size = Pt(18)
    run.font.bold = True
    for _ in range(9):
        doc.add_paragraph()
    for line in (
        "工作总结与技术总结合并稿",
        "统计截止  2026 年 9 月 17 日",
        "事实复核  2026 年 9 月 21 日",
        "最终性能主口径  V7 模型与完整 V7 注册验证集",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(10)
        p.add_run(line).font.size = Pt(11)
    doc.add_page_break()


def build() -> Path:
    doc = Document()
    configure(doc)
    add_cover(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_appendix = False
    after_heading = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if (
            line.startswith("# ")
            or line.startswith("**编制状态：")
            or line.startswith("**统计截止：")
            or line.startswith("**事实复核：")
            or line.startswith("**最终性能主口径：")
        ):
            continue
        if line.startswith("## "):
            heading = line[3:]
            in_appendix = heading.startswith("附录")
            doc.add_paragraph(heading, style="Heading 1")
            after_heading = True
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
            after_heading = True
            continue
        if line.startswith("|"):
            rows = []
            while True:
                if not (line.startswith("|") and line.endswith("|")):
                    raise ValueError(f"表格格式错误：{line}")
                cells = [item.strip() for item in line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", item) for item in cells):
                    rows.append(cells)
                if i >= len(lines) or not lines[i].strip().startswith("|"):
                    break
                line = lines[i].strip()
                i += 1
            add_table(doc, rows)
            after_heading = False
            continue
        if match := re.match(r"\*\*图\s*(\d+)\s*预留\*\*\s*(.+)", line):
            add_reserved_figure(doc, int(match.group(1)), match.group(2))
            after_heading = False
            continue
        if match := re.match(r"\*\*图\s*(\d+)\*\*\s*(.+)", line):
            number = int(match.group(1))
            figure_path = REPORT_FIGURES[number]
            if not figure_path.is_file():
                raise FileNotFoundError(figure_path)
            image = doc.add_paragraph()
            image.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image.paragraph_format.first_line_indent = Cm(0)
            image.paragraph_format.space_before = Pt(6)
            image.paragraph_format.space_after = Pt(3)
            image.paragraph_format.keep_with_next = True
            image.add_run().add_picture(str(figure_path), width=Cm(15.2))
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.first_line_indent = Cm(0)
            caption.paragraph_format.space_after = Pt(9)
            run = caption.add_run(f"图 {number}  {match.group(2)}")
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(68, 75, 82)
            after_heading = False
            continue
        paragraph = doc.add_paragraph()
        if after_heading:
            paragraph.paragraph_format.keep_together = True
            after_heading = False
        if in_appendix or line.startswith("**说明："):
            paragraph.paragraph_format.first_line_indent = Cm(0)
        if in_appendix:
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.30
        add_formatted_text(paragraph, line)
        if not in_appendix:
            for run in paragraph.runs:
                if run.font.name == "Consolas":
                    run.font.size = Pt(11)
        if in_appendix:
            for run in paragraph.runs:
                run.font.size = Pt(14)

    doc.core_properties.title = "基于多源异构数据融合的端对端适变跟踪方法研究结题报告初稿"
    doc.core_properties.subject = "工作总结与技术总结合并 V7 最终验证口径"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
