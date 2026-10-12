from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "software_delivery" / "端到端协同跟踪算法软件_测试报告.docx"
RUN_DIR = ROOT / "artifacts" / "reports" / "software_test" / "IT-03-02" / "20260924_run01"
FIG_DIR = RUN_DIR / "figures"
RESULT = json.loads((RUN_DIR / "result.json").read_text(encoding="utf-8"))
BACKUP = RUN_DIR / "端到端协同跟踪算法软件_测试报告_写入IT-03-02前备份.docx"


def set_run_font(run, bold: bool | None = None) -> None:
    run.font.name = "楷体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "楷体")
    run.font.size = Pt(14)
    if bold is not None:
        run.bold = bold


def add_paragraph(cell, text: str = "", *, bold: bool = False, center: bool = False, after: float = 5.0):
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    if text:
        set_run_font(paragraph.add_run(text), bold=bold)
    return paragraph


def add_figure(cell, filename: str, caption: str) -> None:
    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(FIG_DIR / filename), width=Cm(10.2))
    add_paragraph(cell, caption, center=True, after=7)


assert RESULT["overall_passed"] is True
assert RESULT["passed_count"] == RESULT["check_count"]
assert RESULT["strong_interference"]["all_cells_passed"] is True
for figure in RESULT["figures"]:
    assert (FIG_DIR / figure).is_file(), figure

if not BACKUP.exists():
    shutil.copy2(REPORT, BACKUP)

doc = Document(REPORT)
table = doc.tables[13]
table.autofit = False

label_cell = table.cell(5, 0)
label_cell._element.clear_content()
label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
add_paragraph(label_cell, "测试结果", bold=True, center=True, after=0)

cell = table.cell(5, 1)
cell._element.clear_content()
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

add_paragraph(
    cell,
    "2026年9月24日在Windows环境使用Python 3.11.5执行IT-03-02。正式全量证据绑定批次"
    "46fada4af5ef4c70a03fa56ef521f8c8，并从V7原始MAT重新运行measurement_only_hungarian_with_"
    "temporal_memory关联流程；强干扰稳定性使用平衡抽取的V7场景生成确定性内存扰动副本，未修改模型权重、"
    "训练数据或磁盘原始数据。全部验收检查均通过。",
)

add_paragraph(cell, "1. V7完整集正式关联正确率", bold=True)
add_paragraph(
    cell,
    "完整覆盖2000个V7场景，按correct/total计算：正确关联4238696次，总审计事件4257909次，关联正确率为"
    "99.548769%，高于85%阈值14.548769个百分点。雷达、光电和电子侦察分项分别为99.556518%"
    "（2354882/2365372）、99.644260%（1259349/1263845）和99.327652%（624465/628692）。"
    "所有分母均大于0，分项之和与总体计数一致。",
)
add_paragraph(
    cell,
    "独立证据的job_id、job_signature、模型、数据集、序列长度、预热帧和场景清单均与本轮完整批测一致。"
    "truth_used_for_observation_grouping=false，files_flagged_for_truth_leakage=0，证据状态为"
    "valid_for_assessment=true。",
)
add_figure(cell, "figure_1_full_v7_association.png", "图4-12-1  V7完整集总体及分传感器关联正确率")

add_paragraph(cell, "2. 多目标强干扰分层稳定性", bold=True)
add_paragraph(
    cell,
    "从CV、CA、CT和CSSW各选8个场景，共32个场景，其中16个为典型难场景。每个场景分别执行无干扰、"
    "雷达闪烁、70%软压制、100 m距离抖动和复合干扰，共160次关联重算；再按雷达、雷达+光电、雷达+电子侦察、"
    "雷达+光电+电子侦察四种组合形成80个独立判定格。80格分母均大于0、均未使用真值分组，且全部达到85%。",
)
add_paragraph(
    cell,
    "全传感器组合在无干扰、雷达闪烁、70%软压制、100 m距离抖动和复合干扰下的汇总正确率依次为"
    "99.016572%、99.023041%、95.512096%、99.010700%和95.008092%。各条件所有场景类型和传感器组合中的"
    "最低正确率为91.744285%，出现在70%软压制/CV/雷达组合，仍高于阈值6.744285个百分点。",
)
add_paragraph(
    cell,
    "雷达闪烁平均保留91.89%的有效雷达观测，70%软压制平均保留29.74%，复合干扰平均保留27.51%；"
    "距离抖动对活动观测数量不作删除。观测抑制程度单独记录，关联正确率只按实际进入审计的关联事件计算，"
    "不把缺失观测伪记为正确关联。",
)
add_figure(cell, "figure_2_interference_conditions.png", "图4-12-2  五种干扰条件的汇总正确率与最低单格结果")
add_figure(cell, "figure_3_condition_class_heatmap.png", "图4-12-3  全传感器组合下运动类型与干扰条件分层结果")

add_paragraph(cell, "3. 阈值、分母和真值泄漏防护", bold=True)
add_paragraph(
    cell,
    "阈值边界测试中，84.99%判为不通过，85.00%和85.01%判为通过；total=0返回无效结论，不得按100%处理。"
    "构造99.9%但truth_used_for_observation_grouping=true的证据时，判定仍为不通过。源码核查确认先完成"
    "基于测量的关联分组，再读取Data_Label用于事后映射和正确性审计，且AssociationStats显式记录真值未用于观测分组。",
)

add_paragraph(cell, "4. 可复现性和适用范围", bold=True)
add_paragraph(
    cell,
    "当前代码对V7完整原始MAT连续重算两次，两次correct/total均为4238696/4257909，结果完全一致。"
    "2026年9月17日旧证据为4142948/4162065；两者雷达计数相同，当前流程增加了EO和ESM关联事件覆盖，"
    "因此旧口径只作范围对照，不作为本轮复现基准。当前代码对160次强干扰审计也连续执行两次，两次在80格"
    "correct/total上逐格一致；旧强干扰证据与当前证据均为80格通过，但因EO/ESM事件覆盖增加，最低格由"
    "90.963455%变为91.744285%。V8当前只有1个复合强干扰场景，仅作为典型案例保留，"
    "未混入V7大样本统计，也不据此给出V8总体百分比结论。",
)
add_figure(cell, "figure_4_reproducibility_and_guards.png", "图4-12-4  全量计数复现及阈值、空分母与泄漏防护")

add_paragraph(cell, "5. 判定", bold=True)
add_paragraph(
    cell,
    "本用例通过。V7完整2000场景的目标关联正确率为99.548769%，满足不低于85%的技术指标；32个平衡场景、"
    "5种干扰条件和4种传感器组合形成的80个强干扰审计格全部通过，最差格为91.744285%。证据与正式批测范围"
    "绑定，分母非零，无真值泄漏，边界和异常防护符合测试大纲要求。",
)

for paragraph in cell.paragraphs:
    for run in paragraph.runs:
        if run._element.xpath(".//a:blip"):
            continue
        set_run_font(run)

doc.save(REPORT)
print(REPORT)
