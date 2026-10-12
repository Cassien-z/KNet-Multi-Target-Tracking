"""Build the software user manual from the client-provided Word template."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.build_software_docs import (
    add_numbered_steps,
    add_table,
    prevent_row_split,
    set_repeat_table_header,
    set_table_borders,
)


TEMPLATE = ROOT / "docs" / "software_delivery" / "资料模板" / "OODA-6-使用说明书.docx"
OUTPUT = ROOT / "docs" / "software_delivery" / "端到端协同跟踪原型软件_使用说明书.docx"
WORK_DIR = ROOT / ".codex_work" / "user_manual"

EAST_ASIA_FONT = "楷体"
LATIN_FONT = "Times New Roman"
GRAY = "D9D9D9"


def _set_run_font(run, *, bold: bool | None = None) -> None:
    run.font.name = LATIN_FONT
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:cs"), LATIN_FONT)
    run.font.size = Pt(14)
    if bold is not None:
        run.bold = bold


def _replace_cover_text(doc: Document) -> None:
    replacements = {
        "面向OODA闭环的“协同感知-决策”误差分析与联合优化项目": "基于多源异构数据融合的端对端适变跟踪方法研究",
        "使 用 说 明 书": "使 用 说 明 书",
        "共8页": "共XX页",
        "承担单位：上海交通大学": "承担单位：华东理工大学",
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text not in replacements:
            continue
        replacement = replacements[text]
        if paragraph.runs:
            paragraph.runs[0].text = replacement
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(replacement)


def _set_toc_levels(doc: Document) -> None:
    for node in doc._element.xpath(".//w:instrText"):
        if node.text and "TOC" in node.text.upper():
            node.text = ' TOC \\o "1-2" \\h \\z \\u '


def _remove_body_from(doc: Document, heading_text: str) -> int:
    body = doc._element.body
    boundary = None
    for index, child in enumerate(body):
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, doc._body)
        if paragraph.text.strip() == heading_text:
            boundary = index
            break
    if boundary is None:
        raise RuntimeError(f"Cannot locate template heading: {heading_text}")
    for child in list(body)[boundary:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return boundary


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.22)
        paragraph.add_run("● ")
        paragraph.add_run(item)


def _command(doc: Document, command: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.34)
    paragraph.paragraph_format.right_indent = Inches(0.2)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.add_run(command)


def _lead(doc: Document, lead: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run(f"{lead}：").bold = True
    paragraph.add_run(text)


def _build_body(doc: Document) -> None:
    doc.add_heading("1 文档适用范围", level=1)
    doc.add_heading("1.1 使用对象", level=2)
    doc.add_paragraph(
        "本说明书面向端到端协同跟踪原型软件的操作人员、算法验证人员、测试人员和交付验收人员，说明软件运行前的准备工作、模型和数据选择、单场景回放、批量测试、性能评估、报告归档、跨语言验证以及异常处理方法。读者按照本文步骤操作，可以完成从运行环境检查到正式评测材料导出的完整流程。"
    )
    doc.add_paragraph(
        "普通操作人员主要使用系统总览、模型管理、跟踪验证、批量测试、性能评估和报告中心六个页面。维护人员还需要了解环境变量、模型与数据目录、日志和报告位置。MATLAB或C集成人员只需使用固定ONNX接口和黄金样例，不需要进入Python训练流程。"
    )

    doc.add_heading("1.2 软件边界", level=2)
    doc.add_paragraph(
        "软件以训练完成的端到端协同跟踪模型为核心。模型接收已经完成时间对齐、目标关联和空间注册的多源序列，直接输出每个时刻的六维目标状态。前端和后端用于配置模型、组织数据、调用推理、回放结果、统计指标和保存证据，不参与模型参数训练，也不生成飞行控制或武器控制指令。"
    )
    doc.add_paragraph(
        "单场景页面在开始回放前完成一次完整100帧模型前向，随后按界面时间轴发送和显示已经计算的结果。页面约20帧每秒的播放节奏只用于观看，不代表模型以该频率在线预测。V8源数据按40 Hz生成并保留，现有权重使用同步抽取的约10 Hz兼容视图，不需要重新训练。"
    )

    doc.add_heading("1.3 推荐使用顺序", level=2)
    add_numbered_steps(doc, [
        "检查模型、数据、Python环境、Node.js环境和可用计算设备。",
        "启动后端，访问健康检查地址确认模型加载成功。",
        "启动前端并在系统总览中确认服务、设备和任务状态。",
        "在模型管理中选择模型和设备，再通过跟踪验证执行单场景链路检查。",
        "在批量测试中执行抽样测试或完整测试，并在性能评估中生成独立证据。",
        "在报告中心复核任务条件和判定结果，下载JSON和CSV制品后归档。",
    ])
    doc.add_paragraph(
        "单场景回放适合确认输入、模型和显示链路是否正常；抽样批测适合检查配置和估计耗时；完整批测及与任务绑定的独立证据用于形成正式技术指标结论。三类操作的用途不同，不应以少量场景结果代替完整数据集结论。"
    )

    doc.add_heading("2 软件组成与交付目录", level=1)
    doc.add_heading("2.1 功能组成", level=2)
    doc.add_paragraph(
        "软件由模型运行内核、数据解析与批测服务、Web服务、浏览器端操作界面和跨语言推理组件组成。模型运行内核负责加载CoopTrackerModel权重并执行固定窗口推理；数据解析模块把注册场景转换为18项输入张量；批测服务负责任务队列、指标统计和报告持久化；浏览器端负责用户操作和结果展示；portable目录提供MATLAB和C固定尺寸调用入口。"
    )
    add_table(doc, ["组成部分", "主要位置", "使用目的"], [
        ["核心模型与数据解析", "src", "构造网络、解析注册数据并执行张量预处理"],
        ["后端服务", "tracking_system/backend", "模型管理、单场景推理、批量测试、评估和报告接口"],
        ["前端界面", "tracking_system/frontend", "系统总览、模型管理、跟踪验证、批测、评估和报告操作"],
        ["模型权重", "results", "保存V5、V6、V7训练运行及best_model.pt"],
        ["注册数据", "data", "保存V5至V8注册验证集及V8兼容视图"],
        ["输出制品", "artifacts", "保存指标、批测报告、证据和必要的图表"],
        ["跨语言组件", "portable", "提供ONNX模型、MATLAB入口、C接口和黄金样例"],
        ["交付文档", "docs/software_delivery", "保存设计说明、需求规格和使用说明"],
    ], widths=[1.55, 2.25, 3.25])

    doc.add_heading("2.2 关键文件", level=2)
    doc.add_paragraph(
        "操作前应确认模型权重、注册验证数据和运行入口完整。模型仓库登记的是V5、V6、V7最佳权重；单场景验证默认使用V7注册验证集；批测页面还可以选择V5、V6、V7、V8强干扰验证集以及按场景类别组织的验收想定。"
    )
    add_table(doc, ["文件或目录", "作用", "使用注意事项"], [
        ["tracking_system/backend/server.py", "后端服务入口", "应从项目根目录以模块方式启动"],
        ["tracking_system/frontend/package.json", "前端依赖和启动命令", "首次运行或依赖变化后执行npm install"],
        ["results/.../best_model.pt", "模型权重", "正式评测前记录实际文件、大小和修改时间"],
        ["data/val_data_v7_registered", "默认单场景与V7批测数据", "保持只读，不在原目录写入报告"],
        ["data/val_data_v8_interference_registered", "V8 40 Hz注册源数据", "保留原始400帧序列，不直接交给现有100帧权重"],
        ["data/val_data_v8_interference_registered_10hz", "V8模型兼容数据", "由统一索引同步抽取为约10 Hz、100帧"],
        ["artifacts/reports/batch_tests", "批测任务报告", "已完成报告可在后端重启后重新载入"],
        ["artifacts/metrics/acceptance", "独立指标证据", "必须与对应批测任务来源签名一致"],
    ], widths=[2.6, 2.1, 2.35])

    doc.add_heading("2.3 输出状态与文件格式", level=2)
    doc.add_paragraph(
        "单场景验证输出逐帧时间、真值、预测位置、位置误差和运行RMSE；批量测试输出任务配置、实际场景清单、稳定阶段指标、时延、逐航迹结果和异常信息；性能评估在批测结果上增加阈值、证据状态和综合结论。输出主要采用JSON和UTF-8 CSV，浏览器打印功能可将报告详情保存为PDF。"
    )
    doc.add_paragraph(
        "JSON适合保存完整配置和嵌套证据，CSV适合查看逐航迹指标或横向对比。归档时应同时保留批测原始JSON、综合评估JSON、逐航迹CSV、必要的比较CSV以及运行环境记录，不能只保存页面截图。"
    )

    doc.add_heading("3 使用前准备", level=1)
    doc.add_heading("3.1 运行环境", level=2)
    add_table(doc, ["项目", "最低或推荐条件", "说明"], [
        ["操作系统", "Windows 10或Windows 11", "当前启动命令和交付路径按Windows环境编写"],
        ["Python环境", "项目现有Conda pytorch环境", "应包含PyTorch、FastAPI、Uvicorn、SciPy、h5py和WebSocket依赖"],
        ["Node.js和npm", "能够安装并运行前端依赖", "用于启动Vue和Vite前端"],
        ["浏览器", "支持现代JavaScript、WebSocket和WebGL", "建议启用硬件加速以显示三维航迹"],
        ["内存", "不少于8 GB，建议16 GB及以上", "完整数据集批测需要额外内存和磁盘缓存"],
        ["计算设备", "CPU可运行，正式时延测试建议CUDA GPU", "CUDA设备必须与当前PyTorch和驱动兼容"],
        ["磁盘空间", "为数据、权重、依赖和报告预留空间", "node_modules、多个数据集和批测报告会持续占用空间"],
    ], widths=[1.35, 2.55, 3.15])

    doc.add_heading("3.2 目录与权限检查", level=2)
    doc.add_paragraph(
        "运行账户应具备项目目录的读取权限，并能够向artifacts/reports和artifacts/metrics写入文件。数据集和已验收模型建议设置为只读或通过备份保护，防止操作中误覆盖。报告目录必须可写，否则批测任务即使完成计算也无法形成可追溯制品。"
    )
    _bullets(doc, [
        "确认项目根目录能够访问src、tracking_system、data、results、artifacts和portable。",
        "确认默认V7权重存在，或在启动前设置有效的TRACKER_MODEL_PATH。",
        "确认默认V7注册验证集存在，V8测试时同时确认40 Hz源目录和10 Hz兼容目录。",
        "确认8000和5173端口未被其他程序占用，浏览器可访问本机回环地址。",
        "确认报告目录可创建和修改文件，正式归档目录具备备份条件。",
    ])

    doc.add_heading("3.3 端口与访问地址", level=2)
    add_table(doc, ["项目", "默认地址", "用途"], [
        ["后端服务", "http://127.0.0.1:8000", "提供REST接口和WebSocket回放"],
        ["健康检查", "http://127.0.0.1:8000/api/health", "确认服务、活动模型和批测管理器状态"],
        ["接口文档", "http://127.0.0.1:8000/docs", "查看并试用FastAPI接口"],
        ["前端界面", "http://127.0.0.1:5173", "进入六个主要操作页面"],
        ["跟踪回放", "ws://127.0.0.1:8000/ws/tracking", "发送场景请求和播放控制，接收逐帧结果"],
    ], widths=[1.45, 3.2, 2.4])
    doc.add_paragraph(
        "后端默认只监听本机127.0.0.1，前端跨域来源限制为本机5173端口。若部署到其他主机、端口或代理环境，需要同时调整监听地址、前端API基址和跨域配置，并重新进行安全与功能检查。"
    )

    doc.add_heading("4 软件安装与启动", level=1)
    doc.add_heading("4.1 后端启动", level=2)
    doc.add_paragraph(
        "后端应从transformer_v1项目根目录启动，以保证模型、数据和报告相对路径解析一致。首次启动会检查登记模型并加载默认V7权重；模型较大或CUDA首次初始化时，健康检查可能需要等待一段时间。"
    )
    add_numbered_steps(doc, [
        "打开Anaconda Prompt或已经完成Conda初始化的PowerShell。",
        "进入transformer_v1项目根目录。",
        "使用pytorch环境启动后端服务。",
        "看到Uvicorn启动信息后访问健康检查地址。",
        "确认status为ready，并记录活动模型和推理设备。",
    ])
    _command(doc, "conda run --no-capture-output -n pytorch python -m tracking_system.backend.server")
    _command(doc, "http://127.0.0.1:8000/api/health")
    doc.add_paragraph(
        "如需在已激活的pytorch环境中运行，可执行python -m tracking_system.backend.server。不要直接在其他工作目录运行server.py，否则Python模块搜索路径和相对资源位置可能与交付配置不一致。"
    )

    doc.add_heading("4.2 前端安装与启动", level=2)
    doc.add_paragraph(
        "前端使用Vue和Vite。首次使用、package.json变化或node_modules缺失时需要安装依赖；依赖未变化时可直接启动。前端服务启动后，终端会显示本地访问地址。"
    )
    add_numbered_steps(doc, [
        "另开一个PowerShell窗口。",
        "进入tracking_system/frontend目录。",
        "首次运行或依赖变化后执行npm install。",
        "执行npm run dev启动前端开发服务。",
        "在浏览器中打开终端显示的本地地址。",
    ])
    _command(doc, "cd tracking_system/frontend")
    _command(doc, "npm install")
    _command(doc, "npm run dev")

    doc.add_heading("4.3 启动后检查", level=2)
    doc.add_paragraph(
        "进入前端后，左侧底部应显示后端服务在线。系统总览应能够读取服务版本、活动模型、推理设备、CUDA状态、运行环境、数据集概况和批测队列。如果页面显示后端未连接，应先直接访问健康检查地址，区分后端未启动、端口冲突、模型加载失败和浏览器访问问题。"
    )
    add_table(doc, ["检查项", "正常现象", "异常时优先检查"], [
        ["健康检查", "status为ready", "后端终端、模型路径和Python依赖"],
        ["活动模型", "显示模型标识、设备和加载耗时", "权重文件、d_model结构和CUDA设备"],
        ["模型仓库", "V5至V7登记模型可见", "results目录及模型登记路径"],
        ["数据集目录", "数据集显示可用状态和场景数", "data目录、V8兼容目录和文件权限"],
        ["报告目录", "历史已完成任务可查询", "artifacts/reports/batch_tests写权限"],
        ["三维显示", "跟踪验证页面能够建立图形区域", "浏览器WebGL和硬件加速"],
    ], widths=[1.4, 2.55, 3.05])

    doc.add_heading("5 界面导航与系统总览", level=1)
    doc.add_heading("5.1 界面导航", level=2)
    add_table(doc, ["页面", "地址标识", "主要用途"], [
        ["系统总览", "#overview", "检查服务、模型、设备、队列、环境和技术指标"],
        ["模型管理", "#models", "查看登记权重、刷新仓库并切换活动模型"],
        ["跟踪验证", "#tracking", "选择场景和目标，查看完整序列推理回放"],
        ["批量测试", "#batch", "配置、提交、查看和取消批测任务"],
        ["性能评估", "#evaluation", "按四项指标判定并管理独立证据"],
        ["报告中心", "#reports", "检索归档、复核判定并导出制品"],
    ], widths=[1.4, 1.55, 4.05])
    doc.add_paragraph(
        "桌面端使用左侧导航切换页面，窄屏设备通过顶部菜单按钮打开导航。浏览器地址中的Hash可以直接定位页面。左侧底部显示后端连接状态，旁边的设置按钮用于查看API地址、服务信息、当前模型、推理设备、PyTorch和CUDA状态、报告目录以及后端启动命令。"
    )

    doc.add_heading("5.2 系统总览", level=2)
    doc.add_paragraph(
        "系统总览是每次操作的起点。页面从后端读取运行状态并每30秒自动同步，用户也可以手动刷新。刷新结果应以当前后端返回为准；浏览器页面上的历史卡片或模型摘要只用于快速识别，不替代本次批测报告。"
    )
    add_numbered_steps(doc, [
        "确认服务状态为在线，最近更新时间持续更新。",
        "核对活动模型、推理设备、CUDA状态和活动任务数量。",
        "查看Python、PyTorch、操作系统和显卡信息是否符合评测环境。",
        "查看数据集和模型数量，确认本次需要的资源可用。",
        "核对四项技术指标阈值及其统计口径。",
    ])
    doc.add_paragraph(
        "四项指标包括稳定阶段三维位置RMSE不超过8 m、数据关联正确率不低于85%、相对同条件IMM-JPDA基线的RMSE降幅不低于15%、完整序列模型前向时间折算到单帧不超过100 ms。缺少独立证据时页面应显示数据不足，而不能显示通过。"
    )

    doc.add_heading("5.3 页面状态与提示", level=2)
    doc.add_paragraph(
        "页面操作会使用成功或失败提示说明模型部署、任务创建、证据生成和文件下载结果。出现错误时，应记录页面提示、任务ID或请求ID，并结合后端终端定位。刷新页面不会使运行中的批测变为完成，也不会改变活动模型；实际状态由后端任务记录和模型管理器维护。"
    )
    doc.add_paragraph(
        "浏览器关闭后，后端批测仍可继续运行；后端进程退出则未完成任务会被中断。重新启动后端时，已完成报告会重新载入，先前处于排队或运行状态而无法续跑的任务不会被伪装为正常完成。"
    )

    doc.add_heading("6 模型管理", level=1)
    doc.add_heading("6.1 查看模型仓库", level=2)
    doc.add_paragraph(
        "模型管理页面列出登记的V5、V6、V7权重及其训练集、验证集、权重文件、文件大小、修改时间和可用历史摘要。ACTIVE MODEL标记表示当前单场景推理使用的模型。批量测试可以独立选择其他登记模型，不要求先切换活动模型。"
    )
    doc.add_paragraph(
        "模型卡片显示可用状态时，只能证明登记路径和基本文件检查通过。正式使用前仍需执行部署或批测，以验证权重与当前CoopTrackerModel结构能够严格匹配。历史RMSE和时延字段来自已有评估文件，只能作为参考。"
    )

    doc.add_heading("6.2 切换活动模型", level=2)
    add_numbered_steps(doc, [
        "在部署设备区域选择auto、cpu或可用的cuda设备。",
        "找到目标模型卡片并单击部署。当前活动模型的部署按钮不可用。",
        "在确认窗口核对模型标识和目标设备。",
        "单击确认部署并等待权重严格加载完成。",
        "检查ACTIVE MODEL标记、设备和加载耗时是否更新。",
        "进入跟踪验证执行一个已知场景，确认输出数值有限且回放正常。",
    ])
    doc.add_paragraph(
        "切换采用先加载候选模型、校验成功后再替换活动模型的方式。新权重缺失、键或形状不匹配、CUDA不可用或显存不足时，切换失败且旧模型继续服务。已经开始的WebSocket回放继续使用建立连接时的模型，新连接使用切换后的模型。"
    )

    doc.add_heading("6.3 刷新模型仓库", level=2)
    doc.add_paragraph(
        "新增、恢复或替换登记权重后，单击刷新模型仓库重新检查文件。刷新只更新仓库信息，不自动切换活动模型。若替换了同一路径下的权重，应重新部署并执行链路检查；正式归档时还应记录权重文件摘要，避免仅凭模型名称判断文件一致。"
    )
    _lead(doc, "注意", "模型权重必须与当前d_model为96的网络结构一致。不要通过关闭严格加载或忽略缺失键的方式强行部署不匹配权重。")

    doc.add_heading("7 单场景跟踪验证", level=1)
    doc.add_heading("7.1 场景与目标选择", level=2)
    doc.add_paragraph(
        "跟踪验证用于检查一条目标航迹的输入解析、模型前向、误差计算、回放和导出链路。页面从活动模型的数据目录读取场景，场景ID与后端稳定排序后的文件对应。机动类型筛选包括CV、CA、CT、CSSW和全部；目标索引从0开始。"
    )
    add_numbered_steps(doc, [
        "确认活动模型和单场景数据目录符合本次检查目的。",
        "选择机动类型，必要时选择全部以查看完整场景列表。",
        "选择具体场景并查看场景名称、机动类别和目标数量。",
        "输入目标索引，第一个目标使用0。",
        "单击重新推理，等待完整序列计算完成并开始回放。",
    ])
    doc.add_paragraph(
        "场景不存在、目标索引越界、注册目标为空、有效帧为空或输入包含非有限值时，后端会拒绝当前请求。此时应根据页面提示重新选择资源或检查数据，不应通过修改前端参数范围绕过后端校验。"
    )

    doc.add_heading("7.2 回放结果", level=2)
    doc.add_paragraph(
        "三维区域同时绘制真值航迹和预测航迹，并用独立标记显示当前时刻位置。为避免ECEF百万米量级绝对坐标压缩局部运动，图形以首帧真值位置为显示原点，并在三个方向使用一致空间比例。图下读数和导出文件仍保留原始绝对坐标。"
    )
    _bullets(doc, [
        "当前位置误差为预测位置与真值位置之间的三维欧氏距离。",
        "当前RMSE从最近一次开始或跳转后累计的误差计算。",
        "单帧推理时间为完整序列模型前向时间除以有效帧数的折算值。",
        "完整序列推理时间表示开始回放前一次100帧前向的总耗时。",
        "误差曲线用于观察误差随时间的变化，不单独代表完整数据集指标。",
    ])
    doc.add_paragraph(
        "页面提供俯视、侧视和自由视角。观察误差峰值时，应同时查看场景机动、传感器有效性和当前时间位置，避免只根据轨迹线的视觉距离判断模型精度。"
    )

    doc.add_heading("7.3 播放与跳转", level=2)
    add_numbered_steps(doc, [
        "使用圆形播放按钮暂停或继续回放。",
        "暂停后拖动时间轴到目标时间并结束拖动。",
        "后端使用最近时间点定位新帧，并从跳转位置重新累计运行RMSE。",
        "需要重新执行模型前向时单击重新推理，而不是只把时间轴拖回起点。",
    ])
    doc.add_paragraph(
        "播放、暂停和跳转只控制已经计算结果的发送位置，不改变模型输入、模型参数或推理结果。服务端默认每0.05秒发送一帧用于界面回放，这个间隔也不表示传感器采样周期。"
    )

    doc.add_heading("7.4 导出当前序列", level=2)
    doc.add_paragraph(
        "收到至少一帧结果后可以导出当前序列CSV。文件包含时间、真值位置、预测位置、位置误差和运行RMSE。导出前应确认场景和目标索引正确，导出后抽查首末时间、行数和坐标数量。该CSV用于问题定位和单航迹复核，不替代批量测试生成的逐航迹统计。"
    )
    add_table(doc, ["字段类别", "典型内容", "解释"], [
        ["标识", "场景和目标索引", "确定当前导出对应的注册对象"],
        ["时间", "t", "来自有效时间轴，单位为秒"],
        ["真值位置", "truth_x、truth_y、truth_z", "ECEF绝对坐标，单位为米"],
        ["预测位置", "pred_x、pred_y、pred_z", "模型输出中的位置分量，单位为米"],
        ["误差", "position_error、running_rmse", "三维位置欧氏误差及累计RMSE"],
    ], widths=[1.25, 2.4, 3.7])

    doc.add_heading("8 批量测试", level=1)
    doc.add_heading("8.1 数据集选择", level=2)
    doc.add_paragraph(
        "批量测试页面列出技术数据集和验收想定。每个条目给出名称、状态、路径、场景数、机动类别、传感器模态、干扰模式和质量说明。选择前应确认数据集可用，并阅读其适用边界。V7用于8 m稳定阶段RMSE口径，V8用于强干扰鲁棒性验证，两者不能混用结论。"
    )
    add_table(doc, ["数据集", "主要用途", "结果解释"], [
        ["V5标准验证集", "标准噪声条件下的历史模型比较", "用于查看不同模型和数据条件的变化"],
        ["V6降噪验证集", "相对标准条件降低噪声的验证", "用于比较噪声变化下的模型表现"],
        ["V7低噪验证集", "合同稳定阶段精度评测", "8 m结论应基于完整范围和固定预热口径"],
        ["V8强干扰验证集", "40 Hz源数据兼容后的鲁棒性评测", "用于强干扰证据，不声明40 Hz在线模型能力"],
        ["协同编队目标跟踪想定", "V7中CV和CA子集的想定化检查", "只使用验证数据，不进入训练"],
        ["高机动强干扰想定", "V8复合干扰条件下的想定化检查", "同时包含调幅闪烁、软压制和距离抖动"],
    ], widths=[1.75, 2.55, 2.75])

    doc.add_heading("8.2 创建批测任务", level=2)
    add_numbered_steps(doc, [
        "输入便于识别的任务名称，留空时系统自动生成名称。",
        "选择模型。批测会加载独立模型实例，不需要先切换活动模型。",
        "选择auto、cpu或可用的cuda设备。",
        "选择目标数据集或验收想定，并确认资源状态为可用。",
        "选择抽样测试或完整测试。抽样测试需要填写场景数量。",
        "设置预热帧、批大小和数据线程数。",
        "提交任务并记录系统返回的任务ID。",
    ])
    add_table(doc, ["参数", "允许范围或默认值", "使用说明"], [
        ["任务名称", "最多100个字符", "用于检索和归档，不影响计算"],
        ["模型", "登记的V5至V7模型", "报告保存实际模型标识和来源摘要"],
        ["设备", "auto、cpu、可用cuda设备", "正式时延比较应保持设备一致"],
        ["预热帧", "0至99，默认5", "每条航迹前若干帧不计入稳定阶段精度"],
        ["批大小", "1至256，默认64", "过大可能导致显存不足，过小可能降低吞吐率"],
        ["数据线程", "0至8，默认0", "Windows环境下默认0更便于稳定复现"],
        ["场景数量", "至少1", "仅抽样测试使用，不能超过可用场景数"],
    ], widths=[1.35, 2.0, 3.7])
    doc.add_paragraph(
        "抽样任务用于快速验证配置，系统保存实际文件清单和选择策略。正式评测应选择完整测试，并保持模型、数据集、设备、预热帧、批大小和线程数不变。若修改其中任一条件，应创建新任务，不应覆盖或手工修改既有报告。"
    )

    doc.add_heading("8.3 任务状态与进度", level=2)
    doc.add_paragraph(
        "批测管理器使用单执行槽依次处理GPU密集任务。创建成功的任务先进入queued，开始加载和计算后进入running，最终进入completed、failed或cancelled。运行队列显示活动任务和百分比，任务历史可以按全部、活动、已完成和异常筛选。"
    )
    add_table(doc, ["状态", "含义", "允许的后续操作"], [
        ["queued", "任务已登记，等待执行", "查看配置或请求取消"],
        ["running", "正在加载数据或执行推理", "查看进度或请求取消"],
        ["completed", "计算完成且报告已持久化", "评估、下载和归档"],
        ["failed", "加载、推理、数值检查或写入失败", "查看错误，修复后新建任务"],
        ["cancelled", "取消请求在安全检查点生效", "保留状态记录，按需新建任务"],
    ], widths=[1.25, 3.0, 2.8], center_cols=(0,))
    doc.add_paragraph(
        "取消请求不是强制终止进程，系统在当前安全检查点停止后续工作。关闭后端会请求取消未完成任务；后端重启后不能续跑的旧任务会保留中断信息。已完成报告写入artifacts/reports/batch_tests，服务重启后仍可查询。"
    )

    doc.add_heading("8.4 批测结果", level=2)
    doc.add_paragraph(
        "任务详情显示配置、实际场景清单、样本量、有效帧数、误差统计、时延、吞吐率、关联摘要和逐航迹结果。进入统计的预测、真值和误差必须是有限值；出现NaN或Inf时任务应失败，不能在忽略异常样本后继续给出正式结论。"
    )
    add_table(doc, ["结果类别", "主要字段", "用途"], [
        ["精度", "RMSE、MAE、中位数、P95、最大误差", "描述稳定阶段总体误差水平和尾部风险"],
        ["航迹", "逐航迹有效帧、RMSE、MAE、P95、最大误差", "定位高误差场景和目标"],
        ["时延", "批次、航迹、完整序列和单帧折算时延", "检查计算性能并区分统计口径"],
        ["吞吐率", "有效帧每秒", "比较相同配置下的整体处理能力"],
        ["关联", "正确数、总数、正确率和分传感器统计", "在数据包含审计信息时提供关联证据"],
        ["来源", "模型、数据文件、代码摘要和任务签名", "判断独立证据是否适用于当前任务"],
    ], widths=[1.2, 3.05, 2.8])

    doc.add_heading("9 V8四十赫兹数据使用", level=1)
    doc.add_heading("9.1 数据口径", level=2)
    doc.add_paragraph(
        "V8数据用于接近真实雷达40 Hz量测周期的强干扰场景。源注册数据包含10秒、400帧序列，源采样周期约0.025秒。现有V5至V7权重按100帧、约10 Hz窗口训练；直接把400帧序列截断为前100帧会把覆盖时长缩短到约2.5秒，并改变模型看到的运动尺度，因此不能作为兼容处理。"
    )
    doc.add_paragraph(
        "在不重新训练的条件下，软件保留40 Hz源数据，并从0.1秒目标网格为每个目标时刻选择最近源帧，形成100帧兼容视图。雷达、光电、电子侦察、平台状态、快慢链路、掩码、真值和时间戳必须使用同一组索引同步抽取。处理不执行四帧平均，也不对缺测值做额外插值。"
    )

    doc.add_heading("9.2 生成兼容数据", level=2)
    doc.add_paragraph(
        "交付数据已包含V8 10 Hz兼容目录时，普通操作人员无需重新生成。只有源数据更新、兼容目录缺失或需要复核索引时，才运行兼容构建脚本。脚本不会修改40 Hz源文件，并会在输出目录生成兼容清单。"
    )
    _command(doc, "python scripts/data_generation/build_v8_10hz_compat.py --input-dir data/val_data_v8_interference_registered --output-dir data/val_data_v8_interference_registered_10hz")
    add_numbered_steps(doc, [
        "确认输入目录是V8 40 Hz注册数据，不是原始MAT目录。",
        "确认输出目录与输入目录不同，并具备写权限。",
        "运行脚本并检查完成文件数、目标周期和异常信息。",
        "查看compatibility_manifest.json，确认源文件、目标文件和同步索引记录。",
        "在批量测试页面刷新数据集状态并执行少量场景链路检查。",
    ])

    doc.add_heading("9.3 V8结果解释", level=2)
    doc.add_paragraph(
        "V8页面和报告应同时说明source_frame_period_s为0.025、source_measurement_rate_hz为40、model_input_rate_hz为10以及兼容方式为同步选帧。模型实际接收的是10 Hz兼容视图，因此结果证明现有模型在V8强干扰场景兼容输入上的表现，不证明模型以40 Hz在线输出。"
    )
    doc.add_paragraph(
        "V8同时施加调幅闪烁、软压制和距离抖动，平台链路噪声也高于V7。其结果用于鲁棒性和适用边界分析，不应直接替代V7稳定阶段8 m精度结论。需要比较V7和V8时，应分别保留数据口径、有效航迹数、预热帧和设备条件。"
    )
    add_table(doc, ["检查项", "V8源数据", "模型兼容视图"], [
        ["采样频率", "40 Hz", "约10 Hz"],
        ["采样周期", "约0.025 s", "约0.1 s"],
        ["10秒帧数", "约400帧", "100帧"],
        ["文件用途", "保留真实量测时间分辨率", "供现有100帧权重推理"],
        ["处理方法", "不修改", "统一索引选择最近源帧"],
        ["结论边界", "源场景具有40 Hz量测", "模型不声明40 Hz在线推理"],
    ], widths=[1.5, 2.75, 2.75])

    doc.add_heading("10 性能评估", level=1)
    doc.add_heading("10.1 选择评估任务", level=2)
    doc.add_paragraph(
        "性能评估只能针对已完成批测任务。选择任务后，页面读取任务配置、结果、阈值和独立证据，并分别显示四项指标。若任务未完成、报告缺失或来源签名不一致，评估接口会拒绝形成正式判定。"
    )
    add_numbered_steps(doc, [
        "选择已完成且配置正确的批测任务。",
        "核对模型、数据集、设备、场景范围、预热帧和有效航迹数。",
        "确认位置RMSE和时延已经由当前任务产生。",
        "检查关联正确率和IMM-JPDA降幅证据状态。",
        "必要时生成缺失证据，然后重新判定。",
    ])

    doc.add_heading("10.2 指标口径", level=2)
    add_table(doc, ["指标", "判定阈值", "统计与证据要求"], [
        ["稳定阶段三维位置RMSE", "不超过8 m", "只统计target_valid_mask为真且排除每条航迹预热帧的样本"],
        ["数据关联正确率", "不低于85%", "使用与任务场景范围绑定的关联审计证据"],
        ["相对IMM-JPDA的RMSE降幅", "不低于15%", "基线与模型必须使用同一数据、场景、序列和预热条件"],
        ["单帧结果生成时间", "不超过100 ms", "完整序列模型前向时间除以帧数，不含回放、控制律和通信"],
    ], widths=[2.2, 1.35, 3.45])
    doc.add_paragraph(
        "指标状态包括passed、failed和unavailable。只有实测值或有效证据存在时才判定通过或失败；缺少关联审计、缺少同条件基线或证据签名失效时显示unavailable。总体状态为passed、failed或incomplete，其中incomplete表示仍有指标缺少可用证据。"
    )

    doc.add_heading("10.3 生成独立证据", level=2)
    doc.add_paragraph(
        "关联正确率和相对IMM-JPDA降幅由独立脚本生成。页面提供运行按钮和可复制命令，后端只允许调用登记的白名单脚本。脚本完成后，评估服务校验任务ID、数据范围、模型和配置形成的来源签名；任何影响结果的任务变化都会使旧证据失效。"
    )
    _command(doc, "python scripts/evaluation/test_association_accuracy.py --batch-report artifacts/reports/batch_tests/任务ID.json")
    _command(doc, "python scripts/evaluation/test_rmse_improvement.py --batch-report artifacts/reports/batch_tests/任务ID.json")
    doc.add_paragraph(
        "完整数据集基线计算可能耗时较长。运行期间应保持后端和计算设备稳定，不要修改任务报告或移动关联数据。证据生成失败时先查看错误信息和对应命令，不要手工把缺失指标填为通过。"
    )

    doc.add_heading("10.4 多维结果与比较", level=2)
    doc.add_paragraph(
        "性能评估页面按精度、时延、航迹分布和机动类型展示结果，并可查看V5至V7模型在多类数据集上的历史交叉评估矩阵。跨数据集矩阵和历史基线用于趋势比较，正式结论仍以当前任务及其同条件证据为准。"
    )
    _bullets(doc, [
        "精度维度用于查看总体误差、尾部误差和航迹间差异。",
        "时延维度用于区分批次耗时、序列耗时、单帧折算时延和吞吐率。",
        "航迹分布用于识别少量高误差航迹是否抬高总体结果。",
        "机动类型汇总用于比较CV、CA、CT、CSSW和其他类别。",
        "横向比较前必须核对任务配置，不能只按模型名称比较数值。",
    ])

    doc.add_heading("11 报告中心", level=1)
    doc.add_heading("11.1 检索与复核", level=2)
    doc.add_paragraph(
        "报告中心汇总已经持久化的批测任务。用户可以按任务名称、任务ID、模型或数据集搜索，并按模型和状态筛选。打开详情后，应先核对任务配置、实际场景范围、完成状态和样本数量，再查看核心指标和技术指标判定。"
    )
    add_numbered_steps(doc, [
        "打开报告中心并同步归档列表。",
        "输入任务名称、任务ID、模型或数据集关键字。",
        "使用状态和模型筛选缩小范围。",
        "打开目标任务详情，核对配置、指标、证据和制品。",
        "确认正式证据齐全后再执行导出和归档。",
    ])

    doc.add_heading("11.2 重新判定", level=2)
    doc.add_paragraph(
        "报告详情可以使用同条件IMM-JPDA基线RMSE重新计算降幅。手工输入基线只适用于已经确认测试条件一致的情况，并应在归档说明中记录来源。清空手工基线并重新判定，可恢复使用独立证据文件的结果。重新判定不会修改原始批测误差和时延。"
    )
    doc.add_paragraph(
        "若任务配置、模型文件、数据场景或预热帧发生变化，应新建批测任务和证据，不应反复使用旧任务重新判定。报告中心显示的来源摘要和证据状态用于帮助发现这种不一致。"
    )

    doc.add_heading("11.3 报告导出", level=2)
    add_table(doc, ["导出操作", "输出内容", "使用条件"], [
        ["下载原始JSON", "完整批测任务配置、结果和来源摘要", "批测报告文件存在"],
        ["下载综合评估JSON", "任务、四项判定、阈值、证据和总体结论", "任务已完成并可评估"],
        ["下载逐航迹CSV", "每条航迹的有效帧和误差指标", "任务已完成且存在逐航迹结果"],
        ["导出比较CSV", "选中任务的统一横向指标", "选择1至50个已完成任务"],
        ["打印或保存PDF", "当前报告详情的打印视图", "浏览器支持打印或保存PDF"],
    ], widths=[1.55, 3.15, 2.3])
    doc.add_paragraph(
        "下载后应使用任务ID关联各制品，并检查文件能够打开。正式归档建议采用只读目录，记录模型权重、数据集、设备、代码状态和生成时间。浏览器PDF便于阅读，但可复核信息以JSON和CSV为准。"
    )

    doc.add_heading("12 正式评测与验收操作", level=1)
    doc.add_heading("12.1 评测前冻结", level=2)
    doc.add_paragraph(
        "正式评测前应冻结模型权重、数据集、代码、运行环境和统计参数。冻结的目的不是阻止备份，而是保证从批测到证据生成和报告导出的条件保持一致。任何影响输入、模型或统计口径的变化都应重新执行任务。"
    )
    add_table(doc, ["冻结对象", "需要记录的内容", "复核方法"], [
        ["模型", "模型标识、权重路径、大小和文件摘要", "模型仓库和报告来源信息"],
        ["数据", "数据集标识、路径、实际文件清单和兼容口径", "数据集详情和任务报告"],
        ["环境", "操作系统、Python、PyTorch、CUDA和设备", "系统总览和系统接口"],
        ["参数", "预热帧、批大小、线程数和场景范围", "任务配置"],
        ["阈值", "四项指标阈值和单位", "技术指标配置接口"],
        ["代码", "模型、数据解析和批测关键文件状态", "任务来源签名或交付清单"],
    ], widths=[1.15, 3.4, 2.45])

    doc.add_heading("12.2 正式执行步骤", level=2)
    add_numbered_steps(doc, [
        "启动软件并在系统总览记录运行环境。",
        "在模型管理中确认待验收权重可用，并在目标设备上完成部署检查。",
        "在跟踪验证中选择一个已知场景，确认输入、推理、显示和导出链路正常。",
        "在批量测试中选择待验收模型、目标数据集和完整测试，固定预热帧、批大小和线程数。",
        "提交任务并等待completed，检查有效航迹数、有效帧数和异常信息。",
        "在性能评估中生成关联正确率和IMM-JPDA基线证据。",
        "确认四项指标均有有效证据，并查看总体结论。",
        "在报告中心下载原始JSON、综合评估JSON、逐航迹CSV和必要的比较CSV。",
        "将报告、证据、模型、数据清单和环境记录一并归档。",
    ])
    doc.add_paragraph(
        "若任一指标失败，应保留失败报告和证据，再进行原因分析；若任一指标为数据不足，应补齐相应证据或修复任务条件后重新评测。不得删除失败任务并只保留通过结果，也不得把抽样任务包装为完整评测。"
    )

    doc.add_heading("12.3 验收结果检查", level=2)
    _bullets(doc, [
        "任务状态为completed，报告文件存在且任务ID一致。",
        "场景范围为完整数据集或合同明确的验收范围。",
        "稳定阶段RMSE使用固定预热规则和有效帧掩码。",
        "关联正确率和IMM-JPDA降幅证据状态为有效。",
        "时延来自模型前向计时，设备和批大小与报告一致。",
        "V8结果标明40 Hz源数据和10 Hz模型兼容视图。",
        "下载制品可以打开，并与页面显示的数值一致。",
    ])
    doc.add_paragraph(
        "软件还提供python -m tracking_system.backend.acceptance用于黑盒检查后端健康、模型、数据、场景、WebSocket、批测、评估和报告下载。该检查证明软件链路可用，默认少量场景结果不替代完整算法指标评测。"
    )

    doc.add_heading("13 跨语言推理使用", level=1)
    doc.add_heading("13.1 固定接口", level=2)
    doc.add_paragraph(
        "portable目录提供同一份固定尺寸ONNX模型的MATLAB和C入口。接口接收已经完成时间对齐、目标关联和空间注册的单目标序列，不包括原始MAT解析、数据集生成或训练。固定尺寸为批量1、序列100帧、平台5个，输出为1×100×6。"
    )
    add_table(doc, ["项目", "固定值", "说明"], [
        ["批量大小", "B等于1", "一次处理一条目标序列"],
        ["序列长度", "T等于100", "需要其他长度时重新导出并验证ONNX"],
        ["平台数量", "N等于5", "主机和四个协同平台槽位"],
        ["输入数量", "18项float32张量", "字段、尺寸和掩码含义见portable/INTERFACE.md"],
        ["输出状态", "px、vx、py、vy、pz、vz", "位置为ECEF绝对坐标，速度单位为m/s"],
        ["存储顺序", "C按行优先", "MATLAB加载函数负责列优先转换"],
    ], widths=[1.4, 2.2, 3.6])

    doc.add_heading("13.2 MATLAB验证", level=2)
    doc.add_paragraph(
        "MATLAB运行需要R2025a或更新版本、Deep Learning Toolbox以及ONNX模型转换组件。将portable/matlab加入路径后，使用交付ONNX和黄金样例执行一致性验证。"
    )
    _command(doc, 'addpath("portable/matlab")')
    _command(doc, 'verify_golden_case("portable/artifacts/coop_tracker_v7.onnx", "portable/artifacts/golden_case")')
    doc.add_paragraph(
        "业务调用入口为run_tracker.m。调用方应按接口文档构造18个字段，并保证尺寸、单位、掩码、ref_pos和dt与Python实现一致。验证失败时先检查输入展开顺序和float32类型，再检查ONNX转换环境。"
    )

    doc.add_heading("13.3 C验证", level=2)
    doc.add_paragraph(
        "C版本使用ONNX Runtime C API。交付目录包含Windows x64 CPU运行包和公共头文件tracker.h。调用方创建Tracker，填充TrackerInputs中的18个连续float32缓冲区，调用tracker_run得到600个状态值。"
    )
    _command(doc, "cmake -S portable/c -B portable/c/build -DONNXRUNTIME_ROOT=portable/third_party/onnxruntime-win-x64-1.18.1")
    _command(doc, "cmake --build portable/c/build --config Release")
    _command(doc, "portable/c/build/Release/verify_golden_case.exe portable/artifacts/coop_tracker_v7.onnx portable/artifacts/golden_case")
    doc.add_paragraph(
        "MATLAB或C相对Python参考输出的最大位置分量差应不超过1 m，最大速度分量差应不超过0.2 m/s。黄金样例用于证明跨语言数值一致性，不替代完整验证集上的合同指标测试。"
    )

    doc.add_heading("14 高级配置与数据管理", level=1)
    doc.add_heading("14.1 启动环境变量", level=2)
    doc.add_paragraph(
        "以下环境变量在启动后端前设置，修改后需要重启服务。普通操作优先使用登记模型和页面设备选择；只有默认资源路径不适用或需要固定启动设备时才使用环境变量。"
    )
    add_table(doc, ["环境变量", "示例", "作用和注意事项"], [
        ["TRACKER_MODEL_PATH", "D路径/best_model.pt", "覆盖默认模型路径，文件仍需与登记模型及网络结构匹配"],
        ["TRACKER_DATA_DIR", "D路径/val_data_v7_registered", "覆盖单场景验证数据目录，不改变批测数据集目录"],
        ["TRACKER_DEVICE", "cuda:0或cpu", "指定默认推理设备，CUDA不可用时显式请求会失败"],
        ["TRACKER_SEQ_LEN", "100", "指定单场景序列长度；当前权重和跨语言接口按100帧验证"],
    ], widths=[1.85, 2.2, 3.15])
    doc.add_paragraph(
        "Windows PowerShell可以在当前终端使用环境变量，关闭终端后是否保留取决于设置方式。环境变量中包含空格的路径应正确引用。修改序列长度、输入字段或模型结构属于接口变化，不能只改环境变量后继续沿用原验收结论。"
    )

    doc.add_heading("14.2 数据保护", level=2)
    doc.add_paragraph(
        "注册验证数据、V8源数据和已验收权重应作为受控输入管理。兼容数据、缓存、批测报告和证据写入独立目录，不覆盖源数据。正式操作前建议记录目录大小和文件摘要，评测后将输入清单与输出制品一并备份。"
    )
    _bullets(doc, [
        "不要在数据集目录手工重命名部分文件后继续使用旧报告。",
        "不要用新的权重覆盖已验收best_model.pt而保留原任务名称。",
        "不要手工编辑批测JSON中的指标、状态或来源签名。",
        "不要移动独立证据文件后再用人工数值替代其适用性校验。",
        "V8兼容生成必须写入独立输出目录并保留compatibility_manifest.json。",
    ])

    doc.add_heading("14.3 报告备份与清理", level=2)
    add_table(doc, ["内容", "默认位置", "管理建议"], [
        ["批测报告", "artifacts/reports/batch_tests", "正式任务完成后复制到只读归档目录"],
        ["验收证据", "artifacts/metrics/acceptance/任务ID", "与对应批测JSON一起归档"],
        ["交叉评估", "artifacts/metrics/cross_dataset_rmse.json", "标注为历史参考并保留生成条件"],
        ["模型权重", "results/训练运行/best_model.pt", "保留只读副本和文件摘要"],
        ["V8兼容清单", "10 Hz兼容目录/compatibility_manifest.json", "与40 Hz源目录和兼容数据一并保存"],
    ], widths=[1.45, 3.0, 2.75])
    doc.add_paragraph(
        "清理空间前先确认文件是否属于正式归档。删除批测报告会使报告中心无法重新载入任务，删除证据会使相应指标变为数据不足。需要保留审计链时，应先复制完整任务目录和证据目录，再执行清理。"
    )

    doc.add_heading("15 常见问题与故障处理", level=1)
    doc.add_heading("15.1 启动与连接问题", level=2)
    add_table(doc, ["现象", "可能原因", "处理方法"], [
        ["页面显示后端未连接", "后端未启动、端口占用或模型冷启动失败", "直接访问/api/health，检查后端终端和8000端口"],
        ["后端提示找不到模型", "默认路径或TRACKER_MODEL_PATH无效", "恢复登记权重或修正路径后重启"],
        ["前端无法启动", "Node.js不可用、依赖缺失或5173端口占用", "检查node和npm，重新安装依赖或释放端口"],
        ["接口文档可访问但页面报错", "前端API基址、跨域或浏览器缓存异常", "核对访问主机和端口，刷新页面并查看请求信息"],
        ["CUDA不可用", "驱动、PyTorch CUDA版本或设备索引不匹配", "在系统总览核对CUDA，临时改用cpu或修复环境"],
    ], widths=[1.7, 2.45, 2.85])

    doc.add_heading("15.2 模型与数据问题", level=2)
    add_table(doc, ["现象", "可能原因", "处理方法"], [
        ["模型部署失败", "权重缺失、结构不匹配或显存不足", "查看后端错误，刷新仓库，核对d_model和设备"],
        ["场景列表为空", "数据目录不存在、无PKL文件或解析失败", "核对TRACKER_DATA_DIR和注册数据完整性"],
        ["目标索引越界", "所选场景的注册目标数少于输入索引", "重新查看场景摘要，索引从0开始"],
        ["V8无法批测", "10 Hz兼容目录缺失或清单不完整", "重新生成兼容视图并刷新数据集状态"],
        ["V8帧数仍为400", "误选40 Hz源目录作为模型输入", "改用val_data_v8_interference_registered_10hz"],
        ["结果出现NaN或Inf", "输入异常、数值溢出或模型不匹配", "保留失败任务，定位文件和批次，不得忽略异常继续判定"],
    ], widths=[1.7, 2.45, 2.85])

    doc.add_heading("15.3 批测评估与报告问题", level=2)
    add_table(doc, ["现象", "可能原因", "处理方法"], [
        ["任务长时间排队", "已有GPU密集任务占用单执行槽", "查看活动队列，等待完成或取消前一任务"],
        ["任务在后端重启后失败", "任务运行期间服务退出且不支持续跑", "保留中断记录并重新提交任务"],
        ["证据显示数据不足", "脚本未运行、运行失败或来源签名不一致", "重新生成对应证据并查看任务条件"],
        ["逐航迹CSV无法下载", "任务未完成或没有逐航迹结果", "确认completed且evaluated_tracks大于0"],
        ["报告中心找不到历史任务", "报告被移动、删除或JSON损坏", "恢复batch_tests目录中的原报告并重启后端"],
        ["页面与下载数值不一致", "页面缓存、重新判定条件或文件不属于同一任务", "按任务ID核对并重新同步归档"],
    ], widths=[1.7, 2.45, 2.85])
    doc.add_paragraph(
        "故障处理时应保留任务ID、请求ID、后端错误、输入资源和操作步骤。无法定位时，可以先使用一个已知V7场景在CPU上执行单场景验证，以区分CUDA、数据集、批测队列和报告写入问题。"
    )

    doc.add_heading("16 维护与关闭", level=1)
    doc.add_heading("16.1 日常维护", level=2)
    doc.add_paragraph(
        "日常维护包括检查模型和数据目录、确认报告空间、备份正式制品、清理非正式测试输出以及验证启动命令。依赖或代码发生变化后，应至少完成健康检查、模型严格加载、单场景回放、小批量测试和报告下载。"
    )
    _bullets(doc, [
        "定期检查artifacts目录剩余空间和正式报告备份。",
        "模型、数据解析、统计口径或V8兼容规则变化后重新执行相关测试。",
        "前端依赖变化后重新安装并执行生产构建或启动检查。",
        "Python或CUDA环境变化后重新记录系统信息和时延基线。",
        "跨语言模型重新导出后重新运行Python、MATLAB和C黄金样例。",
    ])

    doc.add_heading("16.2 正常关闭", level=2)
    add_numbered_steps(doc, [
        "确认没有需要继续运行的批测或证据生成任务。",
        "在报告中心下载并备份本次正式制品。",
        "在前端终端按Ctrl+C停止Vite服务。",
        "在后端终端按Ctrl+C停止Uvicorn服务。",
        "确认后端已经退出，并检查未完成任务的最终状态。",
    ])
    doc.add_paragraph(
        "不要直接结束进程后立即删除报告目录。异常断电或强制结束后，重新启动服务并检查任务历史；已完成任务应能够恢复，未完成任务应显示中断或失败，不应手工改为完成。"
    )

    doc.add_heading("附录A 操作速查", level=1)
    doc.add_heading("A.1 常用命令", level=2)
    add_table(doc, ["用途", "命令"], [
        ["启动后端", "conda run --no-capture-output -n pytorch python -m tracking_system.backend.server"],
        ["进入前端", "cd tracking_system/frontend"],
        ["安装前端依赖", "npm install"],
        ["启动前端", "npm run dev"],
        ["后端黑盒验收", "python -m tracking_system.backend.acceptance"],
        ["后端单元测试", "python -m unittest discover -s tracking_system/backend/tests -v"],
        ["生成V8兼容数据", "python scripts/data_generation/build_v8_10hz_compat.py"],
        ["关联正确率证据", "python scripts/evaluation/test_association_accuracy.py --batch-report 报告路径"],
        ["IMM-JPDA证据", "python scripts/evaluation/test_rmse_improvement.py --batch-report 报告路径"],
    ], widths=[1.65, 5.35])

    doc.add_heading("A.2 状态速查", level=2)
    add_table(doc, ["对象", "状态", "说明"], [
        ["批测任务", "queued", "已提交，等待单执行槽"],
        ["批测任务", "running", "正在加载数据或计算"],
        ["批测任务", "completed", "完成并持久化报告"],
        ["批测任务", "failed", "发生错误，保留失败原因"],
        ["批测任务", "cancelled", "取消请求已生效"],
        ["单项指标", "passed", "有有效数据或证据且达到阈值"],
        ["单项指标", "failed", "有有效数据或证据但未达到阈值"],
        ["单项指标", "unavailable", "缺少数据、证据或证据失效"],
        ["总体结论", "passed", "所有指标均通过"],
        ["总体结论", "failed", "至少一个指标失败"],
        ["总体结论", "incomplete", "没有失败项但存在数据不足"],
    ], widths=[1.45, 1.55, 4.0], center_cols=(1,))

    doc.add_heading("A.3 输入输出速查", level=2)
    add_table(doc, ["类别", "输入或输出", "关键约束"], [
        ["模型输入", "18项多源和平台张量", "B为1、T为100、N为5，float32，掩码与时间同步"],
        ["模型输出", "逐时刻六维状态", "px、vx、py、vy、pz、vz"],
        ["单场景输入", "活动模型目录中的场景和目标", "目标索引从0开始"],
        ["批测输入", "登记模型、数据集、设备和统计参数", "正式任务使用完整范围"],
        ["V8输入", "40 Hz源数据的10 Hz同步兼容视图", "使用统一索引，不做四帧平均"],
        ["报告输出", "批测JSON、评估JSON、逐航迹CSV和比较CSV", "按任务ID和来源签名关联"],
    ], widths=[1.4, 2.75, 2.85])


def _format_document(doc: Document, body_boundary: int) -> None:
    for style_name in ("Normal", "Body Text", "List Bullet", "List Number", "List Paragraph"):
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        style.font.name = LATIN_FONT
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        fonts.set(qn("w:ascii"), LATIN_FONT)
        fonts.set(qn("w:hAnsi"), LATIN_FONT)
        fonts.set(qn("w:cs"), LATIN_FONT)
        style.font.size = Pt(14)

    for style_name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = doc.styles[style_name]
        style.font.name = LATIN_FONT
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        fonts.set(qn("w:ascii"), LATIN_FONT)
        fonts.set(qn("w:hAnsi"), LATIN_FONT)
        fonts.set(qn("w:cs"), LATIN_FONT)
        style.font.size = Pt(14)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.italic = False
        style.paragraph_format.keep_with_next = True

    body = doc._element.body
    for child in list(body)[body_boundary:]:
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, doc._body)
        style_name = paragraph.style.name if paragraph.style is not None else ""
        is_heading = style_name in {"Heading 1", "Heading 2", "Heading 3", "Heading 4"}
        if is_heading:
            paragraph.paragraph_format.keep_with_next = True
        else:
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.widow_control = True
        for run in paragraph.runs:
            _set_run_font(run, bold=True if is_heading else None)
            if is_heading:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.italic = False

    for table in doc.tables:
        set_table_borders(table)
        if table.rows:
            set_repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell in row.cells:
                tc_pr = cell._tc.get_or_add_tcPr()
                for shading in list(tc_pr.findall(qn("w:shd"))):
                    tc_pr.remove(shading)
                if row_index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), GRAY)
                    shading.set(qn("w:val"), "clear")
                    tc_pr.append(shading)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.2
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        _set_run_font(run, bold=True if row_index == 0 else None)
                        run.font.color.rgb = RGBColor(0, 0, 0)

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    doc.core_properties.title = "端到端协同跟踪原型软件使用说明书"
    doc.core_properties.subject = "软件配套交付文档"
    doc.core_properties.author = "项目组"


def build() -> Path:
    if not TEMPLATE.exists():
        raise FileNotFoundError(TEMPLATE)
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    working = WORK_DIR / "user_manual_working.docx"
    shutil.copy2(TEMPLATE, working)
    doc = Document(working)
    _replace_cover_text(doc)
    _set_toc_levels(doc)
    boundary = _remove_body_from(doc, "1 使用前准备")
    _build_body(doc)
    _format_document(doc, boundary)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
