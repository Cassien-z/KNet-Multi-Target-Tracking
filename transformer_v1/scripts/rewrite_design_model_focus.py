"""Rewrite the design-report body around the tracker model while preserving front matter.

The cover and the manually edited table of contents are kept byte-for-byte at
the document-body element level.  Only content beginning at the real
``1 前言`` heading is replaced.
"""

from __future__ import annotations

import argparse
import hashlib
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.build_software_docs import add_bullets, add_numbered_steps, add_table
from scripts.design_model_expansions import add_model_expansions


DEFAULT_DOC = ROOT / "docs" / "software_delivery" / "端到端协同跟踪原型软件_设计说明书.docx"
MODEL_FIGURE = ROOT / "artifacts" / "figures" / "架构图.png"
BACKUP_DIR = ROOT / ".codex_work" / "model_focus_backup"

EAST_ASIA_FONT = "楷体"
LATIN_FONT = "Times New Roman"


def _assert_source_contracts() -> None:
    contracts = {
        ROOT / "src" / "model.py": [
            "class LearnedKFrontEnd",
            "nn.Linear(30, feature_dim)",
            "class CrossAttentionBlock",
            "class TemporalTransformer",
            "class CollaborativeGAT",
            "dist <= 20000.0",
            "pred_rel = k_state + residual",
        ],
        ROOT / "src" / "dataset.py": [
            "def _parse_registered_pkl",
            "dt = float(np.median(valid_steps))",
            '"global_time": np.zeros((seq_len,), dtype=np.float32)',
        ],
        ROOT / "tracking_system" / "backend" / "runtime.py": [
            "class TransformerEngine",
            "torch.inference_mode()",
            "model.load_state_dict(checkpoint, strict=True)",
        ],
        ROOT / "train.py": [
            "from src.dataset_pkl import RegisteredPklDataset",
            "criterion = nn.SmoothL1Loss()",
            "optimizer = optim.AdamW",
        ],
    }
    for path, tokens in contracts.items():
        text = path.read_text(encoding="utf-8")
        missing = [token for token in tokens if token not in text]
        if missing:
            raise RuntimeError(f"Source contract changed in {path}: {missing}")


def _front_boundary(doc: Document) -> int:
    body = doc._element.body
    for index, child in enumerate(body):
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, doc._body)
        if (
            paragraph.style is not None
            and paragraph.style.name == "Heading 1"
            and paragraph.text.strip() == "1 前言"
        ):
            return index
    raise RuntimeError("Cannot locate the real '1 前言' body heading")


def _prefix_digest(doc: Document, boundary: int) -> str:
    body = doc._element.body
    data = b"".join(
        etree.tostring(child, with_tail=False) for child in list(body)[:boundary]
    )
    return hashlib.sha256(data).hexdigest()


def _remove_existing_body(doc: Document, boundary: int) -> None:
    body = doc._element.body
    for child in list(body)[boundary:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _extract_heading_bookmarks(source: Document) -> list[tuple[str, str]]:
    records: list[tuple[str, str]] = []
    for paragraph in source.paragraphs:
        if paragraph.style is None or not paragraph.style.name.startswith("Heading"):
            continue
        names = paragraph._p.xpath(".//w:bookmarkStart/@w:name")
        for name in names:
            if name.startswith("_Toc"):
                records.append((paragraph.text.strip(), name))
    return records


def _restore_toc_bookmarks(
    doc: Document,
    boundary: int,
    records: list[tuple[str, str]],
) -> None:
    """Keep the user's existing PAGEREF fields resolvable without editing the TOC."""
    new_headings = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.style is not None
        and paragraph.style.name.startswith("Heading")
        and paragraph._p.getparent() is doc._element.body
    ]
    body_heading_ids = {id(p._p) for p in new_headings}
    new_headings = [p for p in new_headings if id(p._p) in body_heading_ids]
    if not new_headings:
        raise RuntimeError("No new headings available for TOC bookmark restoration")

    by_text = {paragraph.text.strip(): paragraph for paragraph in new_headings}
    by_chapter: dict[str, Paragraph] = {}
    for paragraph in new_headings:
        token = paragraph.text.strip().split(maxsplit=1)[0]
        chapter = token.split(".", 1)[0]
        by_chapter.setdefault(chapter, paragraph)

    existing_ids = []
    for start in doc._element.body.xpath(".//w:bookmarkStart"):
        value = start.get(qn("w:id"))
        if value is not None and value.isdecimal():
            existing_ids.append(int(value))
    next_id = max(existing_ids, default=0) + 1

    targets: dict[int, list[tuple[int, str]]] = {}
    for old_text, name in records:
        target = by_text.get(old_text)
        if target is None:
            token = old_text.split(maxsplit=1)[0] if old_text else ""
            chapter = token.split(".", 1)[0]
            target = by_chapter.get(chapter, new_headings[0])
        targets.setdefault(id(target._p), []).append((next_id, name))
        next_id += 1

    for paragraph in new_headings:
        items = targets.get(id(paragraph._p), [])
        if not items:
            continue
        insert_at = 1 if paragraph._p.pPr is not None else 0
        for bookmark_id, name in items:
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(bookmark_id))
            start.set(qn("w:name"), name)
            paragraph._p.insert(insert_at, start)
            insert_at += 1
        for bookmark_id, _ in reversed(items):
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(bookmark_id))
            paragraph._p.append(end)


def _set_run_font(run, *, bold: bool | None = None) -> None:
    run.font.name = LATIN_FONT
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    fonts.set(qn("w:ascii"), LATIN_FONT)
    fonts.set(qn("w:hAnsi"), LATIN_FONT)
    fonts.set(qn("w:cs"), LATIN_FONT)
    run.font.size = Pt(14)
    if bold is not None:
        run.bold = bold


def _format_new_body(doc: Document, boundary: int) -> None:
    body = doc._element.body
    heading_styles = {"Heading 1", "Heading 2", "Heading 3"}
    for child in list(body)[boundary:]:
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc._body)
            style_name = paragraph.style.name if paragraph.style is not None else ""
            is_heading = style_name in heading_styles
            if not is_heading:
                paragraph.paragraph_format.line_spacing = 1.5
                paragraph.paragraph_format.space_after = Pt(6)
            for run in paragraph.runs:
                _set_run_font(run, bold=True if is_heading else None)
                if is_heading:
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.italic = False
        elif child.tag == qn("w:tbl"):
            # Find the python-docx wrapper for this table and format its runs.
            for table in doc.tables:
                if table._tbl is not child:
                    continue
                for row_index, row in enumerate(table.rows):
                    for cell in row.cells:
                        tc_pr = cell._tc.get_or_add_tcPr()
                        for shading in list(tc_pr.findall(qn("w:shd"))):
                            tc_pr.remove(shading)
                        if row_index == 0:
                            shading = OxmlElement("w:shd")
                            shading.set(qn("w:fill"), "D9D9D9")
                            shading.set(qn("w:val"), "clear")
                            tc_pr.append(shading)
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.line_spacing = 1.2
                            paragraph.paragraph_format.space_before = Pt(0)
                            paragraph.paragraph_format.space_after = Pt(0)
                            for run in paragraph.runs:
                                _set_run_font(run, bold=True if row_index == 0 else None)
                                if row_index == 0:
                                    run.font.color.rgb = RGBColor(0, 0, 0)
                break


def _add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_lead_paragraph(doc: Document, lead: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run(f"{lead}：").bold = True
    paragraph.add_run(text)


def _build_body(doc: Document) -> None:
    doc.add_heading("1 前言", level=1)
    doc.add_heading("1.1 编写目的", level=2)
    doc.add_paragraph(
        "本说明书以端到端协同跟踪模型为核心，说明模型所解决的多无人机多源异构传感器目标状态估计问题，给出输入输出、坐标与时间约定、网络结构、状态递推、多模态融合、时序建模、图协同、残差回归、训练约束和推理验证方法。前后端原型系统仅作为模型加载、结果展示、批量验证和报告归档的支撑环境。"
    )
    doc.add_heading("1.2 项目背景", level=2)
    doc.add_paragraph(
        "项目面向多架无人机协同探测条件下的目标连续跟踪。不同平台能够获得雷达、光电和电子支援量测，并通过数据链共享局部融合结果。模型需要同时处理量测噪声、传感器缺失、平台运动、异步链路和目标机动，在统一的时间窗口内直接输出目标三维位置与速度状态。"
    )
    doc.add_heading("1.3 设计依据", level=2)
    add_bullets(doc, [
        "以src/model.py中的CoopTrackerModel、LearnedKFrontEnd、CrossAttentionBlock、TemporalTransformer和CollaborativeGAT实现为模型结构依据。",
        "以src/dataset.py中的注册数据解析、参考坐标转换、时间间隔推导和掩码构造为模型输入依据。",
        "以train.py中的训练入口、损失函数、优化参数、僚机输入约束和消融方式为权重形成依据。",
        "以tracking_system/backend/runtime.py中的完整序列推理、有效帧过滤、耗时测量和缓存逻辑为部署推理依据。",
    ])
    doc.add_heading("1.4 适用范围与读者", level=2)
    doc.add_paragraph(
        "本说明书适用于模型研发、模型复核、软件联调、测试验收和后续维护。读者应重点关注第3章至第10章的模型数据、网络结构、训练与验证设计；第11章以后仅说明展示验证软件为模型提供的工程支撑。"
    )
    doc.add_heading("1.5 名词与符号约定", level=2)
    add_table(doc, ["术语", "含义"], [
        ["B", "批次大小；单场景回放时B等于1"],
        ["T", "时间序列长度，当前训练与部署配置为100帧"],
        ["N", "协同无人机数量，当前配置为5"],
        ["D", "隐特征维度d_model，当前配置为96"],
        ["Learned-K", "由神经网络生成六维对角增益并执行类卡尔曼递推的物理先验状态前端"],
        ["主流与被动流", "雷达为主动主流，EO与ESM经置信融合后构成被动辅助流"],
        ["参考位置", "每条目标航迹首帧绝对ECEF位置，用于将大坐标平移为局部相对量"],
    ], widths=[1.55, 5.25])

    doc.add_heading("2 系统总体设计", level=1)
    doc.add_heading("2.1 系统定位", level=2)
    doc.add_paragraph(
        "系统主体是CoopTrackerModel及其数据预处理、训练权重和推理运行时。浏览器页面、REST接口、WebSocket回放和批测报告不参与模型内部状态估计，只负责向模型提供受控输入、展示模型输出并形成可复核的验证结果。"
    )
    doc.add_heading("2.2 总体技术架构", level=2)
    add_table(doc, ["层次", "核心组成", "设计定位"], [
        ["模型输入层", "注册PKL、坐标转换、掩码、时间与链路字段", "把场景数据整理为模型约定张量"],
        ["物理状态前端", "几何反演、信息矩阵融合、Learned-K递推", "形成连续且可解释的六维粗状态"],
        ["深度特征主干", "通道编码、被动置信融合、双向跨注意力", "学习主动与被动量测之间的互补关系"],
        ["时空协同层", "时序Transformer、20 km邻接、CollaborativeGAT", "提取单平台时间特征并聚合多平台信息"],
        ["状态输出层", "位置头、速度头、残差叠加、绝对坐标恢复", "输出每帧绝对位置与速度"],
        ["验证支撑层", "推理运行时、批测、前端回放、评估报告", "验证模型，不改变模型计算图"],
    ], widths=[1.25, 2.7, 2.85])
    doc.add_heading("2.3 模型推理主流程", level=2)
    add_numbered_steps(doc, [
        "读取注册场景，按目标拆分样本并构造100帧、5平台的多源输入张量。",
        "以首帧目标绝对位置建立参考坐标，统一平台位置、链路位置、真值和模型输出的坐标口径。",
        "由雷达与快慢链路构造几何粗观测，再由Learned-K逐帧执行匀速预测和自适应观测修正。",
        "分别编码雷达、EO和ESM，通过被动置信融合、时间对齐、姿态门控和双向跨注意力形成单平台融合序列。",
        "由时序Transformer提取整段运动上下文，再在每个时刻依据传感器有效性和平台距离执行图协同。",
        "聚合有效平台特征，经位置与速度残差头修正Learned-K状态，并加回参考位置得到绝对ECEF状态。",
        "运行时过滤无真值帧、检查非有限数值，并将结果交给回放或批量指标计算。",
    ])
    doc.add_heading("2.4 模型与支撑软件的关系", level=2)
    doc.add_paragraph(
        "模型源码和权重决定跟踪结果；数据解析器决定模型看到的张量含义；推理运行时决定权重加载、设备、完整序列调用和结果过滤。前端仅消费TrackingSequence和批测报告，后端接口仅封装上述对象。因此，系统设计评审应首先检查模型与数据口径，其次检查运行时是否忠实调用模型，最后检查展示与报告是否正确表达输出。"
    )
    doc.add_heading("2.5 数据流", level=2)
    _add_lead_paragraph(
        doc,
        "源数据读取与目标级拆分",
        "数据流起点是已经完成传感器注册和目标关联的场景PKL文件。注册数据解析器首先读取雷达、EO、ESM、无人机平台位置、机体系到世界系旋转矩阵、快慢链路位置及其龄期和更新标志、统一时间轴、三类传感器掩码以及目标真值。一个场景中存在多个目标时，解析器按目标索引拆分为相互独立的样本，禁止在批次维上共享目标状态。各字段随后被裁剪或补齐为统一的T帧和N个平台槽位；可选字段缺失时补为规定形状的零张量，同时利用mask、age或update明确标记其无效性，避免仅凭数值为零判断观测是否存在。",
    )
    _add_lead_paragraph(
        doc,
        "坐标、时间与采样口径统一",
        "每条目标样本以首帧目标绝对ECEF位置建立ref_pos，平台位置、链路位置和位置标签均减去该参考位置后进入模型，从而把大尺度绝对坐标回归转化为局部相对状态估计；角度、速度、旋转矩阵和时间戳不执行平移。解析器从t_grid的有效相邻差分中取中位数形成样本dt，使状态转移步长与实际数据一致。V8源数据保留400帧、40 Hz记录，但在不重新训练的部署方案中同步对所有字段执行固定四帧取一，形成100帧、约10 Hz的兼容输入；原始帧索引、时间戳和来源路径随兼容数据保留，以便验证抽样过程。",
    )
    _add_lead_paragraph(
        doc,
        "物理粗状态形成",
        "标准张量首先进入LearnedKFrontEnd。雷达az、el、r依据机体系几何关系还原为相对三维向量，再由rotmat_b2e旋转到世界方向并叠加平台相对位置，形成参考坐标系中的雷达几何位置。每个时刻汇集全部有效雷达、快链和慢链位置：雷达权重来自径向与切向异性精度矩阵，链路权重同时受固定位置精度、age指数新鲜度和update标志约束。各来源的信息矩阵和信息向量累加后，通过带单位阵正则项的三阶安全求解器得到融合位置；无观测时沿用上一粗位置，并由相邻位置和dt构造限幅速度，最终组成B×T×6的粗观测z_obs。",
    )
    _add_lead_paragraph(
        doc,
        "Learned-K递推与多模态编码",
        "Learned-K以z_obs首帧初始化后验状态，后续时刻先按六维匀速转移矩阵生成预测，再以预测状态、当前粗观测、创新量和96维观测质量特征共同生成六维对角增益，对创新逐分量修正，形成连续粗状态k_state。与此同时，雷达主流将原始雷达量测、几何位置和k_state拼接编码；EO与ESM把角度量测和经姿态旋转后的世界单位方向编码为被动特征，并根据各自mask和学习到的置信度进行融合。链路龄期、平台位置、IMU、传感器有效性和ESM模式共同进入观测质量特征，使深度支路能够区分真实量测、陈旧链路与缺失通道。",
    )
    _add_lead_paragraph(
        doc,
        "时空特征融合与状态回归",
        "雷达主流和被动流先在每个平台的100帧序列内执行自适应对齐和姿态门控，再通过雷达到被动、被动到雷达两个独立方向的跨注意力交换上下文。融合门逐时刻决定两类交互结果的相对比例，并与未经过交互的雷达基准特征共同投影为96维序列。TemporalTransformer在单个平台内部提取整段时间上下文；随后模型按每个时刻重排为N节点图，以传感器有效性和平台间不超过20000 m的距离构造邻接矩阵，由CollaborativeGAT完成跨平台聚合。有效节点特征按平台维加权平均得到B×T×96的全局特征，位置头和速度头输出六维修正量，与k_state逐元素相加后，仅对位置分量加回ref_pos，生成B×T×6的绝对ECEF位置与世界系速度。",
    )
    _add_lead_paragraph(
        doc,
        "结果过滤、统计与归档",
        "模型输出先返回CPU，并与真值、t_grid共同按target_valid_mask保留有效帧，保证时间、预测和标签一一对应；随后计算三维位置误差并检查NaN或Inf，任何非有限结果都会终止该场景而不会写入成功缓存。单场景推理将通过检查的时间轴、预测状态、真值状态、逐帧误差和模型前向耗时封装为TrackingSequence，供界面回放使用。批量测试复用同一模型输出定义，在逐航迹层面统计有效点数量、RMSE、MAE、P95和最大误差，并在总体层面汇总稳定阶段位置误差、速度误差及重复前向测得的模型时延。评估服务最后读取批测结果和独立生成的关联及基线证据，形成JSON和CSV制品；该数据流是单向的，展示、统计和证据判定均不会向模型回写状态或改变权重。",
    )
    doc.add_heading("2.6 控制流", level=2)
    _add_lead_paragraph(
        doc,
        "模型发现、校验与原子激活",
        "服务启动或用户切换模型时，ModelManager先扫描已注册权重及其说明信息，再创建候选TransformerEngine。候选引擎按配置选择CPU或CUDA设备，实例化d_model为96的CoopTrackerModel，并使用strict方式加载state_dict；键名、张量形状、模型结构或设备初始化任一环节失败，候选对象立即废弃。只有模型构造、权重校验和设备迁移全部成功后，管理器才在受保护区内一次性替换活动引擎引用，因此失败切换不会让半初始化模型进入可见状态，也不会中断原活动模型继续服务。",
    )
    _add_lead_paragraph(
        doc,
        "单场景推理调度",
        "单场景请求先解析模型标识、数据集、场景路径和目标索引，并以场景绝对路径与目标索引组合作为缓存键。运行时在当前引擎的互斥锁内依次执行缓存查询、注册数据解析、可选字段补齐、批次维构造、张量迁移和完整序列前向，防止多个请求同时操作同一模型实例或写入同一缓存。缓存命中时直接返回已经通过验证的TrackingSequence；缓存未命中时仅在前向、有效帧过滤和有限数检查全部成功后提交结果。模型切换会创建新的引擎和新的生命周期缓存，旧模型产生的序列不会被新模型复用。",
    )
    _add_lead_paragraph(
        doc,
        "模型内部执行顺序",
        "前向调用在torch.inference_mode上下文中一次处理完整100帧窗口，而不是由界面逐帧触发模型。Learned-K内部按照时间从首帧向末帧递推，使后一时刻粗状态依赖前一后验；多模态编码、跨注意力和TemporalTransformer则在固定长度张量上整体计算，当前实现未传入因果或padding mask。use_temporal和use_gat只控制消融路径，传感器mask、链路age和update只改变观测权重、缺失特征及有效图节点，不改变B、T、N张量形状。该顺序保证正常模式、无图模式和雷达主流模式共享同一输入契约，同时明确界面按帧播放不等于严格在线因果推理。",
    )
    _add_lead_paragraph(
        doc,
        "异常终止与结果提交",
        "推理过程中若场景或目标不存在、输入维度不一致、权重无法strict匹配、设备不可用、矩阵计算失败、有效帧为空或误差包含NaN和Inf，控制流立即进入失败分支并返回明确异常，不把部分帧包装成成功结果。完成模型前向后，运行时同步过滤时间轴、预测和真值，并验证长度一致；只有这些条件均满足时才构造TrackingSequence、写入内存缓存并交给回放层。由此形成“准备—计算—校验—提交”的事务边界，避免前端读取正在生成或未经验证的中间状态。",
    )
    _add_lead_paragraph(
        doc,
        "批量任务状态机",
        "批量测试创建后首先处于queued状态，由单工作线程依次领取计算密集任务并转为running，以控制GPU显存占用和模型实例并发。任务运行期间逐场景、逐目标调用相同的推理与指标口径，并周期性检查取消请求；全部样本和报告制品成功写入后才进入completed。用户取消、计算异常或服务重启造成的中断分别进入cancelled或failed，恢复逻辑会把无法安全续跑的queued和running历史任务终结化，禁止把缺少结果文件、汇总统计或来源签名的任务标记为完成。",
    )
    _add_lead_paragraph(
        doc,
        "回放控制与模型计算隔离",
        "WebSocket会话将结果发送和控制命令接收分为相互独立的协程。play和pause只改变已生成TrackingSequence的播放游标及发送节奏；seek根据请求时间选择最近帧，并以该帧为截止位置重新计算界面需要显示的累计误差，不会重新读取PKL、再次调用模型或改变缓存内容。连接关闭时仅终止该会话的发送任务，模型、权重和批测任务不受影响。该隔离保证多次暂停、跳转或改变播放速度得到的仍是同一条预测序列。",
    )
    _add_lead_paragraph(
        doc,
        "证据有效性控制",
        "正式验收在任务完成后单独校验模型权重、数据范围、运行代码、关键参数和任务配置的来源摘要，并要求关联准确率证据与IMM-JPDA基线证据分别由受控脚本生成。证据只有在任务ID、指标ID和任务签名与当前批测一致时才进入正式判定；模型、数据、采样方式、运行代码或阈值发生变化后，旧结果仍可保留为历史记录，但formal_evidence_complete不得沿用。评估服务只读取已提交的模型结果和证据，不回写网络状态，也不会因为前端展示值变化而重新判定模型性能。",
    )
    doc.add_heading("2.7 设计原则", level=2)
    add_bullets(doc, [
        "模型优先：所有展示、测试和报告口径围绕同一模型输入输出定义建立。",
        "物理与学习结合：几何反演和状态递推提供可解释基线，深度网络学习难以显式建模的修正量。",
        "多源缺失可运行：使用掩码、缺失向量、链路龄期和分母保护维持张量结构。",
        "离线序列推理边界明确：当前Transformer使用整段序列上下文，界面回放不等同严格因果在线跟踪。",
        "结果可复核：权重、数据、运行环境、任务配置和指标证据均保留来源信息。",
    ])

    doc.add_heading("3 模型输入数据与坐标约定", level=1)
    doc.add_heading("3.1 基础维度与参数", level=2)
    add_table(doc, ["参数", "当前值", "含义"], [
        ["B", "运行时1；训练和批测按配置", "批次维度"],
        ["T", "100", "单次模型前向的序列长度"],
        ["N", "5", "协同无人机数量"],
        ["D", "96", "各神经特征分支的隐维度"],
        ["注意力头数", "4", "跨注意力和时序Transformer头数"],
        ["时序编码层数", "2", "TransformerEncoder层数"],
        ["图邻接距离", "20000 m", "平台协同边的最大欧氏距离"],
        ["可训练参数量", "496623", "d_model为96时的CoopTrackerModel参数总数"],
    ], widths=[1.6, 1.55, 3.65])
    doc.add_heading("3.2 状态向量与参考位置", level=2)
    doc.add_paragraph(
        "目标状态统一按px、vx、py、vy、pz、vz交错排列。位置单位为米，速度单位为米每秒。每条航迹以首帧绝对ECEF位置作为ref_pos，训练标签的位置分量和平台、链路位置均减去ref_pos后进入模型；模型末端只对位置分量加回ref_pos，速度分量保持原世界坐标语义。"
    )
    doc.add_heading("3.3 传感器输入", level=2)
    add_table(doc, ["字段", "形状", "实际分量顺序", "主要用途"], [
        ["radar", "B×T×N×3", "az、el、r", "几何反演、信息融合和主动主流编码"],
        ["eo", "B×T×N×2", "az、el", "构造世界方向并形成被动流"],
        ["esm", "B×T×N×2", "az、el；缺失分量按零处理", "构造被动方向特征"],
        ["radar_mask", "B×T×N", "0或1", "雷达观测权重、通道缺失和图节点有效性"],
        ["eo_mask、esm_mask", "B×T×N", "0或1", "被动置信融合、缺失处理和图节点有效性"],
    ], widths=[1.45, 1.5, 2.0, 1.85])
    doc.add_paragraph(
        "源码中的雷达解析会将源文件的r、az、el重新排列为模型实际使用的az、el、r。设计和测试必须以模型实际顺序为准，避免仅依据旧注释解释字段。"
    )
    doc.add_heading("3.4 平台位置、姿态和旋转矩阵", level=2)
    doc.add_paragraph(
        "uav_pos形状为B×T×N×3，采用减去ref_pos后的参考坐标。rotmat_b2e形状为B×T×N×3×3，用于把机体系雷达相对位置、EO方向和ESM方向转换到世界方向。当PKL包含UAV_Rotmat时直接使用；缺失时，解析器根据UAV_Pos和实际t_grid计算速度方向，并构造前、右、上三个正交轴。"
    )
    doc.add_heading("3.5 IMU与链路字段", level=2)
    add_table(doc, ["字段", "形状", "设计含义"], [
        ["imu", "B×T×N×6", "姿态与角速度特征，参与Learned-K观测编码和双流姿态门控"],
        ["fast_link、slow_link", "B×T×N×3", "参考坐标下的协同位置观测"],
        ["fast_age、slow_age", "B×T×N", "距离最近一次链路更新的时间，用于新鲜度衰减"],
        ["fast_update、slow_update", "B×T×N", "本帧是否收到链路更新"],
        ["esm_mode", "B×T×N", "ESM模式编码，进入Learned-K的30维观测特征"],
    ], widths=[1.8, 1.7, 3.1])
    doc.add_paragraph(
        "注册PKL未提供的可选字段由TransformerEngine补为固定形状零张量。模型会把全零链路位置判为无效，将对应龄期修正为999秒并清除更新标志，使其不参与粗观测融合。"
    )
    doc.add_heading("3.6 僚机输入口径", level=2)
    doc.add_paragraph(
        "当前统一训练入口train.py保留所有平台的雷达、EO和ESM原始有效掩码，与部署和批测入口一致。有效僚机原始观测可形成GAT节点贡献，快慢链路同时进入Learned-K前端。旧MAT训练入口曾屏蔽僚机原始传感器，但该入口已删除，不能将其规则套用于当前训练或部署。调用方应忠实传递实际有效掩码；若改为仅允许僚机链路，需要共同调整节点编码、有效性、图结构和训练配置并重新验证。"
    )
    doc.add_heading("3.7 时间采样与V8兼容", level=2)
    doc.add_paragraph(
        "数据解析器从PKL的t_grid相邻正时间差中取中位数作为dt，不再固定假设0.1秒。V5至V7训练数据按10 Hz组织，当前权重以100帧序列学习约10秒上下文。V8源数据保持40 Hz和400帧，用于保存真实量测节拍；进入当前权重前通过固定四帧取一帧生成100帧、10 Hz兼容数据，并保留源路径、原始帧索引和时间戳。该处理改变的是推理输入采样，不重新训练权重。"
    )
    doc.add_heading("3.8 有效帧与缺失数据", level=2)
    add_bullets(doc, [
        "三类传感器掩码控制通道替代、观测权重、图邻接和最终节点汇聚。",
        "target_valid_mask只用于真值与指标过滤，不直接作为Transformer的注意力掩码。",
        "ChannelMaskEncoder用可学习mask token替换无效通道编码，EO和ESM均无效时使用passive_missing。",
        "当前时序Transformer未传入因果掩码或padding mask，因此按固定长度整段处理。",
        "所有样本必须保持字段维度一致；补零表示缺失时必须同步设置无效掩码或无效龄期。",
    ])

    doc.add_heading("4 端到端协同跟踪模型总体设计", level=1)
    doc.add_heading("4.1 跟踪任务定义", level=2)
    doc.add_paragraph(
        "CoopTrackerModel实现从多平台、多源传感器时间序列到单目标六维状态序列的端到端映射。输入是已经完成目标级配准和时间对齐的张量，输出覆盖输入窗口中的每个时刻。模型不在forward内部完成原始点迹检测、目标关联或航迹起始终止管理，这些环节属于数据准备与验证边界。"
    )
    doc.add_heading("4.2 总体结构", level=2)
    if MODEL_FIGURE.is_file():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(MODEL_FIGURE), width=Inches(6.7))
        _add_caption(doc, "图 1 CoopTrackerModel端到端协同跟踪结构")
    doc.add_paragraph(
        "网络由五个逻辑层组成：输入与空间配准层、物理先验Learned-K层、主被动双流融合层、时序与图协同层、残差状态输出层。物理支路提供逐帧连续粗状态，深度支路学习多模态、长时序和跨平台关系，最终只预测相对于粗状态的修正量。"
    )
    doc.add_heading("4.3 前向张量主线", level=2)
    add_table(doc, ["阶段", "主要张量变化", "输出"], [
        ["样本输入", "B×T×N的多源张量", "传感器、平台、链路、时间和掩码"],
        ["Learned-K", "平台维观测汇聚为逐帧状态", "k_state和z_obs，均为B×T×6"],
        ["分平台编码", "B×T×N重排为B·N条长度T序列", "三类B·N×T×D特征"],
        ["双流融合", "雷达主流与被动流双向跨注意力", "B·N×T×D融合特征"],
        ["时序建模", "每个平台独立编码整段时间上下文", "B·N×T×D时序特征"],
        ["图协同", "按时刻重排为B·T个N节点图", "B×T×N×D节点特征"],
        ["节点汇聚", "按传感器有效性加权平台维", "B×T×D全局特征"],
        ["状态输出", "位置、速度残差与k_state相加", "B×T×6绝对状态"],
    ], widths=[1.25, 3.45, 2.1])
    doc.add_heading("4.4 物理支路与深度支路的分工", level=2)
    doc.add_paragraph(
        "物理支路利用可解释几何和匀速状态转移建立观测基准，适合承担坐标恢复、量测精度加权和短时连续性；深度支路不重新生成绝对状态，而是基于原始量测、粗状态和上下文估计残差，重点修正观测噪声、目标机动和多平台互补关系。两支路在模型末端相加，使网络学习目标从完整状态回归转为粗状态修正。"
    )
    doc.add_heading("4.5 离线序列推理边界", level=2)
    doc.add_paragraph(
        "TemporalTransformer没有使用因果注意力掩码，因此某一时刻的输出可以利用同一100帧窗口内前后时刻的信息。当前原型软件在完整序列前向完成后再进行WebSocket回放，适用于离线验证、数据复盘和批量评估。若后续要求严格在线逐帧跟踪，应重新设计因果掩码、状态缓存和滑动窗口，并重新验证或训练权重，不能仅提高回放刷新频率替代。"
    )

    doc.add_heading("5 物理先验与Learned-K状态前端", level=1)
    doc.add_heading("5.1 模块作用", level=2)
    doc.add_paragraph(
        "LearnedKFrontEnd先从雷达和异步链路得到粗观测z_obs，再按照六维匀速模型生成预测状态，并由k_net输出逐状态分量的增益。该结构保留传统预测—创新—修正框架，但增益不由固定协方差递推给出，而是由当前预测、观测、创新和多传感器特征共同学习。"
    )
    doc.add_heading("5.2 雷达机体系到参考坐标转换", level=2)
    add_numbered_steps(doc, [
        "从radar的az、el、r计算机体系相对三维向量，三个分量分别对应前向、侧向和竖直方向。",
        "使用rotmat_b2e左乘机体系向量，将其旋转到ECEF方向。",
        "叠加已经减去ref_pos的uav_pos，得到参考位置平移后的目标几何位置。",
        "同一旋转矩阵用于EO和ESM单位方向的世界系转换，保证主被动方向特征采用一致姿态基准。",
    ])
    doc.add_heading("5.3 信息矩阵几何融合", level=2)
    add_table(doc, ["观测来源", "有效条件", "精度与新鲜度模型"], [
        ["雷达", "radar_mask有效", "径向标准差48 m；切向方差由距离乘0.002形成，并增加1平方米保护项"],
        ["快链", "本帧更新或龄期小于2 s", "60 m各向同性标准差；新鲜度按0.3 s时间尺度指数衰减"],
        ["慢链", "本帧更新或龄期小于5 s", "60 m各向同性标准差；新鲜度按2.0 s时间尺度指数衰减"],
    ], widths=[1.2, 2.0, 3.6])
    doc.add_paragraph(
        "每个雷达根据视线方向构造径向和切向不同的三维精度矩阵，链路位置使用各向同性精度矩阵。系统累加全部有效观测的精度矩阵和精度加权位置，并加入1e-6单位阵正则项，通过固定三阶伴随矩阵公式求解融合位置。该求解路径避免torch.linalg.solve对目标ONNX导出环境的兼容限制。"
    )
    doc.add_heading("5.4 无观测与速度构造", level=2)
    doc.add_paragraph(
        "某帧没有任何有效雷达或链路位置时，融合位置沿用前一帧结果，避免由缺失值产生零点跳变。速度由相邻融合位置除以样本dt得到，首帧速度置零，每轴限制在负400 m/s至400 m/s。位置和速度随后按px、vx、py、vy、pz、vz排列为六维z_obs。"
    )
    doc.add_heading("5.5 三十维观测质量特征", level=2)
    add_table(doc, ["特征组", "维数", "内容"], [
        ["原始传感器", "7", "雷达3、EO 2、ESM 2"],
        ["IMU", "6", "姿态与角速度"],
        ["有效性与模式", "4", "三类mask及ESM模式"],
        ["链路龄期", "2", "快链与慢链age"],
        ["链路位置", "6", "快链3和慢链3"],
        ["链路更新", "2", "快链与慢链update"],
        ["平台位置", "3", "参考坐标下的uav_pos"],
    ], widths=[1.6, 1.0, 4.2])
    doc.add_paragraph(
        "30维特征经两组线性层、LayerNorm和GELU映射到96维。仅对至少一种传感器有效的平台特征进行平均，得到逐帧pooled_obs；无有效平台时分母增加1e-6并得到零汇聚结果。"
    )
    doc.add_heading("5.6 Learned-K递推", level=2)
    add_numbered_steps(doc, [
        "以z_obs首帧初始化上一后验状态。",
        "根据每个样本的dt构造六维匀速状态转移矩阵，后续帧由上一后验状态预测当前状态。",
        "用当前粗观测减去预测状态得到六维创新量。",
        "拼接预测状态、粗观测、创新量和96维pooled_obs，形成114维k_net输入。",
        "k_net经两层96维隐藏映射和Sigmoid输出六维对角增益，范围限制在0至1。",
        "按逐元素增益修正创新量并叠加到预测状态，形成当前后验状态，再递推至序列末尾。",
    ])
    doc.add_heading("5.7 数值保护", level=2)
    add_bullets(doc, [
        "信息矩阵加入1e-6单位阵，固定三阶求解器将行列式下限限制为1e-24。",
        "dt在差分速度计算时设置1e-6下限，避免除零。",
        "无观测帧沿用前一位置，差分速度逐轴限幅。",
        "全零链路位置被判为无效并设置999秒龄期，不进入链路权重。",
    ])

    doc.add_heading("6 多模态特征编码与双流融合", level=1)
    doc.add_heading("6.1 雷达主动主流", level=2)
    doc.add_paragraph(
        "雷达分支输入由原始az、el、r三维量测、姿态校正后的三维几何位置和扩展到每个平台的六维k_state组成，共12维。ChannelMaskEncoder使用两层线性映射、LayerNorm和GELU输出96维特征；radar_mask无效时以可学习mask token替代该时刻的通道特征。"
    )
    doc.add_heading("6.2 EO与ESM被动流", level=2)
    doc.add_paragraph(
        "EO和ESM分别使用两维角度与三维世界方向组成5维输入，并通过各自的ChannelMaskEncoder映射到96维。两类被动特征分别经线性层和Sigmoid得到置信权重，再乘对应mask进行加权平均；两者均无效时使用可学习passive_missing向量。"
    )
    doc.add_heading("6.3 时间维自适应对齐", level=2)
    doc.add_paragraph(
        "雷达主流和被动流分别沿长度为T的时间维计算均值与标准差，完成序列内标准化，再通过可学习align_gamma和align_beta调整到统一分布。标准差设置1e-5下限。该对齐在每个平台序列内部完成，不混合不同无人机。"
    )
    doc.add_heading("6.4 姿态门控", level=2)
    doc.add_paragraph(
        "六维IMU特征经线性层投影到96维，再经Sigmoid形成pose_gate。雷达和被动特征均乘以1加pose_gate，使姿态与角速度以幅值调制方式进入两条流。注册PKL缺少IMU时该输入补零，门控仍由线性层偏置形成可学习基线。"
    )
    doc.add_heading("6.5 双向跨注意力", level=2)
    add_table(doc, ["方向", "Query", "Key与Value", "作用"], [
        ["雷达到被动", "雷达序列", "被动序列", "以主动量测查询被动方向上下文"],
        ["被动到雷达", "被动序列", "雷达序列", "以被动方向查询主动距离与粗状态上下文"],
    ], widths=[1.45, 1.5, 1.65, 2.2])
    doc.add_paragraph(
        "每个CrossAttentionBlock使用4头MultiheadAttention，随后执行注意力残差、LayerNorm、两倍通道前馈网络、GELU、0.1 Dropout和第二次残差归一化。注意力在每个平台的时间序列内部计算，输入形状为B·N×T×96。"
    )
    doc.add_heading("6.6 融合门与输出投影", level=2)
    doc.add_paragraph(
        "两路跨注意力上下文拼接后由fusion_gate输出逐时刻标量alpha。模型使用alpha在两路上下文之间加权，再把加权结果与radar_ctx拼接，经线性层投影回96维。该设计保留雷达主流基准，同时允许被动量测在可靠时改变融合结果。"
    )

    doc.add_heading("7 时序建模与多无人机协同", level=1)
    doc.add_heading("7.1 时序Transformer", level=2)
    doc.add_paragraph(
        "融合后的每个平台序列加入96维正弦位置编码，再进入两层TransformerEncoder。每层包含4头自注意力、384维前馈网络、GELU激活和0.1 Dropout，并采用batch_first布局。该模块提取目标机动与量测变化的整段上下文。"
    )
    doc.add_heading("7.2 非因果注意力说明", level=2)
    doc.add_paragraph(
        "模型未配置因果mask，时序编码器可同时访问当前帧前后的窗口特征。因此输出应解释为固定窗口离线序列估计，而不是只依赖历史帧的实时滤波结果。前端播放间隔和雷达量测频率不会改变该计算边界。"
    )
    doc.add_heading("7.3 图节点与邻接矩阵", level=2)
    add_numbered_steps(doc, [
        "把B·N×T×96的时序结果重排为B×T×N×96，使每个时刻形成一个N节点图。",
        "当平台至少有一种雷达、EO或ESM有效时，将该平台标记为有效传感器节点。",
        "计算平台参考坐标之间的欧氏距离，仅保留距离不超过20000 m的节点对。",
        "传感器有效性邻接与距离邻接逐元素相乘，再强制加入对角自环。",
    ])
    doc.add_heading("7.4 CollaborativeGAT", level=2)
    doc.add_paragraph(
        "CollaborativeGAT分别用线性层生成Q、K和V，以96维平方根缩放点积得到节点间得分。无邻接边被填充为负大数，有效边经Softmax归一化和0.1 Dropout后聚合V。聚合结果与原节点特征残差相加，并通过LayerNorm输出协同节点特征。"
    )
    doc.add_heading("7.5 有效节点汇聚", level=2)
    doc.add_paragraph(
        "图输出按valid_sensor在平台维加权平均为B×T×96的global_feat。无有效传感器时分母增加1e-6，输出保持有限。虽然图中为所有节点加入自环，但最终全局汇聚仍会排除无有效传感器的平台。"
    )
    doc.add_heading("7.6 消融开关", level=2)
    add_table(doc, ["模式", "控制方式", "用途"], [
        ["完整模型", "use_temporal与use_gat均启用", "正式模型推理"],
        ["无图协同", "use_gat为false", "验证跨平台图聚合贡献"],
        ["仅雷达主流", "EO和ESM置零并关闭其mask", "验证被动流贡献"],
        ["无时序", "use_temporal为false", "源码支持的时序模块消融入口"],
    ], widths=[1.5, 2.5, 2.8])

    doc.add_heading("8 状态回归与完整推理流程", level=1)
    doc.add_heading("8.1 位置与速度残差头", level=2)
    doc.add_paragraph(
        "global_feat分别进入位置头和速度头。两个回归头均由96维线性映射到48维、GELU和三维输出层构成，分别产生三维位置残差与三维速度残差。两类输出按状态顺序交错写入六维residual。"
    )
    doc.add_heading("8.2 粗状态与深度残差叠加", level=2)
    doc.add_paragraph(
        "六维residual与Learned-K输出k_state逐元素相加得到参考坐标下的pred_rel。该叠加关系使模型在物理粗估计附近进行修正，而不必完全从零回归大尺度ECEF绝对坐标。"
    )
    doc.add_heading("8.3 绝对状态恢复", level=2)
    add_bullets(doc, [
        "对pred_rel的第0、2、4分量分别加回ref_pos的x、y、z。",
        "第1、3、5速度分量不叠加参考位置。",
        "最终pred_abs形状为B×T×6，状态顺序保持px、vx、py、vy、pz、vz。",
    ])
    doc.add_heading("8.4 部署推理步骤", level=2)
    add_numbered_steps(doc, [
        "TransformerEngine创建d_model为96的模型，将权重映射到指定CPU或CUDA设备并strict加载。",
        "读取指定PKL和目标索引，通过_parse_registered_pkl得到标准样本，并补齐可选字段。",
        "所有张量移动到活动设备并增加B等于1的批次维，dt整理为一维批张量。",
        "在torch.inference_mode中执行一次100帧完整前向；CUDA计时前后同步设备。",
        "按target_valid_mask同时过滤真值、预测和时间轴，并以位置分量计算三维欧氏误差。",
        "若误差中出现NaN或Inf则中止场景；否则封装TrackingSequence并写入内存缓存。",
    ])
    doc.add_heading("8.5 推理缓存与模型切换", level=2)
    doc.add_paragraph(
        "缓存键由场景绝对路径和目标索引组成，只在当前TransformerEngine生命周期内有效。模型切换通过新建候选引擎完成，因此新模型不会复用旧模型缓存；不同目标也不会混用结果。缓存不写入磁盘，正式指标由批测任务独立计算。"
    )
    doc.add_heading("8.6 推理输出边界", level=2)
    add_table(doc, ["包含内容", "不包含内容"], [
        ["已配准多源张量到目标状态的模型前向", "原始点迹检测和传感器信号处理"],
        ["100帧窗口内的位置与速度序列", "严格因果逐帧在线状态更新"],
        ["模型前向耗时", "文件读取、网络传输和界面绘制耗时"],
        ["有效真值帧的位置误差", "目标关联算法自身的正确率"],
    ], widths=[3.35, 3.45])

    doc.add_heading("9 模型训练与权重形成", level=1)
    doc.add_heading("9.1 训练数据组织", level=2)
    doc.add_paragraph(
        "RegisteredPklDataset将注册PKL中的每个目标整理为100帧、5平台样本，并以首帧目标位置建立相对坐标。训练入口默认随机种子42，使用独立训练与验证目录；仅显式将val_data_dir置空时按样本索引随机8:2切分。DataLoader训练批次默认为4，验证集不打乱；历史V7日志采用批次32。"
    )
    doc.add_heading("9.2 输入约束", level=2)
    doc.add_paragraph(
        "当前统一train.py保留所有平台的实际原始雷达、EO和ESM掩码，使有效僚机原始观测可参与编码和GAT；快慢链路同时进入Learned-K。部署与批测采用相同口径，不能套用已删除旧MAT入口的僚机屏蔽规则。"
    )
    doc.add_heading("9.3 损失函数", level=2)
    add_table(doc, ["损失项", "计算对象", "权重"], [
        ["位置回归损失", "标准化后的相对位置预测与相对位置标签之间的SmoothL1", "1.0"],
        ["序列平滑损失", "预测与真值相邻位置增量经位置标准差缩放后的SmoothL1", "0.1"],
        ["速度损失", "标准化速度预测与速度标签之间的SmoothL1", "前期0，第二阶段按配置在10轮内渐增到目标权重"],
    ], widths=[1.55, 4.0, 1.25])
    doc.add_paragraph(
        "默认第二阶段从第5轮开始，速度损失权重在10轮内线性增至0.08；命令行或自动阶段参数可覆盖默认值。位置与速度均值和标准差最多从训练加载器前50批估计，标准差设置下限，再将三轴标准差平均为共同尺度。warmup_ignore只用于评估统计排除初始过渡帧，不从训练损失中删除前5帧。"
    )
    doc.add_heading("9.4 优化与权重选择", level=2)
    add_table(doc, ["项目", "默认值"], [
        ["优化器", "AdamW"],
        ["学习率", "1e-4，可由已保存的第二阶段搜索结果覆盖"],
        ["权重衰减", "1e-4"],
        ["训练轮数", "40"],
        ["梯度裁剪", "全局范数1.0"],
        ["最佳权重条件", "验证集位置RMSE最低"],
    ], widths=[2.0, 4.8])
    doc.add_heading("9.5 训练与部署一致性", level=2)
    add_bullets(doc, [
        "部署模型必须使用d_model为96的同构CoopTrackerModel并strict加载。",
        "传感器字段顺序、参考坐标、dt、旋转矩阵和僚机mask应与训练入口一致。",
        "V8兼容输入保持100帧和约0.1秒dt，以避免改变权重学习到的时间尺度。",
        "切换权重不会自动改变数据集和输入规则，模型注册信息应同时记录训练与验证数据口径。",
    ])

    doc.add_heading("10 模型验证与技术指标", level=1)
    doc.add_heading("10.1 验证目标", level=2)
    doc.add_paragraph(
        "模型验证用于回答四个问题：输出位置是否达到精度阈值，目标关联输入是否达到正确率要求，相对IMM-JPDA基线是否具有规定改善，单帧模型前向是否满足时延约束。前端图表只辅助观察，正式结论来自可复核批测报告和独立证据。"
    )
    doc.add_heading("10.2 位置误差口径", level=2)
    doc.add_paragraph(
        "位置误差使用pred_abs和label_abs的第0、2、4分量计算三维欧氏距离。批测按target_valid_mask保留真值有效帧，并对每条航迹排除配置的初始预热帧。总体RMSE以全部有效点的距离平方均值开根号，同时保留MAE、中位数、P95、最大误差和逐航迹分布。"
    )
    doc.add_heading("10.3 推理时延口径", level=2)
    doc.add_paragraph(
        "时延计量范围从已构造张量进入CoopTrackerModel到六维状态输出，不包含文件I/O、原始检测、目标关联、HTTP传输和前端绘制。CUDA设备在计时前后同步。完整批测另外对单航迹序列进行3次预热和10次实测，并将序列时间除以100得到单帧折算时延。"
    )
    doc.add_heading("10.4 关联与基线证据", level=2)
    doc.add_paragraph(
        "关联准确率和IMM-JPDA相对改善率由独立白名单脚本生成。证据必须记录任务ID、指标ID、任务签名和值有效性；AssociationStats和人工基线只能用于技术复核。只有两份独立证据均被接受且与当前模型、数据、运行代码和关键参数一致时，formal_evidence_complete才为true。"
    )
    doc.add_heading("10.5 场景覆盖", level=2)
    add_bullets(doc, [
        "V5、V6和V7用于验证不同噪声与数据处理条件下的模型表现。",
        "V8保留40 Hz原始量测，并使用10 Hz兼容数据验证真实节拍场景在当前权重下的适配能力。",
        "抽样批测用于快速回归，完整批测用于正式指标结论。",
        "批测报告保存实际场景清单、类别数量、航迹数量、有效点数量和运行来源哈希。",
    ])
    doc.add_heading("10.6 消融与解释", level=2)
    doc.add_paragraph(
        "训练脚本支持完整模型、无图协同和仅雷达主流三种对比，用于观察被动传感器与协同图模块对位置误差和时延的影响。模型解释应结合粗状态、残差输出、场景类别和逐航迹误差分布，不能只用单一平均值推断某个子模块在所有场景中均有效。"
    )

    doc.add_heading("11 展示与验证支撑系统", level=1)
    doc.add_heading("11.1 支撑功能边界", level=2)
    doc.add_paragraph(
        "支撑系统负责发现模型和场景、调用完整序列推理、控制结果回放、提交批量测试、展示技术指标并导出报告。它不修改网络结构、权重参数或模型输出，也不承担飞行控制和在线任务规划。"
    )
    doc.add_heading("11.2 前端展示", level=2)
    add_table(doc, ["功能", "与模型的关系"], [
        ["模型选择", "显示已注册权重和活动设备，发起原子切换"],
        ["跟踪回放", "消费TrackingSequence，绘制预测、真值和误差"],
        ["批量测试", "提交模型、数据集、预热帧和批大小等验证配置"],
        ["性能评估", "显示四项指标、证据完整性和场景分布"],
        ["报告中心", "查询已持久化任务并下载JSON或CSV"],
    ], widths=[1.7, 5.1])
    doc.add_heading("11.3 后端服务", level=2)
    doc.add_paragraph(
        "FastAPI服务把ModelManager、TransformerEngine、BatchTestManager和EvaluationService封装为受控接口。模型管理负责权重发现与切换，运行时负责单场景推理，批测管理负责完整数据集计算，评估服务负责证据校验和结论生成。"
    )
    doc.add_heading("11.4 接口分类", level=2)
    add_table(doc, ["接口类别", "主要操作"], [
        ["系统与模型", "健康检查、环境信息、模型列表、活动模型切换"],
        ["数据与场景", "数据集列表、场景列表、目标索引解析"],
        ["回放", "WebSocket连接、pause、play和seek"],
        ["批测", "任务创建、查询、取消和任务结果读取"],
        ["评估与报告", "四项指标判定、证据生成、JSON与CSV下载"],
    ], widths=[1.65, 5.15])
    doc.add_heading("11.5 状态与异常", level=2)
    doc.add_paragraph(
        "HTTP中间件生成或透传X-Request-ID并记录方法、路径、状态和耗时。资源不存在、状态冲突、参数不合法和未处理异常分别映射为404、409、422和500。模型数值异常、权重结构不匹配和设备不可用均在进入展示层前明确失败。"
    )

    doc.add_heading("12 数据与运行制品", level=1)
    doc.add_heading("12.1 数据类别", level=2)
    add_table(doc, ["类别", "内容", "用途"], [
        ["模型权重", "各训练运行的best_model.pt", "部署与对比"],
        ["注册场景", "V5至V8的PKL文件", "单场景推理与批测"],
        ["兼容清单", "V8 40 Hz到10 Hz的索引和来源记录", "时间口径追溯"],
        ["批测报告", "任务配置、结果、场景清单和来源哈希", "指标统计与复核"],
        ["独立证据", "关联准确率与IMM-JPDA对比", "正式验收"],
        ["导出制品", "综合JSON、逐航迹CSV和任务对比CSV", "交付归档"],
    ], widths=[1.4, 3.2, 2.2])
    doc.add_heading("12.2 完整性与追溯", level=2)
    doc.add_paragraph(
        "正式任务记录模型权重、数据文件和运行代码的SHA-256摘要。证据读取时同时核对任务签名和必要来源字段；旧报告缺少来源摘要时显示为证据不完整，而不是自动沿用历史结论。JSON写入禁止NaN和Inf，任务更新使用临时文件与原子替换。"
    )
    doc.add_heading("12.3 生命周期", level=2)
    doc.add_paragraph(
        "模型与数据目录只读挂载，批测报告和证据目录单独授予写权限。后端启动时恢复历史终态任务；无法安全续跑的queued或running任务标记为failed并回写。正式验收报告、证据、环境信息和场景清单应作为同一归档单元保存。"
    )

    doc.add_heading("13 部署测试与维护", level=1)
    doc.add_heading("13.1 部署要求", level=2)
    add_bullets(doc, [
        "模型权重与注册数据目录必须存在，活动权重能够strict加载。",
        "计算设备支持CPU或可用CUDA设备，正式测试记录PyTorch、CUDA和GPU信息。",
        "前端仅用于受控环境展示，后端默认监听127.0.0.1，不直接作为公网服务。",
        "V8验证默认注册10 Hz兼容目录，同时保留40 Hz源数据路径。",
    ])
    doc.add_heading("13.2 测试层级", level=2)
    add_table(doc, ["层级", "重点"], [
        ["模型纯函数测试", "坐标反演、三阶安全求解、掩码、图邻接和张量形状"],
        ["模型运行时测试", "权重strict加载、完整序列前向、有限数和有效帧过滤"],
        ["批测测试", "稳定阶段统计、时延测量、任务恢复和来源签名"],
        ["展示验证测试", "模型选择、WebSocket回放、任务状态与报告下载"],
        ["正式指标测试", "完整数据集和两份独立验收证据"],
    ], widths=[1.8, 5.0])
    doc.add_heading("13.3 模型维护约束", level=2)
    add_bullets(doc, [
        "修改输入字段时同步更新dataset解析、运行时补齐、模型forward、训练入口和测试。",
        "修改d_model、编码层或回归头时生成新权重，不允许用strict为false掩盖结构不匹配。",
        "修改序列长度、采样率或因果性时重新评估位置精度和时延，必要时重新训练。",
        "修改邻接距离、掩码或僚机输入口径时同步更新训练与部署调用方。",
        "新增证据指标时同步更新阈值、签名、页面展示和导出结构。",
    ])
    doc.add_heading("13.4 交付检查", level=2)
    add_numbered_steps(doc, [
        "确认活动模型权重、训练数据说明和验证数据目录一致。",
        "执行模型运行时和批测自动化测试，确认所有输出为有限数。",
        "完成前端生产构建和后端健康检查。",
        "针对正式任务生成关联准确率和IMM-JPDA两份独立证据。",
        "归档任务JSON、综合评估JSON、逐航迹CSV、证据与环境信息。",
    ])

    doc.add_heading("附录A 模型关键参数", level=1)
    add_table(doc, ["参数", "数值或规则", "源码位置"], [
        ["d_model", "96", "CoopTrackerModel"],
        ["可训练参数", "496623", "CoopTrackerModel(d_model=96)"],
        ["Learned-K观测特征", "30维", "LearnedKFrontEnd.obs_encoder"],
        ["K网络输入", "114维", "预测6＋观测6＋创新6＋特征96"],
        ["跨注意力", "4头", "CrossAttentionBlock"],
        ["时序Transformer", "2层、4头、前馈384维", "TemporalTransformer"],
        ["协同距离", "不超过20000 m", "CoopTrackerModel.forward"],
        ["雷达径向标准差", "48 m", "LearnedKFrontEnd._build_obs_state"],
        ["雷达角度尺度", "0.002 rad", "LearnedKFrontEnd._build_obs_state"],
        ["链路位置标准差", "60 m", "LearnedKFrontEnd._build_obs_state"],
        ["速度限幅", "每轴±400 m/s", "LearnedKFrontEnd._build_obs_state"],
        ["序列长度", "100帧", "训练与TransformerEngine"],
    ], widths=[2.0, 2.35, 2.45])

    doc.add_heading("附录B 模型张量速查", level=1)
    add_table(doc, ["张量", "形状", "说明"], [
        ["radar", "B×T×N×3", "az、el、r"],
        ["eo、esm", "B×T×N×2", "被动角度量测"],
        ["uav_pos", "B×T×N×3", "相对ref_pos的平台位置"],
        ["rotmat_b2e", "B×T×N×3×3", "机体系到世界系旋转矩阵"],
        ["k_state、z_obs", "B×T×6", "Learned-K状态与粗观测"],
        ["radar_feat", "B·N×T×96", "雷达主流编码"],
        ["passive_feat", "B·N×T×96", "EO与ESM融合特征"],
        ["temporal", "B·N×T×96", "单平台时序特征"],
        ["graph_out", "B×T×N×96", "多平台协同节点特征"],
        ["global_feat", "B×T×96", "有效节点汇聚特征"],
        ["pred_abs", "B×T×6", "绝对位置与世界速度状态"],
    ], widths=[1.8, 1.75, 3.25])


def rewrite(path: Path) -> None:
    _assert_source_contracts()
    doc = Document(path)
    boundary = _front_boundary(doc)
    before_digest = _prefix_digest(doc, boundary)
    bookmark_records = _extract_heading_bookmarks(doc)

    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    backup = BACKUP_DIR / "设计说明书_本轮修改前备份.docx"
    shutil.copy2(path, backup)

    _remove_existing_body(doc, boundary)
    _build_body(doc)
    add_model_expansions(doc)
    _restore_toc_bookmarks(doc, boundary, bookmark_records)
    _format_new_body(doc, boundary)

    # Keep the mandatory paragraph after the last table compact.
    trailing = doc.paragraphs[-1]
    if not trailing.text.strip():
        trailing.paragraph_format.space_before = Pt(0)
        trailing.paragraph_format.space_after = Pt(0)
        trailing.paragraph_format.line_spacing = Pt(1)
        marker = trailing.add_run("\u200b")
        marker.font.size = Pt(1)
        marker.font.hidden = True

    doc.save(path)

    check = Document(path)
    after_boundary = _front_boundary(check)
    after_digest = _prefix_digest(check, after_boundary)
    if before_digest != after_digest or boundary != after_boundary:
        shutil.copy2(backup, path)
        raise RuntimeError("Front matter changed; original document restored")

    print(f"updated={path}")
    print(f"backup={backup}")
    print(f"front_boundary={boundary}")
    print(f"front_sha256={after_digest}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("document", nargs="?", default=str(DEFAULT_DOC))
    args = parser.parse_args()
    rewrite(Path(args.document).resolve())


if __name__ == "__main__":
    main()
