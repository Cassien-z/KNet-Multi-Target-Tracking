from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "docs" / "software_delivery"
ASSET_DIR = ROOT / "artifacts" / "documentation_build"
MODEL_ARCH = ROOT / "artifacts" / "figures" / "架构图.png"
DOC_DATE = "2026年9月13日"
SOFTWARE_NAME = "端到端协同跟踪原型软件"
SOFTWARE_VERSION = "前端 0.0.0  后端 API 1.4.0"

NAVY = "1F4E78"
PALE_BLUE = "DDEBF7"
PALE_GRAY = "F2F2F2"
BORDER = "D9D9D9"
TEXT_GRAY = "666666"


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_mixed_run_font(
    run,
    *,
    east_asia="楷体",
    latin="Times New Roman",
    size=14,
):
    """Apply the client-requested mixed Chinese/Latin body typography."""
    run.font.name = latin
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:cs"), latin)
    run.font.size = Pt(size)


def set_mixed_style_font(
    style,
    *,
    east_asia="楷体",
    latin="Times New Roman",
    size=14,
):
    style.font.name = latin
    r_fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), east_asia)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:cs"), latin)
    style.font.size = Pt(size)


def apply_design_body_typography(doc):
    """Set design-report text to the client-requested size-four mixed typography."""
    for style_name in ("Normal", "List Bullet", "List Number"):
        set_mixed_style_font(doc.styles[style_name], size=14)

    heading_styles = {"Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"}
    for style_name in heading_styles:
        style = doc.styles[style_name]
        set_mixed_style_font(style, size=14)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.italic = False
        remove_style_paragraph_borders(style)

    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if paragraph.text.strip().startswith("图 "):
            continue
        if style_name in heading_styles:
            for run in paragraph.runs:
                set_mixed_run_font(run, size=14)
                run.bold = True
                run.italic = False
                run.font.color.rgb = RGBColor(0, 0, 0)
            continue
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(6)
        for run in paragraph.runs:
            set_mixed_run_font(run, size=14)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1.2
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        set_mixed_run_font(run, size=14)


def set_style_font(style, name="Microsoft YaHei", size=None, bold=None, color="000000"):
    style.font.name = name
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)


def remove_style_paragraph_borders(style):
    p_pr = style._element.find(qn("w:pPr"))
    if p_pr is None:
        return
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "5")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_width(cell, width_in):
    cell.width = Inches(width_in)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_in * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def style_cell_text(cell, *, header=False, center=False, small=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            set_run_font(
                run,
                size=9 if small else 9.5,
                bold=header,
                color="FFFFFF" if header else "000000",
            )


def add_table(doc, headers, rows, widths=None, center_cols=(), small=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    prevent_row_split(header_row)
    for index, value in enumerate(headers):
        cell = header_row.cells[index]
        cell.text = str(value)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        if widths:
            set_cell_width(cell, widths[index])
        style_cell_text(cell, header=True, center=True, small=small)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = "" if value is None else str(value)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[index])
            style_cell_text(cell, center=index in center_cols, small=small)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    return table


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=TEXT_GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    number = OxmlElement("w:t")
    number.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, number, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=TEXT_GRAY)


def configure_document(doc, short_title):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    section.different_first_page_header_footer = True

    styles = doc.styles
    set_style_font(styles["Normal"], size=10.8)
    styles["Normal"].paragraph_format.line_spacing = 1.45
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.widow_control = True

    set_style_font(styles["Title"], size=25, bold=True)
    remove_style_paragraph_borders(styles["Title"])
    styles["Title"].paragraph_format.space_before = Pt(0)
    styles["Title"].paragraph_format.space_after = Pt(16)
    styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    set_style_font(styles["Subtitle"], size=12, bold=False, color="000000")
    styles["Subtitle"].font.italic = False
    remove_style_paragraph_borders(styles["Subtitle"])
    styles["Subtitle"].paragraph_format.space_after = Pt(8)

    for name, size, before, after in (
        ("Heading 1", 16, 18, 8),
        ("Heading 2", 13.5, 14, 6),
        ("Heading 3", 11.5, 10, 4),
    ):
        set_style_font(styles[name], size=size, bold=True)
        styles[name].paragraph_format.space_before = Pt(before)
        styles[name].paragraph_format.space_after = Pt(after)
        styles[name].paragraph_format.keep_with_next = True

    set_style_font(styles["List Bullet"], size=10.8)
    set_style_font(styles["List Number"], size=10.8)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.86))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = table.rows[0].cells
    set_cell_width(left, 5.4)
    set_cell_width(right, 1.46)
    left_p = left.paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    left_r = left_p.add_run(short_title)
    set_run_font(left_r, size=8.5, color=TEXT_GRAY)
    right_p = right.paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(right_p)
    for cell in (left, right):
        set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    doc.core_properties.title = short_title
    doc.core_properties.subject = f"{SOFTWARE_NAME}配套交付文档"
    doc.core_properties.author = "项目组"
    doc.core_properties.keywords = "端到端跟踪 Transformer 原型软件 配套文档"


def add_cover(
    doc,
    doc_name,
    doc_code,
    *,
    doc_version="V1.0",
    doc_status="初稿",
    include_version_metadata=True,
):
    doc.add_paragraph("软件配套交付文档", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(52)
    title = doc.add_paragraph(style="Title")
    title.add_run(f"{SOFTWARE_NAME}\n{doc_name}")
    for run in title.runs:
        set_run_font(run, size=25, bold=True)
    subtitle = doc.add_paragraph("面向模型部署 批量测试 跟踪验证 性能评估和报告归档", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        set_run_font(run, size=12, color=TEXT_GRAY)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(48)
    metadata = [("文档编号", doc_code)]
    if include_version_metadata:
        metadata.extend(
            [
                ("文档版本", doc_version),
                ("软件版本", SOFTWARE_VERSION),
            ]
        )
    metadata.extend(
        [
            ("文档状态", doc_status),
            ("编制日期", DOC_DATE),
        ]
    )
    table = doc.add_table(rows=len(metadata), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, (key, value) in enumerate(metadata):
        left, right = table.rows[index].cells
        left.text = key
        right.text = value
        set_cell_width(left, 1.5)
        set_cell_width(right, 4.5)
        for cell in (left, right):
            set_cell_margins(cell, top=100, bottom=100, start=120, end=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, size=10.5, bold=(cell is left))
    doc.add_page_break()


def add_document_control(
    doc,
    scope_text,
    contents,
    *,
    version_rows=None,
    basis_items=None,
    include_version_history=True,
):
    doc.add_heading("文档说明", level=1)
    doc.add_paragraph(scope_text)
    if include_version_history:
        doc.add_heading("版本记录", level=2)
        add_table(
            doc,
            ["版本", "日期", "变更说明"],
            version_rows
            or [["V1.0", DOC_DATE, "依据现有前后端代码 运行配置和技术指标形成配套文档初稿"]],
            widths=[0.9, 1.35, 4.55],
            center_cols=(0, 1),
        )
    doc.add_heading("编制依据", level=2)
    for text in basis_items or (
        "前端源代码 tracking_system frontend",
        "后端源代码 tracking_system backend",
        "核心模型及数据处理代码 src",
        "项目说明 README 和 tracking_system README",
        "技术指标文件 docs 技术指标 txt",
    ):
        doc.add_paragraph(text, style="List Bullet")
    doc.add_page_break()
    doc.add_heading("目录", level=1)
    for item in contents:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18 if item[0] == 1 else 0.48)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item[1])
        set_run_font(run, size=11 if item[0] == 1 else 10.5, bold=True)
    doc.add_page_break()


def add_paragraph(doc, text, lead=None):
    p = doc.add_paragraph()
    if lead:
        r = p.add_run(lead)
        set_run_font(r, bold=True)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered_steps(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.26)
        p.paragraph_format.first_line_indent = Inches(-0.26)
        number = p.add_run(f"{index}. ")
        set_run_font(number, bold=True)
        body = p.add_run(item)
        set_run_font(body)


def add_command(doc, command):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.34)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(command)
    set_run_font(r, name="Consolas", size=9.3)


def make_software_architecture_image():
    from PIL import Image, ImageDraw, ImageFont

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / "software_architecture.png"
    width, height = 2200, 1320
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path("C:/Windows/Fonts/msyh.ttc")
    bold_path = Path("C:/Windows/Fonts/msyhbd.ttc")
    title_font = ImageFont.truetype(str(bold_path if bold_path.exists() else font_path), 56)
    box_font = ImageFont.truetype(str(bold_path if bold_path.exists() else font_path), 34)
    small_font = ImageFont.truetype(str(font_path), 28)

    draw.text((80, 42), "端到端协同跟踪原型软件总体架构", font=title_font, fill="#000000")

    def rounded_box(x1, y1, x2, y2, title, lines, fill):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=26, fill=fill, outline="#7F8C8D", width=4)
        tw = draw.textbbox((0, 0), title, font=box_font)[2]
        draw.text(((x1 + x2 - tw) / 2, y1 + 28), title, font=box_font, fill="#000000")
        y = y1 + 92
        for line in lines:
            tw = draw.textbbox((0, 0), line, font=small_font)[2]
            draw.text(((x1 + x2 - tw) / 2, y), line, font=small_font, fill="#202020")
            y += 44

    def arrow(x1, y1, x2, y2, label):
        draw.line((x1, y1, x2, y2), fill="#1F4E78", width=8)
        angle = 18
        if x2 > x1:
            draw.polygon([(x2, y2), (x2 - 24, y2 - angle), (x2 - 24, y2 + angle)], fill="#1F4E78")
        else:
            draw.polygon([(x2, y2), (x2 + 24, y2 - angle), (x2 + 24, y2 + angle)], fill="#1F4E78")
        if label:
            tw = draw.textbbox((0, 0), label, font=small_font)[2]
            draw.text(((x1 + x2 - tw) / 2, y1 - 44), label, font=small_font, fill="#1F4E78")

    def down_arrow(x, y1, y2, label):
        draw.line((x, y1, x, y2), fill="#1F4E78", width=8)
        draw.polygon([(x, y2), (x - 18, y2 - 26), (x + 18, y2 - 26)], fill="#1F4E78")
        tw = draw.textbbox((0, 0), label, font=small_font)[2]
        draw.text((x - tw / 2, y1 + 22), label, font=small_font, fill="#1F4E78")

    rounded_box(80, 190, 570, 510, "浏览器前端", ["Vue 3 和 TypeScript", "六个业务页面", "ECharts 三维与统计图表"], "#EAF2F8")
    rounded_box(850, 170, 1380, 530, "FastAPI 服务层", ["REST 接口", "WebSocket 回放", "请求追踪和错误处理"], "#DDEBF7")
    rounded_box(1630, 170, 2120, 530, "业务服务", ["ModelManager", "BatchTestManager", "EvaluationService"], "#E8F8F5")
    rounded_box(180, 760, 710, 1130, "模型推理层", ["CoopTrackerModel", "Learned K 和 Transformer", "时序建模与协同 GAT"], "#FCF3CF")
    rounded_box(850, 760, 1380, 1130, "数据资源", ["注册 PKL 数据集", "V5 V6 V7 V8 验证集", "场景和目标索引"], "#FDEBD0")
    rounded_box(1520, 760, 2050, 1130, "持久化制品", ["批测任务 JSON", "逐航迹 CSV", "评估证据和对比报告"], "#F4ECF7")

    arrow(570, 350, 850, 350, "HTTP 和 WebSocket")
    arrow(1380, 350, 1630, 350, "服务调用")
    draw.line((1875, 530, 1875, 650), fill="#1F4E78", width=8)
    draw.line((445, 650, 1875, 650), fill="#1F4E78", width=8)
    down_arrow(445, 650, 760, "模型加载与推理")
    down_arrow(1115, 650, 760, "场景与批测数据")
    down_arrow(1785, 650, 760, "报告与证据")
    image.save(path, quality=95)
    return path


def build_requirements_doc():
    doc = Document()
    title = f"{SOFTWARE_NAME} 需求规格说明书"
    configure_document(doc, title)
    add_cover(doc, "需求规格说明书", "TS DOC SRS 001")
    contents = [
        (1, "1 文档范围"), (1, "2 产品概述"), (1, "3 业务流程"),
        (1, "4 功能需求"), (1, "5 数据需求"), (1, "6 外部接口需求"),
        (1, "7 非功能需求"), (1, "8 技术指标与验收要求"),
        (1, "9 需求追踪关系"), (1, "10 验收建议"), (1, "附录 术语说明"),
    ]
    add_document_control(
        doc,
        "本说明书定义端到端协同跟踪原型软件的业务范围 功能需求 数据要求 外部接口和验收准则 供需求确认 软件验收 测试设计和后续维护使用",
        contents,
    )

    doc.add_heading("1 文档范围", level=1)
    doc.add_heading("1.1 编写目的", level=2)
    doc.add_paragraph("本说明书将合同及任务书中的原型验证软件要求落实为可检查的软件需求。需求覆盖模型部署 批量测试 单场景跟踪验证 多维性能评估和报告归档，并给出与现有界面 接口和代码模块的对应关系。")
    doc.add_heading("1.2 软件边界", level=2)
    doc.add_paragraph("软件以本地原型验证平台方式运行。浏览器负责交互和可视化，FastAPI 服务负责模型加载 推理 批测 评估和报告输出，核心模型由 CoopTrackerModel 提供。训练 数据生成和数据注册脚本属于支撑工具，不纳入日常前端操作范围。")
    doc.add_heading("1.3 术语约定", level=2)
    add_table(doc, ["术语", "含义"], [
        ["场景", "一个注册后的 PKL 数据文件及其包含的一个或多个目标航迹"],
        ["批测任务", "针对指定模型 数据集和计算设备提交的后台评估任务"],
        ["测试证据", "与具体批测配置和结果签名严格匹配的关联正确率或 IMM JPDA 基线比较结果"],
        ["完整", "批测范围未设置最大场景数并覆盖所选数据集全部可用场景"],
    ], widths=[1.35, 5.45])

    doc.add_heading("2 产品概述", level=1)
    doc.add_heading("2.1 产品定位", level=2)
    doc.add_paragraph("本软件用于验证多无人机、多源异构传感器条件下端到端协同跟踪模型的部署能力、推理效果、性能指标和报告可追溯性。软件面向研发人员、测试人员和验收人员，不直接承担飞行控制、武器控制或在线任务规划。")
    doc.add_heading("2.2 用户角色", level=2)
    add_table(doc, ["角色", "主要任务", "典型权限"], [
        ["系统维护人员", "启动服务 配置模型和数据目录 检查运行环境", "本机文件和进程管理"],
        ["模型研发人员", "切换模型 进行场景回放 发起抽样或完整批测", "全部前端页面和接口"],
        ["测试验收人员", "核对四项指标 生成测试证据 导出报告", "评估和报告相关功能"],
    ], widths=[1.25, 3.35, 2.2])
    doc.add_heading("2.3 运行环境", level=2)
    add_bullets(doc, [
        "后端运行于 Windows 本机 Python 环境，依赖 FastAPI Uvicorn NumPy PyTorch h5py 和 SciPy。",
        "前端运行于现代桌面浏览器，由 Vue 3 TypeScript Vite ECharts 和 ECharts GL 构建。",
        "默认后端地址为 http://127.0.0.1:8000，前端开发服务默认为 http://127.0.0.1:5173。",
        "推理设备支持 CPU 和可用的 CUDA 设备。自动设备模式优先选择 cuda:0，不可用时使用 CPU。",
    ])
    doc.add_heading("2.4 约束与依赖", level=2)
    add_bullets(doc, [
        "模型文件和注册数据集必须位于配置目录并可被后端读取。",
        "模型注册表当前固定包含 V5 V6 和 V7 三个最佳权重。",
        "实时跟踪验证使用当前激活模型和当前激活模型的数据目录。",
        "批测任务一次执行一个 GPU 密集型任务，其余任务进入队列。",
        "关联正确率和相对 IMM JPDA 降幅在缺少有效测试证据时必须显示为数据不足，不得判定为通过。",
    ])

    doc.add_heading("3 业务流程", level=1)
    add_numbered_steps(doc, [
        "维护人员启动后端，确认健康检查返回 ready，并启动前端开发服务。",
        "用户在系统总览中检查服务 模型 设备 数据集和任务书指标。",
        "用户在模型管理中选择 V5 V6 或 V7 模型及计算设备并完成部署。",
        "用户可先在跟踪验证页面选择机动类型 场景和目标，检查三维航迹与实时误差。",
        "用户在批量测试页面选择模型 数据集 设备和测试范围，提交后台任务并查看进度。",
        "任务完成后，用户在性能评估页面核对位置误差 推理时延 关联正确率和基线降幅，并生成缺失证据。",
        "用户在报告中心检索任务，下载原始 JSON 综合评估 JSON 逐航迹 CSV 或多任务对比 CSV。",
    ])

    doc.add_heading("4 功能需求", level=1)
    functional_rows = [
        ["FR 001", "系统总览", "显示服务状态 版本 运行时长 当前模型 推理设备 CUDA 状态 批测队列和四项技术指标", "页面可刷新且后台每 30 秒自动同步"],
        ["FR 002", "模型发现", "列出已注册的 V5 V6 V7 权重及版本 训练集 验证集 文件大小 修改时间和历史指标摘要", "缺失权重不得显示为可部署模型"],
        ["FR 003", "模型切换", "允许用户选择模型和 CPU 或 CUDA 设备，经二次确认后加载并切换", "加载失败时当前模型保持不变 已开始的回放不受影响"],
        ["FR 004", "场景目录", "列出当前数据目录中按文件名排序的 PKL 场景，并支持 ID 或文件名解析", "非法路径 越界场景和越界目标必须拒绝"],
        ["FR 005", "单场景推理", "对选定场景和目标执行完整序列推理并通过 WebSocket 按时间轴发送结果", "输出真实位置 预测位置 当前误差 运行 RMSE 和推理时延"],
        ["FR 006", "回放控制", "支持播放 暂停 时间轴跳转 重新推理以及俯视 侧视和自由相机", "跳转后运行 RMSE 从新位置重新累计"],
        ["FR 007", "当前序列导出", "将已接收的真实坐标 预测坐标 误差和时间信息导出为 CSV", "导出数据保留绝对坐标"],
        ["FR 008", "数据集目录", "展示两个验收想定和 V5 V6 V7 V8 技术验证集的可用性 场景数 类型 模态和干扰说明", "不可用数据集不得提交批测"],
        ["FR 009", "批测配置", "允许配置任务名称 模型 数据集 设备 预热帧 批大小 数据线程以及抽样或完整范围", "预热帧 0 到 99 批大小 1 到 256 数据线程 0 到 8"],
        ["FR 010", "任务队列", "批测提交后立即返回任务 ID，后台显示 queued running completed failed cancelled 状态和进度", "同一时刻最多运行一个批测任务"],
        ["FR 011", "任务取消和历史", "允许取消排队中或运行中的任务，并在后端重启后恢复可读取的历史结果", "重启前未结束任务转为失败并记录中断原因"],
        ["FR 012", "批测指标", "计算位置 RMSE MAE 中位数 P95 最大误差 平均航迹 RMSE 吞吐率和多类时延", "忽略无效目标帧和指定预热帧"],
        ["FR 013", "多维判定", "按四项技术指标生成 passed failed unavailable 三态结果及总体 passed failed incomplete 状态", "任何失败项使总体失败 仅缺证据时总体不完整"],
        ["FR 014", "测试证据", "支持为已完成任务生成关联正确率和相对 IMM JPDA 降幅证据", "证据必须通过任务 ID 指标 ID 任务签名和值有效性校验"],
        ["FR 015", "跨数据集评估", "展示 V5 V6 V7 模型在 V5 V6 V7 V8 数据集上的 3 乘 4 RMSE 矩阵", "历史矩阵仅作为参考结果"],
        ["FR 016", "报告中心", "支持按状态 模型 名称 ID 和数据集筛选报告并查看详情", "最多列出 200 条持久化任务"],
        ["FR 017", "制品导出", "支持原始批测 JSON 综合评估 JSON 逐航迹 CSV 和最多 50 个任务的对比 CSV", "仅在相应制品可用时允许下载"],
        ["FR 018", "运行配置查看", "侧栏设置显示 API 服务版本 当前模型 设备 PyTorch CUDA 和报告目录，并提供后端启动命令复制", "连接失败时给出重新连接入口"],
    ]
    add_table(doc, ["编号", "需求名称", "需求描述", "验收要点"], functional_rows, widths=[0.75, 1.25, 3.05, 1.75], center_cols=(0,), small=True)

    doc.add_heading("5 数据需求", level=1)
    doc.add_heading("5.1 输入数据", level=2)
    add_table(doc, ["数据类别", "主要字段", "用途"], [
        ["雷达", "距离 方位角 俯仰角及有效掩码", "主动主流观测和几何位置反演"],
        ["光电", "方位角 俯仰角及有效掩码", "被动辅助流特征"],
        ["电子侦察", "方位角 俯仰角 模式及有效掩码", "被动辅助流特征和工作模式信息"],
        ["平台状态", "无人机位置 姿态旋转矩阵 IMU 和时间间隔", "坐标转换 运动预测和姿态调制"],
        ["协同数据链", "快链 慢链 更新标志和数据龄期", "跨平台协同观测"],
        ["真值与掩码", "绝对状态 label_abs 和 target_valid_mask", "误差计算和有效帧筛选"],
    ], widths=[1.25, 2.75, 2.8])
    doc.add_heading("5.2 输出数据", level=2)
    add_bullets(doc, [
        "模型输出绝对坐标状态 px vx py vy pz vz。",
        "实时回放输出时间 真值位置 预测位置 单帧位置误差 运行 RMSE 和推理耗时。",
        "批测输出总体指标 时延指标 关联摘要 场景清单和逐航迹指标。",
        "评估输出四项判定 证据状态 航迹误差分布和按机动类型汇总。",
    ])
    doc.add_heading("5.3 数据完整性", level=2)
    doc.add_paragraph("系统应拒绝空场景 无有效真值帧 非有限模型输出 目标索引越界和无效测试证据。批测报告采用临时文件写入后原子替换，减少中途写入导致的报告损坏。")

    doc.add_heading("6 外部接口需求", level=1)
    endpoint_rows = [
        ["GET", "/api/health", "健康状态 当前模型和批测摘要"],
        ["GET", "/api/system", "Python PyTorch CUDA 设备和运行状态"],
        ["GET", "/api/config/acceptance-criteria", "四项技术指标"],
        ["GET", "/api/devices", "可用 CPU 和 CUDA 设备"],
        ["GET POST", "/api/models  /api/models/refresh", "模型目录查询和刷新"],
        ["GET POST", "/api/models/active  /api/models/activate", "当前模型查询和模型切换"],
        ["GET", "/api/scenarios  /api/datasets", "场景和数据集目录"],
        ["GET POST", "/api/batch-tests", "任务列表和任务创建"],
        ["GET POST", "/api/batch-tests/{job_id}  /cancel", "任务详情和取消"],
        ["GET POST", "/api/batch-tests/{job_id}/assessment  /evidence/{metric_id}", "评估和证据生成"],
        ["GET", "/api/evaluations/references  /cross-dataset  /compare", "历史参考 跨集矩阵和任务对比"],
        ["GET", "/api/reports  /api/reports/{job_id}/assessment.json", "报告目录和综合评估下载"],
        ["GET", "/api/batch-tests/{job_id}/report.json  /tracks.csv", "原始报告和逐航迹指标下载"],
        ["GET", "/api/reports/comparison.csv", "多任务对比 CSV 下载"],
        ["WS", "/ws/tracking", "单场景推理结果回放及播放控制"],
    ]
    add_table(doc, ["方法", "路径", "说明"], endpoint_rows, widths=[0.85, 3.45, 2.5], center_cols=(0,), small=True)
    doc.add_heading("6.1 接口通用约定", level=2)
    add_bullets(doc, [
        "REST 接口使用 JSON 交换结构化数据，下载接口按 JSON 或 UTF 8 CSV 返回文件。",
        "请求失败时返回合适的 HTTP 状态码和 detail，未处理异常返回请求 ID 以便定位。",
        "每个 HTTP 响应附带 X Request ID 和 X Process Time Ms 响应头。",
        "WebSocket 控制消息支持 pause play 和带 time 参数的 seek。",
    ])

    doc.add_heading("7 非功能需求", level=1)
    nonfunctional_rows = [
        ["NFR 001", "性能", "合同要求的单帧前向推理时间不超过 100 ms。本软件将其定义为已对准和关联的多源张量输入至跟踪状态输出的模型前向耗时，不含控制律和执行机构；正式报告须记录硬件与软件环境。"],
        ["NFR 002", "可靠性", "模型切换采用候选模型先加载后替换。批测报告原子写入，任务异常必须记录类型和原因。"],
        ["NFR 003", "可追溯性", "每个任务保存模型 数据集 设备 参数 场景清单 时间 指标和报告路径。测试证据绑定任务签名。"],
        ["NFR 004", "可用性", "主要操作在六个页面内完成，状态和错误以中文提示，危险的模型切换操作需二次确认。"],
        ["NFR 005", "兼容性", "前端支持现代浏览器。后端支持 CPU 和 CUDA，并对 CUDA 不可用或索引无效给出明确错误。"],
        ["NFR 006", "可维护性", "前后端按视图 接口 业务服务 运行时和算法模块分层，模型 数据集和指标目录采用集中注册表。"],
        ["NFR 007", "安全边界", "服务默认仅监听 127.0.0.1，CORS 仅允许本机 5173 开发地址，场景解析禁止目录穿越。"],
    ]
    add_table(doc, ["编号", "类别", "要求"], nonfunctional_rows, widths=[0.9, 1.15, 4.75], center_cols=(0, 1))

    doc.add_heading("8 技术指标与验收要求", level=1)
    add_table(doc, ["指标", "判定条件", "软件取值来源", "缺失处理"], [
        ["跟踪稳定阶段位置 RMSE", "不超过 8 m", "排除每条航迹预热帧后的 position_rmse_m", "无有效点时任务失败"],
        ["目标正确关联率", "不低于 85%", "独立关联测试证据或注册数据 AssociationStats", "标记 unavailable"],
        ["相对 IMM JPDA 降幅", "不低于 15%", "同数据集 同场景范围 同序列长度和同预热设置的基线证据", "标记 unavailable"],
        ["单帧跟踪结果生成时间", "不超过 100 ms", "预热后模型前向同步实测并记录测试平台", "回退到历史有效点折算值"],
    ], widths=[1.45, 1.35, 2.85, 1.15], center_cols=(1,))
    doc.add_paragraph("完整算法验收应使用完整数据集任务。单场景或少量抽样任务仅用于检查软件链路和操作流程，不替代完整数据集上的合同指标结论。")

    doc.add_heading("9 需求追踪关系", level=1)
    trace_rows = [
        ["原型软件具备模型部署", "FR 002 FR 003", "模型管理页面  ModelManager"],
        ["原型软件具备批量测试", "FR 008 至 FR 012", "批量测试页面  BatchTestManager"],
        ["原型软件具备多维性能评估", "FR 013 至 FR 015", "性能评估页面  EvaluationService"],
        ["稳定阶段位置 RMSE 不超过 8 m", "FR 012 FR 013", "排除预热帧后的 position_rmse_m 判定"],
        ["关联正确率不低于 85%", "FR 013 FR 014", "association_accuracy 证据"],
        ["相对基线降幅不低于 15%", "FR 013 FR 014", "rmse_improvement_pct 证据"],
        ["单帧模型前向时间不超过 100 ms", "FR 012 FR 013", "measured_frame_ms 及环境记录判定"],
        ["不少于三种传感器模态", "FR 008  第 5 章", "雷达 光电 电子侦察并含协同数据链"],
        ["不少于两个典型想定和三种干扰", "FR 008", "协同编队想定 高机动强干扰想定及三类干扰目录"],
    ]
    add_table(doc, ["合同及任务书要求", "软件需求", "实现对应"], trace_rows, widths=[2.45, 1.55, 2.8])

    doc.add_heading("10 验收建议", level=1)
    add_numbered_steps(doc, [
        "执行前端生产构建，确认 TypeScript 类型检查和 Vite 构建成功。",
        "执行后端单元测试，确认场景分层抽样 评估判定 证据签名和报告目录功能通过。",
        "启动后端并运行 tracking_system backend acceptance，完成健康 模型切换 WebSocket 批测 评估和报告下载黑盒检查。",
        "以完整数据集提交正式批测，并生成关联正确率和 IMM JPDA 基线两份独立证据。",
        "从报告中心导出原始报告 综合评估报告 逐航迹数据和横向对比表，核对任务配置和结论。",
    ])
    doc.add_paragraph("当前代码核对结果显示，后端 10 项单元测试通过，前端生产构建通过。该结果证明现有测试范围内的软件模块可构建且核心评估逻辑工作正常，正式验收仍应按上述完整流程执行。")

    doc.add_heading("附录 术语说明", level=1)
    add_table(doc, ["缩写", "说明"], [
        ["RMSE", "均方根误差"], ["MAE", "平均绝对误差"], ["P95", "误差或时延的第 95 百分位"],
        ["IMM JPDA", "交互式多模型联合概率数据关联基线算法"], ["EO", "光电传感器"],
        ["ESM", "电子支援或电子侦察传感器"], ["GAT", "图注意力网络"], ["PKL", "Python Pickle 序列化数据文件"],
    ], widths=[1.4, 5.4])

    path = OUTPUT_DIR / "端到端协同跟踪原型软件_需求规格说明书_V1.0.docx"
    doc.save(path)
    return path


def build_design_doc(software_arch_path):
    doc = Document()
    title = f"{SOFTWARE_NAME} 设计说明书"
    configure_document(doc, title)
    add_cover(
        doc,
        "设计说明书",
        "TS DOC DDS 001",
        doc_status="交付稿",
        include_version_metadata=False,
    )
    contents = [
        (1, "1 前言"), (2, "1.1 编写目的"), (2, "1.2 项目背景"),
        (2, "1.3 设计依据"), (2, "1.4 适用范围与读者"), (2, "1.5 名词术语与缩略语"),
        (1, "2 系统总体设计"), (2, "2.1 系统功能概述"), (2, "2.2 总体架构"),
        (2, "2.3 系统运行流程总览"), (2, "2.4 主要模块划分与关系"),
        (2, "2.5 数据流"), (2, "2.6 控制流"), (2, "2.7 设计目标与原则"),
        (1, "3 公共数据结构与坐标约定"), (2, "3.1 张量维度与基础参数"),
        (2, "3.2 坐标系与单位"), (2, "3.3 目标状态与参考坐标"),
        (2, "3.4 传感器输入张量"), (2, "3.5 平台链路与时间张量"),
        (2, "3.6 掩码与缺失数据"), (2, "3.7 跟踪序列结构"),
        (2, "3.8 传输与持久化结构"), (2, "3.9 数据有效性约束"),
        (1, "4 技术选型"),
        (1, "5 模块详细设计"), (2, "5.1 模块划分与依赖约束"),
        (2, "5.2 前端应用模块"), (2, "5.3 接口适配与传输模块"),
        (2, "5.4 模型管理模块"), (2, "5.5 单场景推理运行时"),
        (2, "5.6 批量测试管理模块"), (2, "5.7 性能评估与报告模块"),
        (2, "5.8 场景目录与数据访问模块"), (2, "5.9 模块协作流程"),
        (1, "6 前端界面详细设计"), (2, "6.1 应用框架与导航"),
        (2, "6.2 数据访问与公共状态"), (2, "6.3 系统总览界面"),
        (2, "6.4 模型管理界面"), (2, "6.5 跟踪验证界面"),
        (2, "6.6 批量测试界面"), (2, "6.7 性能评估界面"),
        (2, "6.8 报告中心界面"), (2, "6.9 响应式、可用性与异常反馈"),
        (1, "7 后端接口详细设计"), (2, "7.1 应用初始化与公共处理"),
        (2, "7.2 接口分层与资源模型"), (2, "7.3 系统与模型接口"),
        (2, "7.4 场景与数据集接口"), (2, "7.5 批量测试接口"),
        (2, "7.6 评估与报告接口"), (2, "7.7 WebSocket 跟踪接口"),
        (2, "7.8 请求响应约定"), (2, "7.9 错误处理与安全边界"),
        (1, "8 模型推理详细设计"), (2, "8.1 模型输入与输出"),
        (2, "8.2 几何预处理与观测状态构造"), (2, "8.3 Learned-K 状态前端"),
        (2, "8.4 多模态特征编码"), (2, "8.5 主被动双流融合"),
        (2, "8.6 时序建模与协同图交互"), (2, "8.7 残差回归与绝对坐标恢复"),
        (2, "8.8 运行时推理流程与缓存"), (2, "8.9 数值稳定性与降级策略"),
        (1, "9 批量测试详细设计"), (2, "9.1 功能边界与任务配置"),
        (2, "9.2 任务对象与状态机"), (2, "9.3 创建 排队 执行与取消"),
        (2, "9.4 场景发现与分层抽样"), (2, "9.5 数据装载与批量推理"),
        (2, "9.6 精度指标计算"), (2, "9.7 时延与吞吐测量"),
        (2, "9.8 关联统计"), (2, "9.9 持久化与重启恢复"),
        (2, "9.10 失败处理与可复现信息"),
        (1, "10 性能评估与报告详细设计"), (2, "10.1 评估输入与总体状态"),
        (2, "10.2 四项技术指标判定"), (2, "10.3 独立证据生成"),
        (2, "10.4 证据完整性与任务签名"), (2, "10.5 航迹分布与机动类型分析"),
        (2, "10.6 历史参考与跨数据集结果"), (2, "10.7 多任务对比"),
        (2, "10.8 报告目录与下载制品"), (2, "10.9 审慎判定规则"),
        (1, "11 数据与持久化详细设计"), (2, "11.1 数据分类与目录边界"),
        (2, "11.2 注册 PKL 场景数据"), (2, "11.3 运行时数据对象"),
        (2, "11.4 模型权重与指标摘要"), (2, "11.5 批测报告与验收证据"),
        (2, "11.6 命名 索引与完整性约束"), (2, "11.7 数据生命周期与备份"),
        (1, "12 接口设计"), (2, "12.1 WebSocket 帧结构"),
        (1, "13 并发与状态一致性详细设计"), (2, "13.1 共享状态与同步对象"),
        (2, "13.2 服务单例初始化"), (2, "13.3 模型激活一致性"),
        (2, "13.4 批测任务并发控制"), (2, "13.5 持久化一致性"),
        (2, "13.6 证据生成互斥"), (2, "13.7 WebSocket 会话状态"),
        (2, "13.8 故障隔离与资源释放"),
        (1, "14 异常处理与安全边界详细设计"), (2, "14.1 异常分类与对外映射"),
        (2, "14.2 参数与状态校验"), (2, "14.3 文件路径与目录边界"),
        (2, "14.4 模型与数值安全"), (2, "14.5 证据脚本执行边界"),
        (2, "14.6 日志与错误信息"), (2, "14.7 部署安全边界"),
        (1, "15 配置与部署详细设计"), (2, "15.1 配置来源与优先级"),
        (2, "15.2 运行目录结构"), (2, "15.3 后端启动与健康检查"),
        (2, "15.4 前端构建与启动"), (2, "15.5 端口发现与跨域配置"),
        (2, "15.6 部署前检查"), (2, "15.7 交付环境调整"),
        (1, "16 测试与维护详细设计"), (2, "16.1 测试层级与职责"),
        (2, "16.2 后端单元测试"), (2, "16.3 前端静态检查与构建"),
        (2, "16.4 后端黑盒验收"), (2, "16.5 完整指标验收"),
        (2, "16.6 日常维护"), (2, "16.7 扩展与变更控制"),
        (2, "16.8 交付检查"),
        (1, "附录 关键设计参数"),
    ]
    add_document_control(
        doc,
        "本说明书规定端到端协同跟踪原型软件的系统边界、总体架构、运行流程、公共数据约定、模块职责、接口、持久化、并发控制、异常处理和部署设计，作为软件开发、联调、维护和验收的技术依据。",
        contents,
        basis_items=[
            "项目合同书及任务书中关于原型软件功能和技术指标的要求",
            "端到端协同跟踪原型软件需求规格说明书",
            "技术指标文件 docs 技术指标 txt",
            "前端源代码 tracking_system frontend",
            "后端源代码 tracking_system backend",
            "核心模型和数据处理代码 src",
            "项目说明 README 和 tracking_system README",
        ],
        include_version_history=False,
    )

    doc.add_heading("1 前言", level=1)
    doc.add_heading("1.1 编写目的", level=2)
    doc.add_paragraph(
        "本说明书用于明确端到端协同跟踪原型软件的设计边界、总体结构、运行流程和公共数据约定，使前端实现、后端服务、模型推理、批量测试、性能评估与报告归档采用一致的技术口径。文档既用于开发人员理解模块职责和接口关系，也用于测试、维护及验收人员核查软件实现是否覆盖合同和任务书要求。"
    )
    doc.add_heading("1.2 项目背景", level=2)
    doc.add_paragraph(
        "本软件面向多无人机、多源异构传感器条件下的目标协同跟踪验证。核心 CoopTrackerModel 已完成训练，能够综合雷达、光电、电子侦察、平台运动状态和协同链路信息，输出目标位置与速度状态。为使训练结果能够被部署、演示、批量验证和正式评估，项目建设了浏览器前端、FastAPI 后端、模型运行时、批测管理、证据生成和报告归档组成的本地原型系统。"
    )
    doc.add_paragraph(
        "软件将单场景跟踪回放与完整数据集验收分开处理。单场景回放用于检查模型链路、三维轨迹和误差变化；批量测试用于形成稳定阶段精度、推理时延和逐航迹统计；性能评估进一步结合关联正确率与 IMM-JPDA 同条件基线证据，对合同规定的四项指标给出可追溯判定。"
    )
    doc.add_heading("1.3 设计依据", level=2)
    add_bullets(doc, [
        "项目合同书及任务书中关于原型软件功能、交付物和技术指标的要求。",
        "端到端协同跟踪原型软件需求规格说明书。",
        "docs 技术指标 txt 中关于位置 RMSE、关联正确率、相对 IMM-JPDA 降幅和单帧推理时间的规定。",
        "tracking_system frontend 与 tracking_system backend 的现有实现。",
        "src model.py、src dataset.py 和 src dataset_pkl.py 中的模型与数据接口。",
        "README、tracking_system README 及项目运行环境配置。",
    ])
    doc.add_heading("1.4 适用范围与读者", level=2)
    add_table(doc, ["项目", "说明"], [
        ["软件范围", "浏览器前端、FastAPI 后端、模型部署与推理、单场景回放、批量测试、性能评估和报告归档"],
        ["数据范围", "注册 PKL 场景、模型权重、批测任务 JSON、验收证据 JSON 和导出 CSV"],
        ["范围边界", "不包含模型训练、原始数据生成、飞行控制、武器控制和在线任务规划"],
        ["主要读者", "软件开发人员、模型研发人员、测试人员、系统维护人员和验收人员"],
    ], widths=[1.35, 5.45])
    doc.add_heading("1.5 名词术语与缩略语", level=2)
    add_table(doc, ["缩写", "英文或字段名", "说明"], [
        ["RMSE", "Root Mean Square Error", "均方根误差，本软件主要使用三维位置误差，单位为 m"],
        ["MAE", "Mean Absolute Error", "位置误差绝对值的平均值，单位为 m"],
        ["P95", "95th Percentile", "误差或时延样本的第 95 百分位"],
        ["EO", "Electro Optical", "光电传感器，在软件中提供方位角和俯仰角"],
        ["ESM", "Electronic Support Measures", "电子支援或电子侦察传感器"],
        ["GAT", "Graph Attention Network", "图注意力网络，用于无人机节点之间的协同信息交互"],
        ["IMM-JPDA", "Interacting Multiple Model Joint Probabilistic Data Association", "合同对比使用的交互式多模型联合概率数据关联基线"],
        ["ECEF", "Earth Centered Earth Fixed", "地心地固直角坐标系"],
        ["PKL", "Python Pickle", "注册场景使用的 Python 序列化数据文件"],
        ["T N B C", "Time Agent Batch Channel", "分别表示时间步、无人机数量、批大小和特征通道数"],
    ], widths=[1.1, 2.55, 3.15], small=True)

    doc.add_heading("2 系统总体设计", level=1)
    doc.add_heading("2.1 系统功能概述", level=2)
    doc.add_paragraph(
        "系统以六个前端业务页面组织操作流程，由后端统一维护模型、数据集、任务和报告状态。用户可以在同一软件内完成运行环境检查、模型切换、单场景验证、完整或抽样批测、四项指标判定以及报告下载。"
    )
    add_table(doc, ["功能", "主要输入", "主要输出"], [
        ["系统总览", "健康状态、运行时、活动模型、设备和批测队列", "服务状态、CUDA 状态、技术指标和任务摘要"],
        ["模型管理", "模型编号和计算设备", "模型清单、加载结果和当前活动模型"],
        ["跟踪验证", "场景、目标索引和播放控制", "真值与预测三维轨迹、实时误差和当前序列 CSV"],
        ["批量测试", "模型、数据集、设备、预热帧、批大小和测试范围", "任务进度、总体指标、时延和逐航迹结果"],
        ["性能评估", "已完成批测、关联证据和同条件基线", "四项指标判定、多维统计和任务对比"],
        ["报告中心", "任务筛选条件和任务编号", "原始 JSON、综合评估 JSON、逐航迹 CSV 和对比 CSV"],
    ], widths=[1.25, 2.75, 2.8])

    doc.add_heading("2.2 总体架构", level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(software_arch_path), width=Inches(6.7))
    cap = doc.add_paragraph("图 1 软件总体架构")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        set_run_font(run, size=9.5, color=TEXT_GRAY)
    doc.add_paragraph(
        "浏览器前端通过 REST 接口读取系统状态、模型、数据集、批测和评估信息，通过 WebSocket 接收单场景跟踪帧。FastAPI 服务将请求分派给模型管理、批测管理和评估服务。业务服务复用 TransformerEngine 与 CoopTrackerModel，并将批测和评估制品写入项目 artifacts 目录。系统当前按可信本机原型部署，后端监听 127.0.0.1 的 8000 端口，前端开发服务默认使用 5173 端口。"
    )

    doc.add_heading("2.3 系统运行流程总览", level=2)
    doc.add_heading("2.3.1 离线初始化阶段", level=3)
    add_numbered_steps(doc, [
        "后端启动 FastAPI 应用并建立模型管理、批测管理和评估服务的延迟初始化入口。",
        "ModelManager 扫描 V5、V6 和 V7 注册权重，检查模型文件和对应数据目录。",
        "首次健康检查触发默认 V7 模型加载，并按 auto、CPU 或指定 CUDA 设备创建 TransformerEngine。",
        "前端读取健康状态、运行环境、模型目录、设备目录和四项验收指标，建立页面初始状态。",
    ])
    doc.add_heading("2.3.2 单场景跟踪验证阶段", level=3)
    add_numbered_steps(doc, [
        "用户选择场景和目标索引，前端使用 dataset 与 target 参数建立 WebSocket 连接。",
        "运行时解析注册 PKL 文件，将选定目标转换为固定长度、固定无人机数量的模型输入字典。",
        "缺失的可选输入按统一形状补齐，模型一次完成整段序列推理，运行时仅保留 target_valid_mask 标记的有效帧。",
        "后端按 50 ms 间隔发送真值位置、预测位置、当前误差、运行 RMSE 和推理时延，独立接收暂停、继续和跳转控制。",
        "同一模型引擎内重复请求相同场景和目标时复用 TrackingSequence 缓存，模型切换后创建新引擎并自然失效。",
    ])
    doc.add_heading("2.3.3 批量测试与评估归档阶段", level=3)
    add_numbered_steps(doc, [
        "批测接口校验模型、数据集、设备、预热帧、批大小、工作进程和场景范围，生成 UUID 任务并先以 queued 状态持久化。",
        "单工作线程执行场景分层抽样、数据集索引、模型加载和分批推理，持续更新批次、航迹和百分比进度。",
        "任务完成后写入总体精度、逐航迹指标、推理时延、关联摘要和实际场景清单；失败或取消时写入终止状态及原因。",
        "EvaluationService 对已完成任务生成或读取签名匹配的关联正确率及 IMM-JPDA 基线证据，形成四项指标判定。",
        "报告中心将任务、评估和逐航迹结果导出为 JSON 或 CSV，供验收复核和归档。",
    ])

    doc.add_heading("2.4 主要模块划分与关系", level=2)
    add_table(doc, ["层次", "模块", "关系与职责"], [
        ["表现层", "Vue 前端六个业务页面", "调用 REST 和 WebSocket，展示运行状态、轨迹、任务、评估和报告"],
        ["接口层", "server.py", "提供路由、参数校验、请求追踪、错误映射、下载响应和 WebSocket 回放"],
        ["模型管理", "ModelManager", "发现权重与设备，维护活动模型，完成候选引擎加载和原子切换"],
        ["推理运行时", "TransformerEngine", "解析场景、补齐输入、调用 CoopTrackerModel、过滤有效帧并缓存结果"],
        ["批测管理", "BatchTestManager", "创建任务、分层抽样、单线程调度、指标计算和任务 JSON 持久化"],
        ["评估服务", "EvaluationService", "校验证据签名，执行四项判定，生成多维汇总和报告目录"],
        ["核心算法", "CoopTrackerModel", "完成 Learned-K、双流跨注意力、时序 Transformer、协同 GAT 和残差回归"],
        ["资源与制品", "模型 数据集 artifacts", "保存权重、注册场景、批测任务、证据和导出报告"],
    ], widths=[1.1, 1.9, 3.8])

    doc.add_heading("2.5 数据流", level=2)
    add_numbered_steps(doc, [
        "注册 PKL 场景经 dataset 解析为传感器、平台、链路、时间、掩码和目标真值张量。",
        "TransformerEngine 将单样本增加批维度并移动到目标设备，CoopTrackerModel 输出绝对位置与速度状态。",
        "单场景链路将有效帧组织为 TrackingSequence，并转换为 WebSocket JSON 帧供三维图和误差图使用。",
        "批测链路将多个目标样本组成批次，累计稳定阶段点误差、航迹统计、推理时延和关联摘要。",
        "评估链路读取任务结果与独立证据，输出综合评估 JSON；报告接口再生成逐航迹和横向对比 CSV。",
    ])

    doc.add_heading("2.6 控制流", level=2)
    add_numbered_steps(doc, [
        "模型切换先在活动锁内加载候选模型和候选引擎，只有完整加载成功后才替换活动引用；失败时原模型继续服务。",
        "批测任务按 queued、running、completed、failed 或 cancelled 状态迁移，GPU 密集任务由单工作线程串行执行。",
        "WebSocket 发送循环与控制接收协程相互独立，pause 停止推进，play 恢复推进，seek 修改索引并重新累计运行 RMSE。",
        "验收证据仅针对已完成任务生成，白名单脚本串行执行，任务签名不匹配时旧证据不得参与正式判定。",
        "HTTP 中间件为请求生成或透传 X-Request-ID，业务异常映射为 404、409 或 422，未处理异常统一返回 500。",
    ])

    doc.add_heading("2.7 设计目标与原则", level=2)
    add_bullets(doc, [
        "将训练完成的 CoopTrackerModel 封装为可操作、可回放、可批测和可评估的本地原型系统。",
        "把合同及任务书的四项量化指标映射为统一的测试、证据和报告流程。",
        "保持模型加载、单场景推理、批量推理和报告生成之间的数据口径一致。",
        "在本地单机环境中保证模型切换失败可回退、批测状态可追踪、报告可恢复。",
    ])
    add_table(doc, ["原则", "落实方式"], [
        ["分层", "前端表现、接口适配、业务服务、推理运行时、核心算法和数据制品相互分离"],
        ["单一事实来源", "模型、数据集、技术指标和报告目录均由后端集中维护"],
        ["安全切换", "新模型完整加载成功后才替换当前推理引擎"],
        ["结果可复现", "批测保存场景清单、配置、时间、设备、模型和任务签名"],
        ["证据审慎", "缺失或过期证据显示 unavailable，不使用历史值替代正式判定"],
    ], widths=[1.4, 5.4])

    doc.add_heading("3 公共数据结构与坐标约定", level=1)
    doc.add_paragraph(
        "前端、后端运行时、批测管理和模型代码必须使用本章定义的数据顺序、张量维度、坐标基准和有效性规则。未经同步修改数据解析、模型接口和报告字段，不得改变公共字段含义。"
    )
    doc.add_heading("3.1 张量维度与基础参数", level=2)
    add_table(doc, ["符号或参数", "当前值", "说明"], [
        ["B", "运行时决定", "批大小；单场景推理时 B 等于 1"],
        ["T", "100", "序列时间步数，由 TRACKER_SEQ_LEN 或批测 sequence_length 确定"],
        ["N", "5", "协同无人机数量"],
        ["C", "按字段确定", "传感器或状态特征通道数"],
        ["dt", "由 t_grid 推导，兼容输入为 0.1 s", "相邻时间步间隔；模型内部对非正值按最小正数保护"],
        ["dtype", "float32", "传感器、状态、掩码和时间张量的默认数据类型"],
    ], widths=[1.3, 1.45, 4.05])

    doc.add_heading("3.2 坐标系与单位", level=2)
    add_table(doc, ["对象", "坐标或顺序", "单位与说明"], [
        ["绝对世界坐标", "ECEF x y z", "位置使用 m，速度使用 m/s"],
        ["相对参考坐标", "绝对 ECEF 减 ref_pos", "用于模型内部位置、无人机位置和链路量测，单位 m"],
        ["机体坐标", "前向轴 右向轴 上向轴", "rotmat_b2e 的三列给出机体轴在世界坐标中的方向"],
        ["雷达球坐标", "az el r", "方位角和俯仰角使用 rad，斜距使用 m"],
        ["EO 与 ESM 角度", "az el", "方位角和俯仰角使用 rad"],
        ["时间", "global_time times dt age", "统一使用 s；推理耗时和接口处理耗时使用 ms"],
        ["误差", "三维欧氏位置误差", "单位 m，按状态位置分量 0 2 4 计算"],
    ], widths=[1.35, 2.05, 3.4])
    doc.add_paragraph(
        "雷达机体系量测按 az、el、r 存储。模型先根据斜距和角度构造机体系相对向量，再使用 rotmat_b2e 转换到世界坐标，并叠加无人机相对位置。EO 和 ESM 量测按相同角度约定构造机体系方向向量，再转换为世界坐标方向。"
    )

    doc.add_heading("3.3 目标状态与参考坐标", level=2)
    add_table(doc, ["字段", "形状", "定义"], [
        ["label_abs", "T × 6", "绝对目标状态，顺序为 px vx py vy pz vz"],
        ["label_pos_abs", "T × 3", "绝对目标位置 px py pz"],
        ["ref_pos", "3", "目标首个真值帧的绝对 ECEF 位置"],
        ["label_pos_rel", "T × 3", "label_pos_abs 减 ref_pos"],
        ["label", "T × 6", "兼容字段；注册 PKL 解析时位置分量采用相对参考坐标"],
        ["模型输出", "B × T × 6", "绝对目标状态，顺序与 label_abs 一致"],
    ], widths=[1.45, 1.25, 4.1])
    doc.add_paragraph(
        "模型的 Learned-K 和残差回归在相对参考坐标中工作。输出阶段只对位置分量 px、py、pz 加回 ref_pos，速度分量保持世界坐标速度。运行时和批测均以 label_abs 作为精度计算的真值来源。"
    )

    doc.add_heading("3.4 传感器输入张量", level=2)
    add_table(doc, ["字段", "形状", "分量顺序", "有效性"], [
        ["radar", "T × N × 3", "az el r", "radar_mask 等于 1 时有效"],
        ["eo", "T × N × 2", "az el", "eo_mask 等于 1 时有效"],
        ["esm", "T × N × 2", "az el", "esm_mask 等于 1 时有效；部分注册数据仅写入 az"],
        ["radar_mask", "T × N", "0 或 1", "控制雷达编码和粗观测融合"],
        ["eo_mask", "T × N", "0 或 1", "控制 EO 编码和被动流融合"],
        ["esm_mask", "T × N", "0 或 1", "控制 ESM 编码和被动流融合"],
    ], widths=[1.15, 1.35, 1.45, 2.85], small=True)

    doc.add_heading("3.5 平台链路与时间张量", level=2)
    add_table(doc, ["字段", "形状", "说明"], [
        ["uav_pos", "T × N × 3", "无人机绝对 ECEF 位置减 ref_pos 后的相对位置，单位 m"],
        ["rotmat_b2e", "T × N × 3 × 3", "机体坐标到 ECEF 世界坐标的旋转矩阵"],
        ["imu", "T × N × 6", "roll pitch yaw p q r；角度使用 rad，角速度使用 rad/s"],
        ["euler", "T × N × 3", "辅助姿态字段 roll pitch yaw，当前后端推理不直接传入模型"],
        ["quat", "T × N × 4", "辅助四元数字段 w x y z，当前后端推理不直接传入模型"],
        ["fast_link", "T × N × 3", "快链目标位置量测的相对 ECEF 坐标，单位 m"],
        ["slow_link", "T × N × 3", "慢链目标位置量测的相对 ECEF 坐标，单位 m"],
        ["fast_age slow_age", "T × N", "链路数据龄期，单位 s"],
        ["fast_update slow_update", "T × N", "本时间步是否更新，语义值为 0 或 1"],
        ["esm_mode", "T × N", "ESM 工作模式代码"],
        ["global_time", "T", "数据记录的全局时间；缺失时运行时按 dt 构造时间轴"],
    ], widths=[1.65, 1.55, 3.6], small=True)

    doc.add_heading("3.6 掩码与缺失数据", level=2)
    add_table(doc, ["规则", "处理方式"], [
        ["传感器无效", "量测张量保留形状，掩码置 0；编码器使用可学习 mask token 或零权重屏蔽无效观测"],
        ["目标真值无效", "target_valid_mask 置 0；单场景回放和批测精度统计均排除该帧"],
        ["可选字段缺失", "imu、链路、龄期、更新标志和 esm_mode 按固定形状补零"],
        ["链路量测缺失", "全零链路位置视为无效，模型将对应数据龄期修正为 999 s 并清除更新标志"],
        ["整帧无位置来源", "Learned-K 粗观测位置沿用上一有效帧，首帧保持初始化结果"],
        ["序列不足或超长", "解析器按 T 补齐或裁剪，保持所有字段时间维一致"],
    ], widths=[1.55, 5.25])
    doc.add_paragraph(
        "链路有效性除全零检查外，还结合更新标志和龄期判断。快链以更新标志大于 0.5 或龄期小于 2 s 为有效条件，慢链以更新标志大于 0.5 或龄期小于 5 s 为有效条件，并分别按龄期衰减权重。"
    )

    doc.add_heading("3.7 跟踪序列结构", level=2)
    add_table(doc, ["TrackingSequence 字段", "类型或形状", "说明"], [
        ["scenario_id", "字符串", "从 100 开始的场景编号"],
        ["scenario_name", "字符串", "PKL 文件名去除扩展名后的场景名称"],
        ["target_index", "整数", "场景内目标索引，从 0 开始"],
        ["times", "M", "仅包含有效真值帧的时间轴，单位 s"],
        ["truth_states", "M × 6", "有效帧绝对真值状态"],
        ["predicted_states", "M × 6", "有效帧绝对预测状态"],
        ["position_errors", "M", "逐帧三维位置误差，单位 m"],
        ["sequence_inference_ms", "浮点数", "完整序列模型前向耗时"],
        ["per_frame_inference_ms", "浮点数", "完整序列耗时按帧数折算的单帧值"],
    ], widths=[2.0, 1.55, 3.25], small=True)

    doc.add_heading("3.8 传输与持久化结构", level=2)
    add_table(doc, ["结构", "主要字段", "用途"], [
        ["WebSocket 跟踪帧", "time max_time is_reset scenario target_index truth_pos pred_pos metrics", "前端逐帧轨迹、误差和时延展示"],
        ["批测任务 JSON", "id name status created_at config progress result error", "任务恢复、状态查询和原始结果归档"],
        ["批测 config", "model_id dataset_id device warmup_frames batch_size workers max_scenarios sequence_length", "固定批测范围和复现条件"],
        ["批测 result", "metrics latency association scenario_manifest tracks report_path", "保存总体指标、时延、关联和逐航迹结果"],
        ["综合评估 JSON", "schema_version kind generated_at job assessment provenance", "保存四项判定及阈值和基线来源"],
        ["逐航迹 CSV", "track valid_frames rmse_m mae_m p95_m max_error_m", "离线分析和人工复核"],
    ], widths=[1.45, 3.6, 1.75], small=True)

    doc.add_heading("3.9 数据有效性约束", level=2)
    add_bullets(doc, [
        "场景必须解析出至少一个注册目标，目标索引必须在场景目标数量范围内。",
        "单场景结果必须包含至少一个 target_valid_mask 有效帧，否则推理请求失败。",
        "模型输出位置误差必须全部为有限数；出现 NaN 或 Inf 时中止该场景。",
        "场景参数只能是数字编号或不含目录组件的 PKL 文件名，禁止访问注册数据目录之外的文件。",
        "批测精度只统计有效目标帧，并按任务配置排除每条航迹前 warmup_frames 个初始过渡帧。",
        "正式评估使用的证据必须同时匹配指标编号、任务编号、任务签名和 valid_for_assessment 标志。",
    ])

    doc.add_heading("4 技术选型", level=1)
    add_table(doc, ["层次", "技术", "作用"], [
        ["前端框架", "Vue 3.5 和 TypeScript 5.9", "组件化界面 类型约束和响应式状态"],
        ["构建工具", "Vite 8", "开发服务和生产构建"],
        ["可视化", "ECharts 5.5 和 ECharts GL 2.0", "三维航迹 误差曲线 分布和矩阵展示"],
        ["后端框架", "FastAPI 和 Uvicorn", "REST WebSocket 生命周期和接口文档"],
        ["推理计算", "PyTorch NumPy", "模型加载 张量计算和指标统计"],
        ["数据读取", "Pickle h5py SciPy", "注册 PKL 和相关实验数据支持"],
        ["并发", "asyncio threading ThreadPoolExecutor", "异步接口 原子切换和单任务后台批测"],
        ["持久化", "JSON CSV 和文件系统", "任务恢复 报告下载和证据归档"],
    ], widths=[1.15, 2.4, 3.25])

    doc.add_heading("5 模块详细设计", level=1)
    doc.add_paragraph(
        "本章在总体架构基础上进一步规定各软件模块的边界、入口、输入输出、内部状态和异常行为。模块之间以明确定义的数据对象或接口调用协作，不跨层直接修改其他模块的内部状态。"
    )

    doc.add_heading("5.1 模块划分与依赖约束", level=2)
    add_table(doc, ["模块", "实现位置", "主要入口", "输出或持久状态"], [
        ["前端应用", "frontend src App.vue、views", "页面导航、用户操作和生命周期钩子", "页面状态、图表、提示和下载动作"],
        ["前端数据访问", "frontend src api.ts", "request、loadDashboard 及业务 API 函数", "类型化响应、WebSocket 地址和报告 URL"],
        ["接口适配", "backend server.py", "REST 路由与 ws tracking", "JSON、CSV、文件响应和跟踪帧"],
        ["模型管理", "backend model_manager.py", "refresh、activate、get_engine", "模型目录、活动引擎和设备信息"],
        ["单场景运行时", "backend runtime.py", "infer_scenario", "TrackingSequence 及内存缓存"],
        ["批测管理", "backend batch_testing.py", "create_job、cancel_job、get_job", "任务状态、批测结果和任务 JSON"],
        ["评估与报告", "backend evaluation.py", "assess_job、generate_evidence、report_catalog", "指标判定、证据和报告索引"],
        ["场景与数据访问", "backend scenario_catalog.py、src dataset.py", "数据集枚举、分层选样、PKL 解析", "目标样本、场景清单和类别统计"],
        ["核心算法", "src model.py", "CoopTrackerModel.forward", "B × T × 6 绝对目标状态"],
    ], widths=[1.15, 2.15, 1.8, 1.7], small=True)
    add_bullets(doc, [
        "表现层仅通过 api.ts 访问后端，不直接读取模型权重、PKL 数据或 artifacts 目录。",
        "server.py 负责协议适配和错误码映射；模型加载、批测执行与指标判定分别交由业务模块完成。",
        "ModelManager 是在线推理引擎的唯一持有者，单场景运行时和批测管理通过受控入口取得模型或模型路径。",
        "核心算法仅处理张量，不依赖 HTTP、界面状态或文件报告格式，便于训练、推理和跨语言导出共用。",
    ])

    doc.add_heading("5.2 前端应用模块", level=2)
    doc.add_paragraph(
        "App.vue 作为应用壳维护 currentPage、dashboard、loading、errorMessage、lastUpdated、移动导航、设置抽屉和全局提示。总览与模型管理页面随首包加载，其余四个业务页面采用异步组件，进入页面时再加载相应代码。页面标识写入 URL Hash，使刷新浏览器后仍能恢复当前工作区。"
    )
    add_table(doc, ["事件", "处理", "状态变化"], [
        ["应用挂载", "解析 URL Hash，调用 loadDashboard，并注册 Hash、键盘和定时刷新事件", "默认进入总览；成功记录更新时间，失败保留错误信息"],
        ["导航切换", "校验页面是否可用后修改 currentPage 并替换 Hash", "关闭移动导航，不触发整页刷新"],
        ["全局刷新", "先读取 health，再并行读取 system、models、devices 和 criteria", "非静默刷新显示加载态；成功清除连接错误"],
        ["后台轮询", "每 30 s 静默刷新总览数据", "保持当前页面和用户输入不变"],
        ["Escape 键", "关闭设置抽屉和移动端导航", "避免遮罩层滞留"],
        ["组件卸载", "清理刷新与提示定时器并移除事件监听", "不保留悬空回调"],
    ], widths=[1.25, 3.6, 1.95])

    doc.add_heading("5.3 接口适配与传输模块", level=2)
    doc.add_paragraph(
        "前端 api.ts 定义服务健康、模型、场景、批测任务、评估和报告的 TypeScript 接口。request 函数统一拼接 API 基址、设置 JSON 请求头并解析错误响应中的 detail；业务页面只调用语义化函数。后端 server.py 以 Pydantic 模型校验写请求参数，并将文件下载和 WebSocket 流与普通 JSON 查询分离。"
    )
    add_table(doc, ["通道", "适用数据", "设计约束"], [
        ["REST JSON", "健康状态、模型、数据集、任务、评估和报告目录", "请求失败抛出包含后端 detail 的异常；列表接口返回 items"],
        ["WebSocket JSON", "单场景逐帧真值、预测、误差、RMSE 和时延", "连接参数为 dataset 与 target；控制消息仅接受 pause、play、seek"],
        ["文件响应", "原始任务 JSON 与模型相关静态报告", "下载前确认文件存在并使用安全文件名"],
        ["流式 CSV", "逐航迹结果和多任务对比结果", "仅允许已完成且具备结果的任务导出；UTF-8 文本输出"],
    ], widths=[1.35, 2.65, 2.8])

    doc.add_heading("5.4 模型管理模块", level=2)
    doc.add_paragraph(
        "ModelManager 使用固定目录登记 V5、V6 和 V7 三个模型版本。refresh 仅登记实际存在的权重文件；list_models 同步读取权重大小、修改时间和同目录 metrics_summary.json。available_devices 始终返回 CPU，并按 PyTorch 检测结果补充可用 CUDA 设备及显存和计算能力。"
    )
    add_numbered_steps(doc, [
        "根据 model_id 在登记表中解析权重路径，根据 device 的 auto、cpu、cuda 或 cuda:n 解析计算设备。",
        "在 activation_lock 内构造候选 TransformerEngine，完成路径检查、模型实例化、权重严格加载和 eval 模式设置。",
        "只有候选引擎完整加载成功后，才在 state_lock 内原子替换活动引擎、模型编号、激活时间和加载耗时。",
        "模型和设备均未变化时直接返回 changed 为 false；加载异常时保持原活动引擎不变。",
    ])
    doc.add_paragraph(
        "该设计保证模型切换失败不会破坏正在使用的引擎。切换成功后新建引擎自身携带空的单场景缓存，因此旧模型的预测序列不会被新模型复用。"
    )

    doc.add_heading("5.5 单场景推理运行时", level=2)
    add_table(doc, ["处理阶段", "输入", "处理结果"], [
        ["场景解析", "数字编号或安全 PKL 文件名", "已登记目录内的绝对路径及从 100 开始的场景编号"],
        ["目标样本构造", "PKL 内容、目标索引、T 和 N", "固定长度目标样本；可选字段按约定补零"],
        ["整段前向", "增加批维的输入张量和 dt", "B 为 1 的绝对状态预测及完整序列耗时"],
        ["有效帧筛选", "target_valid_mask、label_abs 和预测", "有效时间、真值、预测、逐帧三维位置误差"],
        ["结果封装", "场景元数据和推理结果", "不可变 TrackingSequence，含序列和折算单帧时延"],
        ["结果缓存", "场景绝对路径与目标索引", "同一引擎内重复请求直接返回缓存序列"],
    ], widths=[1.2, 2.65, 2.95])
    doc.add_paragraph(
        "infer_scenario 在可重入锁内完成缓存检查与首次推理，避免同一引擎并发构造重复结果。真值无有效帧或误差出现 NaN、Inf 时不写入缓存，并向接口层抛出明确异常。"
    )

    doc.add_heading("5.6 批量测试管理模块", level=2)
    doc.add_paragraph(
        "BatchTestManager 管理任务表、取消事件、后台 Future 和单工作线程执行器。create_job 对模型、数据集、设备、批大小、预热帧及场景范围进行校验，生成 UUID 后立即以 queued 状态原子持久化，再提交后台执行。单工作线程使 GPU 密集任务串行运行，避免多个批测任务竞争显存。"
    )
    add_table(doc, ["内部对象", "用途", "一致性措施"], [
        ["jobs", "保存全部任务快照", "RLock 保护读写；对外返回副本"],
        ["cancel_events", "通知排队或运行任务终止", "任务阶段与批次循环均检查事件"],
        ["futures", "记录后台任务句柄", "排队任务可直接取消，运行任务采用协作式取消"],
        ["单线程 executor", "执行数据索引、模型推理和统计", "worker_count 固定为 1，降低显存争用"],
        ["任务 JSON", "服务重启后的任务恢复和报告来源", "临时文件写完后 os.replace 原子替换"],
    ], widths=[1.35, 2.75, 2.7])

    doc.add_heading("5.7 性能评估与报告模块", level=2)
    doc.add_paragraph(
        "EvaluationService 将已完成批测结果转换为合同技术指标判定，并同时提供历史参考、跨数据集矩阵、任务对比和报告目录。position_rmse_m 与 inference_time_ms 直接读取批测结果；association_accuracy 与 imm_jpda_reduction 需要独立证据或同条件基线。证据不可用时返回 unavailable，不以界面演示值代替正式结论。"
    )
    add_table(doc, ["功能", "入口条件", "输出"], [
        ["任务评估", "任务状态为 completed", "四项 criterion、总体状态和多维统计"],
        ["证据生成", "指标在白名单内且任务可评估", "带任务签名和 valid_for_assessment 的 JSON"],
        ["技术复核", "用户提供大于 0 的 IMM-JPDA RMSE", "计算相对降幅供复核，但不能替代正式独立证据"],
        ["任务对比", "选择不超过 50 个已完成任务", "统一字段的横向比较记录"],
        ["报告目录", "已持久化任务存在", "制品可用性、大小、摘要和判定概览"],
    ], widths=[1.35, 2.75, 2.7])

    doc.add_heading("5.8 场景目录与数据访问模块", level=2)
    doc.add_paragraph(
        "scenario_catalog.py 维护验收想定与技术数据集登记信息，向界面提供数据集名称、路径、可用性、场景数量、场景类别、传感器模态和干扰说明。dataset_files 仅枚举登记目录内指定后缀文件；scenario_class 根据文件名识别 CV、CA、CT、CSSW 及复合干扰类别。"
    )
    add_bullets(doc, [
        "完整批测保留排序后的全部场景文件；抽样批测按场景类别进行确定性轮转选取。",
        "分层选样结果在 scenario_manifest 中保存 sampling_strategy、文件清单和类别数量。",
        "src dataset.py 负责将单个注册 PKL 解析为按目标组织的张量样本；运行时和批测复用同一解析口径。",
        "V8 保留 40 Hz 原始注册数据，模型侧使用按 0.1 s 时间点同步抽取且不做量测平均的 10 Hz 兼容视图。",
        "场景文件名只作为登记目录内的逻辑标识，接口层不接受带目录组件的任意路径。",
    ])

    doc.add_heading("5.9 模块协作流程", level=2)
    add_table(doc, ["业务流程", "调用链", "关键边界"], [
        ["启动与总览", "App.vue → api.ts → server.py → ModelManager、BatchTestManager", "前端只接收汇总状态；首次健康检查触发模型管理初始化"],
        ["模型切换", "ModelManagementView → activateModel → ModelManager.activate → TransformerEngine", "候选引擎成功后原子切换，失败时保持原模型"],
        ["单场景回放", "TrackingValidationView → ws tracking → TransformerEngine → CoopTrackerModel", "模型一次计算完整序列，WebSocket 只负责节拍与播放控制"],
        ["批量测试", "BatchTestingView → create_job → 数据集解析 → 模型前向 → 任务 JSON", "界面轮询状态，后台单线程执行，取消采用协作事件"],
        ["评估归档", "PerformanceEvaluationView 或 ReportCenterView → EvaluationService → 证据与报告", "仅完成任务参与正式判定和制品导出"],
    ], widths=[1.2, 3.45, 2.15], small=True)

    doc.add_heading("6 前端界面详细设计", level=1)
    doc.add_heading("6.1 应用框架与导航", level=2)
    doc.add_paragraph(
        "界面采用左侧固定工作区导航与右侧内容区布局，移动端切换为顶部栏、抽屉式导航和遮罩。主导航依次提供系统总览、模型管理、跟踪验证、批量测试、性能评估和报告中心。App.vue 根据 currentPage 条件装载单个页面组件，以 URL Hash 保留页面定位。"
    )
    add_table(doc, ["界面区域", "显示内容", "交互规则"], [
        ["品牌与工作区", "软件名称、原型标识和 API 版本", "信息只读"],
        ["主导航", "六个业务入口、图标和功能说明", "当前项显示 page 状态；移动端选择后自动关闭抽屉"],
        ["后端状态", "在线或未连接、默认服务地址", "错误时显示警示条和重新连接按钮"],
        ["设置抽屉", "API、服务版本、模型、设备、PyTorch、CUDA 和报告目录", "支持刷新、复制后端启动命令和 Escape 关闭"],
        ["全局提示", "成功或错误消息", "显示约 3.6 s；新消息覆盖旧定时器"],
    ], widths=[1.25, 2.85, 2.7])

    doc.add_heading("6.2 数据访问与公共状态", level=2)
    doc.add_paragraph(
        "API 基址由浏览器当前协议和主机名组合为 8000 端口，WebSocket 根据页面协议自动选择 ws 或 wss。loadDashboard 先请求 health，确保默认模型已初始化，再并行请求 system、models、devices 和 acceptance criteria；任一关键请求失败时保留可读错误信息。"
    )
    add_bullets(doc, [
        "应用级状态由 App.vue 持有，并通过 props 向总览、模型、批测、评估和报告页面下发。",
        "页面通过 refresh 或 notice 事件请求父组件刷新数据或显示全局提示，不直接修改父组件状态。",
        "请求期间使用 loading、submitting、activating、assessing 或 cancelling 禁用重复操作。",
        "空列表、尚无结果、连接失败和数据不足均设置文字占位，不显示伪造的零值。",
    ])

    doc.add_heading("6.3 系统总览界面", level=2)
    add_table(doc, ["区域", "信息来源", "显示与操作"], [
        ["页面标题", "lastUpdated", "显示最近成功更新时间，提供手动刷新"],
        ["运行状态带", "health", "服务状态、API 版本、运行时长、设备和场景数量"],
        ["指标卡", "active_model、devices、batch_testing", "当前模型、参数量、活动队列和可用加速设备"],
        ["环境与模型摘要", "system、models", "运行环境、活动模型信息和快捷进入模型管理"],
        ["验收指标", "criteria", "名称、比较运算符、阈值、单位和数据来源"],
    ], widths=[1.25, 2.2, 3.35])
    doc.add_paragraph(
        "当 dashboard 尚未取得时，各卡片以破折号和等待连接文字占位；不会将缺失状态误显示为正常。用户刷新时按钮显示同步中并暂时禁用。"
    )

    doc.add_heading("6.4 模型管理界面", level=2)
    doc.add_paragraph(
        "模型管理页以活动模型摘要、计算设备选择和模型版本卡片组织部署操作。设备列表仅展示后端标记 available 的项目；若已有活动设备，则自动将其作为当前选择。模型卡展示版本、训练与验证数据集、文件大小、修改时间和离线训练摘要。"
    )
    add_numbered_steps(doc, [
        "用户选择目标设备并在模型卡上发起部署，页面将目标模型保存为 pendingModel。",
        "确认对话框明确显示模型版本与设备；用户确认后调用 activateModel。",
        "激活期间关闭重复提交入口；成功后清除 pendingModel、显示通知并请求全局静默刷新。",
        "刷新仓库调用 refreshModels，只重新扫描登记权重；失败时保留当前活动模型并显示错误通知。",
    ])

    doc.add_heading("6.5 跟踪验证界面", level=2)
    doc.add_paragraph(
        "跟踪验证页独立加载场景目录，并维护 idle、connecting、live、paused、completed 和 error 六种连接状态。用户可按 CA、CT、CV、CSSW 或全部类型过滤场景，选择目标索引后建立 WebSocket。连接关闭、切换场景或组件卸载时主动释放 Socket。"
    )
    add_table(doc, ["区域", "内容", "交互或更新规则"], [
        ["场景控制", "机动类型、场景和目标索引", "改变机动类型时重置到筛选后的首个场景"],
        ["实时指标", "当前时间、位置误差、运行 RMSE 和单帧时延", "每收到一帧更新；重置帧清空历史序列"],
        ["三维轨迹", "真值与预测相对首个真值点的轨迹", "采用相同比例轴；支持顶视、侧视和自由视角"],
        ["误差曲线", "逐帧位置误差和运行 RMSE", "与时间轴同步追加并随容器尺寸调整"],
        ["播放控制", "播放、暂停、进度拖动和重播", "向后端发送 pause、play、seek；跳转后重新累计运行 RMSE"],
        ["数据导出", "已接收帧的时间、真值、预测和误差", "在浏览器生成 CSV；无帧时禁用"],
    ], widths=[1.2, 2.55, 3.05], small=True)
    doc.add_paragraph(
        "页面初始化时检查 WebGL。WebGL 不可用时显示告警，但仍保持 WebSocket 数据接收和二维误差曲线；该降级不影响原始数据与 CSV 导出。"
    )

    doc.add_heading("6.6 批量测试界面", level=2)
    doc.add_paragraph(
        "批量测试页同时提供任务配置、实时队列和历史任务。表单默认使用当前活动模型、协同编队验收想定、当前活动设备或首个可用设备、5 个预热帧、64 批大小和 0 个数据加载工作进程。CPU 环境不会保留不可用的 cuda:0。sample 模式使用 max_scenarios，full 模式将其置为空以覆盖全部场景。"
    )
    add_table(doc, ["状态或动作", "前端处理", "后端交互"], [
        ["页面加载", "并行读取数据集与最近 100 个任务", "listDatasets、listBatchJobs"],
        ["持续监控", "每 2.5 s 静默刷新；活动任务变化时同步详情", "listBatchJobs、getBatchJob"],
        ["创建任务", "校验选择并锁定提交按钮，成功后打开任务详情", "createBatchJob"],
        ["历史筛选", "按全部、活动、完成、失败或取消分类", "前端 computed，不重复请求"],
        ["查看详情", "显示配置、进度、总体指标、时延、抽样清单和错误", "getBatchJob"],
        ["取消任务", "仅 queued 或 running 可操作，显示取消中状态", "cancelBatchJob"],
    ], widths=[1.25, 2.85, 2.7])

    doc.add_heading("6.7 性能评估界面", level=2)
    doc.add_paragraph(
        "性能评估页只将已完成批测列入可选任务。页面首次加载并行获取批测列表、跨数据集矩阵、历史参考和已完成任务对比，再对首个任务执行四项指标判定。手工输入的同条件 IMM-JPDA RMSE 只生成技术复核值；正式通过必须具备两份与当前任务签名一致的独立证据。"
    )
    add_bullets(doc, [
        "四项指标卡同时显示阈值、实测值、passed、failed 或 unavailable 状态及说明。",
        "关联正确率和基线降幅缺少正式证据时提供生成入口；生成期间只禁用对应动作。",
        "跨数据集矩阵固定展示 V5、V6、V7 模型与 V5、V6、V7、V8 强干扰数据条件，并按 RMSE 分级着色。",
        "历史参考明确标识为参考信息，关联审计与 IMM-JPDA 历史值不自动替代当前任务证据。",
        "任务对比表展示模型、数据集、场景与航迹数、误差、时延、吞吐率和完成时间。",
    ])

    doc.add_heading("6.8 报告中心界面", level=2)
    doc.add_paragraph(
        "报告中心从持久化任务建立档案列表，支持按任务名称或编号、模型、数据集文本检索，并按全部、已完成或异常状态筛选。只有 completed 记录可以勾选进行横向对比导出；筛选变化不会清除仍然有效的选择。"
    )
    add_table(doc, ["区域", "主要内容", "设计规则"], [
        ["归档摘要", "记录数、完整报告数、可判定数和制品体积", "数据来自 report_catalog 汇总"],
        ["档案表", "实验组合、RMSE、航迹数、结论、时间和制品数量", "最低 RMSE 仅作视觉标识，不改变判定"],
        ["选择栏", "当前可导出记录的全选与计数", "自动排除未完成任务；最多由后端限制 50 项"],
        ["详情抽屉", "任务配置、总体指标、四项判定和错误", "任务与评估并行读取；异常任务不请求评估"],
        ["导出区", "原始 JSON、综合评估 JSON、逐航迹 CSV、对比 CSV", "按制品可用性显示；打印使用浏览器打印样式"],
    ], widths=[1.25, 2.65, 2.9])

    doc.add_heading("6.9 响应式、可用性与异常反馈", level=2)
    add_bullets(doc, [
        "桌面端使用侧栏和多列信息面板，窄屏改为可关闭导航与纵向内容，表格区域允许横向滚动。",
        "对话框和抽屉设置 role、aria-modal、标题关联和关闭按钮；导航项标识 aria-current。",
        "颜色只表达辅助状态，在线、失败、取消、通过和证据不足均同时提供中文文字。",
        "所有异步操作在 finally 中恢复 loading 状态；错误优先显示后端返回 detail，无法解析时使用页面级默认说明。",
        "连接警示、空状态、禁用状态和执行中状态均在原操作附近显示，避免用户重复提交。",
        "三维图以相对坐标提升可读性，但实时数值、接口数据和下载制品仍保留绝对坐标，防止展示口径改变计算结果。",
    ])

    doc.add_heading("7 后端接口详细设计", level=1)
    doc.add_paragraph(
        "后端以 FastAPI 提供 REST 与 WebSocket 两类接口。REST 接口面向系统状态、模型管理、场景目录、批测任务、评估结果和报告制品；WebSocket 接口面向单场景逐帧回放。接口层只负责参数校验、状态映射和传输封装，模型加载、批量计算与证据判定分别委托给专用服务，避免路由代码直接持有计算状态。"
    )

    doc.add_heading("7.1 应用初始化与公共处理", level=2)
    doc.add_paragraph(
        "server.py 使用带缓存的工厂函数延迟创建 ModelManager、BatchTestManager 和 EvaluationService。单例初始化由可重入锁保护；EvaluationService 复用同一个 BatchTestManager，确保评估读取的任务状态与报告目录一致。应用退出时调用批测管理器 shutdown，请求取消尚未结束的任务并关闭执行器。"
    )
    add_table(doc, ["处理环节", "设计行为", "对接口调用方的影响"], [
        ["请求标识", "读取请求头 X-Request-ID；缺失时生成 UUID", "响应回传相同标识，便于前后端和日志关联"],
        ["处理耗时", "HTTP 中间件按请求计算端到端处理毫秒数", "响应头返回 X-Process-Time-Ms"],
        ["跨域访问", "仅允许 localhost 和 127.0.0.1 的 5173 端口", "允许 GET、POST 及全部请求头，不开放凭据"],
        ["全局异常", "记录请求标识、方法、路径和堆栈后返回 500", "响应 detail 不暴露内部堆栈"],
        ["字段校验", "由 Pydantic 在进入业务逻辑前校验类型、长度和数值范围", "非法请求统一返回 422"],
    ], widths=[1.2, 3.25, 2.35])

    doc.add_heading("7.2 接口分层与资源模型", level=2)
    add_table(doc, ["接口组", "资源前缀", "主要服务", "设计职责"], [
        ["运行状态", "/api/health 和 /api/system", "ModelManager 与 BatchTestManager", "汇总健康状态、版本、设备、活动模型和队列"],
        ["模型管理", "/api/models", "ModelManager", "刷新仓库、查询活动模型和切换模型设备"],
        ["场景数据", "/api/scenarios 和 /api/datasets", "ScenarioCatalog", "返回已注册想定和验证数据集"],
        ["批量测试", "/api/batch-tests", "BatchTestManager", "创建、查询、取消任务并下载原始结果"],
        ["评估报告", "/api/evaluations 和 /api/reports", "EvaluationService", "判定、证据生成、对比、目录和导出"],
        ["跟踪回放", "/ws/tracking", "TransformerEngine", "发送元数据和逐帧航迹数据，接收播放控制"],
    ], widths=[1.05, 1.8, 1.6, 2.35], small=True)
    doc.add_paragraph(
        "接口使用 JSON 作为控制与查询数据的默认表示；原始报告和综合评估仍为 JSON 文件，逐航迹与多任务对比采用 CSV。下载接口设置 attachment 响应头，使浏览器按任务标识形成稳定文件名。"
    )

    doc.add_heading("7.3 系统与模型接口", level=2)
    add_table(doc, ["方法与路径", "请求", "成功响应", "关键处理"], [
        ["GET /api/health", "无", "服务状态、API 版本、运行时长、活动模型摘要", "触发模型管理器初始化；失败时由统一异常处理返回 500"],
        ["GET /api/system", "无", "运行环境、设备、活动模型、批测队列和报告目录", "组合多个服务的只读状态"],
        ["GET /api/config/acceptance-criteria", "无", "四项技术指标的运算符、阈值和单位", "直接返回评估服务的固定判定定义"],
        ["GET /api/devices", "无", "CPU 与已检测到的 CUDA 设备", "仅 available 为真时允许模型激活"],
        ["GET /api/models", "无", "可用模型列表及活动标记", "只登记权重文件真实存在的模型"],
        ["GET /api/models/active", "无", "当前模型、设备、激活时间和加载耗时", "无活动模型时返回空状态而不是虚构结果"],
        ["POST /api/models/refresh", "无", "刷新后的模型列表", "重新扫描固定模型目录和指标摘要"],
        ["POST /api/models/activate", "model_id 与 device", "新的活动模型状态", "先完成设备和模型校验，再原子替换推理引擎"],
    ], widths=[2.0, 1.2, 1.8, 1.8], small=True)
    doc.add_paragraph(
        "模型切换属于互斥写操作。新引擎完成权重读取、结构构建和设备迁移后才替换活动引用；加载失败时保留原活动模型。由于场景缓存属于引擎实例，模型切换同时自然清除旧模型的回放缓存。"
    )

    doc.add_heading("7.4 场景与数据集接口", level=2)
    add_table(doc, ["方法与路径", "返回内容", "用途与约束"], [
        ["GET /api/scenarios", "当前活动目录中文件的 id、name、file 和 size_bytes", "供跟踪验证页选择；仅返回目录中已发现的 PKL 文件"],
        ["GET /api/datasets", "数据集 ID、名称、说明、路径、类别、模态、质量配置和可用性", "供批测配置页选择；路径由后端注册表控制"],
    ], widths=[2.15, 2.75, 1.9])
    doc.add_paragraph(
        "浏览器通过 WebSocket 的 dataset 参数提交数字场景 ID 或 PKL 文件名，通过 target 参数提交目标索引，不能提交任意服务器文件路径。后端根据当前活动注册目录解析场景并校验目标索引，从接口边界限制可访问的数据范围。"
    )

    doc.add_heading("7.5 批量测试接口", level=2)
    add_table(doc, ["方法与路径", "主要参数", "响应与状态约束"], [
        ["POST /api/batch-tests", "name、model_id、dataset_id、device、warmup_frames、batch_size、workers、max_scenarios", "校验后立即持久化并以 202 返回 queued 任务"],
        ["GET /api/batch-tests", "limit", "按创建时间倒序返回紧凑任务摘要"],
        ["GET /api/batch-tests/{job_id}", "任务 ID", "返回配置、进度、结果或错误；不存在返回 404"],
        ["POST /api/batch-tests/{job_id}/cancel", "任务 ID", "queued 或 running 发起取消；终态任务原样返回；不存在返回 404"],
        ["GET /api/batch-tests/{job_id}/assessment", "可选 baseline_rmse_m", "仅 completed 且有结果的任务可评估"],
        ["POST /api/batch-tests/{job_id}/evidence/{metric_id}", "白名单指标 ID", "执行独立证据脚本并返回更新后的评估"],
        ["GET /api/batch-tests/{job_id}/report.json", "任务 ID", "下载持久化原始任务报告；文件缺失返回 404"],
        ["GET /api/batch-tests/{job_id}/tracks.csv", "任务 ID", "完成且含逐航迹结果时生成并下载 CSV"],
    ], widths=[2.65, 2.1, 2.05], small=True)
    doc.add_paragraph(
        "创建请求中 name 最长 100 个字符；warmup_frames 允许 0 至 99；batch_size 允许 1 至 256；workers 允许 0 至 8；max_scenarios 为空表示完整测试，非空时必须大于等于 1。model_id 为空时使用活动模型，device 为 auto 时由服务选择 CUDA 0 或 CPU。"
    )

    doc.add_heading("7.6 评估与报告接口", level=2)
    add_table(doc, ["方法与路径", "输入", "输出"], [
        ["GET /api/evaluations/references", "无", "IMM-JPDA 历史对比、关联审计和来源路径，并附非正式判定提示"],
        ["GET /api/evaluations/cross-dataset", "无", "模型与数据集组合的历史 RMSE 矩阵；源文件缺失时 available 为 false"],
        ["GET /api/evaluations/compare", "逗号分隔 job_ids，可省略", "已完成任务的标准化指标行；显式任务最多 50 项"],
        ["GET /api/reports", "limit", "报告记录、汇总统计、评估可用性和制品标志"],
        ["GET /api/reports/comparison.csv", "逗号分隔 job_ids", "选中已完成任务的横向比较 CSV"],
        ["GET /api/reports/{job_id}/assessment.json", "可选 baseline_rmse_m", "包含原始任务、判定、证据状态和来源信息的综合评估包"],
    ], widths=[2.65, 1.35, 2.8], small=True)
    doc.add_paragraph(
        "历史参考与当前任务判定严格分离。references 和 cross-dataset 用于展示既有试验结果；除非存在与当前任务签名匹配的独立证据，历史数值不自动写入正式判定。"
    )

    doc.add_heading("7.7 WebSocket 跟踪接口", level=2)
    doc.add_paragraph(
        "客户端以 /ws/tracking 建立连接，并通过 dataset 和 target 查询参数提交场景标识与目标索引。服务接受连接后取得当时的活动 TransformerEngine，在工作线程中完成一次完整序列推理，再按时间顺序每 50 ms 发送一帧 JSON。播放暂停只停止发送进度，不重复执行模型推理。"
    )
    add_table(doc, ["方向", "动作或内容", "关键字段", "处理规则"], [
        ["服务端到客户端", "逐帧结果", "time、max_time、is_reset、scenario、target_index、truth_pos、pred_pos、metrics", "发送绝对坐标和当前误差、运行 RMSE、序列及单帧折算时延"],
        ["客户端到服务端", "pause", "action", "保持当前索引并停止逐帧发送"],
        ["客户端到服务端", "play", "action", "从当前索引恢复发送"],
        ["客户端到服务端", "seek", "action 与 time", "按时间轴二分查找最近帧，跳转后清空运行 RMSE 累计"],
    ], widths=[1.35, 1.25, 2.45, 1.75], small=True)
    doc.add_paragraph(
        "场景、模型、数据或目标索引错误作为策略错误关闭连接，关闭码为 1008；客户端主动断开按正常结束处理。WebSocket 不承担批测任务提交和文件下载，防止长时任务状态与回放连接生命周期耦合。"
    )

    doc.add_heading("7.8 请求响应约定", level=2)
    add_bullets(doc, [
        "所有时间字段使用带时区的 ISO 8601 字符串；运行时长和处理耗时使用秒或毫秒并在字段名中标明单位。",
        "资源不存在返回 404；任务状态冲突返回 409；参数、数据、模型或运行环境不满足执行条件返回 422。",
        "列表接口返回 total 与 items；任务详情保留 config、progress、result 和 error，前端不得根据缺失字段自行推断完成状态。",
        "可下载制品的文件名包含任务 ID；JSON 以 UTF-8 输出，CSV 使用固定列顺序，便于归档脚本稳定读取。",
        "接口字段使用 snake_case，枚举状态使用小写英文；界面负责将其映射为中文，不改变传输值。",
    ])

    doc.add_heading("7.9 错误处理与安全边界", level=2)
    add_table(doc, ["异常类别", "HTTP 或 WS 结果", "设计处理"], [
        ["未知资源或任务", "404", "不返回目录遍历信息，只给出明确资源不存在说明"],
        ["非法任务状态", "409", "例如评估未完成任务或下载尚不可用的逐航迹结果"],
        ["请求字段或业务输入非法", "422", "保留可读 detail，拒绝进入计算流程"],
        ["未处理服务异常", "500", "服务器记录完整堆栈，响应只含通用说明与请求 ID"],
        ["回放策略错误", "WebSocket 1008", "关闭当前连接，不影响其他连接和批测任务"],
    ], widths=[1.55, 1.35, 3.9])
    doc.add_paragraph(
        "后端当前面向本机原型联调，不包含账号鉴权和公网暴露设计。部署边界应保持在受控网络或本机；如扩展为多用户服务，应另行增加身份认证、授权、限流、TLS、审计日志和文件下载权限控制。"
    )

    doc.add_heading("8 模型推理详细设计", level=1)
    if MODEL_ARCH.is_file():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(MODEL_ARCH), width=Inches(6.7))
        cap = doc.add_paragraph("图 2 CoopTrackerModel 内部结构示意")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cap.runs:
            set_run_font(run, size=9.5, color=TEXT_GRAY)
    doc.add_paragraph(
        "CoopTrackerModel 采用“物理可解释粗估计＋深度网络残差修正”的组合结构。模型先利用雷达与协同链路构造粗观测状态，再通过 Learned-K 完成递推更新；随后对雷达、EO 和 ESM 进行多模态编码、跨注意力融合、时序建模和多无人机图交互，最终输出绝对 ECEF 状态。"
    )

    doc.add_heading("8.1 模型输入与输出", level=2)
    add_table(doc, ["接口项", "形状", "主要内容", "使用位置"], [
        ["radar", "B × T × N × 3", "az、el、r", "几何反演、粗观测和雷达主流"],
        ["eo、esm", "B × T × N × 2", "az、el", "被动方向编码与融合"],
        ["三类 mask", "B × T × N", "0 或 1", "通道替代、观测加权、图邻接和节点汇聚"],
        ["uav_pos", "B × T × N × 3", "相对参考位置", "几何反演、特征输入和邻接距离"],
        ["rotmat_b2e", "B × T × N × 3 × 3", "机体系到世界系旋转矩阵", "雷达位置及 EO、ESM 方向转换"],
        ["imu", "B × T × N × 6", "姿态与角速度", "Learned-K 观测编码和姿态门控"],
        ["快慢链输入", "B × T × N × 3 或 B × T × N", "位置、龄期和更新标志", "异步协同观测融合"],
        ["ref_pos", "B × 3", "参考绝对 ECEF 位置", "输出阶段恢复绝对位置"],
        ["dt", "B", "样本时间间隔", "匀速状态转移和差分速度"],
        ["模型输出", "B × T × 6", "px、vx、py、vy、pz、vz", "运行时回放及批测指标计算"],
    ], widths=[1.15, 1.5, 2.1, 2.05], small=True)
    doc.add_paragraph(
        "缺失的 imu、链路位置、龄期、更新标志和 esm_mode 在运行时按固定形状补零。模型输入默认采用 float32；单场景运行时增加 B 等于 1 的批维，批测则由 DataLoader 形成实际批大小。"
    )

    doc.add_heading("8.2 几何预处理与观测状态构造", level=2)
    doc.add_heading("8.2.1 雷达坐标反演", level=3)
    doc.add_paragraph(
        "radar_body_to_world_xyz 按雷达分量 az、el、r 计算机体系相对位置：x 等于 r cos(el) cos(az)，y 等于 r cos(el) sin(az)，z 等于 r sin(el)。相对向量左乘 rotmat_b2e 后叠加 uav_pos，得到参考坐标系下的目标位置。"
    )
    doc.add_heading("8.2.2 信息矩阵融合", level=3)
    doc.add_paragraph(
        "LearnedKFrontEnd._build_obs_state 将各平台雷达位置与快、慢链位置统一到三维观测空间。雷达径向标准差取 48 m，角度引起的切向标准差按 r × 0.002 计算并增加 1 m² 方差保护；链路位置采用 60 m 各向同性标准差。每项观测以精度矩阵形式累加，最后求解三阶线性方程得到融合位置。"
    )
    add_table(doc, ["来源", "有效条件", "时间新鲜度权重", "空间权重"], [
        ["雷达", "radar_mask 等于 1", "不额外衰减", "径向与切向不同方差的信息矩阵"],
        ["快链", "fast_update 大于 0.5 或 fast_age 小于 2 s", "exp 负 fast_age 除以 0.3", "60 m 标准差的各向同性矩阵"],
        ["慢链", "slow_update 大于 0.5 或 slow_age 小于 5 s", "exp 负 slow_age 除以 2.0", "60 m 标准差的各向同性矩阵"],
    ], widths=[1.0, 2.75, 1.55, 1.5], small=True)
    doc.add_paragraph(
        "总信息矩阵加入 1e-6 的单位阵正则项，再由固定三阶伴随矩阵公式求解。若某帧没有任何有效位置来源，则沿用上一帧融合位置；首帧使用正则化求解结果。速度由相邻位置除以 dt 得到，首帧速度置零，并将每轴速度限制在负 400 m/s 至 400 m/s。随后将位置和速度交错组成 z_obs。"
    )

    doc.add_heading("8.3 Learned-K 状态前端", level=2)
    doc.add_paragraph(
        "Learned-K 前端将每个平台每帧的雷达、EO、ESM、IMU、三类掩码、ESM 模式、快慢链龄期与位置、更新标志和无人机位置拼接为 30 维特征，经两层线性映射、LayerNorm 和 GELU 编码为 d_model 维。只对至少一种传感器有效的平台特征求平均，得到逐帧 pooled_obs。"
    )
    add_numbered_steps(doc, [
        "以 z_obs 首帧作为 x_prev；根据 dt 构造六维匀速状态转移矩阵 F。",
        "首帧令 x_pred 等于 x_prev，后续帧计算 x_pred 等于 F 乘以前一帧 x_post。",
        "计算 innovation 等于 z_obs 减 x_pred，并拼接 x_pred、z_obs、innovation 和 pooled_obs。",
        "k_net 使用两层隐藏映射和 Sigmoid 输出六维对角增益 K，取值限制在 0 至 1。",
        "按 x_post 等于 x_pred 加 K 逐元素乘 innovation 更新状态，并递推至序列末尾。",
    ])
    doc.add_paragraph(
        "该前端输出 k_state 作为物理粗状态，同时保留 z_obs 供内部诊断。匀速预测提供连续性，学习增益根据当前创新量和多传感器质量决定观测修正强度。"
    )

    doc.add_heading("8.4 多模态特征编码", level=2)
    add_table(doc, ["分支", "编码器输入", "输入维数", "缺失处理"], [
        ["雷达主流", "原始 radar 3 维＋世界几何位置 3 维＋k_state 6 维", "12", "radar_mask 无效时替换为可学习 mask token"],
        ["EO 被动流", "原始角度 2 维＋世界方向向量 3 维", "5", "eo_mask 无效时替换为可学习 mask token"],
        ["ESM 被动流", "原始角度 2 维＋世界方向向量 3 维", "5", "esm_mask 无效时替换为可学习 mask token"],
    ], widths=[1.1, 3.25, 0.8, 1.65], small=True)
    doc.add_paragraph(
        "三个 ChannelMaskEncoder 均采用“线性层、LayerNorm、GELU、线性层、LayerNorm、GELU”结构，输出 d_model 等于 96 的特征。EO 和 ESM 先由方位、俯仰角生成机体系单位方向，再乘 rotmat_b2e 转换到世界方向。为了以单平台序列为单位建模，输入由 B × T × N 重排为 B×N 条长度为 T 的序列。"
    )

    doc.add_heading("8.5 主被动双流融合", level=2)
    doc.add_heading("8.5.1 被动通道置信融合", level=3)
    doc.add_paragraph(
        "EO 与 ESM 特征分别经线性层和 Sigmoid 得到逐时刻置信度，再乘各自 mask。被动融合特征为两类特征的置信度加权平均；两类被动量测均无效时，使用可学习 passive_missing 向量。"
    )
    doc.add_heading("8.5.2 自适应对齐与姿态门控", level=3)
    doc.add_paragraph(
        "雷达和被动特征分别沿时间维标准化，并使用可学习 align_gamma 与 align_beta 对齐分布。IMU 通过线性投影和 Sigmoid 形成 pose_gate，两条流均乘以 1 加 pose_gate，使平台姿态和角速度参与特征调制。标准差下限取 1e-5。"
    )
    doc.add_heading("8.5.3 双向跨注意力", level=3)
    doc.add_paragraph(
        "系统分别计算 radar 对 passive 的四头跨注意力和 passive 对 radar 的四头跨注意力。每个 CrossAttentionBlock 包含注意力残差、LayerNorm、两倍通道前馈网络、GELU、0.1 Dropout 和第二次残差归一化。fusion_gate 对两路上下文拼接后输出标量门值 alpha，并按 alpha 加权；融合结果再与 radar_ctx 拼接投影回 d_model。"
    )

    doc.add_heading("8.6 时序建模与协同图交互", level=2)
    doc.add_heading("8.6.1 时序 Transformer", level=3)
    doc.add_paragraph(
        "每个平台的融合序列首先加入正弦位置编码，再进入两层 TransformerEncoder。每层使用 d_model 96、4 个注意力头、4 倍通道前馈网络、GELU 激活、0.1 Dropout 和 batch_first 布局。use_temporal 为 false 时可跳过该模块，用于消融验证。"
    )
    doc.add_heading("8.6.2 协同 GAT", level=3)
    doc.add_paragraph(
        "时序结果重排为每个时刻 N 个无人机节点。邻接矩阵同时要求两个节点存在至少一种有效传感器，并且平台间欧氏距离不超过 20 km；对角线始终置为有效，保证每个节点保留自身信息。CollaborativeGAT 对节点分别生成 Q、K、V，以缩放点积计算注意力，对无邻接边填充负大数，经 Softmax 和 0.1 Dropout 后聚合 V，并与原特征残差相加后 LayerNorm。"
    )
    doc.add_paragraph(
        "use_gat 为 false 时直接使用时序节点特征。图输出按 valid_sensor 加权汇聚为逐帧 global_feat；无有效传感器时分母增加 1e-6，防止除零。"
    )

    doc.add_heading("8.7 残差回归与绝对坐标恢复", level=2)
    doc.add_paragraph(
        "global_feat 分别进入位置头和速度头。两个回归头均由 d_model 到 d_model 除以 2 的线性层、GELU 和三维输出层构成。三维位置残差与三维速度残差按 px、vx、py、vy、pz、vz 交错写入六维 residual，并与 k_state 相加形成 pred_rel。"
    )
    add_numbered_steps(doc, [
        "对 pred_rel 的第 0、2、4 分量分别加回 ref_pos 的 x、y、z。",
        "速度分量 1、3、5 不叠加参考位置，保持世界坐标速度语义。",
        "返回 pred_abs，形状为 B × T × 6，供运行时和批测使用。",
    ])

    doc.add_heading("8.8 运行时推理流程与缓存", level=2)
    add_numbered_steps(doc, [
        "TransformerEngine 初始化时创建 d_model 等于 96 的 CoopTrackerModel，将 checkpoint 映射到指定设备并以 strict 等于 true 加载权重。",
        "infer_scenario 将 PKL 解析为选定目标样本，补齐可选字段，并把全部张量移动到推理设备后增加批维。",
        "在 torch.inference_mode 中执行一次完整序列前向；CUDA 模式在计时前后同步设备，得到可比较的序列耗时。",
        "按 target_valid_mask 过滤真值和预测，通过位置分量 0、2、4 计算三维欧氏误差。",
        "将完整序列耗时除以有效帧数得到回放显示使用的折算单帧时延，封装为 TrackingSequence。",
        "以场景绝对路径和目标索引为键写入引擎内存缓存；重复回放直接返回缓存。",
    ])
    doc.add_paragraph(
        "缓存只存在于当前 TransformerEngine 生命周期内，不写入磁盘。模型切换通过新建引擎自然清空缓存；不同目标索引分别缓存，避免混用航迹结果。"
    )

    doc.add_heading("8.9 数值稳定性与降级策略", level=2)
    add_table(doc, ["风险", "保护措施", "失败结果"], [
        ["信息矩阵奇异或近奇异", "加入 1e-6 单位阵；三阶伴随公式的行列式下限为 1e-24", "继续得到有限粗观测"],
        ["无有效观测", "位置沿用上一有效帧；图节点保留自环", "保持序列连续，不用零点突变替代"],
        ["链路字段为全零", "视为无效，龄期改为 999 s，更新标志清零", "不参与快慢链观测融合"],
        ["非正 dt", "观测速度计算将 dt 下限设为 1e-6", "避免除零"],
        ["差分速度尖峰", "逐轴限制在正负 400 m/s", "限制异常观测对递推的影响"],
        ["标准化方差过小", "标准差下限为 1e-5", "避免特征放大为非有限值"],
        ["掩码汇聚无有效项", "分母增加 1e-6并使用缺失向量", "保持张量形状和网络可执行"],
        ["输出出现 NaN 或 Inf", "运行时检查全部位置误差是否有限", "中止场景并返回推理错误，不缓存结果"],
    ], widths=[1.45, 3.65, 1.7], small=True)

    doc.add_heading("9 批量测试详细设计", level=1)
    doc.add_paragraph(
        "BatchTestManager 将一次批测定义为可持久化任务，以单线程后台执行器串行调度 GPU 重负载。接口线程只完成校验、建档和提交，不等待模型加载或数据集计算；前端通过任务查询接口读取进度。该设计避免多个模型同时占用显存，并允许任务完成后独立执行评估和报告下载。"
    )

    doc.add_heading("9.1 功能边界与任务配置", level=2)
    add_table(doc, ["配置项", "默认或范围", "设计用途"], [
        ["name", "为空时自动生成；最长 100 字符", "标识模型与数据集组合，供报告中心检索"],
        ["model_id", "为空时取活动模型", "确定批测权重；创建任务时即验证存在性"],
        ["dataset_id", "默认 v7", "从 DATASET_CATALOG 选择受控数据目录"],
        ["device", "默认 auto", "解析为实际 CPU 或 CUDA 设备并写入任务配置"],
        ["warmup_frames", "默认 5，范围 0 至 99", "排除每条航迹初始过渡帧，不改变模型输入"],
        ["batch_size", "默认 64，范围 1 至 256", "控制 DataLoader 批量和显存占用"],
        ["workers", "默认 0，范围 0 至 8", "控制数据装载子进程数"],
        ["max_scenarios", "为空表示全部；非空至少 1", "控制抽样测试规模，不限制每文件目标数量"],
        ["sequence_length", "固定 100", "统一模型输入、预热约束和实测单帧折算口径"],
    ], widths=[1.3, 2.05, 3.45], small=True)
    doc.add_paragraph(
        "任务创建前验证模型 ID、数据集注册项、数据目录和设备。解析后的模型、数据集与设备写入 config，后续执行不依赖界面当前选择；即使活动模型在排队期间发生变化，任务仍按建档时配置运行。"
    )

    doc.add_heading("9.2 任务对象与状态机", level=2)
    add_table(doc, ["状态", "进入条件", "主要字段", "允许后续操作"], [
        ["queued", "任务已写入报告目录并提交执行器", "created_at、config、初始 progress", "查询、取消"],
        ["running", "后台线程取得任务并记录 started_at", "阶段、批次数、航迹数和完成百分比", "查询、请求取消"],
        ["completed", "推理、统计和最终持久化均成功", "finished_at、result、100% progress", "评估、对比、下载"],
        ["failed", "数据、模型、环境或运行异常", "finished_at、异常类型与信息", "查询、保留失败档案"],
        ["cancelled", "排队 Future 取消成功或运行线程检测到取消事件", "finished_at、cancelled 阶段", "查询、保留取消档案"],
    ], widths=[0.9, 2.2, 2.15, 1.55], small=True)
    doc.add_paragraph(
        "progress 包含 stage、message、completed_batches、total_batches、processed_tracks、total_tracks 和 percent。百分比仅在 total_batches 已知后按完成批次数计算；状态字段是任务终态判定的唯一依据，界面不得仅凭 percent 等于 100 推断 completed。"
    )

    doc.add_heading("9.3 创建 排队 执行与取消", level=2)
    add_numbered_steps(doc, [
        "create_job 解析模型、数据集和设备，生成 32 位十六进制 UUID，构造 queued 任务并立即写入 JSON。",
        "任务提交到 max_workers 等于 1 的 ThreadPoolExecutor；多个任务按执行器队列顺序串行运行。",
        "后台线程将任务置为 running，记录 started_at，并依次进入 preparing、loading_model 和 running 阶段。",
        "数据集索引建立后写入总批次数和总航迹数；每批完成后更新已处理批次、航迹和百分比。",
        "取消操作设置独立 threading.Event。排队任务若 Future.cancel 成功则立即转 cancelled；运行任务在准备、模型加载前后及每批开始处检查事件。",
        "成功时先组装 result 与终态字段，再执行最终原子写入；失败或取消均写入 finished_at 并持久化。",
    ])
    doc.add_paragraph(
        "取消属于协作式取消，不强制中断正在执行的一次模型前向。因此请求取消后可能短暂显示“正在取消”，直至下一个检查点到达；这一行为可避免在 CUDA 内核执行中破坏进程状态。"
    )

    doc.add_heading("9.4 场景发现与分层抽样", level=2)
    doc.add_paragraph(
        "批测只读取 DATASET_CATALOG 注册目录中的 PKL 文件。max_scenarios 为空时按确定顺序使用全部文件；指定上限时，selected_dataset_files 根据文件名识别 CV、CA、CT、CSSW 和复合干扰类别，并按类别轮转选择，使小样本测试尽量覆盖不同运动或干扰类型。"
    )
    add_table(doc, ["输出信息", "保存位置", "复核用途"], [
        ["sampling_strategy", "config 和 scenario_manifest", "确认完整测试或分层抽样的固定策略"],
        ["files", "scenario_manifest", "记录实际参与批测的文件主名"],
        ["class_counts", "scenario_manifest", "统计各场景类别入选数量"],
        ["evaluated_scenarios", "metrics", "记录最终装载的场景文件数"],
    ], widths=[1.55, 2.2, 3.05])
    doc.add_paragraph(
        "抽样不使用随机数，因此相同数据目录、分类规则和 max_scenarios 可得到相同文件清单。报告保存实际清单而不只保存上限，防止数据目录内容变化后无法解释历史结果。"
    )

    doc.add_heading("9.5 数据装载与批量推理", level=2)
    add_numbered_steps(doc, [
        "_BatchRegisteredPklDataset 逐文件读取已配准样本，将每个有效目标展开为独立航迹记录，并按 100 帧及 5 架无人机整理张量。",
        "DataLoader 使用 shuffle 等于 false；batch_size 和 workers 取任务配置，CUDA 设备启用 pin_memory。",
        "加载模型权重后构造独立 TransformerEngine，批测不复用单场景回放缓存，也不改变当前活动引擎。",
        "每批张量移动到目标设备，在 torch.inference_mode 中调用模型前向；CUDA 在计时前后同步，避免异步执行低估耗时。",
        "从输出和 truth 中取位置分量 0、2、4，计算每帧三维欧氏误差，并结合稳定阶段掩码累计点级与航迹级统计。",
        "模型输出、稳定阶段误差、逐航迹误差、总体误差和时延数组均执行有限数检查，出现 NaN 或 Inf 时任务失败。",
        "每 10 个批次或最后一个批次写入一次中间进度，平衡故障可见性和磁盘写入开销。",
    ])

    doc.add_heading("9.6 精度指标计算", level=2)
    doc.add_paragraph(
        "stable_stage_mask 先复制 target_valid_mask，再将每条航迹前 warmup_frames 帧置为无效。位置误差只统计真值有效且已过初始过渡阶段的点；如掩码后没有任何有效点，任务失败而不生成零值指标。"
    )
    add_table(doc, ["结果字段", "计算口径", "统计层级"], [
        ["position_rmse_m", "全部有效点三维距离平方均值开根号", "全任务点级"],
        ["position_mae_m", "全部有效点三维距离算术平均", "全任务点级"],
        ["position_median_m", "全部有效点三维距离中位数", "全任务点级"],
        ["position_p95_m", "全部有效点三维距离第 95 百分位", "全任务点级"],
        ["position_max_m", "全部有效点三维距离最大值", "全任务点级"],
        ["mean_track_rmse_m", "先对每条航迹计算 RMSE，再对航迹 RMSE 求平均", "航迹级"],
        ["tracks 列表", "每条航迹保存 valid_frames、RMSE、MAE、P95 和最大误差", "逐航迹"],
    ], widths=[1.75, 3.7, 1.35], small=True)
    doc.add_paragraph(
        "metrics 同时记录 evaluated_scenarios、evaluated_tracks 和 evaluated_points。metric_definitions 保存 RMSE 的坐标分量、单位、掩码和预热排除规则，使报告使用者可从文件本身复核统计口径。"
    )

    doc.add_heading("9.7 时延与吞吐测量", level=2)
    add_table(doc, ["指标", "测量方法", "解释边界"], [
        ["total_inference_ms", "全部批次模型前向耗时求和", "包含第一批冷启动，不含数据文件读取"],
        ["mean_batch_ms 和 p95_batch_ms", "多批时排除第一批后统计；单批时保留该批", "描述配置批大小下的稳态批处理时延"],
        ["mean_track_ms", "总前向耗时除以有效逐航迹结果数", "批量执行下的航迹平均分摊值"],
        ["estimated_frame_ms", "总前向耗时除以稳定阶段有效点数", "历史兼容折算值"],
        ["measured_frame_ms", "取首个单航迹样本预热 3 次、实测 10 次，平均序列时延除以 100", "四项指标判定优先采用的单帧模型前向实测值"],
        ["p95_measured_frame_ms", "10 次单航迹序列耗时 P95 除以 100", "反映重复实测波动"],
        ["throughput_tracks_per_second", "逐航迹结果数除以总前向秒数", "批处理吞吐，不等同端到端界面刷新率"],
    ], widths=[1.65, 3.35, 1.8], small=True)
    doc.add_paragraph(
        "metric_definitions 明确时延范围为“已对齐和关联的多源张量输入到目标状态输出”，排除文件 I/O、原始检测提取、目标关联、界面绘制和飞控指令生成。正式报告应同时保留模型、设备及软件环境，避免跨硬件直接比较。"
    )

    doc.add_heading("9.8 关联统计", level=2)
    doc.add_paragraph(
        "批测管理器调用 _association_summary 读取入选场景中的 AssociationStats。仅当数据文件包含可识别的关联统计时返回 association 对象，内容可包含正确数、总数和准确率；缺失时返回空值，不从位置误差推导关联准确率。评估服务可优先采用与任务签名严格匹配的独立关联证据。"
    )

    doc.add_heading("9.9 持久化与重启恢复", level=2)
    doc.add_paragraph(
        "每个任务以任务 ID 命名，写入 artifacts/reports/batch_tests 下的 JSON。_persist 在锁内复制任务快照，在锁外写入包含线程标识的临时文件，再调用 os.replace 原子替换目标文件，避免读取方观察到半写入内容。"
    )
    add_bullets(doc, [
        "启动时扫描报告目录中的 JSON；无法解析或缺少任务 ID 的文件被忽略，不阻止服务启动。",
        "历史终态任务原样装入内存，继续支持查询、评估和下载。",
        "历史 queued 或 running 任务无法安全续跑，统一标记 failed，记录新的 finished_at 和“后端重启中断”错误。",
        "恢复得到的 failed 状态立即原子回写原任务 JSON，确保内存查询与磁盘报告一致。",
        "任务列表返回紧凑 result，只保留核心指标、时延、关联和报告路径；任务详情保留完整逐航迹数据。",
    ])

    doc.add_heading("9.10 失败处理与可复现信息", level=2)
    add_table(doc, ["失败点", "任务结果", "保留信息"], [
        ["模型、数据集或设备校验失败", "任务不创建，接口返回 404 或 422", "接口 detail 与请求 ID"],
        ["数据集中无已注册目标", "failed", "配置、阶段、异常类型与信息"],
        ["预热与掩码后无有效点", "failed", "warmup_frames、数据集与错误原因"],
        ["模型前向或 CUDA 异常", "failed", "model_id、model_path、device、当前进度和异常"],
        ["用户取消", "cancelled", "原配置、停止阶段、完成批次和结束时间"],
    ], widths=[2.0, 1.45, 3.35])
    doc.add_paragraph(
        "成功报告保存 model_path、dataset_path、任务配置、指标定义、实际场景清单、逐航迹结果、时延方法及模型、数据和运行代码的 SHA-256 摘要；失败报告至少保存输入配置、进度和错误。JSON 写入禁止 NaN 和 Inf。任务结束后释放局部模型引用，并在 CUDA 可用时清理缓存。"
    )

    doc.add_heading("10 性能评估与报告详细设计", level=1)
    doc.add_paragraph(
        "EvaluationService 在不修改原始批测结果的前提下生成技术指标判定、证据状态、多维统计和可下载评估包。评估只接受 completed 且包含 result 的任务；原始任务 JSON 始终作为事实来源，评估结果可按同一任务重新生成。"
    )

    doc.add_heading("10.1 评估输入与总体状态", level=2)
    add_table(doc, ["输入类别", "主要字段", "作用"], [
        ["任务配置", "model_id、dataset_id、device、warmup_frames、batch_size、workers、max_scenarios、sequence_length", "说明实验条件并参与证据签名"],
        ["精度结果", "position_rmse_m、点数、航迹数及逐航迹统计", "形成位置精度判定和分布分析"],
        ["时延结果", "measured_frame_ms、estimated_frame_ms、吞吐与测量方法", "形成时延判定并说明测量边界"],
        ["关联复核结果", "注册数据 AssociationStats", "仅形成关联准确率技术复核值，不构成正式证据"],
        ["外部证据", "association_accuracy 和 rmse_improvement_pct JSON", "正式验收必须同时具备两份签名匹配的独立证据"],
        ["可选复核值", "baseline_rmse_m", "由调用方显式提供时计算改善率，但不构成正式证据"],
    ], widths=[1.25, 3.35, 2.2], small=True)
    doc.add_paragraph(
        "每项 criterion 的状态为 passed、failed 或 unavailable。总体状态按保守顺序汇总：任一已计算指标失败则 overall_status 为 failed；没有失败但存在 unavailable 或两份正式证据未同时 accepted 时为 incomplete；只有四项均通过且关联正确率与 IMM-JPDA 改善率两份独立证据均与当前任务签名一致时，overall_status 才为 passed。"
    )

    doc.add_heading("10.2 四项技术指标判定", level=2)
    add_table(doc, ["指标 ID", "判定规则", "取值优先级", "结果说明"], [
        ["position_rmse_m", "小于等于 8 m", "批测稳定阶段 position_rmse_m", "附 warmup_frames 与 target_valid_mask 口径"],
        ["association_accuracy", "大于等于 0.85", "独立证据优先；AssociationStats 仅供技术复核", "正式通过必须有 accepted 独立证据"],
        ["rmse_improvement_pct", "大于等于 15%", "独立证据优先；手工 baseline_rmse_m 仅供技术复核", "正式通过必须有 accepted 独立证据"],
        ["inference_time_ms", "小于等于 100 ms", "measured_frame_ms，缺失时 estimated_frame_ms", "注明实测或历史折算方法"],
    ], widths=[1.55, 1.4, 2.6, 1.25], small=True)
    doc.add_paragraph(
        "_criterion 根据固定运算符和阈值计算布尔 passed，并保留 value、threshold、operator、unit、status 和 note。缺少值时不以零代替，而是返回 unavailable，防止证据不足被误判为通过。"
    )

    doc.add_heading("10.3 独立证据生成", level=2)
    doc.add_paragraph(
        "关联准确率和相对 IMM-JPDA 改善率由白名单脚本生成。接口只接受 association_accuracy 与 rmse_improvement_pct 两个 metric_id，服务端用当前 Python 解释器执行 scripts/evaluation 下的固定脚本，并强制传入当前任务报告路径和受控输出路径。"
    )
    add_table(doc, ["证据指标", "脚本", "必须验证的内容"], [
        ["association_accuracy", "test_association_accuracy.py", "metric_id、job_id、job_signature、valid_for_assessment 和有限数 value"],
        ["rmse_improvement_pct", "test_rmse_improvement.py", "上述字段以及大于 0 的有限数 imm_jpda_rmse_m"],
    ], widths=[1.65, 2.35, 2.8])
    doc.add_paragraph(
        "证据生成使用互斥锁，避免同一证据目录被并发脚本写入；单次执行超时为 3600 秒。脚本非零退出时截取标准错误或标准输出中的末尾信息返回，且不接受未通过完整性校验的文件。"
    )

    doc.add_heading("10.4 证据完整性与任务签名", level=2)
    doc.add_paragraph(
        "任务签名把证据绑定到具体批测事实。签名输入包括任务 ID、模型、数据集、设备、批大小、数据加载进程数、最大场景数、预热帧、序列长度、已评估场景数、位置 RMSE、模型路径、数据路径、抽样策略、实际场景文件清单，以及模型权重、入选数据文件和运行代码的 SHA-256 摘要。全部字段规范序列化后再次计算 SHA-256。"
    )
    add_numbered_steps(doc, [
        "读取 evidence_root 下 job_id/metric_id.json；文件不存在、不是 JSON 对象或字段类型非法时标记不可用。",
        "核对文件内 metric_id 与 job_id，拒绝跨任务或跨指标复用。",
        "重新计算当前任务签名并与 job_signature 比较；配置、结果或场景清单变化均使旧证据失效。",
        "确认 valid_for_assessment 严格等于 true，并校验 value 为有限数。",
        "改善率证据额外校验 imm_jpda_rmse_m 为有限正数；全部通过后才将 accepted 置为 true。",
    ])

    doc.add_heading("10.5 航迹分布与机动类型分析", level=2)
    doc.add_paragraph(
        "评估从 result.tracks 汇总航迹 RMSE 分布，包括航迹数、均值、标准差、最小值、中位数、P75、P90、P95、最大值以及 RMSE 不超过 8 m 的航迹数和比例。该分布用于解释总体 RMSE，不能替代点级 position_rmse_m 判定。"
    )
    add_table(doc, ["机动类别", "识别依据", "输出统计"], [
        ["CV", "航迹名中的 CV 标识", "航迹数、平均、中位、P95 和最大航迹 RMSE"],
        ["CA", "航迹名中的 CA 标识", "航迹数、平均、中位、P95 和最大航迹 RMSE"],
        ["CT", "航迹名中的 CT 标识", "航迹数、平均、中位、P95 和最大航迹 RMSE"],
        ["CSSW", "航迹名中的 CSSW 标识", "航迹数、平均、中位、P95 和最大航迹 RMSE"],
        ["OTHER", "不匹配上述规则", "保留未分类航迹，避免静默丢失"],
    ], widths=[1.0, 2.35, 3.45])

    doc.add_heading("10.6 历史参考与跨数据集结果", level=2)
    add_table(doc, ["数据源", "接口表示", "使用规则"], [
        ["IMM-JPDA 历史比较", "references.imm_jpda_comparison", "只展示来源和既有结果，不自动充当当前任务基线"],
        ["历史关联审计", "references.association_audit", "用于背景复核，不覆盖当前任务证据"],
        ["跨数据集 RMSE", "cross-dataset 的 model_id、dataset_id、position_rmse_m 矩阵", "用于观察泛化趋势；源文件缺失时返回 available 为 false"],
    ], widths=[1.7, 2.45, 2.65])
    doc.add_paragraph(
        "references 响应明确声明历史结果仅用于展示，正式判定需要与当前批测相同数据集和设置的基线。该提示属于评估逻辑的一部分，不由前端自行省略。"
    )

    doc.add_heading("10.7 多任务对比", level=2)
    doc.add_paragraph(
        "compare_jobs 对已完成任务归一化为统一行，包含任务名称、模型、数据集、创建和完成时间、RMSE、MAE、P95、单帧时延、吞吐率、关联准确率以及评估场景和航迹数。未完成或无结果任务被跳过，结果按创建时间倒序排列。"
    )
    add_bullets(doc, [
        "未指定 job_ids 时从最近 200 个任务中选择已完成项，适用于报告中心总览。",
        "显式指定 job_ids 时逐项读取；接口层将数量限制为最多 50，避免生成过大的比较响应或 CSV。",
        "紧凑任务列表可能不含完整 result，因此模型和数据集字段以 result 为先、config 为后备。",
        "comparison.csv 使用固定核心列，便于外部表格工具横向比较，不改变各任务原始 JSON。",
    ])

    doc.add_heading("10.8 报告目录与下载制品", level=2)
    doc.add_paragraph(
        "report_catalog 以批测任务为主索引，汇总完成、失败或取消、可评估数量和原始报告总体积。每项记录包含实验组合、配置范围、核心指标、总体判定摘要、错误和制品可用性，供报告中心决定按钮状态。"
    )
    add_table(doc, ["制品", "内容", "生成与可用条件", "用途"], [
        ["原始批测 JSON", "完整任务配置、进度、指标、时延、关联、清单和逐航迹结果", "任务创建后即有文件，执行中持续更新", "复现与机器处理"],
        ["综合评估 JSON", "原始任务、四项判定、证据状态、阈值和来源", "completed 且包含 result", "正式评估归档"],
        ["逐航迹 CSV", "track、valid_frames、RMSE、MAE、P95 和最大误差", "completed 且 evaluated_tracks 大于 0", "离线分布分析"],
        ["任务对比 CSV", "最多 50 个任务的统一核心指标", "至少选中一个可对比的已完成任务", "横向比较"],
        ["浏览器打印或 PDF", "报告中心当前详情视图", "由前端打印样式生成", "人工审阅和提交"],
    ], widths=[1.25, 2.5, 1.75, 1.3], small=True)
    doc.add_paragraph(
        "assessment_bundle 的 schema_version 为 1.0，kind 为 tracking_batch_assessment，并记录 generated_at、job、assessment 和 provenance。provenance 包含固定判定条件、调用方提供的基线以及“IMM-JPDA 基线必须使用相同数据集和设置”的要求。"
    )

    doc.add_heading("10.9 审慎判定规则", level=2)
    add_bullets(doc, [
        "原始批测结果与派生评估分离；生成评估不会回写或覆盖 result。",
        "只对 completed 且包含 result 的任务评估，failed、cancelled、queued 和 running 均不生成结论。",
        "历史参考不自动视为当前证据；独立证据必须通过任务签名、有效标志和有限数校验。",
        "手动 baseline_rmse_m 必须大于 0，只用于技术复核；即使计算结果达标，也不能替代签名匹配的改善率独立证据。",
        "注册数据中的 AssociationStats 只用于技术复核；即使数值达标，也不能替代签名匹配的关联正确率独立证据。",
        "任何已计算指标失败时总体结论为 failed；证据不足时为 incomplete，不以其余指标通过掩盖缺项。",
        "时延优先采用重复实测 measured_frame_ms；降级使用 estimated_frame_ms 时必须在 note 中说明口径。",
    ])

    doc.add_heading("11 数据与持久化详细设计", level=1)
    doc.add_paragraph(
        "系统不使用关系数据库。模型权重和注册场景以文件形式只读访问，运行中的活动模型与回放缓存保存在进程内存，批测任务和验收证据以 JSON 或 CSV 持久化。数据设计通过固定目录、注册表、受控文件名和原子写入保持输入来源与评估结果之间的可追溯关系。"
    )

    doc.add_heading("11.1 数据分类与目录边界", level=2)
    add_table(doc, ["数据类别", "默认位置", "读写方式", "主要责任模块"], [
        ["原始场景数据", "data/val_data_v5 至 v8 对应目录", "离线生成，运行时只读", "数据生成与配准脚本"],
        ["注册场景数据", "data/val_data_*_registered", "运行时只读", "TransformerEngine 与 BatchTestManager"],
        ["V8 模型兼容数据", "data/val_data_v8_interference_registered_10hz", "从40 Hz注册数据同步抽取，运行时只读", "兼容数据生成脚本与场景目录"],
        ["模型权重", "results 对应训练运行目录/best_model.pt", "运行时只读", "ModelManager 与 TransformerEngine"],
        ["训练指标摘要", "权重同目录/metrics_summary.json", "运行时只读，可缺省", "ModelManager"],
        ["批测任务报告", "artifacts/reports/batch_tests/{job_id}.json", "创建后持续原子更新", "BatchTestManager"],
        ["验收证据", "artifacts/metrics/acceptance/{job_id}/{metric_id}.json", "白名单脚本生成，评估只读", "EvaluationService"],
        ["历史评估参考", "artifacts/metrics 下固定 JSON", "离线生成，接口只读", "EvaluationService"],
    ], widths=[1.25, 2.65, 1.45, 1.45], small=True)
    doc.add_paragraph(
        "DATASET_CATALOG 和 MODEL_CATALOG 是可访问数据的注册边界。接口接收逻辑 ID，不接受任意本机路径；服务在注册项内解析真实目录和文件。原始数据仅作为数据来源说明，在线推理和批测只读取已完成目标配准的 PKL。"
    )

    doc.add_heading("11.2 注册 PKL 场景数据", level=2)
    doc.add_paragraph(
        "每个注册 PKL 表示一个场景，可包含多条 RegisteredTrack。场景级数据保存 t_grid、平台位置与姿态等公共信息；航迹级数据保存目标真值、各传感器和数据链观测及有效标志。_parse_registered_pkl 将每条目标航迹展开为统一模型样本，并从有限、递增的 t_grid 差值中取中位数作为 dt；时间轴不足时才回退为 0.1 s。"
    )
    add_table(doc, ["数据层级", "关键内容", "完整性要求", "使用方式"], [
        ["场景文件", "文件名、t_grid、UAV_Pos、姿态或旋转矩阵", "文件可反序列化；时间轴与平台维一致", "场景列表、分类和公共张量构造"],
        ["RegisteredTrack", "target_pos、target_vel、Radar、EO、ESM、FastLink 等", "至少存在一条可解析目标；字段长度与时间轴一致", "每条目标展开为一个模型样本"],
        ["观测记录", "量测、Meas_ECEF、valid、源时间戳", "无效项由 valid 标识；字段缺失或结构不符时当前场景或任务失败", "时空对齐、目标关联和模型输入"],
        ["标签与掩码", "label_abs、target_valid_mask", "掩码与序列长度一致；有效点产生的非有限结果由运行时或批测拒绝", "模型监督、回放过滤和稳定阶段统计"],
        ["采样间隔", "dt 或由 t_grid 推导", "必须为正；不足时间轴时用于重建显示时间", "状态转移、速度差分和时间显示"],
    ], widths=[1.15, 2.35, 2.1, 1.2], small=True)
    doc.add_paragraph(
        "单场景运行时对缺失的 IMU、快慢链位置、龄期、更新标志和 ESM 模式补同形零张量；这种补齐只用于保持模型接口稳定，不改写源 PKL。批测数据集以固定 100 帧和 5 架无人机组织样本，场景与目标的逻辑来源通过文件名和目标索引保留。"
    )
    doc.add_paragraph(
        "V8 原始注册数据采用 40 Hz、400 帧和 0.025 s 时间间隔。现有模型不重新训练，运行时改用 build_v8_10hz_compat.py 生成的 10 Hz 兼容视图：按 0、0.1 至 9.9 s 选择原始索引 0、4 至 396，所有首维等于时间轴长度的真值、平台姿态、量测、链路和有效标志使用同一索引，不进行量测平均。原始 40 Hz 文件保留不变，compatibility_manifest.json 记录转换规则和源文件对应关系。"
    )

    doc.add_heading("11.3 运行时数据对象", level=2)
    add_table(doc, ["对象", "主要字段", "存储位置", "生命周期"], [
        ["TrackingSequence", "scenario_id、scenario_name、target_index、times、truth_states、predicted_states、position_errors、两类时延", "TransformerEngine 内存缓存", "随活动引擎存在；模型切换或进程退出后释放"],
        ["活动模型状态", "engine、model_id、activated_at、last_load_ms", "ModelManager 内存", "成功激活后原子替换"],
        ["批测任务", "id、name、status、时间、config、progress、result、error", "内存任务表和同名 JSON", "创建后长期保留，服务启动时恢复"],
        ["评估结果", "criteria、overall_status、evidence、dimensions", "按请求派生；综合包可下载", "不回写原始批测 result"],
    ], widths=[1.35, 3.0, 1.4, 1.45], small=True)
    doc.add_paragraph(
        "TrackingSequence 为冻结数据类，frame_count 和 max_time 由数组计算。缓存键为场景绝对路径与 target_index，防止不同目标共享结果。批测任务对外返回深拷贝，调用方不能通过获得的字典修改管理器内部状态。"
    )

    doc.add_heading("11.4 模型权重与指标摘要", level=2)
    add_table(doc, ["模型 ID", "权重位置", "训练数据", "验证数据"], [
        ["v5", "results/runs_standard(v5)/run_20260715_194636/best_model.pt", "train_data_v5_registered", "val_data_v5_registered"],
        ["v6", "results/runs_30off(v6)/run_20260717_152359/best_model.pt", "train_data_v6_registered", "val_data_v6_registered"],
        ["v7", "results/runs_40off(v7)/run_20260728_103728/best_model.pt", "train_data_v7_registered", "val_data_v7_registered"],
    ], widths=[0.7, 3.1, 1.5, 1.5], center_cols=(0,), small=True)
    doc.add_paragraph(
        "ModelManager.refresh 只登记目录中实际存在的 best_model.pt。权重同目录可提供 metrics_summary.json，系统优先读取 metrics_wo_warmup，其次读取 metrics_all，并提取 RMSE、MAE、P95、最大误差、完整模式时延、训练轮数和 d_model。摘要缺失或 JSON 无效时仅返回空指标，不影响权重登记。"
    )
    add_bullets(doc, [
        "权重通过 torch.load 的 weights_only 模式加载，并以 strict 等于 true 校验状态字典。",
        "模型登记信息同时包含文件大小、UTC 修改时间、训练运行编号和当前活动标志。",
        "默认权重必须属于 MODEL_CATALOG 且文件存在；否则 ModelManager 初始化失败。",
        "模型目录和权重按只读制品管理，正式交付后不得用同名文件覆盖已验收权重。",
    ])

    doc.add_heading("11.5 批测报告与验收证据", level=2)
    add_table(doc, ["文件", "顶层结构", "关联键", "完整性机制"], [
        ["批测任务 JSON", "任务基本信息、config、progress、result、error", "文件名与 id 相同", "临时文件写入后 os.replace；UTF-8、缩进 JSON 且禁止 NaN 和 Inf"],
        ["关联证据 JSON", "metric_id、job_id、job_signature、value、valid_for_assessment 和关联明细", "job_id 与 metric_id", "白名单脚本生成；加载时重算任务签名"],
        ["改善率证据 JSON", "上述公共字段、IMM-JPDA RMSE 和改善率", "job_id 与 rmse_improvement_pct", "基线必须为有限正数；签名必须一致"],
        ["综合评估 JSON", "schema_version、kind、generated_at、job、assessment、provenance", "内含 job_id", "由当前原始任务按请求重新生成"],
        ["逐航迹 CSV", "track、valid_frames、rmse_m、mae_m、p95_m、max_error_m", "任务下载路径", "只从 completed 任务的 tracks 生成"],
    ], widths=[1.35, 2.6, 1.4, 1.45], small=True)
    doc.add_paragraph(
        "任务报告是批测事实记录，证据和综合评估属于派生制品。result.provenance 保存模型权重、入选数据文件和运行代码的 SHA-256 摘要。修改执行配置、路径、位置 RMSE、场景清单、模型文件、数据内容或运行代码会改变 job_signature，使旧证据变为不可接受；评估服务不尝试自动修复或迁移旧证据。"
    )

    doc.add_heading("11.6 命名 索引与完整性约束", level=2)
    add_table(doc, ["对象", "命名或索引规则", "边界检查"], [
        ["单场景 ID", "按当前数据目录排序，从 100 起顺序编号", "数字索引必须落在场景列表范围内"],
        ["场景文件名", "仅接受当前目录中的 PKL 主名或文件名", "Path.name 必须等于输入，拒绝目录组件"],
        ["目标索引", "从 0 开始", "必须小于解析出的目标样本数"],
        ["批测任务 ID", "uuid4 十六进制字符串", "作为内存键、报告文件名和证据目录名"],
        ["数据集与模型 ID", "固定注册表键", "未知 ID 返回错误，不从路径猜测"],
        ["场景类别", "由文件名正则识别 CV、CA、CT、CSSW 或干扰类别", "无法识别时归为 OTHER 并保留复核"],
    ], widths=[1.25, 3.25, 2.3], small=True)
    add_bullets(doc, [
        "模型输出位置误差必须全部为有限数，出现 NaN 或 Inf 时场景失败且不写入缓存。",
        "批测模型输出、稳定阶段误差、逐航迹误差、汇总误差和时延必须全部为有限数，否则任务失败且不写入完成结果。",
        "批测在稳定阶段掩码后必须至少保留一个有效目标点，否则任务失败。",
        "历史 JSON 无法解析或缺少任务 ID 时在启动恢复中忽略，不让单个坏文件阻止服务启动。",
        "列表接口提供摘要，详情接口提供完整任务；下载制品只能由后端从受控对象生成。",
    ])

    doc.add_heading("11.7 数据生命周期与备份", level=2)
    add_numbered_steps(doc, [
        "数据准备阶段生成原始场景并完成配准，将注册 PKL 放入 DATASET_CATALOG 指定目录。",
        "模型训练阶段把 best_model.pt 和可选 metrics_summary.json 固化到 MODEL_CATALOG 指定运行目录。",
        "运行阶段只读模型和注册数据；单场景结果仅缓存在内存，批测任务从创建时开始写入持久化 JSON。",
        "验收阶段基于 completed 任务生成独立证据、综合评估和 CSV，并保持任务 ID 对应关系。",
        "归档阶段整体备份权重、注册数据清单、批测 JSON、证据 JSON 和导出文件；备份后不覆盖原文件。",
    ])
    doc.add_paragraph(
        "系统没有自动删除任务、证据或模型的逻辑。运维人员如需清理，应先确认制品未用于正式验收，并按任务 ID 同步处理任务报告与证据目录；正式交付记录应优先采用只读介质或具备权限控制的归档位置。"
    )

    doc.add_heading("12 接口设计", level=1)
    api_rows = [
        ["GET", "/api/health", "200", "触发模型初始化并返回健康摘要"],
        ["GET", "/api/system", "200", "返回运行时 设备 活动模型和队列"],
        ["POST", "/api/models/activate", "200 404 422", "加载模型并原子切换"],
        ["GET", "/api/scenarios", "200", "按当前活动数据目录列出场景"],
        ["POST", "/api/batch-tests", "202 404 422", "验证配置 持久化并排队"],
        ["POST", "/api/batch-tests/{id}/cancel", "200 404", "请求取消任务"],
        ["GET", "/api/batch-tests/{id}/assessment", "200 404 409", "评估已完成任务"],
        ["POST", "/api/batch-tests/{id}/evidence/{metric}", "200 404 409 422", "运行白名单独立脚本"],
        ["GET", "/api/reports", "200", "生成持久化任务索引"],
        ["WS", "/ws/tracking", "1008", "推理失败时以策略违规代码关闭连接"],
    ]
    add_table(doc, ["方法", "路径", "状态", "设计说明"], api_rows, widths=[0.7, 3.1, 1.0, 2.0], center_cols=(0, 2), small=True)
    doc.add_heading("12.1 WebSocket 帧结构", level=2)
    add_table(doc, ["字段", "说明"], [
        ["time  max_time", "当前时间和序列最大时间"], ["is_reset", "运行 RMSE 是否从当前帧开始重新累计"],
        ["scenario  target_index", "场景名和目标索引"], ["truth_pos  pred_pos", "三维绝对真值位置和预测位置"],
        ["metrics.err_pos", "当前帧三维位置误差"], ["metrics.rmse_current", "自最近重置点开始的运行 RMSE"],
        ["metrics.inference_time_ms", "完整序列时延折算的单帧值"],
    ], widths=[2.2, 4.6])

    doc.add_heading("13 并发与状态一致性详细设计", level=1)
    doc.add_paragraph(
        "系统同时存在 HTTP 请求线程、WebSocket 异步会话、模型推理工作线程、批测后台线程和证据脚本子进程。共享状态仅集中在服务单例、ModelManager、BatchTestManager 和 EvaluationService 中，并分别使用锁、事件、不可变结果对象和原子文件替换控制一致性。"
    )

    doc.add_heading("13.1 共享状态与同步对象", level=2)
    add_table(doc, ["共享对象", "同步机制", "保护范围", "读取方式"], [
        ["三个服务单例", "MANAGER_INIT_LOCK 可重入锁与 lru_cache", "首次创建和依赖顺序", "工厂函数返回同一实例"],
        ["活动模型状态", "_state_lock 可重入锁", "模型表、引擎、活动 ID、激活时间和加载耗时", "在锁内复制引用或标量"],
        ["模型激活过程", "_activation_lock 普通互斥锁", "候选引擎构建与最终替换的完整流程", "同一时刻只允许一个激活请求"],
        ["批测任务表", "BatchTestManager._lock 可重入锁", "jobs、cancel_events 和 futures", "详情和列表返回深拷贝"],
        ["任务取消", "每任务 threading.Event", "排队或运行任务的协作式停止", "执行阶段在固定检查点读取"],
        ["验收证据生成", "EvaluationService._evidence_lock", "白名单脚本执行和目标证据文件", "生成完成后重新读取校验"],
        ["回放缓存", "TransformerEngine._lock", "场景解析、推理和缓存写入", "缓存命中时返回冻结结果对象"],
    ], widths=[1.3, 1.65, 2.45, 1.4], small=True)

    doc.add_heading("13.2 服务单例初始化", level=2)
    doc.add_paragraph(
        "get_model_manager、get_batch_test_manager 和 get_evaluation_service 均使用 maxsize 等于 1 的 lru_cache，并在缓存未命中时进入 MANAGER_INIT_LOCK。创建顺序固定为模型管理器、批测管理器、评估服务，后两者复用前一层实例。"
    )
    add_numbered_steps(doc, [
        "第一个需要模型状态的请求触发 ModelManager 创建并加载默认活动模型。",
        "首次访问批测接口时创建 BatchTestManager，并读取报告目录恢复历史任务。",
        "首次访问评估接口时创建 EvaluationService，并注入同一个 BatchTestManager。",
        "应用关闭时仅在批测管理器已实际创建的情况下调用 shutdown，避免退出阶段反向触发初始化。",
    ])

    doc.add_heading("13.3 模型激活一致性", level=2)
    doc.add_paragraph(
        "模型激活采用先构建候选引擎、后替换活动引用的两阶段方式。_activation_lock 将整个操作串行化，_state_lock 只保护短时间状态读取和提交。候选权重加载、结构严格校验或设备迁移失败时不会修改当前活动模型。"
    )
    add_table(doc, ["阶段", "锁与操作", "一致性保证"], [
        ["读取当前状态", "activation_lock 内短暂获取 state_lock", "取得模型路径、当前引擎和活动 ID 的一致快照"],
        ["幂等判断", "同模型且同设备直接返回 changed 为 false", "避免重复加载和无意义缓存失效"],
        ["候选构建", "保持 activation_lock，不持有 state_lock", "读取请求不被长时间模型加载阻塞"],
        ["提交切换", "state_lock 内同时替换 engine、model_id、时间和加载耗时", "活动信息不会出现新旧字段混合"],
    ], widths=[1.2, 2.75, 2.85])
    doc.add_paragraph(
        "已开始的 WebSocket 会话在连接建立时取得引擎引用，因此继续使用原模型完成当前回放；切换后的新连接取得新引擎。批测任务创建时固化 model_id 和 device，排队期间的活动模型变化不改变任务配置。"
    )

    doc.add_heading("13.4 批测任务并发控制", level=2)
    add_bullets(doc, [
        "ThreadPoolExecutor 的 max_workers 固定为 1，GPU 密集任务串行执行，任务创建接口仍可并发接收。",
        "jobs、cancel_events 和 futures 在同一可重入锁下登记，避免任务可见但取消事件尚未建立。",
        "状态和进度更新在锁内完成，磁盘持久化使用锁内深拷贝后的快照，缩短临界区。",
        "取消事件在数据索引、模型加载前后和每批推理开始处检查；单次前向完成后再响应取消。",
        "终态任务不会重新进入执行器；重复取消终态任务返回当前任务快照。",
    ])
    doc.add_paragraph(
        "任务状态按 queued、running 和终态单向变化。completed、failed、cancelled 均为终态；progress.percent 只用于展示，不能作为状态迁移条件。"
    )

    doc.add_heading("13.5 持久化一致性", level=2)
    doc.add_paragraph(
        "_persist 先在任务锁内对目标任务执行深拷贝，再将 JSON 写入带线程标识的临时文件，最后通过 os.replace 替换正式报告。读取者只会观察到旧快照或新快照，不会读到半个 JSON。"
    )
    add_bullets(doc, [
        "任务创建后先持久化再提交后台执行器，服务异常退出时仍保留 queued 记录。",
        "执行阶段每 10 个批次或最后一个批次持久化进度，关键阶段切换也单独写入。",
        "服务重启将非终态历史任务标记为 failed，不自动续跑无法确认上下文的计算。",
        "评估读取 BatchTestManager 返回的深拷贝，证据生成后再次读取任务并重新校验签名。",
    ])

    doc.add_heading("13.6 证据生成互斥", level=2)
    doc.add_paragraph(
        "关联准确率和改善率脚本共享一个 evidence_lock。锁覆盖子进程调用，防止两个生成请求同时写入证据目录或争用大量计算资源。脚本完成后，服务不直接相信退出码，而是重新读取证据并验证 metric_id、job_id、任务签名、有效标志和有限数。"
    )

    doc.add_heading("13.7 WebSocket 会话状态", level=2)
    doc.add_paragraph(
        "每个 WebSocket 连接建立独立 control 字典、接收协程和运行 RMSE 累计列表，不与其他连接共享播放进度。接收协程只更新 paused 和 seek_target；发送循环读取这些值并维护 index。asyncio 事件循环内不存在跨线程写入 control 的操作。"
    )
    add_table(doc, ["控制动作", "状态变化", "发送循环行为"], [
        ["pause", "paused 置为 true", "每 100 ms 检查一次，不推进 index"],
        ["play", "paused 置为 false", "从当前 index 继续按 50 ms 间隔发送"],
        ["seek", "seek_target 记录目标时间", "按 times 二分定位，清空 RMSE 累计并从新位置发送"],
        ["断开", "接收或发送出现 WebSocketDisconnect", "结束循环并取消接收协程"],
    ], widths=[1.1, 2.5, 3.2])

    doc.add_heading("13.8 故障隔离与资源释放", level=2)
    add_bullets(doc, [
        "单个 HTTP 请求异常由中间件转换为 500，不终止服务进程。",
        "单个 WebSocket 场景解析或推理失败只关闭当前连接，不改变活动模型和批测任务。",
        "批测异常只把当前任务置为 failed，后台执行器可继续处理后续排队任务。",
        "批测 finally 删除局部模型引用并在 CUDA 可用时调用 empty_cache。",
        "应用关闭请求取消所有非终态批测，并以 cancel_futures 关闭执行器。",
        "WebSocket finally 取消接收协程并抑制预期的 CancelledError，避免遗留任务。",
    ])

    doc.add_heading("14 异常处理与安全边界详细设计", level=1)
    doc.add_paragraph(
        "异常处理以明确失败为原则。系统拒绝不可用模型、越界设备、未知场景和不完整证据，不通过静默降级生成看似有效的结果。接口响应保留定位所需的 detail 或 request_id，完整堆栈只进入服务器日志。"
    )

    doc.add_heading("14.1 异常分类与对外映射", level=2)
    add_table(doc, ["异常来源", "典型异常", "接口结果", "处理原则"], [
        ["资源查找", "未知模型、数据集或任务 ID", "HTTP 404", "说明目标不存在，不创建替代对象"],
        ["任务状态", "未完成任务评估、逐航迹结果不可用", "HTTP 409", "保留现有状态，要求调用方等待或改选任务"],
        ["参数与运行条件", "目录缺失、设备不可用、非法基线、模型加载失败", "HTTP 422 或 409", "在进入计算或证据使用前拒绝"],
        ["未处理异常", "路由外的意外运行错误", "HTTP 500", "返回通用说明和 request_id，日志记录堆栈"],
        ["回放输入或推理", "场景、目标、模型或数值错误", "WebSocket 1008", "关闭当前会话并截断 reason 到 120 字符"],
        ["批测内部异常", "数据为空、有效点为空、CUDA 或前向异常", "任务 failed", "写入异常类型、信息和终止阶段"],
    ], widths=[1.25, 2.35, 1.2, 2.0], small=True)

    doc.add_heading("14.2 参数与状态校验", level=2)
    add_table(doc, ["输入", "校验位置", "规则"], [
        ["模型激活请求", "Pydantic 与 ModelManager", "model_id 和 device 非空；模型已登记；CUDA 可用且索引存在"],
        ["批测创建请求", "Pydantic 与 BatchTestManager", "名称长度、预热帧、批大小、装载进程和场景上限满足范围；模型和数据集存在"],
        ["列表 limit", "FastAPI Query", "按各接口定义限制最小值和最大值"],
        ["baseline_rmse_m", "FastAPI Query 与 EvaluationService", "必须大于 0；只用于 completed 任务"],
        ["target", "FastAPI Query 与场景解析", "必须为非负整数且小于目标数"],
        ["证据 metric_id", "EvaluationService 白名单", "只允许 association_accuracy 和 rmse_improvement_pct"],
    ], widths=[1.4, 2.2, 3.4], small=True)

    doc.add_heading("14.3 文件路径与目录边界", level=2)
    add_bullets(doc, [
        "场景只能通过当前活动数据目录的数字 ID 或 PKL 文件名访问；Path.name 与原输入不一致时拒绝。",
        "模型切换只接受 MODEL_CATALOG 中已发现的模型 ID，不允许接口直接提交权重路径。",
        "批测数据集只接受 DATASET_CATALOG 注册项，报告目录由后端常量确定。",
        "证据脚本名由 metric_id 白名单映射，调用方不能提交命令、脚本路径或输出路径。",
        "下载文件名由任务 ID 和模型 ID构造，服务端不采用用户提交的名称作为物理路径。",
    ])
    doc.add_paragraph(
        "这些约束限制了原型接口的文件访问面，但不等同于操作系统级沙箱。部署账户仍应只授予模型、注册数据和报告目录所需的最小读写权限。"
    )

    doc.add_heading("14.4 模型与数值安全", level=2)
    add_table(doc, ["风险", "检测或保护", "失败行为"], [
        ["权重结构不匹配", "strict 等于 true 加载状态字典", "候选引擎创建失败，活动模型不变"],
        ["非受控反序列化范围", "权重使用 weights_only；场景仅从注册目录读取", "文件异常时当前场景或任务失败"],
        ["CUDA 不可用或索引越界", "_select_device 显式检查", "拒绝请求，不自动改用 CPU"],
        ["模型误差出现 NaN 或 Inf", "单场景检查全部 position_errors", "抛出运行错误且不缓存"],
        ["批测出现 NaN 或 Inf", "逐级检查模型输出、入选误差、航迹统计、总体统计和时延；JSON 禁止非有限数", "任务标记 failed，不生成 completed 报告"],
        ["批测无有效稳定点", "stable_stage_mask 后检查 point_errors", "任务 failed，不生成零值指标"],
        ["模型内部数值不稳定", "正则项、分母下限、速度限幅和缺失向量", "尽量保持有限计算；最终仍由输出检查兜底"],
    ], widths=[1.55, 3.2, 2.05], small=True)

    doc.add_heading("14.5 证据脚本执行边界", level=2)
    add_bullets(doc, [
        "脚本由服务端固定字典选择，使用 sys.executable 和参数列表启动，不拼接用户命令。",
        "工作目录固定为项目根目录，输出固定到当前 job_id 与 metric_id 对应路径。",
        "单次执行超时 3600 秒，标准输出和错误均被捕获，进程返回码非零即判失败。",
        "证据必须与当前任务签名一致并标记 valid_for_assessment，不能仅凭文件存在使用。",
        "改善率证据中的 IMM-JPDA RMSE 必须是有限正数，防止除零或无意义比较。",
    ])

    doc.add_heading("14.6 日志与错误信息", level=2)
    doc.add_paragraph(
        "HTTP 中间件为每个请求生成或透传 X-Request-ID，并在响应中返回处理耗时。每个请求完成后以信息级别记录方法、路径、响应状态、处理耗时和请求标识；未处理异常通过 LOGGER.exception 记录请求标识和完整堆栈。对外 500 响应只包含 Internal server error 与 request_id。业务错误保留可读 detail，便于前端在原操作位置提示。"
    )
    add_bullets(doc, [
        "日志和错误响应不得输出权重内容、场景数据内容或完整证据载荷。",
        "路径错误可用于本机联调，但面向不可信网络部署时应进一步隐藏绝对路径。",
        "批测 error 保存异常类型与消息，不保存 Python 堆栈；详细诊断依赖服务器日志。",
        "前端只显示后端 detail 或通用提示，不将异常文本作为 HTML 执行。",
    ])

    doc.add_heading("14.7 部署安全边界", level=2)
    doc.add_paragraph(
        "当前软件定位为本机或受控内网原型，没有账户、身份认证、角色授权、传输加密、多租户隔离、接口限流和远程文件上传能力。后端默认监听 127.0.0.1，CORS 只允许本机 5173 端口，符合单机联调边界。"
    )
    add_table(doc, ["扩展场景", "必须补充的控制"], [
        ["跨主机内网访问", "HTTPS 或受控反向代理、身份认证、明确 CORS 白名单和网络访问控制"],
        ["多用户并发", "角色权限、任务归属、资源配额、操作审计和报告访问控制"],
        ["公网访问", "TLS、强认证、限流、输入大小限制、安全响应头、依赖漏洞管理和持续监控"],
        ["外部数据上传", "文件类型与大小校验、隔离解析、恶意内容检测和独立存储区"],
    ], widths=[1.7, 5.1])

    doc.add_heading("15 配置与部署详细设计", level=1)
    doc.add_paragraph(
        "软件采用前后端分离的本机部署方式。后端从项目根目录读取模型、数据和报告路径，前端由 Vite 提供开发服务或生成 dist 静态文件。启动顺序为后端、健康检查、前端，避免界面加载后长时间处于不可用状态。"
    )

    doc.add_heading("15.1 配置来源与优先级", level=2)
    add_table(doc, ["配置项", "默认值", "读取位置", "作用与约束"], [
        ["TRACKER_MODEL_PATH", "V7 best_model.pt", "ModelManager 与 TransformerEngine", "覆盖默认权重，但路径必须对应注册模型"],
        ["TRACKER_DATA_DIR", "data/val_data_v7_registered", "ModelManager 与 TransformerEngine", "覆盖单场景目录；必须存在且含注册 PKL"],
        ["TRACKER_DEVICE", "auto", "ModelManager 与 TransformerEngine", "auto 优先 cuda:0，否则 CPU；显式 CUDA 不可用时失败"],
        ["TRACKER_SEQ_LEN", "100", "TransformerEngine", "转换为整数并控制单场景解析长度"],
        ["后端主机与端口", "127.0.0.1:8000", "server.py 启动入口", "限制本机访问并提供 REST 与 WebSocket"],
        ["前端开发端口", "5173", "Vite", "与后端 CORS 白名单保持一致"],
    ], widths=[1.55, 1.65, 2.0, 1.8], small=True)
    doc.add_paragraph(
        "显式构造参数优先于环境变量，环境变量优先于代码默认值。ModelManager 构造活动引擎时把已解析的权重、数据目录和设备显式传入，避免同一实例内重复读取变化中的环境变量。环境变量修改后需要重启后端。"
    )

    doc.add_heading("15.2 运行目录结构", level=2)
    add_table(doc, ["目录", "必须内容", "运行权限"], [
        ["tracking_system/backend", "server.py、runtime.py、model_manager.py、batch_testing.py、evaluation.py", "只读"],
        ["tracking_system/frontend", "前端源代码、package.json；生产交付可使用 dist", "开发时可构建，运行时源代码只读"],
        ["src", "CoopTrackerModel 与注册数据解析实现", "只读"],
        ["results", "MODEL_CATALOG 指向的 best_model.pt 和可选 metrics_summary.json", "只读"],
        ["data", "DATASET_CATALOG 指向的注册 PKL", "只读"],
        ["artifacts/reports/batch_tests", "批测任务 JSON", "后端进程可创建和替换"],
        ["artifacts/metrics/acceptance", "按任务分组的验收证据 JSON", "证据脚本与后端可写"],
    ], widths=[2.05, 3.25, 1.7], small=True)

    doc.add_heading("15.3 后端启动与健康检查", level=2)
    add_numbered_steps(doc, [
        "进入项目根目录并使用已包含 PyTorch、CUDA、FastAPI、Uvicorn、WebSocket、SciPy 和 h5py 的 pytorch 环境。",
        "运行 python -m tracking_system.backend.server；启动入口监听 127.0.0.1 的 8000 端口。",
        "访问 /api/health，确认 status 为 ready，并检查 active_model、device 和 batch_testing。",
        "访问 /api/system，复核 Python、PyTorch、CUDA Runtime、设备列表、活动模型和报告目录。",
        "访问 /docs 查看 FastAPI 自动接口文档，必要时执行只读接口联通检查。",
    ])
    doc.add_paragraph(
        "首次健康检查会触发 ModelManager 初始化和默认权重加载，耗时可能明显高于普通查询。默认模型或数据目录不存在、权重结构不匹配或显式 CUDA 不可用时，健康检查返回失败信息。"
    )

    doc.add_heading("15.4 前端构建与启动", level=2)
    add_table(doc, ["方式", "命令", "输出或用途"], [
        ["开发服务", "npm run dev", "Vite 启动本地页面并支持热更新"],
        ["生产构建", "npm run build", "先执行 vue-tsc 类型检查，再生成 dist"],
        ["构建预览", "npm run preview", "本地检查 dist 的静态部署结果"],
    ], widths=[1.25, 2.0, 3.55])
    doc.add_paragraph(
        "前端依赖 Vue、ECharts 和 ECharts GL，构建工具为 TypeScript、vue-tsc 与 Vite。交付环境如直接托管 dist，不需要运行 Vite 开发服务器，但后端地址发现和 CORS 必须与实际部署地址同步调整。"
    )

    doc.add_heading("15.5 端口发现与跨域配置", level=2)
    add_bullets(doc, [
        "当前前端按浏览器主机名组合 8000 端口，分别构造 HTTP 与 WebSocket 地址。",
        "后端默认只监听 127.0.0.1:8000，外部主机不能直接连接。",
        "CORS 只允许 http://127.0.0.1:5173 和 http://localhost:5173，方法限制为 GET 与 POST。",
        "更改前端端口、域名、协议或反向代理路径时，必须同步修改后端白名单和前端地址构造。",
        "HTTPS 页面不能直接连接明文 ws 或 http；采用 TLS 时应同时提供 https 和 wss。",
    ])

    doc.add_heading("15.6 部署前检查", level=2)
    add_table(doc, ["检查项", "通过条件", "失败处理"], [
        ["目录完整性", "src、backend、frontend 或 dist、注册数据和权重存在", "补齐交付文件或修正注册路径"],
        ["Python 环境", "核心依赖可导入，PyTorch 可识别预期设备", "切换正确环境或修复依赖"],
        ["模型权重", "默认权重属于注册表且 strict 加载成功", "恢复匹配的 best_model.pt"],
        ["数据集", "默认注册目录存在且至少包含一个可解析 PKL", "恢复注册数据或调整受控配置"],
        ["端口", "8000 和所选前端端口未被占用", "停止冲突进程或同步修改配置"],
        ["报告权限", "batch_tests 与 acceptance 目录可创建和原子替换文件", "授予最小写权限"],
        ["接口联通", "health、system、models、datasets 和 WebSocket 可访问", "根据 request_id 和后端日志定位"],
    ], widths=[1.25, 3.55, 2.0], small=True)

    doc.add_heading("15.7 交付环境调整", level=2)
    doc.add_paragraph(
        "现有设计面向单机开发与验收。若交付环境要求前端静态托管或跨主机访问，应把地址、证书、反向代理、CORS、日志目录和服务账户权限作为同一部署变更处理，并重新执行前端构建、接口联通、WebSocket 回放、批测和报告下载检查。"
    )
    add_bullets(doc, [
        "不得只修改监听地址而保留无鉴权接口直接暴露。",
        "模型和数据目录应只读挂载，报告与证据目录单独授予写权限。",
        "部署脚本应固定工作目录，确保相对路径解析到项目根目录。",
        "正式验收前记录操作系统、Python、PyTorch、CUDA、GPU 和活动模型信息。",
    ])

    doc.add_heading("16 测试与维护详细设计", level=1)
    doc.add_paragraph(
        "测试分为纯函数单元测试、后端服务单元测试、前端静态检查与构建、运行中后端黑盒验收以及完整数据集指标验收。前四类验证软件功能和边界，完整指标验收用于形成合同技术指标结论。"
    )

    doc.add_heading("16.1 测试层级与职责", level=2)
    add_table(doc, ["层级", "主要对象", "依赖", "通过含义"], [
        ["算法与数据纯函数", "场景分类、抽样、稳定阶段掩码、关联和机动统计", "NumPy 或小型张量", "基础规则在确定输入下可重复"],
        ["后端服务单元测试", "评估判定、证据校验、报告目录、接口辅助逻辑", "临时目录和模拟任务", "错误分支和持久化规则符合设计"],
        ["前端静态检查", "Vue 组件和 TypeScript 类型", "Node 依赖", "代码可通过类型检查并完成打包"],
        ["后端黑盒验收", "真实 HTTP、WebSocket、模型、数据、批测和下载", "运行中的后端与可用模型数据", "完整软件链路可执行"],
        ["完整指标验收", "四项技术指标及独立证据", "完整注册数据集和同条件 IMM-JPDA 基线", "可生成正式 passed、failed 或 incomplete 结论"],
    ], widths=[1.25, 2.45, 1.6, 1.7], small=True)

    doc.add_heading("16.2 后端单元测试", level=2)
    doc.add_paragraph(
        "现有自动化测试直接覆盖验收判定三态与正式证据门禁、证据签名变化、重启恢复回写、批测非有限数拒绝、最近帧跳转、V8 40 Hz 到 10 Hz 同步抽取、跨数据集矩阵、分层抽样、稳定阶段掩码、目标关联、ONNX 安全求解及评审辅助统计。模型权重发现、完整批测和全部 HTTP 路由主要通过黑盒验收覆盖，不将其表述为已被单元测试完全覆盖。"
    )
    add_table(doc, ["测试集合", "重点验证"], [
        ["tracking_system/backend/tests", "批测恢复与数值安全、评估服务、场景目录及服务端时间跳转辅助逻辑"],
        ["tests/test_review_items_1_3.py", "随机丢测、全源遮挡和链路延迟的确定性"],
        ["tests/test_review_items_4_6.py", "场景分类、干扰施加和传感器组合聚合"],
        ["tests/test_review_items_7_9.py", "稳定阶段误差、判定独立性和单帧批处理"],
        ["tests/test_review_items_10_12.py", "匀速与转弯机动强度统计及非法输入"],
        ["tests/test_track_association.py", "时间对齐、一对一分配、跨传感器和跨平台关联"],
        ["tests/test_onnx_safe_solve.py", "三阶安全求解的数值与导出兼容性"],
        ["tests/test_v8_10hz_compat.py", "V8 同步抽取索引、字段对齐、采样周期和全局时间"],
    ], widths=[2.7, 4.1])

    doc.add_heading("16.3 前端静态检查与构建", level=2)
    add_numbered_steps(doc, [
        "在 tracking_system/frontend 安装锁定依赖，保持 package-lock.json 与 package.json 一致。",
        "执行 npm run build，由 vue-tsc 完成项目类型检查，再由 Vite 生成 dist。",
        "确认构建退出码为 0，dist 包含入口页面和静态资源，且没有未处理的类型错误。",
        "使用 npm run preview 或正式静态服务器检查六个页面、刷新、直接 Hash 访问和资源加载。",
    ])
    doc.add_paragraph(
        "ECharts GL 相关资源体积较大属于已知构建特征。若调整分包或延迟加载，应重新验证三维跟踪图初始化、页面切换和首次进入跟踪验证页的加载状态。"
    )

    doc.add_heading("16.4 后端黑盒验收", level=2)
    doc.add_paragraph(
        "在后端运行状态下执行 python -m tracking_system.backend.acceptance。该入口按客户端方式访问服务，检查健康状态、运行环境、模型与设备、场景读取、真实 WebSocket 推理、指标配置、历史参考、跨数据集结果、小规模批测、多维评估及 JSON 与 CSV 下载。"
    )
    add_bullets(doc, [
        "默认使用 V7 模型、V7 数据集和首个场景，重点验证链路而非完整算法结论。",
        "验收脚本必须以退出码 0 结束，并输出后端功能验收通过。",
        "失败时先按请求 ID、任务 error 和后端日志定位，再决定是否重跑。",
        "真实小批量测试形成的任务报告应标明 max_scenarios，不得冒充完整数据集结果。",
    ])

    doc.add_heading("16.5 完整指标验收", level=2)
    add_numbered_steps(doc, [
        "选择已确认的模型、数据集和设备，提交 max_scenarios 为空的完整批测任务。",
        "等待任务状态变为 completed，核对场景清单、类别数量、预热帧、有效点和硬件环境。",
        "生成 association_accuracy 独立证据，并确认 accepted 为 true。",
        "在相同数据集、场景清单、序列长度和预热设置下运行 IMM-JPDA 基线，生成 rmse_improvement_pct 证据。",
        "复核位置 RMSE、关联准确率、相对改善率和单帧时延四项 criterion 及其 note。",
        "下载原始批测 JSON、综合评估 JSON、逐航迹 CSV 和必要的任务对比 CSV，整体归档。",
    ])
    doc.add_paragraph(
        "EvaluationService 以 formal_evidence_complete 记录两份独立证据是否同时 accepted。只有 overall_status 为 passed、formal_evidence_complete 为 true，且两份证据均与当前任务签名一致时，才能形成四项指标全部通过的结论。AssociationStats 和手工 baseline_rmse_m 即使得到达标数值，也只能用于技术复核，不能将总体状态提升为 passed。incomplete 表示证据不足，failed 表示至少一项已计算指标未达到阈值。"
    )

    doc.add_heading("16.6 日常维护", level=2)
    add_table(doc, ["维护对象", "检查内容", "处理原则"], [
        ["模型权重", "文件存在、大小和修改时间是否异常，能否 strict 加载", "保留已验收权重，不用同名文件覆盖"],
        ["注册数据", "文件数量、可解析性、时间轴和目标数", "源数据与注册数据分开，只读保存"],
        ["任务报告", "JSON 可解析、状态终态、制品可下载", "按任务 ID 归档并定期备份"],
        ["验收证据", "accepted、job_signature 和基线来源", "与对应任务一起保存；任务变化后重新生成"],
        ["运行环境", "Python、PyTorch、CUDA、GPU、端口和磁盘空间", "环境变化后执行黑盒验收和代表性批测"],
        ["前端构建", "依赖锁文件、类型检查和 dist", "依赖升级后完整重建与页面回归"],
    ], widths=[1.25, 3.05, 2.5], small=True)

    doc.add_heading("16.7 扩展与变更控制", level=2)
    add_bullets(doc, [
        "新增模型时扩展 MODEL_CATALOG，提供独立 ID、权重路径、训练与验证数据说明，并验证 strict 加载。",
        "新增数据集或验收想定时扩展 DATASET_CATALOG，说明 raw_path、注册路径、场景类别、传感器和质量边界。",
        "新增场景命名类别时同步修改分类正则、SCENARIO_TAXONOMY、抽样测试和前端筛选。",
        "新增判定指标时同步修改 ACCEPTANCE_CRITERIA、评估逻辑、证据校验、接口类型、页面展示和下载结构。",
        "修改模型输入字段时同步更新解析、单场景运行时、批测 DataLoader、模型 forward、文档和自动化测试。",
        "修改接口字段或状态枚举时保持前后端类型一致，并对旧报告读取策略作出明确决定。",
    ])

    doc.add_heading("16.8 交付检查", level=2)
    add_table(doc, ["检查项", "交付要求"], [
        ["源代码", "前端、后端、核心模型和必要脚本完整，可在交付环境读取"],
        ["模型与数据", "注册权重和注册验证数据路径与目录清单一致"],
        ["构建与启动", "后端健康检查通过，前端生产构建成功，六个页面可访问"],
        ["接口与回放", "REST 查询、WebSocket 回放、模型切换和场景选择正常"],
        ["批测与评估", "任务状态、持久化、取消、四项判定、证据和报告下载正常"],
        ["异常与边界", "未知资源、非法设备、越界目标、不可用制品和证据失效均明确失败"],
        ["文档", "设计说明、需求规格和使用说明与当前软件行为一致"],
        ["归档", "正式报告、证据、环境信息和数据清单已备份，交付文件只保留确认稿"],
    ], widths=[1.4, 5.4])

    doc.add_heading("附录 关键设计参数", level=1)
    add_table(doc, ["参数", "值", "位置"], [
        ["服务版本", "1.4.0", "tracking_system backend server.py"],
        ["模型特征维度", "d_model 96", "runtime.py 和 src model.py"],
        ["序列长度", "100", "TransformerEngine 和批测任务配置"],
        ["无人机数量", "5", "TransformerEngine 和批测数据集"],
        ["协同邻接距离", "20 km", "CoopTrackerModel"],
        ["批测工作线程", "1", "BatchTestManager"],
        ["证据脚本超时", "3600 s", "EvaluationService"],
        ["WebSocket 回放间隔", "50 ms", "server.py"],
        ["总览自动刷新间隔", "30 s", "App.vue"],
    ], widths=[1.8, 1.3, 3.7])

    apply_design_body_typography(doc)

    # Word requires a paragraph after the final table.  Keep that structural
    # paragraph to one point so it does not spill onto a blank trailing page.
    trailing = doc.paragraphs[-1]
    trailing.paragraph_format.space_before = Pt(0)
    trailing.paragraph_format.space_after = Pt(0)
    trailing.paragraph_format.line_spacing = Pt(1)
    marker = trailing.add_run("\u200b")
    marker.font.size = Pt(1)
    marker.font.hidden = True
    path = OUTPUT_DIR / "端到端协同跟踪原型软件_设计说明书.docx"
    doc.save(path)
    return path


def build_user_manual():
    doc = Document()
    title = f"{SOFTWARE_NAME} 使用说明书"
    configure_document(doc, title)
    add_cover(doc, "使用说明书", "TS DOC UM 001")
    contents = [
        (1, "1 软件简介"), (1, "2 使用前准备"), (1, "3 启动与登录"),
        (1, "4 界面导航"), (1, "5 系统总览"), (1, "6 模型管理"),
        (1, "7 跟踪验证"), (1, "8 批量测试"), (1, "9 性能评估"),
        (1, "10 报告中心"), (1, "11 正式指标验收流程"), (1, "12 高级配置"),
        (1, "13 常见问题"), (1, "14 数据与报告管理"), (1, "15 关闭软件"),
        (1, "附录 接口与状态速查"),
    ]
    add_document_control(
        doc,
        "本说明书面向软件操作人员和验收人员，说明端到端协同跟踪原型软件的安装启动 页面操作 批测评估 报告导出和故障处理方法",
        contents,
    )

    doc.add_heading("1 软件简介", level=1)
    doc.add_paragraph("端到端协同跟踪原型软件用于加载 V5 V6 V7 三类训练权重，回放单场景跟踪结果，对 V5 V6 V7 V8 和验收想定执行批量测试，并依据合同及任务书指标生成多维评估和报告。主界面包含系统总览 模型管理 跟踪验证 批量测试 性能评估和报告中心。")
    doc.add_heading("1.1 主要功能", level=2)
    add_bullets(doc, [
        "查看后端运行环境 当前模型 CUDA 状态和批测队列。",
        "刷新模型仓库并在 CPU 或 CUDA 设备上部署指定权重。",
        "按机动类型选择场景和目标，查看三维真值与预测航迹及实时误差。",
        "创建 抽样或完整批测任务，查看进度 历史 结果和逐航迹指标。",
        "按四项指标自动判定，生成关联正确率和 IMM JPDA 基线证据。",
        "检索归档任务并导出 JSON CSV 或浏览器打印 PDF。",
    ])

    doc.add_heading("2 使用前准备", level=1)
    doc.add_heading("2.1 软件与硬件", level=2)
    add_table(doc, ["项目", "要求"], [
        ["操作系统", "Windows 10 或 Windows 11 本机环境"],
        ["Python 环境", "现有 Conda pytorch 环境，或安装 requirements.txt 中的后端依赖"],
        ["Node.js 和 npm", "能够安装并运行 tracking_system frontend 中的前端依赖"],
        ["浏览器", "支持现代 JavaScript WebSocket 和 WebGL 的桌面浏览器"],
        ["计算设备", "CPU 可运行；正式时延测试建议使用已正确安装驱动的 CUDA GPU"],
        ["磁盘", "应为模型 数据集 node_modules 和批测报告预留足够空间"],
    ], widths=[1.5, 5.3])
    doc.add_heading("2.2 必需文件", level=2)
    add_bullets(doc, [
        "results 目录中的 V5 V6 V7 best_model.pt 权重。",
        "data 目录中的注册验证集，例如 val_data_v7_registered。",
        "tracking_system backend 和 frontend 源代码。",
        "src 目录中的模型和数据解析模块。",
    ])
    doc.add_heading("2.3 端口", level=2)
    add_table(doc, ["端口", "用途", "默认地址"], [
        ["8000", "后端 API 和 WebSocket", "http://127.0.0.1:8000"],
        ["5173", "Vite 前端开发服务", "http://127.0.0.1:5173"],
    ], widths=[1.0, 2.8, 3.0], center_cols=(0,))

    doc.add_heading("3 启动与登录", level=1)
    doc.add_paragraph("本软件是本地原型系统，不设置账户登录。操作前需要分别启动后端和前端。所有命令均从 transformer_v1 项目根目录执行。")
    doc.add_heading("3.1 启动后端", level=2)
    add_numbered_steps(doc, [
        "打开 Anaconda Prompt 或已初始化 Conda 的 PowerShell。",
        "进入 transformer_v1 项目根目录。",
        "启动 FastAPI 后端。首次启动会加载默认 V7 权重。",
    ])
    add_command(doc, "conda run --no-capture-output -n pytorch python -m tracking_system.backend.server")
    doc.add_paragraph("看到 Uvicorn 服务启动信息后，在浏览器访问下列地址检查状态。")
    add_command(doc, "http://127.0.0.1:8000/api/health")
    doc.add_paragraph("正常响应中的 status 应为 ready。可访问 http://127.0.0.1:8000/docs 查看交互式 API 文档。")
    doc.add_heading("3.2 启动前端", level=2)
    add_numbered_steps(doc, [
        "另开一个 PowerShell 窗口并进入 tracking_system frontend。",
        "首次使用或依赖变化后执行 npm install。",
        "执行 npm run dev 启动前端。",
        "浏览器打开终端显示的本地地址，通常为 http://127.0.0.1:5173。",
    ])
    add_command(doc, "cd tracking_system/frontend")
    add_command(doc, "npm install")
    add_command(doc, "npm run dev")
    doc.add_heading("3.3 启动检查", level=2)
    add_bullets(doc, [
        "左侧底部显示后端服务在线。",
        "系统总览显示 API 版本 当前模型和推理设备。",
        "模型仓库至少列出一个可用权重，数据集列表显示预期场景数量。",
        "如显示后端未连接，先确认 8000 端口和 pytorch 环境中的后端进程。",
    ])

    doc.add_heading("4 界面导航", level=1)
    add_table(doc, ["导航", "用途", "页面地址"], [
        ["系统总览", "检查运行态势和技术指标", "#overview"],
        ["模型管理", "查看和部署模型", "#models"],
        ["跟踪验证", "单场景实时回放", "#tracking"],
        ["批量测试", "创建和管理批测任务", "#batch"],
        ["性能评估", "技术指标判定和对比", "#evaluation"],
        ["报告中心", "检索 复核和导出报告", "#reports"],
    ], widths=[1.3, 3.65, 1.85])
    doc.add_paragraph("桌面端通过左侧导航切换页面。窄屏设备可使用顶部菜单按钮。浏览器地址中的 Hash 可直接打开指定页面。右下方设置按钮用于查看 API 模型 运行时 CUDA 报告目录和后端启动命令。")

    doc.add_heading("5 系统总览", level=1)
    add_numbered_steps(doc, [
        "打开系统总览，先确认服务状态为在线。",
        "核对当前模型 设备 CUDA 状态和活动任务数量。",
        "查看运行环境中的 Python PyTorch 平台和显卡信息。",
        "查看任务书指标区域，确认阈值为 RMSE 8 m 关联率 85% 基线降幅 15% 和单帧 100 ms。",
        "需要立即更新时单击刷新数据。页面也会每 30 秒自动同步一次。",
    ])
    doc.add_paragraph("总览中的模型历史指标仅用于快速识别模型，不替代本次批测和正式验收报告。")

    doc.add_heading("6 模型管理", level=1)
    doc.add_heading("6.1 查看模型", level=2)
    doc.add_paragraph("模型仓库显示 V5 V6 V7 的名称 版本 训练集 验证集 权重文件 文件大小 修改时间和可用历史指标。ACTIVE MODEL 标记表示当前单场景推理使用的模型。")
    doc.add_heading("6.2 切换模型", level=2)
    add_numbered_steps(doc, [
        "在部署设备区域选择 auto CPU 或可用的 CUDA 设备。",
        "在目标模型卡片上单击部署。当前模型的部署按钮不可用。",
        "在确认窗口核对模型版本和设备，然后单击确认部署。",
        "等待成功提示，检查当前模型和加载耗时是否更新。",
    ])
    doc.add_paragraph("模型切换期间系统先加载新权重。加载失败时旧模型继续工作。已经开始的 WebSocket 回放使用原模型，新建立的回放使用新模型。")
    doc.add_heading("6.3 刷新仓库", level=2)
    doc.add_paragraph("新增或恢复权重文件后单击刷新模型仓库。刷新只重新检查注册路径，不会自动改变当前模型。")

    doc.add_heading("7 跟踪验证", level=1)
    doc.add_heading("7.1 开始回放", level=2)
    add_numbered_steps(doc, [
        "选择机动类型。默认筛选 CT，也可选择全部 CV CA CT CSSW。",
        "选择场景。场景 ID 从 100 开始，与当前数据目录按文件名排序的位置对应。",
        "输入目标索引。第一个目标为 0，越界索引会被后端拒绝。",
        "单击重新推理。后端对完整序列推理完成后开始逐帧回放。",
    ])
    doc.add_heading("7.2 查看结果", level=2)
    add_bullets(doc, [
        "三维图同时显示真实航迹和预测航迹。俯视 侧视和自由按钮用于切换观察角度。",
        "实时指标显示当前位置误差 运行 RMSE 单帧推理时间和完整序列推理时间。",
        "误差时间序列显示随时间变化的三维位置误差。",
        "三维图使用首帧真值作为显示原点，读数和导出仍是绝对坐标。",
    ])
    doc.add_heading("7.3 控制回放", level=2)
    add_numbered_steps(doc, [
        "单击圆形控制按钮暂停或继续。",
        "暂停后拖动时间轴并结束拖动，系统跳转到目标时间。",
        "跳转后运行 RMSE 重新累计，is_reset 对应第一帧。",
        "回放完成或需要重新计算时单击重新推理。",
    ])
    doc.add_heading("7.4 导出当前数据", level=2)
    doc.add_paragraph("接收到至少一帧后单击导出当前数据。浏览器下载 tracking 场景 target 目标索引 csv，内容包括时间 真值坐标 预测坐标 位置误差和运行 RMSE。")

    doc.add_heading("8 批量测试", level=1)
    doc.add_heading("8.1 创建任务", level=2)
    add_numbered_steps(doc, [
        "在任务名称中输入便于识别的名称，名称可留空。",
        "选择测试模型和计算设备。批测使用独立模型实例，不要求先切换活动模型。",
        "选择数据集或验收想定。确认其状态为可用并查看场景数 类型 模态和干扰说明。",
        "选择抽样测试或完整测试。抽样测试填写场景数，系统按场景类型均衡选取。",
        "设置预热帧数 批大小和数据线程。默认建议为 5 64 和 0。",
        "单击创建批测任务。成功后页面显示任务 ID 并进入运行队列。",
    ])
    doc.add_heading("8.2 参数说明", level=2)
    add_table(doc, ["参数", "范围", "说明"], [
        ["任务名称", "0 到 100 个字符", "留空时自动生成模型与数据集名称"],
        ["预热帧数", "0 到 99", "每条航迹前若干帧不计入精度统计"],
        ["批大小", "1 到 256", "一次前向推理包含的航迹数，受显存限制"],
        ["数据线程", "0 到 8", "PyTorch DataLoader 工作进程数量"],
        ["场景数", "至少 1", "仅抽样测试使用，不得超过数据集场景数"],
    ], widths=[1.35, 1.5, 3.95])
    doc.add_heading("8.3 查看和取消任务", level=2)
    add_bullets(doc, [
        "运行队列显示当前 queued 和 running 任务及百分比。",
        "任务历史可按全部 活动 已完成和异常筛选。",
        "单击任务行或队列卡片查看配置 进度 指标 时延和场景清单。",
        "排队中或运行中的任务可在详情中单击取消任务。取消请求在当前安全检查点生效。",
        "任务完成后可直接下载原始 JSON 和逐航迹 CSV。",
    ])

    doc.add_heading("9 性能评估", level=1)
    doc.add_heading("9.1 自动判定", level=2)
    add_numbered_steps(doc, [
        "选择一个已完成的批测任务。",
        "如已有同条件 IMM JPDA 基线 RMSE，可输入大于 0 的数值；否则留空。",
        "单击重新判定。",
        "查看总体状态和四项指标。全部有值且满足阈值时为全部通过；存在未达标项时为存在未达标项；仅缺证据时为证据不完整。",
    ])
    doc.add_heading("9.2 生成独立证据", level=2)
    doc.add_paragraph("关联正确率或基线降幅显示未生成时，可在证据区域单击运行测试。系统使用后端当前 Python 环境运行白名单脚本，完成后自动校验任务签名并刷新判定。完整数据集证据可能耗时较长。也可复制页面给出的命令在项目根目录执行。")
    doc.add_heading("9.3 查看多维结果", level=2)
    add_bullets(doc, [
        "精度维度包括 RMSE MAE 中位数 P95 最大误差和平均航迹 RMSE。",
        "时延维度包括批次 航迹 单帧实测或估算时延和吞吐率。",
        "航迹分布显示中位数 P75 P90 P95 最大值和 8 m 内航迹比例。",
        "机动类型汇总按 CV CA CT CSSW 和 OTHER 展示。",
        "跨数据集矩阵和历史基线标注为参考，正式判定以当前任务和同条件证据为准。",
    ])

    doc.add_heading("10 报告中心", level=1)
    doc.add_heading("10.1 检索和筛选", level=2)
    add_numbered_steps(doc, [
        "打开报告中心并单击同步归档。",
        "可按任务名称 ID 模型或数据集搜索。",
        "可按模型和全部 已完成 异常状态筛选。",
        "单击报告行右侧箭头查看详情。",
    ])
    doc.add_heading("10.2 报告复核", level=2)
    doc.add_paragraph("详情显示任务配置 核心指标 四项判定 证据状态和可用制品。如需用手工提供的同条件 IMM JPDA RMSE 复核，输入基线值后重新判定。清空基线并重新判定可恢复证据驱动的原始状态。")
    doc.add_heading("10.3 导出", level=2)
    add_table(doc, ["操作", "结果", "使用条件"], [
        ["下载原始 JSON", "完整批测任务文件", "任务报告文件存在"],
        ["下载综合评估 JSON", "任务 四项判定 证据和阈值来源", "任务已完成"],
        ["下载逐航迹 CSV", "每条航迹的精度指标", "任务已完成且有逐航迹结果"],
        ["导出对比 CSV", "勾选任务的统一横向指标", "至少选择一个已完成任务 最多 50 个"],
        ["打印或 PDF", "当前报告详情的浏览器打印版", "浏览器支持打印"],
    ], widths=[1.55, 3.15, 2.1])

    doc.add_heading("11 正式指标验收流程", level=1)
    doc.add_paragraph("正式验收应使用完整数据集和与合同一致的设备 环境及测试条件。建议按以下顺序执行。")
    add_numbered_steps(doc, [
        "在模型管理中确认待验收权重可用，并记录设备和运行环境。",
        "在批量测试中选择待验收模型和目标数据集，选择完整测试，提交任务并等待完成。",
        "在性能评估中运行关联正确率测试，确认测试证据 accepted。",
        "运行相对 IMM JPDA 降幅测试，确认基线使用同一数据集 场景清单 序列长度和预热帧设置。",
        "确认位置 RMSE 不超过 8 m 关联正确率不低于 85% 基线降幅不低于 15% 单帧推理时间不超过 100 ms。",
        "在报告中心下载原始 JSON 综合评估 JSON 逐航迹 CSV 和必要的对比 CSV，连同软件版本和环境记录归档。",
    ])
    doc.add_paragraph("如只运行一个或少量场景，结果用于链路检查和问题定位，不作为完整算法指标结论。")

    doc.add_heading("12 高级配置", level=1)
    doc.add_paragraph("可在启动后端前设置以下环境变量。修改后需要重启后端。")
    add_table(doc, ["环境变量", "示例", "说明"], [
        ["TRACKER_MODEL_PATH", "D 路径 best_model.pt", "覆盖默认模型路径，路径仍需匹配注册模型"],
        ["TRACKER_DATA_DIR", "D 路径 val_data_v7_registered", "覆盖单场景验证数据目录"],
        ["TRACKER_DEVICE", "cuda:0 或 cpu", "指定默认推理设备"],
        ["TRACKER_SEQ_LEN", "100", "指定单场景序列长度"],
    ], widths=[1.65, 2.45, 2.7])
    doc.add_paragraph("批测任务的序列长度当前固定为 100。前端后端地址由当前浏览器主机名和 8000 端口组成，如调整端口需同步修改前端 API 适配。")

    doc.add_heading("13 常见问题", level=1)
    faq_rows = [
        ["页面显示后端未连接", "后端未启动 端口被占用或模型冷启动失败", "访问 /api/health，检查后端终端和 8000 端口，修复后单击重新连接"],
        ["启动提示找不到模型权重", "默认或环境变量指定路径不存在", "恢复 V7 best_model.pt 或检查 TRACKER_MODEL_PATH"],
        ["CUDA 不可用", "驱动 PyTorch CUDA 版本或设备索引不匹配", "在设置中检查 CUDA Runtime，改用 cpu 或修复 CUDA 环境"],
        ["模型部署失败", "权重缺失 结构不匹配或显存不足", "查看后端错误，刷新仓库，确认权重与 d_model 96 模型匹配"],
        ["场景或目标越界", "当前数据目录中没有该 ID 或目标索引", "重新选择场景，目标索引从 0 开始"],
        ["三维图无法显示", "浏览器未提供 WebGL", "启用硬件加速或更换支持 WebGL 的浏览器"],
        ["批测长时间排队", "已有 GPU 密集任务运行", "查看运行队列，等待完成或按需取消前一任务"],
        ["证据不完整", "独立脚本未运行 失败或任务签名变化", "在性能评估中重新运行对应测试并查看证据消息"],
        ["下载逐航迹 CSV 失败", "任务未完成或没有逐航迹数据", "等待任务完成并确认 evaluated_tracks 大于 0"],
        ["后端重启后任务变失败", "任务未完成时服务退出", "重新提交任务。已完成任务报告仍可查询"],
    ]
    add_table(doc, ["现象", "可能原因", "处理方法"], faq_rows, widths=[1.65, 2.25, 2.9], small=True)

    doc.add_heading("14 数据与报告管理", level=1)
    add_table(doc, ["内容", "默认位置", "建议"], [
        ["注册验证数据", "data val_data 版本 registered", "保持只读并记录数据集版本"],
        ["模型权重", "results 各训练运行 best_model.pt", "不得覆盖已验收权重"],
        ["批测报告", "artifacts reports batch_tests", "定期备份正式任务 JSON"],
        ["验收证据", "artifacts metrics acceptance 任务 ID", "与对应任务报告一并归档"],
        ["交叉评估", "artifacts metrics cross_dataset_rmse.json", "仅作历史参考并记录生成条件"],
    ], widths=[1.35, 3.15, 2.3])
    doc.add_paragraph("删除或修改批测报告会影响报告中心索引和证据签名。正式验收完成后，建议将任务 JSON 证据 JSON 导出 CSV 软件版本和环境信息整体备份。")

    doc.add_heading("15 关闭软件", level=1)
    add_numbered_steps(doc, [
        "确认没有需要保留的运行中批测任务。关闭后端会请求取消未完成任务。",
        "在前端终端按 Ctrl C 停止 Vite 服务。",
        "在后端终端按 Ctrl C 停止 Uvicorn 服务。",
        "检查正式报告已写入 artifacts reports batch_tests 并完成备份。",
    ])

    doc.add_heading("附录 接口与状态速查", level=1)
    add_table(doc, ["项目", "值"], [
        ["健康检查", "http://127.0.0.1:8000/api/health"],
        ["接口文档", "http://127.0.0.1:8000/docs"],
        ["前端地址", "http://127.0.0.1:5173"],
        ["任务状态", "queued  running  completed  failed  cancelled"],
        ["指标状态", "passed  failed  unavailable"],
        ["总体状态", "passed  failed  incomplete"],
        ["WebSocket 控制", "pause  play  seek"],
    ], widths=[1.6, 5.2])

    path = OUTPUT_DIR / "端到端协同跟踪原型软件_使用说明书_V1.0.docx"
    doc.save(path)
    return path


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    architecture = make_software_architecture_image()
    outputs = [
        build_design_doc(architecture),
        build_requirements_doc(),
        build_user_manual(),
    ]
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
